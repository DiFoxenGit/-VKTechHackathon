"""Full business HTTP acceptance; uses trusted imports, real templates and exports."""
import argparse
from decimal import Decimal
import json
import time
from pathlib import Path
import httpx
from pptx import Presentation
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'samples/business-acceptance/data-v1'
OUT=ROOT/'output/business-review'


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://127.0.0.1:8018/api/v1')
    parser.add_argument('--only',default='');parser.add_argument('--updated',action='store_true');args=parser.parse_args()
    OUT.mkdir(exist_ok=True,parents=True)
    report={'scenario':'full synthetic business acceptance','image_provider_live':False,'runs':[]}
    with httpx.Client(base_url=args.url,timeout=240) as c:
        def req(method,path,**kw):
            r=c.request(method,path,**kw)
            if not r.is_success:raise RuntimeError((r.status_code,r.text[:1000]))
            return r
        templates=list({t.get('sha256',t['id']):t for t in req('GET','/templates').json()['items']}.values())
        if args.only:templates=[t for t in templates if args.only.lower() in t['name'].lower()]
        assert templates,'No matching templates'
        packs=[]
        names=['sales.csv','plan.csv','products.csv','regions.csv','catalog.csv','business-context.txt','content-pack.zip']
        for name in names:
            directory=DATA/'updated' if args.updated and name in {'sales.csv','plan.csv'} else DATA
            r=req('POST','/content-packs',files={'file':(name,(directory/name).read_bytes())}).json()
            packs.append(r['id'])
        brief=(ROOT/'samples/business-acceptance/README.md').read_text(encoding='utf8').split('## Запрос для приложения')[1].split('## Состав')[0]
        brief='\n'.join(line[2:] if line.startswith('> ') else '' if line=='> ' or line=='>' else line for line in brief.splitlines()).strip()
        expected_net='117501720' if args.updated else '116241720'
        for template in templates:
            started=time.monotonic()
            job=req('POST','/generations',json={'brief':brief,'template_id':template['id'],'slide_count':12,
                'content_pack_ids':packs,'narrative':True,'generate_images':True}).json()
            deadline=time.monotonic()+300
            while job['status'] not in {'completed','failed'} and time.monotonic()<deadline:
                time.sleep(.3);job=req('GET','/jobs/'+job['id']).json()
            assert job['status']=='completed',job
            for identifier in job['presentation_ids']:
                record=req('GET','/presentations/'+identifier).json()
                base=OUT/('updated-' if args.updated else '')/(template['name'].replace('.pptx','')+'-'+record['variant'])
                base.parent.mkdir(exist_ok=True,parents=True)
                assert len(record['deck']['slides'])==12
                assert record['image_generation']['status']=='unavailable',record['image_generation']
                results=record['analysis'];assert len(results)==12
                kpi=next(r for r in results if r.get('business_view')=='quarter_kpi')
                assert any(Decimal(f['exact'])==Decimal(expected_net) for f in kpi['facts'])
                for fmt in ('pptx','pdf'):
                    base.with_suffix('.'+fmt).write_bytes(req('GET',f'/presentations/{identifier}/export/{fmt}').content)
                prs=Presentation(base.with_suffix('.pptx'));pdf=PdfReader(base.with_suffix('.pdf'))
                assert len(prs.slides)==len(pdf.pages)==12
                charts=0;tables=0
                for native,planned in zip(prs.slides,record['deck']['slides']):
                    v=planned['content']['visual']
                    if v['kind'] in {'line','bar','pie','grouped_bar','stacked_bar'}:
                        chart=next(s.chart for s in native.shapes if s.has_chart);charts+=1
                        assert [list(s.values) for s in chart.series]==[s['values'] for s in v['series']]
                    if v['kind']=='table':
                        assert any(s.has_table for s in native.shapes);tables+=1
                slides=[s['content'] for s in record['deck']['slides']]
                saved=req('PATCH','/presentations/'+identifier,json={'revision':1,'title':record['title'],'variant':record['variant'],'slides':slides}).json()
                assert saved['revision']==2
                bad=json.loads(json.dumps(slides));bad[0]['title']='Подменённый вывод'
                rejected=c.patch('/presentations/'+identifier,json={'revision':2,'title':record['title'],'variant':record['variant'],'slides':bad})
                assert rejected.status_code==422
                used={e['asset']['id'] for s in record['deck']['slides'] for e in s['elements'] if e.get('asset',{}).get('source')=='pack'}
                base.with_suffix('.json').write_text(json.dumps(record,ensure_ascii=False,indent=2))
                assert len(used)>=3,(template['name'],record['variant'],used)
                item={'template':template['name'],'template_sha256':template.get('sha256'),'variant':record['variant'],'id':identifier,'slides':12,'native_charts':charts,
                    'native_tables':tables,'pack_assets':len(used),'saved_revision':2,'tampering_rejected':True,
                    'audit':record['audit']['counts'],'image_generation':record['image_generation']}
                report['runs'].append(item);print(json.dumps(item,ensure_ascii=False),flush=True)
            print('Template seconds',round(time.monotonic()-started,2),flush=True)
    (OUT/('updated-service-report.json' if args.updated else 'service-report.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
