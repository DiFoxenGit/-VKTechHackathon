"""Inspect real HTTP exports; no synthetic API/model responses or screenshots."""
import argparse
import io
import json
from pathlib import Path
import httpx
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2 as pdfium
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from designer.analytic_export import display_number


def inspect(url, ids, output):
    output.mkdir(parents=True,exist_ok=True)
    report={'url':url,'ids':ids,'runs':[]}
    with httpx.Client(base_url=url,timeout=180) as client:
        for identifier in ids:
            response=client.get('/presentations/'+identifier);response.raise_for_status();record=response.json()
            stem=identifier+'-r'+str(record['revision'])
            (output/(stem+'.json')).write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf8')
            files={}
            for fmt in ('pptx','pdf'):
                r=client.get(f'/presentations/{identifier}/export/{fmt}',params={'revision':record['revision']});r.raise_for_status()
                files[fmt]=r.content;(output/(stem+'.'+fmt)).write_bytes(r.content)
            native=Presentation(io.BytesIO(files['pptx']));pdf=PdfReader(io.BytesIO(files['pdf']))
            planned=record['deck']['slides']
            assert len(native.slides)==len(pdf.pages)==len(planned)==12
            assert planned[0]['content']['role']=='cover' and planned[0]['native_cover']
            assert any(s.has_text_frame and 'Коммерческие результаты' in s.text for s in native.slides[0].shapes)
            charts=tables=0
            for page,slide in zip(native.slides,planned):
                visual=slide['content']['visual'];kind=visual['kind']
                if kind in {'bar','line','pie','grouped_bar','stacked_bar','scatter','histogram'}:
                    chart=next(s.chart for s in page.shapes if s.has_chart);charts+=1
                    assert [list(series.values) for series in chart.series]==[s['values'] for s in visual['series']]
                    if kind not in {'pie','scatter'}:
                        assert chart.category_axis._element.xpath('./c:tickLblSkip/@val')==['1']
                if kind=='table':
                    table=next(s.table for s in page.shapes if s.has_table);tables+=1
                    actual=[[table.cell(i+1,j).text for j in range(len(visual['columns']))] for i in range(len(visual['rows']))]
                    assert actual==[[display_number(cell) for cell in row] for row in visual['rows']]
            assert charts==5 and tables==2
            selected={0,1,next(i for i,s in enumerate(planned) if s['content'].get('analysis_id','').startswith('business:regions:')),
                      next(i for i,s in enumerate(planned) if s['content'].get('analysis_id','').startswith('business:risks:'))}
            document=pdfium.PdfDocument(files['pdf'])
            for index in selected:
                page=document[index];bitmap=page.render(scale=1000/page.get_width())
                bitmap.to_pil().save(output/f'{stem}-page-{index+1}.png');page.close()
            document.close()
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
            assert 'Коммерческие' in text
            result={'id':identifier,'template_id':record['template_id'],'variant':record['variant'],'revision':record['revision'],
                    'slides':len(planned),'native_charts':charts,'native_tables':tables,'exact_chart_series':True,'exact_table_cells':True,
                    'audit':record['audit']['counts'],'image_generation':record.get('image_generation'),
                    'inference':{k:v for k,v in record.get('generation',{}).get('inference',{}).items() if k not in {'shared_inference'}}}
            report['runs'].append(result)
            (output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
            print(identifier,record['variant'],'PASS: 12 pages, native charts/tables, exact values, rendered PDF',flush=True)
    return report


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://127.0.0.1:18078/api/v1')
    parser.add_argument('--ids',nargs='+',required=True);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    inspect(args.url,args.ids,args.output)
