"""Reproducible visual QA, explicitly separate from full business acceptance.

Uses a 1,512-row synthetic probe and independently prepared aggregations to
exercise four native evidence types and mixed template/pack pages. Passing this
does NOT accept D1-D3: production joins, weighted ratios and story remain gates.
"""
import csv
import argparse
import hashlib
import io
import json
import math
import os
import sys
import subprocess
import time
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))
from designer.analytics_bridge import extract, analyse_sources, build_outline, verify_outline
from designer.assets import library, parse_pack_zip, save_assets
from designer.exporting import export_pptx, convert_pdf, render_slides, verify_pptx
from designer.layout import compose
from designer.parsing import parse_template, parse_content
from designer.storage import Store
from pptx import Presentation
from pypdf import PdfReader
from PIL import Image, ImageOps, ImageDraw

OUT = ROOT/'output/next-iteration'

def write_csv(name, columns, rows):
    target=OUT/'input'/name;target.parent.mkdir(parents=True,exist_ok=True)
    with target.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(columns);writer.writerows(rows)
    return target

def prepare():
    rows=[]
    for m in range(21):
        year=2025+m//12;month=m%12+1
        season=1+.13*math.sin((month-1)*math.pi/6)
        for region in range(6):
            for product in range(4):
                for channel in range(3):
                    gross=round((160000+region*19000+product*32000+channel*27000)*season*(1+m*(.009+product*.003)))
                    discount=round(gross*(.04+channel*.015))
                    refund=round(gross*(.007+region*.002))
                    net=gross-discount-refund
                    margin=.51-product*.055-channel*.035-(.07 if product==2 and m>=18 else 0)
                    cogs=round(net*(1-margin))
                    rows.append([f'{year}-{month:02}',f'R{region+1}',f'P{product+1}',['Прямые продажи','Партнёры','Онлайн'][channel],gross,discount,refund,cogs,round(gross/6500)])
    path=write_csv('sales.csv',['month','region_id','product_id','channel','gross_revenue_rub','discount_rub','refund_rub','cogs_rub','orders_count'],rows)
    assert len(rows)==1512 and len({tuple(r[:4]) for r in rows})==1512
    # Independent decimal reference, never imports production calculation helpers.
    monthly=defaultdict(Decimal);quarters=defaultdict(Decimal);products=defaultdict(Decimal);profits=defaultdict(Decimal)
    for r in rows:
        net=Decimal(r[4])-Decimal(r[5])-Decimal(r[6]);profit=net-Decimal(r[7])
        monthly[r[0]]+=net
        quarters[r[0][:4]+'-'+f'{((int(r[0][5:])-1)//3)*3+1:02}']+=net
        if '2026-07'<=r[0]<='2026-09':products[r[2]]+=net;profits[r[2]]+=profit
    names={'P1':'Командная работа','P2':'Облачная инфраструктура','P3':'Аналитика продаж','P4':'Сервисная платформа'}
    q3=sum(products.values());q2=quarters['2026-04'];prior=quarters['2025-07']
    expected={'synthetic':True,'rows':len(rows),'q3_net':str(q3),'q3_profit':str(sum(profits.values())),
        'q3_margin':str(sum(profits.values())/q3),'qoq_percent':str((q3/q2-1)*100),'yoy_percent':str((q3/prior-1)*100),
        'by_product':{names[k]:str(v) for k,v in products.items()}}
    (OUT/'input/expected-probe.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
    plan=[]
    for m in range(1,13):
        for region in range(6):
            for product in range(4):
                base=650000+region*57000+product*100000
                plan.append([f'2026-{m:02}',f'R{region+1}',f'P{product+1}',base,round(base*.4)])
    write_csv('plan.csv',['month','region_id','product_id','net_revenue_plan_rub','gross_profit_plan_rub'],plan)
    paths=[write_csv('kpi-q3.csv',['Период','Выручка после скидок и возвратов'],[['III квартал 2026',str(q3)]]),
        write_csv('dynamics.csv',['Месяц','Выручка'],[[k,str(v)] for k,v in list(monthly.items())[-12:]]),
        write_csv('products-q3.csv',['Продукт','Выручка'],[[names[k],str(v)] for k,v in products.items()]),
        write_csv('economics-q3.csv',['Продукт','Выручка, руб.','Валовая прибыль, руб.','Валовая маржа, %'],[[names[k],str(v),str(profits[k]),str(round(profits[k]/v*100,2))] for k,v in products.items()])]
    sources=[{'id':'brief','text':'Проверка вёрстки на синтетических данных','_story':{'goal':'Развитие продуктов в IV квартале: выручка, динамика и прибыльность'}}]
    for i,path in enumerate(paths):
        table=extract(path.read_bytes(),path.name)[0]
        sources.append({'id':'data'+str(i),'datasets':[table], 'analysis_requests':{table['dataset_id']:{'task':['summary','dynamics','comparison','reference'][i],'unit':'руб.' if i<3 else ''}},'analysis_goal':['Итог','Динамика','Сравнить','Справочная таблица'][i]})
    context='\n\n'.join([
        'Команды коммерции и продукта согласуют критерии отбора направлений на следующий квартал.',
        'Совместная работа команды опирается на полностью закрытый квартал и единый словарь показателей.',
        'Роли команды обсуждают ограничения внедрения до принятия решения о расширении.',
        'Продажи оцениваются после скидок и возвратов. Рост выручки требует отдельной проверки экономики.',
        'Динамика выручки показана за последние двенадцать месяцев. Последний месяц закрыт полностью.',
        'Результаты продаж сами по себе не показывают эффективность маркетинга и не позволяют рассчитать CAC.',
        'Риски оцениваются по продуктам и каналам. Изменение маржи не доказывает причину его возникновения.',
        'Контроль качества требует сверки периода, валюты и источника. Смешивать валюты без курса нельзя.',
        'Безопасность решения зависит от доступности подтверждённых данных и проверки ограничений.',
        'Следующий шаг: коммерческий директор и руководитель продукта проверяют маржу выбранных направлений до утверждения плана. Критерий: согласованный расчёт по каждому направлению.'
    ])
    sources.append({'id':'context','text':context})
    (OUT/'input/context.txt').write_text(context,encoding='utf-8')
    write_csv('products.csv',['product_id','name'],list(names.items()))
    write_csv('regions.csv',['region_id','name'],[[f'R{i+1}',n] for i,n in enumerate(['Москва и Московская область','Санкт-Петербург и Ленинградская область','Центральный федеральный округ','Уральский федеральный округ','Сибирский федеральный округ','Дальний Восток и удалённые территории'])])
    (OUT/'input/PROBE.md').write_text('Синтетический технический пробник: 1512 продаж и 288 строк плана. Это НЕ готовый D1: нет полного делового контекста, каталога и согласованного эталона 10–15 вопросов. Четыре агрегата рассчитаны вне приложения только для визуальной проверки S1/S2. Их загрузка не доказывает поддержку производственных объединений и маржи.',encoding='utf-8')
    return sources,path,expected

def contact_sheet(images,target):
    thumbs=[]
    for p in images:
        thumb=ImageOps.contain(Image.open(p).convert('RGB'),(640,360))
        canvas=Image.new('RGB',(660,400),'#e7ecf4');canvas.paste(thumb,((660-thumb.width)//2,15))
        ImageDraw.Draw(canvas).text((16,378),p.stem,fill='#243149');thumbs.append(canvas)
    result=Image.new('RGB',(1320,400*math.ceil(len(thumbs)/2)),'#e7ecf4')
    for i,im in enumerate(thumbs):result.paste(im,((i%2)*660,(i//2)*400))
    result.save(target)

def main():
    from run_final_server import load_server_env
    load_server_env()
    parser=argparse.ArgumentParser()
    parser.add_argument('--templates',type=Path,default=Path(os.getenv('DESIGNER_TEMPLATE_DIR',str(ROOT/'templates'))))
    parser.add_argument('--only',default='',help='Rerender only matching template names; retain the other validated runs')
    args=parser.parse_args()
    template_paths=list(args.templates.glob('*.pptx'))
    if len(template_paths)<3:
        parser.error('Provide --templates with the three original corporate PPTX templates')
    OUT.mkdir(parents=True,exist_ok=True);started=time.monotonic()
    sources,_,expected=prepare()
    results=analyse_sources(sources,'Анализ')
    (OUT/'analysis-contracts.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    outline=build_outline(results,sources,12).model_dump()
    assert {s['visual']['kind'] for s in outline['slides'] if s.get('analysis_id')}=={'kpi','line','bar','table'}
    report={'synthetic':True,'scope':'S1/S2 visual regression, NOT business acceptance','independent_probe':expected,'runs':[],'open_gates':['D1 complete dataset/oracle','D2 production joins/weighted ratios','D3 editorial story','S4 live image provider','A2 visual approval','S5 full scenario and video after approval']}
    report['base_commit']=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
    report['working_tree_changes']=True
    report['input_sha256']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (OUT/'input').glob('*.csv')}
    if args.only and (OUT/'visual-report.json').exists():
        previous=json.loads((OUT/'visual-report.json').read_text(encoding='utf-8'))
        report['runs']=[r for r in previous['runs'] if args.only.casefold() not in r['template'].casefold()]
    store=Store(OUT/'server-data')
    # Team-owned existing pack is disclosed; no organiser assets are fabricated.
    _,harvested=parse_pack_zip((ROOT/'samples/content-pack/visual-pack.zip').read_bytes(), parse_content)
    owner='f'*32;metas=save_assets(store.directory('content_packs',owner)/'assets',harvested)
    packs=[{'id':owner,'assets':metas}]
    for index,source in enumerate(template_paths):
        if args.only and args.only.casefold() not in source.name.casefold():continue
        digest=hashlib.sha256(source.read_bytes()).hexdigest()
        template=next((t for t in store.list('templates') if t.get('sha256')==digest),None)
        if not template:
            template=parse_template(source.read_bytes(),source.name);template['id']=str(index+1)*32
        base=OUT/f'template-{index+1}';base.mkdir(exist_ok=True)
        for variant in ['classic','split','focus']:
            deck=compose(outline,template,variant,library(template,packs))
            assert verify_outline({'slides':[s['content'] for s in deck['slides']]},results,sources)
            pptx=base/(variant+'.pptx');pdf=pptx.with_suffix('.pdf')
            export_pptx(source,template,deck,pptx,store.root);convert_pdf(pptx,pdf)
            assert len(Presentation(pptx).slides)==len(PdfReader(pdf).pages)==len(deck['slides'])
            reopened=Presentation(pptx)
            for native,planned in zip(reopened.slides,deck['slides']):
                visual=planned['content']['visual']
                if visual['kind'] in {'line','bar'}:
                    chart=next(shape.chart for shape in native.shapes if shape.has_chart)
                    assert [list(s.values) for s in chart.series]==[s['values'] for s in visual['series']]
            assets=[{'slide':i+1,'title':s['content']['title'],'usage':s.get('asset_usage',[])} for i,s in enumerate(deck['slides'])]
            used={e['asset']['id'] for s in deck['slides'] for e in s['elements'] if e.get('asset',{}).get('source')=='pack'}
            pngs=[]
            for n,png in enumerate(render_slides(pdf,960)):
                target=base/f'{variant}-{n+1:02}.png';target.write_bytes(png);pngs.append(target)
            contact_sheet(pngs,base/(variant+'-contact.png'))
            (base/(variant+'-asset-map.json')).write_text(json.dumps(assets,ensure_ascii=False,indent=2),encoding='utf-8')
            (base/(variant+'-deck.json')).write_text(json.dumps(deck,ensure_ascii=False,indent=2),encoding='utf-8')
            check=verify_pptx(pptx,len(deck['slides']))
            item={'template':source.name,'template_sha256':digest,'variant':variant,'slides':len(deck['slides']),'pack_assets':len(used),'native_check':check}
            report['runs'].append(item);print({k:v for k,v in item.items() if k!='native_check'},flush=True)
        if index==0:
            proof=[base/f'classic-{i+1:02}.png' for i,s in enumerate(deck['slides']) if s['content'].get('analysis_id')]
            contact_sheet(proof,OUT/'four-reference-slides.png')
    report['elapsed_seconds']=round(time.monotonic()-started,2)
    (OUT/'visual-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':main()
