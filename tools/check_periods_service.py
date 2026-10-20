"""Replay Q2 and B through HTTP, editable Office objects, PDF and revisions."""
import argparse
from decimal import Decimal
import io
import json
from pathlib import Path
import time

import httpx
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parents[1]
EXT=ROOT/'samples/business-acceptance/period-context-v1'
DATA=ROOT/'samples/business-acceptance/data-v1'


def read(path):return json.loads(path.read_text(encoding='utf8'))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--url',default='http://127.0.0.1:8026/api/v1')
    parser.add_argument('--output',type=Path,default=ROOT/'output/periods-service')
    parser.add_argument('--only',default='')
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    report={'synthetic':True,'image_provider_live':False,'runs':[]}
    with httpx.Client(base_url=args.url,timeout=240) as client:
        def req(method,path,**kw):
            response=client.request(method,path,**kw);response.raise_for_status();return response
        templates=[t for t in req('GET','/templates').json()['items'] if any(n in t['name'].lower() for n in ('vk tech','workspace','education'))]
        if args.only:templates=[t for t in templates if args.only.lower() in t['name'].lower()]
        assert templates,'Upload corporate PPTX templates first'
        for version in ('baseline','version-b'):
            folder=EXT/version;ids=[]
            for name in ('sales.csv','plan.csv','products.csv','regions.csv','context.txt','content-pack.zip'):
                path=(DATA if name in {'products.csv','regions.csv','content-pack.zip'} else folder)/name
                raw=path.read_bytes();pack=req('POST','/content-packs',files={'file':(name,raw)}).json();ids.append(pack['id'])
                assert req('GET',f'/content-packs/{pack["id"]}/source').content==raw
            request={'brief':'Чистая выручка, прибыль, маржа и план-факт. Приоритеты, ресурсы и действия по выбранным периодам.',
                'template_id':templates[0]['id'],'slide_count':12,'content_pack_ids':ids,'generate_images':False,
                'business_settings':read(EXT/'settings.json'),'business_conditions':read(folder/'conditions.json'),'business_context_id':ids[-2]}
            response=req('POST','/analysis',json=request).json()
            assert response['business'] and len(response['items'])==12
            assert all('216' in r['business_explanation']['exclusions'] for r in response['items'])
            # B is replayed on VK Tech; baseline uses each original corporate template.
            selected=templates if version=='baseline' else [next((t for t in templates if 'vk tech' in t['name'].lower()),templates[0])]
            for template in selected:
                request['template_id']=template['id']
                job=req('POST','/generations',json=request).json();deadline=time.monotonic()+300
                while job['status'] not in {'completed','failed'} and time.monotonic()<deadline:
                    time.sleep(.2);job=req('GET','/jobs/'+job['id']).json()
                assert job['status']=='completed',job
                for identifier in job['presentation_ids']:
                    record=req('GET','/presentations/'+identifier).json()
                    assert record['business_request']=={'settings':request['business_settings'],'conditions':request['business_conditions'],'context_source_id':ids[-2]}
                    slides=[s['content'] for s in record['deck']['slides']]
                    kpi=next(s['visual'] for s in slides if s['analysis_id'].startswith('business:kpi:'))
                    expected=['93563698','30580684'] if version=='baseline' else ['94823698','31570684']
                    assert [Decimal(str(s['values'][0])) for s in kpi['series']]==list(map(Decimal,expected))
                    trend=next(s['visual'] for s in slides if s['analysis_id'].startswith('business:trend:'))
                    assert trend['categories'][-1]=='2026-06-01'
                    assert not any('III квартал' in s['title'] for s in slides)
                    edit={'revision':1,'title':f'Q2 {version}: сохранённый отчёт','variant':record['variant'],'slides':slides}
                    changed=json.loads(json.dumps(edit));changed['slides'][1]['title']='Подмена'
                    assert client.patch('/presentations/'+identifier,json=changed).status_code==422
                    saved=req('PATCH','/presentations/'+identifier,json=edit).json();assert saved['revision']==2
                    assert req('GET','/presentations/'+identifier).json()['business_request']==record['business_request']
                    base=args.output/version/(template['name']+'-'+record['variant']);base.parent.mkdir(parents=True,exist_ok=True)
                    raw_pptx=req('GET',f'/presentations/{identifier}/export/pptx?revision=2').content
                    raw_pdf=req('GET',f'/presentations/{identifier}/export/pdf?revision=2').content
                    base.with_suffix('.pptx').write_bytes(raw_pptx);base.with_suffix('.pdf').write_bytes(raw_pdf)
                    native=Presentation(io.BytesIO(raw_pptx));pdf=PdfReader(io.BytesIO(raw_pdf))
                    assert len(native.slides)==len(pdf.pages)==13
                    charts=tables=0
                    for page,content in zip(native.slides,slides):
                        visual=content['visual']
                        if visual['kind'] in {'line','bar','pie','grouped_bar','stacked_bar'}:
                            chart=next(s.chart for s in page.shapes if s.has_chart);charts+=1
                            assert [list(s.values) for s in chart.series]==[s['values'] for s in visual['series']]
                        if visual['kind']=='table':assert any(s.has_table for s in page.shapes);tables+=1
                    assert charts==5 and tables==2
                    thumb=req('GET',f'/presentations/{identifier}/slides/0/thumbnail?revision=2&width=640')
                    assert thumb.content.startswith(b'\x89PNG')
                    base.with_name(base.name+'-preview.png').write_bytes(thumb.content)
                    doc=pdfium.PdfDocument(raw_pdf);contact=Image.new('RGB',(1440,4*300),'#e7ecf4')
                    for index in range(len(doc)):
                        page=doc[index];bitmap=page.render(scale=1280/page.get_width());image=bitmap.to_pil().convert('RGB')
                        image.save(base.with_name(base.name+f'-{index+1:02}.png'))
                        small=ImageOps.contain(image,(470,270));x=(index%3)*480;y=(index//3)*300
                        contact.paste(small,(x,y));ImageDraw.Draw(contact).text((x+10,y+274),f'{index+1:02}',fill='#172b46')
                        bitmap.close();page.close()
                    doc.close();contact.save(base.with_name(base.name+'-contact.png'))
                    item={'version':version,'template':template['name'],'variant':record['variant'],'id':identifier,
                        'net_rub':expected[0],'profit_rub':expected[1],'excluded_sales_rows':216,'slides':12,'native_charts':charts,'native_tables':tables,
                        'revision':2,'preview':True,'tampering_rejected':True,'audit':saved['audit']['counts']}
                    report['runs'].append(item);print(json.dumps(item,ensure_ascii=False),flush=True)
                    (args.output/'service-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')


if __name__=='__main__':main()
