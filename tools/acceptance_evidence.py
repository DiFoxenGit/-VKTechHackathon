"""Shared fail-closed evidence checks for controlled and live acceptance tools.

Only sanitized, server-recorded model telemetry is written to reports. Model names
and environment readiness are deliberately insufficient evidence of a model call.
"""
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import re
import subprocess

from pptx import Presentation
from pypdf import PdfReader


def code_version(root):
    def git(*args):
        return subprocess.check_output(["git", *args], cwd=root, text=True).strip()
    digest = hashlib.sha256()
    paths = subprocess.check_output(["git", "ls-files", "-z", "--cached", "--others", "--exclude-standard",
                                     "--", "backend", "frontend/src", "tools"], cwd=root).split(b'\0')
    for name in sorted(set(paths)):
        if name:
            path = Path(root) / name.decode("utf8")
            if path.is_file():
                digest.update(name + b'\0' + hashlib.sha256(path.read_bytes()).digest())
    return {"commit": git("rev-parse", "HEAD"), "branch": git("branch", "--show-current"),
            "working_tree_dirty": bool(git("status", "--porcelain")),
            "source_tree_sha256": digest.hexdigest(),
            "started_at_utc": datetime.now(timezone.utc).isoformat()}


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def safe_name(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:140]


class EvidenceError(ValueError):
    """A safe diagnostic authored by these tools, never a provider response body."""


def require(condition, message):
    if not condition:
        raise EvidenceError(message)


def sanitized_inference(trace):
    """Whitelist telemetry fields even when recording a failed live run."""
    trace = trace or {}
    fields = ("requested", "attempted", "completed", "provider", "model", "attempts",
              "successful_calls", "duration_seconds", "validation", "fallback_used", "status",
              "error_code", "response_id", "response_sha256", "usage", "content_reused_for_variants",
              "primary_skipped_calls", "content_batch_id", "content_source_job_id", "shared_from_job_id",
              "content_input_sha256", "content_revalidated", "model_attempt", "model_max_attempts", "retry_reason")
    call_fields = ("attempt", "status_code", "duration_seconds", "validation", "response_id",
                   "response_sha256", "usage", "error_code", "provider", "validation_error")
    result = {**{key: trace.get(key) for key in fields},
              "calls": [{key: call.get(key) for key in call_fields} for call in trace.get("calls", [])]}
    if isinstance(trace.get('shared_inference'), dict):
        require(not trace['shared_inference'].get('shared_inference'), 'Nested shared provenance is invalid')
        result['shared_inference'] = sanitized_inference(trace['shared_inference'])
    return result


def validate_inference(trace, *, live, variants=3, provider=None, model=None):
    """Require a validated HTTP call for live, and zero calls for explicit local mode."""
    require(isinstance(trace, dict), "Missing server inference telemetry")
    calls = trace.get("calls")
    require(isinstance(calls, list), "Missing inference calls journal")
    require(isinstance(trace.get("attempts"), int) and not isinstance(trace["attempts"], bool),
            "Invalid inference attempt count")
    require(trace["attempts"] == len(calls), "Inference attempts differ from journal")
    require(isinstance(trace.get("successful_calls"), int) and not isinstance(trace["successful_calls"], bool),
            "Invalid successful call count")
    duration = trace.get("duration_seconds")
    require(isinstance(duration, (int, float)) and not isinstance(duration, bool) and math.isfinite(duration) and duration >= 0,
            "Invalid inference duration")
    require(all(call.get("attempt") == index for index, call in enumerate(calls, start=1)),
            "Inference call attempt sequence is invalid")
    validated = [call for call in calls if call.get("validation") == "passed"
                 and isinstance(call.get("status_code"), int) and 200 <= call["status_code"] < 300]
    require(trace.get("successful_calls") == len(validated), "Successful calls lack validated HTTP evidence")
    if live:
        shared = trace.get('shared_inference')
        if shared is not None:
            require(isinstance(shared, dict) and not shared.get('shared_inference'), 'Invalid shared call evidence')
            validate_inference(shared, live=True, variants=0, provider=provider, model=model)
            require(trace.get('attempted') is False and not calls and trace['attempts'] == trace['successful_calls'] == 0,
                    'Reused content must not claim another HTTP call')
            require(trace.get('content_revalidated') is True and trace.get('shared_from_job_id')
                    and trace['shared_from_job_id'] == trace.get('content_source_job_id') == shared.get('content_source_job_id'),
                    'Missing source job or per-template validation')
            require(bool(re.fullmatch(r'[a-f0-9]{64}', trace.get('content_input_sha256') or ''))
                    and trace['content_input_sha256'] == shared.get('content_input_sha256')
                    and trace.get('content_batch_id') and trace['content_batch_id'] == shared.get('content_batch_id'),
                    'Shared response belongs to another input or batch')
            require(all(trace.get(key) == shared.get(key) for key in ('provider', 'model', 'response_id', 'response_sha256')),
                    'Shared response identity differs from source')
        if shared is None:
            require(trace.get('requested') is True and trace.get('attempted') is True,
                    'Model was configured but never requested/attempted')
        else:
            require(trace.get('requested') is True, 'Shared content was not requested')
        require(trace.get("completed") is True and trace.get("validation") == "passed"
                and trace.get("status") == "completed", "Model response was not validated")
        require(bool(validated) or shared is not None, "No actual validated model call")
        require(trace.get("fallback_used") is False and not trace.get("error_code"),
                "Fallback/error cannot pass live Qwen acceptance")
        if variants:
            require(trace.get("content_reused_for_variants") == variants,
                    "Missing proof of shared model content across variants")
        if provider:
            require(trace.get("provider") == provider, "Unexpected model provider")
        if model:
            require(trace.get("model") == model, "Unexpected model identifier")
    else:
        require(trace.get("requested") is False and trace.get("attempted") is False
                and trace.get("completed") is False and not calls,
                "Explicit deterministic run unexpectedly requested a model")
    return sanitized_inference(trace)


def response_summary(traces):
    calls = [call for trace in traces for call in trace["calls"]
             if call.get("validation") == "passed" and 200 <= (call.get("status_code") or 0) < 300]
    ids = [call["response_id"] for call in calls if call.get("response_id")]
    digests = [call["response_sha256"] for call in calls if call.get("response_sha256")]
    owners = [trace for trace in traces if not trace.get('shared_from_job_id') and trace.get('requested')]
    first_passes = sum(trace.get('completed') is True and trace.get('validation') == 'passed'
                      and (trace.get('model_attempt') == 1 if trace.get('model_attempt') else
                           bool(trace.get('calls') and trace['calls'][0].get('validation') == 'passed'))
                      for trace in owners)
    return {"actual_attempts": sum(t.get("attempts") or 0 for t in traces),
            "validated_model_calls": len(calls), "provider_response_ids": ids,
            "distinct_provider_response_ids": len(set(ids)) if len(ids) == len(calls) else None,
            "distinct_response_content_hashes": len(set(digests)) if len(digests) == len(calls) else None,
            "response_identity_complete": len(ids) == len(calls),
            "fallback_jobs": sum(bool(t["fallback_used"]) for t in traces),
            "first_attempt_passed": first_passes, "independent_editorial_requests": len(owners),
            "first_attempt_pass_rate": first_passes / len(owners) if owners else None}


def verify_shared_jobs(jobs, expected_count):
    require(len(jobs) == expected_count and len({job['job_id'] for job in jobs}) == expected_count,
            'Expected distinct template jobs')
    owners = [job for job in jobs if not job['inference'].get('shared_from_job_id')]
    require(len(owners) == 1, 'Expected exactly one fresh editorial source job per batch')
    owner = owners[0]
    source = validate_inference(owner['inference'], live=True)
    require(source.get('content_source_job_id') == owner['job_id'], 'Editorial source job mismatch')
    for job in jobs:
        trace = validate_inference(job['inference'], live=True)
        if job is owner:
            continue
        require(trace.get('shared_from_job_id') == owner['job_id'], 'Shared content links to a missing job')
        inherited = trace['shared_inference']
        require(all(inherited.get(key) == source.get(key) for key in
                    ('calls', 'response_id', 'response_sha256', 'content_input_sha256', 'content_batch_id')),
                'Shared content does not match the source HTTP journal')
    summary = response_summary([job['inference'] for job in jobs])
    require(summary['validated_model_calls'] == 1, 'Expected one validated model response for all templates')
    return {**summary, 'source_job_id': owner['job_id'], 'templates_reusing_content': expected_count - 1,
            'per_template_validation_verified': True}


def verify_independent_jobs(jobs, expected_count):
    """Fresh jobs plus validated HTTP journals prove independent model calls.

    Provider response IDs are optional metadata. Equal text hashes are legitimate
    for repeated prompts and never establish a hidden cache on their own.
    """
    require(len(jobs) == expected_count and len({job["job_id"] for job in jobs}) == expected_count,
            "Expected distinct independent business jobs")
    traces = [validate_inference(job["inference"], live=True) for job in jobs]
    summary = response_summary(traces)
    if summary["response_identity_complete"]:
        require(summary["distinct_provider_response_ids"] == summary["validated_model_calls"],
                "A provider response ID was reused across independent jobs")
    else:
        validated = [call for trace in traces for call in trace["calls"] if call.get("validation") == "passed"]
        require(all(call.get("response_id") or re.fullmatch(r'[a-f0-9]{64}', call.get("response_sha256") or '')
                    for call in validated), "Responses without provider IDs require content digests")
    return {**summary, "independent_business_jobs": expected_count,
            "independence_evidence": "distinct job IDs and validated actual HTTP call journals",
            "provider_response_identity": "available" if summary["response_identity_complete"] else "unavailable_or_partial"}


def verify_office_exports(pptx, pdf, record):
    """Compare editable native chart/table data with trusted saved slide content."""
    native = Presentation(io.BytesIO(pptx))
    document = PdfReader(io.BytesIO(pdf))
    slides = record["deck"]["slides"]
    require(len(native.slides) == len(document.pages) == len(slides), "PPTX/PDF slide count mismatch")
    charts = tables = titles = kpis = 0
    pdf_text = "\n".join(page.extract_text() or "" for page in document.pages)
    require(bool(pdf_text.strip()), "PDF contains no extractable text")
    def compact(text):
        return ''.join(char.casefold() for char in str(text) if char.isalnum())
    def numeric_text(value):
        return re.sub(r'[\s\u00a0\u202f]', '', str(value)).replace(',', '.').replace('−', '-')
    for index, (page, planned) in enumerate(zip(native.slides, slides)):
        page_pdf = document.pages[index].extract_text() or ""
        native_text = "\n".join(shape.text for shape in page.shapes if shape.has_text_frame)
        title = planned["content"].get("title")
        if title:
            require(compact(title) in compact(native_text) and compact(title) in compact(page_pdf),
                    "Saved slide title missing from native PPTX/PDF")
            titles += 1
        visual = planned["content"]["visual"]
        if visual["kind"] in {"line", "bar", "pie", "scatter", "histogram", "grouped_bar", "stacked_bar"}:
            native_charts = [shape.chart for shape in page.shapes if shape.has_chart]
            require(bool(native_charts), "Planned chart lost its native PPTX object")
            def exact(value):
                return Decimal(str(value)) if value is not None else None
            expected = [[exact(v) for v in series["values"]] for series in visual["series"]]
            actual = [[exact(v) for v in series.values] for series in native_charts[0].series]
            require(actual == expected, "Native PPTX chart numbers differ from saved content")
            if visual["kind"] == "scatter":
                for series in native_charts[0].series:
                    xs = [Decimal(node.text) for node in series._element.xpath('.//c:xVal//c:pt/c:v')]
                    require(xs == [exact(value) for value in visual["x_values"]], "Native PPTX scatter X values differ")
            charts += 1
        if visual["kind"] == "kpi":
            for series in visual["series"]:
                value = series["values"][0]
                if value is None:
                    continue
                rendered = (str(int(value)) if float(value).is_integer() else
                            f'{value:.2f}'.rstrip('0').rstrip('.') if abs(value) >= .01 else str(value))
                require(rendered in numeric_text(native_text) and rendered in numeric_text(page_pdf),
                        "KPI number missing from native PPTX/PDF")
                kpis += 1
        if visual["kind"] == "table":
            native_tables = [shape.table for shape in page.shapes if shape.has_table]
            require(bool(native_tables), "Planned table lost its native PPTX object")
            def normalized(value):
                text = str(value)
                numeric = re.sub(r'[\s\u00a0\u202f]', '', text).replace(',', '.')
                return (str(Decimal(numeric)) if re.fullmatch(r'[-+]?\d+(?:\.\d+)?', numeric)
                        and not re.match(r'[-+]?0\d', numeric) else ' '.join(text.split()))
            expected = [[normalized(cell) for cell in row] for row in visual.get("rows", [])]
            actual = [[normalized(cell.text) for cell in row.cells] for row in native_tables[0].rows]
            require(not expected or actual[-len(expected):] == expected, "Native PPTX table values differ")
            for row in visual.get("rows", []):
                for cell in row:
                    if re.fullmatch(r'[-+]?\d+(?:\.\d+)?', numeric_text(cell)):
                        require(numeric_text(cell) in numeric_text(page_pdf), "Signed table value missing from PDF")
                    else:
                        require(not compact(cell) or compact(cell) in compact(page_pdf), "Table cell missing from PDF")
            tables += 1
    return {"slides": len(slides), "native_charts": charts, "native_tables": tables,
            "native_values_verified": True, "pdf_extractable_text": True,
            "saved_titles_in_pptx_pdf": titles, "kpi_values_in_pptx_pdf": kpis,
            "pptx_sha256": sha256(pptx), "pdf_sha256": sha256(pdf)}


def verify_business(record, expected, settings):
    """Compare all independent Decimal oracle answers and evidence pointers."""
    analysis = record.get("business_proof")
    if not analysis:
        business = next((s["_business"] for s in record.get("sources", []) if s.get("_business")), None)
        analysis = business["analysis"] if business else None
    require(bool(analysis), "Saved presentation has no business proof")
    require(analysis["period_settings"] == settings, "Saved business periods differ from request")
    controls = 0
    verified = []
    for answer in expected["answers"]:
        fact = analysis["facts"][answer["id"]]
        actual, wanted = fact["exact"], answer["exact"]
        require((actual is None and wanted is None) or (actual is not None and wanted is not None
                and abs(Decimal(actual) - Decimal(wanted)) <= Decimal("1e-20")),
                "Independent fact mismatch: " + answer["id"])
        require(fact["period"] == answer["period"], "Wrong fact period: " + answer["id"])
        # The oracle names raw CSV rows; production keeps them in explicit lineage.
        require(bool(fact.get("evidence") or fact.get("source_rows") or fact.get("lineage")),
                "Fact has no evidence pointer: " + answer["id"])
        for table, key in (("sales", "source_rows"), ("plan", "plan_source_rows")):
            if answer.get(key):
                rows = {row for ev in fact.get("evidence", []) if ev.get("table") == table for row in ev["row_ids"]}
                require(set(answer[key]) <= rows, "Oracle source rows absent from fact evidence: " + answer["id"])
        controls += 1
        verified.append({"id": answer["id"], "exact": fact["exact"], "unit": fact["unit"],
                         "period": fact["period"], "source_rows": len(answer["source_rows"])})
    return {"independent_numeric_controls": controls, "periods_verified": True,
            "fact_evidence_verified": True, "verified_facts": verified,
            "excluded_sales_rows": len(analysis["source_scope"]["excluded_sales_rows"])}


def verify_slide_evidence(record, *, live):
    proof = record["business_proof"]
    fact_ids = set(proof["facts"])
    slides = record["deck"]["slides"]
    for slide in slides:
        content = slide["content"]
        require(bool(content["source_refs"]), "Slide lost source evidence links")
        notes = json.loads(content["notes"])
        require(bool(notes.get("fact_ids")) and set(notes["fact_ids"]) <= fact_ids,
                "Slide cites an unknown or missing fact")
        require(bool(notes.get("context_refs")), "Slide lost context evidence")
        if live:
            require(notes.get("editorial", {}).get("origin") == "validated_model",
                    "Saved slide lacks validated model provenance")
    if live:
        comparisons = record.get("generation", {}).get("editorial", [])
        require(sum(item.get("before") != item.get("after") for item in comparisons) >= 3,
                "Model prose changes not demonstrated on at least three pages")
    return {"slide_evidence_links_verified": len(slides), "model_provenance_on_every_slide": live}
