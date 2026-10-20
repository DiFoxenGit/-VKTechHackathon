"""One paired live run per stand, frozen outline, no retries of generation POSTs.

Run inside the primary backend container so its API key stays on that server.
The default is a dry run. --live creates four jobs and requests two images total.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import time

import httpx
from pptx import Presentation
from pypdf import PdfReader

OUTLINE = {
    'title': 'Пилот зелёного офиса',
    'slides': [
        {'title': 'Пилот зелёного офиса', 'bullets': ['Сократить расход энергии', 'Сохранить комфорт команды'],
         'visual': {'kind': 'process', 'steps': ['Наблюдение', 'Пилот', 'Оценка']}},
        {'title': 'Современный офис с солнечными панелями',
         'bullets': ['Использовать естественное освещение', 'Проверить режимы работы оборудования'],
         'visual': {'kind': 'none'}},
        {'title': 'Следующий шаг команды', 'bullets': ['Выбрать пилотную площадку', 'Согласовать условия оценки'],
         'visual': {'kind': 'process', 'steps': ['Площадка', 'Ответственный', 'Проверка']}},
    ],
}
BRIEF = ('Проверка пилота зелёного офиса. Сохрани предоставленный текст без добавления чисел. '
         'Если создание новых изображений включено, создай РОВНО ОДНУ фотографию: '
         'на слайде 2 справа современный офис с солнечными панелями, дневной свет. '
         'На слайдах 1 и 3 новые изображения запрещены. Не добавляй другие картинки. '
         'На фотографии не должно быть текста, цифр, логотипов или водяных знаков.')


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def verify_generated_picture(blob, record):
    expected = {(index + 1, Path(element['asset']['ref'].rsplit(':', 1)[-1]).stem)
                for index, slide in enumerate(record['deck']['slides'])
                for element in slide['elements']
                if element.get('asset', {}).get('ref', '').startswith('generated:')}
    assert len(expected) == 1
    hashes = {digest for _, digest in expected}
    actual = {(index + 1, hashlib.sha256(shape.image.blob).hexdigest()[:24])
              for index, slide in enumerate(Presentation(io.BytesIO(blob)).slides)
              for shape in slide.shapes if shape.shape_type == 13
              and hashlib.sha256(shape.image.blob).hexdigest()[:24] in hashes}
    assert actual == expected
    return sorted(actual)


def poll(client, job_id, events, started, checkpoint):
    deadline = started + 360
    previous = None
    while time.perf_counter() < deadline:
        response = client.get('/jobs/' + job_id)
        response.raise_for_status()
        job = response.json()
        state = (job['status'], job.get('stage'))
        if state != previous:
            events.append({'status': state[0], 'stage': state[1],
                           'observed_seconds': round(time.perf_counter() - started, 3)})
            checkpoint()
            previous = state
        if job['status'] in {'completed', 'failed'}:
            return job
        time.sleep(0.5)
    raise RuntimeError('Polling deadline; do not resubmit the paid task')


def run_case(client, payload, output, label, checkpoint):
    entry = {'label': label, 'requested_images': 1 if payload['generate_images'] else 0,
             'payload_sha256': fingerprint(payload), 'outline_sha256': fingerprint(payload['outline']),
             'events': [], 'status': 'submitting'}
    output.append(entry)
    checkpoint()  # Durable marker before POST: restarting this run must not submit again.
    started = time.perf_counter()
    response = client.post('/generations', json=payload)
    response.raise_for_status()
    job = response.json()
    entry['job_id'] = job['id']
    entry['status'] = 'accepted'
    checkpoint()
    job = poll(client, job['id'], entry['events'], started, checkpoint)
    entry.update(status=job['status'], client_ready_seconds=round(time.perf_counter()-started, 3),
                 server_ready_seconds=job.get('elapsed_seconds'), warnings=job.get('warnings', []))
    checkpoint()
    if job['status'] != 'completed':
        raise RuntimeError('Generation failed; no automatic resubmission')
    entry['variants'] = []
    contents = []
    export_started = time.perf_counter()
    for identifier in job['presentation_ids']:
        response = client.get('/presentations/' + identifier)
        response.raise_for_status()
        record = response.json()
        image = record['image_generation']
        expected = int(payload['generate_images'])
        assert image['status'] == ('ready' if expected else 'not_requested'), image['status']
        assert image.get('count', 0) == image.get('requested_count', 0) == expected
        if expected:
            assert [(x['slide'], x['placement']) for x in image['plan']['images']] == [(2, 'right')]
            assert image['planning']['provider'] == 'vk' and not image['planning'].get('fallback_used')
        content = [slide['content'] for slide in record['deck']['slides']]
        contents.append(fingerprint(content))
        item = {'id': identifier, 'variant': record['variant'], 'content_sha256': contents[-1],
                'image_count': image.get('count', 0), 'illustrations_id': record.get('illustrations_id'),
                'planning': image.get('planning'), 'exports': {}}
        for fmt in ('pptx', 'pdf'):
            tick = time.perf_counter()
            exported = client.get('/presentations/' + identifier + '/export/' + fmt)
            exported.raise_for_status()
            blob = exported.content
            pages = len(Presentation(io.BytesIO(blob)).slides) if fmt == 'pptx' else len(PdfReader(io.BytesIO(blob)).pages)
            assert pages == 3
            item['exports'][fmt] = {'seconds': round(time.perf_counter()-tick, 3),
                                     'pages': pages, 'bytes': len(blob), 'sha256': hashlib.sha256(blob).hexdigest()}
            if fmt == 'pptx':
                images = sum(shape.shape_type == 13 for slide in Presentation(io.BytesIO(blob)).slides for shape in slide.shapes)
                item['pptx_total_picture_shapes'] = images
                if expected:
                    item['native_generated_pictures'] = verify_generated_picture(blob, record)
        entry['variants'].append(item)
        checkpoint()
    assert len(contents) == 3 and len(set(contents)) == 1
    entry.update(export_all_seconds=round(time.perf_counter()-export_started, 3),
                 end_to_end_seconds=round(time.perf_counter()-started, 3))
    if payload['generate_images']:
        assert len({x['illustrations_id'] for x in entry['variants']}) == 1
    checkpoint()
    return entry


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live', action='store_true')
    parser.add_argument('--verify-only', action='store_true', help='GET/export only; verify previously created images')
    parser.add_argument('--output', type=Path, default=Path('/data/dima-image-timing/report.json'))
    args = parser.parse_args()
    if args.verify_only:
        if args.live:
            raise SystemExit('Use either --live or --verify-only')
        report = json.loads(args.output.read_text(encoding='utf-8'))
        for label, stand in report['stands'].items():
            base = 'http://127.0.0.1:8000/api/v1' if label == 'primary' else 'https://auto-sell.site/presentations/api/v1'
            key = os.getenv('DESIGNER_API_KEY') if label == 'primary' else None
            with httpx.Client(base_url=base, headers={'Authorization': 'Bearer '+key} if key else {}, timeout=180) as client:
                for run in stand['runs']:
                    if not run['requested_images']:
                        continue
                    for variant in run['variants']:
                        record_response = client.get('/presentations/'+variant['id']); record_response.raise_for_status()
                        pptx_response = client.get('/presentations/'+variant['id']+'/export/pptx'); pptx_response.raise_for_status()
                        variant['native_generated_pictures'] = verify_generated_picture(pptx_response.content, record_response.json())
        report['native_picture_verification'] = 'passed'
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('Existing images verified in PPTX; no POSTs or new generations')
        return
    if not args.live:
        print(json.dumps({'dry_run': True, 'jobs': 4, 'requested_new_images': 2,
                          'outline_sha256': fingerprint(OUTLINE)}, ensure_ascii=False))
        return
    if args.output.exists():
        raise SystemExit('Report already exists; inspect accepted jobs instead of repeating POSTs')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    try:
        descriptor = os.open(args.output, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError:
        raise SystemExit('Report already exists; inspect accepted jobs instead of repeating POSTs') from None
    with os.fdopen(descriptor, 'w', encoding='utf-8') as marker:
        marker.write('{"status":"reserved; do not repeat live submission"}')
    report = {'version': 1, 'started_at_utc': datetime.now(timezone.utc).isoformat(),
              'outline_sha256': fingerprint(OUTLINE), 'brief_sha256': fingerprint(BRIEF),
              'new_images_budget': 2, 'stands': {}, 'limitations': ['one pair per stand; not a percentile or SLA',
              'frozen outline excludes main text generation; image planning, VLM and assembly remain enabled',
              'server-ready excludes explicit export downloads; export is timed separately']}

    def checkpoint():
        temp = args.output.with_suffix('.pending')
        temp.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        temp.replace(args.output)

    def balance():
        response = httpx.get('https://api.gen-api.ru/api/v1/user',
                            headers={'Authorization': 'Bearer '+os.environ['DESIGNER_IMAGE_API_KEY']}, timeout=30)
        response.raise_for_status()
        return response.json()['balance']

    report['balance_before'] = balance()
    checkpoint()
    definitions = [('primary', 'http://127.0.0.1:8000/api/v1', os.getenv('DESIGNER_API_KEY')),
                   ('autosell', 'https://auto-sell.site/presentations/api/v1', None)]
    clients = []
    try:
        common_sha = None
        for label, base, key in definitions:
            client = httpx.Client(base_url=base, headers={'Authorization': 'Bearer '+key} if key else {}, timeout=180)
            clients.append((label, client))
            response = client.get('/templates'); response.raise_for_status()
            templates = response.json()['items']
            template = next((x for x in templates if common_sha is None or x['sha256'] == common_sha), None)
            assert template is not None, 'Stands must have identical source template'
            common_sha = template['sha256']
            report['stands'][label] = {'template_id': template['id'], 'template_sha256': common_sha, 'runs': []}
        for label, client in clients:
            stand = report['stands'][label]
            for images in (False, True):
                payload = {'brief': BRIEF, 'purpose': 'project', 'language': 'ru', 'slide_count': 3,
                           'content_pack_ids': [], 'template_id': stand['template_id'],
                           'outline': OUTLINE, 'model_content_generation': False, 'narrative': False,
                           'contextual_audit': False, 'generate_images': images}
                run_case(client, payload, stand['runs'], 'with_image' if images else 'without_images', checkpoint)
            a, b = stand['runs']
            assert a['variants'][0]['content_sha256'] == b['variants'][0]['content_sha256']
            stand['ready_delta_seconds'] = round(b['server_ready_seconds'] - a['server_ready_seconds'], 3)
            stand['download_delta_seconds'] = round(b['end_to_end_seconds'] - a['end_to_end_seconds'], 3)
            checkpoint()
    finally:
        for _, client in clients:
            client.close()
        try:
            report['balance_after'] = balance()
            report['balance_delta'] = round(float(report['balance_before'])-float(report['balance_after']), 4)
        except (httpx.HTTPError, ValueError, KeyError):
            report['balance_after_unavailable'] = True
        checkpoint()
    print(json.dumps({'report': str(args.output), 'balance_delta': report.get('balance_delta'),
                      'stands': {k: {'ready_delta_seconds': v.get('ready_delta_seconds'),
                                    'jobs': [x.get('job_id') for x in v['runs']]} for k,v in report['stands'].items()}}))


if __name__ == '__main__':
    main()
