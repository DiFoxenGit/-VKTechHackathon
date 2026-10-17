"""Real HTTP integration: upload, jobs, native PPTX, PDF and independent controls.

Run against a local server. Organiser templates are read in place, never copied
into source control. All fixtures are explicitly synthetic.
"""
import argparse
import io
import json
import time
from pathlib import Path
import httpx
from pptx import Presentation
from run_final_server import ROOT


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--url",default="http://127.0.0.1:8000/api/v1");parser.add_argument("--templates",type=Path);parser.add_argument("--pdf",action="store_true");args=parser.parse_args()
    output=ROOT/"output"/"final";output.mkdir(parents=True,exist_ok=True)
    samples=ROOT/"samples"/"final"
    report={"synthetic":True,"checks":[],"runs":[]}
    with httpx.Client(base_url=args.url,timeout=180) as client:
        def req(method,path,**kw):
            response=client.request(method,path,**kw);response.raise_for_status();return response
        def upload(path,endpoint="/content-packs",name=None):
            return req("POST",endpoint,files={"file":(name or path.name,path.read_bytes())}).json()
        templates=req("GET","/templates").json()["items"]
        if args.templates:
            for path in sorted(args.templates.glob("*.pptx")):
                template=upload(path,"/templates")
                if not any(t["id"]==template["id"] for t in templates):templates.append(template)
        if len(templates)<2:
            for i in range(2-len(templates)):
                deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=f"Шаблон {i+1}";slide.placeholders[1].text="Содержание"
                data=io.BytesIO();deck.save(data)
                templates.append(req("POST","/templates",files={"file":(f"test-{i}.pptx",data.getvalue())}).json())
        controls=json.loads((samples/"controls.json").read_text(encoding="utf-8"))
        packs=[];goals={}
        for control in controls["datasets"]:
            pack=upload(samples/control["file"]);packs.append(pack);goals[pack["id"]]=control["goal"]
            result=req("GET",f"/content-packs/{pack['id']}/analysis",params={"goal":control["goal"]}).json()["items"][0]
            assert result["profile"]["row_count"]==control["rows"]
            assert result["selected"] in control["allowed"]
            if "total" in control and result["facts"]:assert abs(result["facts"][0]["value"]-control["total"])<1e-6
            if "median" in control:assert abs(result["facts"][0]["value"]-control["median"])<1e-6
            report["checks"].append(dict(file=control["file"],rows=control["rows"],selected=result["selected"],passed=True))
        book=upload(samples/"brandbook-blue.json","/brandbooks")
        def generate(name,selected,template,brandbook=None,limit=30):
            started=time.monotonic()
            payload=dict(brief="Синтетический отчёт: проверяемые результаты продаж",template_id=template["id"],slide_count=limit,content_pack_ids=[p["id"] for p in selected],analysis_goals={p["id"]:goals[p["id"]] for p in selected})
            if brandbook:payload["brandbook_id"]=brandbook["id"]
            job=req("POST","/generations",json=payload).json()
            deadline=time.monotonic()+300
            while job["status"] not in {"completed","failed"} and time.monotonic()<deadline:
                time.sleep(.15);job=req("GET",f"/jobs/{job['id']}").json()
            assert job["status"]=="completed",job
            records=[req("GET",f"/presentations/{p}").json() for p in job["presentation_ids"]]
            first=records[0]
            for record in records:
                assert [s["content"]["visual"] for s in record["deck"]["slides"]]==[s["content"]["visual"] for s in first["deck"]["slides"]]
                assert not any(i["code"] in {"analysis_mismatch","unverified_number","chart_without_data"} for i in record["audit"]["issues"])
            pptx=req("GET",f"/presentations/{first['id']}/export/pptx").content
            (output/f"{name}.pptx").write_bytes(pptx)
            reopened=Presentation(io.BytesIO(pptx))
            assert len(reopened.slides)==len(first["deck"]["slides"])
            pdf_bytes=None
            if args.pdf:
                pdf_bytes=req("GET",f"/presentations/{first['id']}/export/pdf").content
                (output/f"{name}.pdf").write_bytes(pdf_bytes)
                from pypdf import PdfReader
                pdf=PdfReader(io.BytesIO(pdf_bytes))
                assert len(pdf.pages)==len(reopened.slides)
                assert all("Источник" in (p.extract_text() or "") for p in pdf.pages)
            (output/f"{name}.json").write_text(json.dumps(first,ensure_ascii=False,indent=2),encoding="utf-8")
            report["runs"].append(dict(name=name,slides=len(reopened.slides),elapsed_seconds=round(time.monotonic()-started,3),pptx_bytes=len(pptx),pdf_bytes=len(pdf_bytes) if pdf_bytes else None,variants=len(records),audit=first["audit"]["counts"]))
            print(name,report["runs"][-1],flush=True)
            return first
        first=generate("all-seven",packs,templates[0])
        second=generate("same-data-other-template",packs[:1],templates[1])
        assert first["analysis"][0]["visual"]==second["analysis"][0]["visual"]
        styled=generate("brandbook",packs[:3],templates[0],book)
        # Demonstrate all chart types through the real API and renderer.
        typed=[]
        for name,filename,goal in [("bar","02-regions.csv","Сравнить"),("line","06-gaps.csv","Динамика"),("pie","03-channels.json","Доли"),("histogram","04-orders.csv","Распределение"),("scatter","04-orders.csv","Взаимосвязь"),("grouped_bar","01-sales.csv","Группированные столбцы"),("stacked_bar","01-sales.csv","Накопительные столбцы"),("kpi","02-regions.csv","KPI")]:
            pack=upload(samples/filename,name=name+Path(filename).suffix);typed.append(pack);goals[pack["id"]]=goal
        types=generate("all-chart-types",typed,templates[0],book)
        kinds=[s["content"]["visual"]["kind"] for s in types["deck"]["slides"]]
        assert set(kinds)=={"bar","line","pie","histogram","scatter","grouped_bar","stacked_bar","kpi"},kinds
        # A real change to input must alter both graph data and verified title.
        updated=req("POST","/content-packs",files={"file":("updated.csv","Регион,Продажи (тыс. руб.)\nСевер,1500\nЮг,2400\nВосток,3600".encode())}).json()
        goals[updated["id"]]="Сравнить"
        changed=generate("changed-input",[updated],templates[0])
        assert changed["analysis"][0]["facts"][0]["value"]==7500
        report["checks"].append(dict(check="three variants + another template + changed input + native chart types",passed=True))
    (output/"acceptance.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print("Acceptance complete:",output,flush=True)


if __name__=="__main__":main()
