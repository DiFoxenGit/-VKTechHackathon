"""Additional synthetic dark template, explicitly not unseen-template acceptance."""
import argparse
import hashlib
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def make(path):
    prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(7.5)
    for index in range(5):
        slide=prs.slides.add_slide(prs.slide_layouts[0 if index==0 else 1])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('192B28')
        title=slide.shapes.title
        title.left=Inches(.8);title.top=Inches(2.7 if index==0 else .55)
        title.width=Inches(9.6);title.height=Inches(1.4 if index==0 else 1.1)
        title.text='Синтетический стресс-шаблон' if index==0 else f'Образец страницы {index}'
        body=slide.placeholders[1];body.left=Inches(.8);body.top=Inches(4.6 if index==0 else 2)
        body.width=Inches(9.6);body.height=Inches(1 if index==0 else 4.3)
        body.text='Период и адресат' if index==0 else 'Вводный блок\nДоказательство\nДействие'
        for shape,size in ((title,42 if index==0 else 30),(body,20)):
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.name='DejaVu Sans';paragraph.font.size=Pt(size)
                paragraph.font.color.rgb=RGBColor.from_string('F3F5ED')
        marker=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(10.9),Inches(.5),Inches(.35),Inches(.35))
        marker.fill.solid();marker.fill.fore_color.rgb=RGBColor.from_string('D9AC60');marker.line.fill.background()
        footer=slide.shapes.add_textbox(Inches(.8),Inches(6.9),Inches(6),Inches(.3))
        footer.text='STRESS · synthetic · 2026'
        footer.text_frame.paragraphs[0].font.size=Pt(11)
        footer.text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string('D9AC60')
    path.parent.mkdir(parents=True,exist_ok=True);prs.save(path)
    return hashlib.sha256(path.read_bytes()).hexdigest()


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('output',type=Path);args=parser.parse_args()
    print(make(args.output))
