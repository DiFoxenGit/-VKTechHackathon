"""Independent final acceptance: real templates, edits, story and unfamiliar data."""
import argparse, io, json, time
from pathlib import Path
import httpx
from pptx import Presentation
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://127.0.0.1:8017/api/v1');args=parser.parse_args()
    out=ROOT/'output/winning';out.mkdir(parents=True,exist_ok=True)
    report={'synthetic_data':True,'templates':[],'edge_cases':[]}
    with httpx.Client(base_url=args.url,timeout=180) as c:
        def req(method,path,**kw):
            r=c.request(method,path,**kw);r.raise_for_status();return r
        def upload(name,data):return req('POST','/content-packs',files={'file':(name,data)}).json()
        def analyse(pack,settings):
            dataset=pack['datasets'][0]['dataset_id'];return req('POST',f"/content-packs/{pack['id']}/analysis",json={'goal':'Проверка','requests':{dataset:settings}}).json()['items'][0]
        # Expected values below are hand-calculated, independent of production code.
        cases=[('small.csv','Категория,Значение\nА,0.000001\nБ,0.000002',{'task':'summary'},[.000003]),
            ('negative.csv','Категория,Значение\nА,-10\nБ,15',{'task':'comparison'},[15,-10]),
            ('zeros.csv','Категория,Значение\nА,0\nБ,0',{'task':'summary'},[0]),
            ('missing.csv','Категория,Значение\nА,10\nБ,\nВ,30',{'task':'summary'},[40]),
            ('duplicate.csv','Категория,Значение\nА,10\nА,20\nБ,5',{'task':'comparison'},[30,5])]
        for name,text,settings,expected in cases:
            result=analyse(upload(name,text.encode()),settings);assert result['status']!='refused',result
            values=[f['value'] for f in result['facts']]
            assert sorted(values)==sorted(expected),(name,values)
            report['edge_cases'].append({'name':name,'values':values,'passed':True})
        rejected=analyse(upload('mixed.csv','Категория,Значение,Валюта\nА,10,RUB\nБ,20,USD'.encode()),{'task':'summary'})
        assert rejected['status']=='refused';report['edge_cases'].append({'name':'mixed-units','refused':True})
        folder=ROOT/'samples/final-demo'
        packs=[upload(name,(folder/name).read_bytes()) for name in ['продажи.csv','контекст.txt']]
        changed=upload('продажи-обновлено.csv',(folder/'продажи-обновлено.csv').read_bytes())
        templates=req('GET','/templates').json()['items']
        def generate(template,data,prefix):
            dataset=data['datasets'][0]['dataset_id']
            payload={'brief':'Результаты пилота: динамика выручки','analysis_goal':'Динамика','narrative':True,'slide_count':10,'template_id':template['id'],
                'content_pack_ids':[data['id'],packs[1]['id']], 'analysis_requests':{dataset:{'task':'dynamics','unit':'тыс. руб.'}}}
            job=req('POST','/generations',json=payload).json()
            deadline=time.monotonic()+240
            while job['status'] not in {'completed','failed'} and time.monotonic()<deadline:
                time.sleep(.2);job=req('GET',f"/jobs/{job['id']}").json()
            assert job['status']=='completed',job
            records=[req('GET',f'/presentations/{id}').json() for id in job['presentation_ids']]
            boxes=[]
            for record in records:
                analytical=next(s for s in record['deck']['slides'] if s['content']['analysis_id'])
                boxes.append(next(e['box'] for e in analytical['elements'] if e['id']=='visual'))
                stem=prefix+'-'+record['variant']
                for fmt in ['pptx','pdf']:
                    (out/(stem+'.'+fmt)).write_bytes(req('GET',f"/presentations/{record['id']}/export/{fmt}").content)
                prs=Presentation(out/(stem+'.pptx'));pdf=PdfReader(out/(stem+'.pdf'))
                assert len(prs.slides)==len(pdf.pages)==len(record['deck']['slides'])
                assert any('Следующий шаг' in s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
                assert not any(i['code']=='analysis_mismatch' for i in record['audit']['issues'])
                (out/(stem+'.json')).write_text(json.dumps(record,ensure_ascii=False,indent=2))
            assert len({tuple(b) for b in boxes})==3
            return records
        for index,template in enumerate(templates):
            records=generate(template,packs[0],f'template-{index+1}')
            record=records[0];url=f"/presentations/{record['id']}"
            slides=[s['content'] for s in record['deck']['slides']];slides[0]['title']='Пилот: результаты и следующий шаг'
            saved=req('PATCH',url,json={'revision':1,'title':'Проверенная история','variant':'split','slides':slides}).json()
            assert saved['revision']==2 and saved['template_id']==template['id']
            (out/f'template-{index+1}-edited.pptx').write_bytes(req('GET',url+'/export/pptx').content)
            (out/f'template-{index+1}-edited.pdf').write_bytes(req('GET',url+'/export/pdf').content)
            report['templates'].append({'name':template['name'],'id':template['id'],'variants':3,'slides':len(slides),'saved_revision':2,'audit':saved['audit']['counts']})
            print(report['templates'][-1],flush=True)
        updated=generate(templates[0],changed,'changed')
        title=next(s['content']['title'] for s in updated[0]['deck']['slides'] if s['content']['analysis_id'])
        assert '200%' in title,title
        report['mutation']={'expected_growth_percent':200,'title':title}
        (out/'readiness.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
        print('PASS: source-grounded story, three templates and variants, edits, exports and independent values',flush=True)

if __name__=='__main__':main()
