"""Measure three independent template jobs together, then inspect all nine variants.

Uploads and pre-analysis finish before the shared wall clock starts. The measured
interval still includes server queueing and any repeated analysis. PDF conversion
is checked after the generation clock; this is a local controlled measurement,
not a production SLA or a substitute for the screen recording.
"""
import argparse,asyncio,hashlib,io,json,subprocess,time
from pathlib import Path
import httpx
from pptx import Presentation
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]


async def main(args):
    output=ROOT/'output/andrey-periods'/args.scenario
    output.mkdir(parents=True,exist_ok=True)
    report={'scenario':args.scenario,'commit':subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip(),
        'environment':'local','image_generation_requested':False,'model_content_generation':False,
        'analysis_preflight_outside_clock':True,'repeated_analysis_and_queue_inside_clock':True,
        'pdf_conversion_outside_clock':True,'checks':[],'runs':[]}
    async with httpx.AsyncClient(base_url=args.url,timeout=240) as client:
        async def req(method,path,**kw):
            r=await client.request(method,path,**kw);r.raise_for_status();return r
        async def upload(path):
            return (await req('POST','/content-packs',files={'file':(path.name,path.read_bytes())})).json()
        templates=list({t['sha256']:t for t in (await req('GET','/templates')).json()['items']}.values())
        assert len(templates)==3,[(t['name'],t['id']) for t in templates]
        if args.scenario=='business':
            ext=ROOT/'samples/business-acceptance/period-context-v1';data=ROOT/'samples/business-acceptance/data-v1'
            packs=[]
            for name in ('sales.csv','plan.csv','products.csv','regions.csv','context.txt','content-pack.zip'):
                path=(data if name in {'products.csv','regions.csv','content-pack.zip'} else ext/'baseline')/name
                packs.append(await upload(path))
            payload={'brief':'Чистая выручка, прибыль, маржа и план-факт; ресурсы и действия по выбранным периодам.',
                'slide_count':12,'content_pack_ids':[p['id'] for p in packs],'generate_images':False,
                'business_settings':json.loads((ext/'settings.json').read_text()),
                'business_conditions':json.loads((ext/'baseline/conditions.json').read_text()),'business_context_id':packs[-2]['id']}
            await req('POST','/analysis',json={**payload,'template_id':templates[0]['id']})
        else:
            folder=ROOT/'samples/analytics/datasets';packs=[];selections={};goals={}
            for spec in sorted(folder.glob('*.request.json')):
                stem=spec.name.removesuffix('.request.json')
                path=next(p for p in folder.glob(stem+'.*') if p.suffix in {'.csv','.md'})
                pack=await upload(path);packs.append(pack)
                selection=json.loads(spec.read_text());dataset=pack['datasets'][0]['dataset_id']
                preview=(await req('POST',f"/content-packs/{pack['id']}/analysis",json={'goal':selection['goal'],'requests':{dataset:selection}})).json()
                expected=json.loads((folder/(stem+'.result.json')).read_text())
                assert preview['analysis_results'][0]==expected,stem
                selections[dataset]=selection;goals[pack['id']]=selection['goal']
                report['checks'].append({'dataset':stem,'rows':len(pack['datasets'][0]['rows']),'exact_contract_match':True})
            assert len(packs)==7
            payload={'brief':'Проверка семи видов данных: динамика, категории, доли, распределение, справочник, пропуски и таблица в тексте.',
                'slide_count':30,'narrative':False,'content_pack_ids':[p['id'] for p in packs],
                'analysis_requests':selections,'analysis_goals':goals,'generate_images':False}
        started=time.monotonic()
        async def generate(template):
            job=(await req('POST','/generations',json={**payload,'template_id':template['id']})).json()
            while job['status'] not in {'completed','failed'} and time.monotonic()-started<300:
                await asyncio.sleep(.15);job=(await req('GET','/jobs/'+job['id'])).json()
            assert job['status']=='completed',job
            assert len(job['presentation_ids'])==3
            return template,job,round(time.monotonic()-started,3)
        groups=await asyncio.gather(*(generate(t) for t in templates))
        report['all_nine_ready_seconds']=round(time.monotonic()-started,3)
        report['within_five_minutes']=report['all_nine_ready_seconds']<=300
        report['jobs']=[{'template':t['name'],'job_id':job['id'],'ready_since_shared_start_seconds':elapsed} for t,job,elapsed in groups]
        for template,job,elapsed in groups:
            for identifier in job['presentation_ids']:
                record=(await req('GET','/presentations/'+identifier)).json()
                base=output/(template['name']+'-'+record['variant'])
                raw=(await req('GET',f'/presentations/{identifier}/export/pptx')).content
                pdf=(await req('GET',f'/presentations/{identifier}/export/pdf')).content
                base.with_suffix('.pptx').write_bytes(raw);base.with_suffix('.pdf').write_bytes(pdf)
                native=Presentation(io.BytesIO(raw));pages=PdfReader(io.BytesIO(pdf)).pages
                assert len(native.slides)==len(pages)==len(record['deck']['slides'])
                for page,planned in zip(native.slides,record['deck']['slides']):
                    v=planned['content']['visual']
                    if v['kind'] in {'line','bar','pie','histogram','grouped_bar','stacked_bar'}:
                        chart=next(s.chart for s in page.shapes if s.has_chart)
                        assert [list(s.values) for s in chart.series]==[s['values'] for s in v['series']]
                item={'template':template['name'],'variant':record['variant'],'id':identifier,'slides':len(native.slides),
                    'visual_types':sorted({s['content']['visual']['kind'] for s in record['deck']['slides']}),
                    'pdf_sha256':hashlib.sha256(pdf).hexdigest(),'pptx_sha256':hashlib.sha256(raw).hexdigest(),
                    'native_values_verified':True,'audit':record['audit']['counts']}
                report['runs'].append(item)
                print(json.dumps(item,ensure_ascii=False),flush=True)
        (output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
        print(json.dumps({'scenario':args.scenario,'all_nine_ready_seconds':report['all_nine_ready_seconds'],'within_five_minutes':report['within_five_minutes']},ensure_ascii=False),flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://127.0.0.1:8019/api/v1')
    parser.add_argument('--scenario',choices=['business','seven-types'],default='business')
    asyncio.run(main(parser.parse_args()))
