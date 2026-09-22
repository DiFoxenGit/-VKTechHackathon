"""Синтетические шаблоны для регрессионного прогона в CI.

Шаблоны кейса (VK Tech, VK WorkSpace, VK Education, ЛЦТ2026) — файлы
организаторов, в репозиторий они не попадают. Чтобы в CI регресс шёл на
настоящих .pptx, здесь собираются два шаблона, которые воспроизводят то, на чём
вёрстка чаще всего ошибается:

* «Светлый 16x9»: логотип на каждой странице (брендинг по повторяемости),
  оглавление, ряд из трёх карточек, раздел и финальная страница «Спасибо».
* «Тёмный 4x3»: другая пропорция, тёмный фон и светлый текст, сетка 2x2.

Файлы детерминированные: тот же код даёт те же страницы.

    python tools/make_fixture_templates.py out/templates
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

BRAND = RGBColor(0x5B, 0x2E, 0xFF)
CARD_LIGHT = RGBColor(0xF1, 0xED, 0xFF)
DARK = RGBColor(0x14, 0x16, 0x1F)
CARD_DARK = RGBColor(0x24, 0x28, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Макеты шаблона python-pptx по умолчанию.
TITLE, CONTENT, SECTION, TWO_CONTENT, TITLE_ONLY = 0, 1, 2, 3, 5


def logo(slide, width, height):
    """Одинаковый знак в углу каждой страницы — так шаблон держит бренд."""
    mark = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(int(width * 0.03)),
        Emu(int(height * 0.9)),
        Emu(int(width * 0.06)),
        Emu(int(height * 0.05)),
    )
    mark.fill.solid()
    mark.fill.fore_color.rgb = BRAND
    mark.line.fill.background()


def paint(shape, color, size=None):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color
            if size:
                run.font.size = Pt(size)


def cards(slide, width, height, texts, columns, fill, color):
    """Сетка одинаковых карточек с текстом — её вёрстка должна распознать."""
    rows = -(-len(texts) // columns)
    left, top = width * 0.06, height * 0.3
    gap = width * 0.02
    card_w = (width * 0.88 - gap * (columns - 1)) / columns
    card_h = (height * 0.55 - gap * (rows - 1)) / rows
    for index, text in enumerate(texts):
        column, row = index % columns, index // columns
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(int(left + column * (card_w + gap))),
            Emu(int(top + row * (card_h + gap))),
            Emu(int(card_w)),
            Emu(int(card_h)),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        shape.text_frame.text = text
        paint(shape, color, 16)


def dark_background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK


def light_16x9(target: Path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    width, height = deck.slide_width, deck.slide_height
    layouts = deck.slide_layouts

    cover = deck.slides.add_slide(layouts[TITLE])
    cover.shapes.title.text = "Название презентации"
    cover.placeholders[1].text = "Имя Фамилия, команда, 2026"

    agenda = deck.slides.add_slide(layouts[CONTENT])
    agenda.shapes.title.text = "Содержание"
    agenda.placeholders[1].text = "Контекст\nПредложение\nРезультаты\nПлан"

    content = deck.slides.add_slide(layouts[CONTENT])
    content.shapes.title.text = "Заголовок слайда формулирует вывод"
    content.placeholders[1].text = (
        "Первый тезис с фактом и цифрой\nВторой тезис с последствием\nТретий тезис"
    )

    grid = deck.slides.add_slide(layouts[TITLE_ONLY])
    grid.shapes.title.text = "Три карточки в ряд"
    cards(
        grid, width, height,
        ["Первая карточка\nКороткое описание", "Вторая карточка\nКороткое описание",
         "Третья карточка\nКороткое описание"],
        columns=3, fill=CARD_LIGHT, color=RGBColor(0x1A, 0x1A, 0x1A),
    )

    two = deck.slides.add_slide(layouts[TWO_CONTENT])
    two.shapes.title.text = "Две колонки"
    two.placeholders[1].text = "Левая колонка\nТезис"
    two.placeholders[2].text = "Правая колонка\nТезис"

    section = deck.slides.add_slide(layouts[SECTION])
    section.shapes.title.text = "Раздел 2"

    closing = deck.slides.add_slide(layouts[TITLE_ONLY])
    closing.shapes.title.text = "Спасибо за внимание"

    for slide in deck.slides:
        logo(slide, width, height)
    deck.save(target)


def dark_4x3(target: Path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    width, height = deck.slide_width, deck.slide_height
    layouts = deck.slide_layouts

    cover = deck.slides.add_slide(layouts[TITLE])
    cover.shapes.title.text = "Тёмный шаблон"
    cover.placeholders[1].text = "Подзаголовок обложки"

    content = deck.slides.add_slide(layouts[CONTENT])
    content.shapes.title.text = "Контентная страница"
    content.placeholders[1].text = "Тезис один\nТезис два\nТезис три"

    grid = deck.slides.add_slide(layouts[TITLE_ONLY])
    grid.shapes.title.text = "Сетка два на два"
    cards(
        grid, width, height,
        ["Карточка 1\nОписание", "Карточка 2\nОписание",
         "Карточка 3\nОписание", "Карточка 4\nОписание"],
        columns=2, fill=CARD_DARK, color=WHITE,
    )

    closing = deck.slides.add_slide(layouts[TITLE_ONLY])
    closing.shapes.title.text = "Вопросы?"

    for slide in deck.slides:
        dark_background(slide)
        logo(slide, width, height)
        for shape in slide.placeholders:
            paint(shape, WHITE)
    deck.save(target)


def main():
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else "out/templates")
    folder.mkdir(parents=True, exist_ok=True)
    light_16x9(folder / "Светлый 16x9.pptx")
    dark_4x3(folder / "Тёмный 4x3.pptx")
    for path in sorted(folder.glob("*.pptx")):
        print(f"{path}  {path.stat().st_size // 1024} КБ")


if __name__ == "__main__":
    main()
