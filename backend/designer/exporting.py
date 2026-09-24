"""Native PowerPoint export; LibreOffice renders the same deck to PDF/HTML."""

import copy
import html
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .assets import color_roles, draw_asset
from .fonts import printable
from .layout import STAT_BAR, STAT_INSET, estimated_text_height
from .visuals import DIAGRAMS, MIN_LABEL_SIZE
from .parsing import (
    background_color,
    best_text_color,
    contrast_ratio,
    is_footer_placeholder,
    preserved_shape,
    shape_key,
    stage_clutter,
)


def _font(font, name, size, color, bold=False):
    font.name, font.size, font.bold = name, Pt(size), bold
    font.color.rgb = RGBColor.from_string(color)
    for tag in ("a:ea", "a:cs"):
        element = font._rPr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            font._rPr.append(element)
        element.set("typeface", name)


def _buried(shape, content):
    """Мелкий объект прототипа, который окажется под нашим текстом.

    Иконка, стрелка или точка, оставшаяся от образца содержания, попадает ровно
    туда, куда ляжет новый блок, и читается как грязь. Рамку-контейнер это
    правило не трогает: она больше блока и обрамляет его, а не лежит под ним.
    """
    left, top = shape.left, shape.top
    right, bottom = left + shape.width, top + shape.height
    area = max(1, shape.width * shape.height)
    for bx, by, bw, bh in content:
        overlap = max(0, min(right, bx + bw) - max(left, bx)) * max(
            0, min(bottom, by + bh) - max(top, by)
        )
        if overlap > 0.4 * area and area < 0.5 * max(1, bw * bh):
            return True
    return False


def _orphan_icon_row(source, width, height, branding):
    """Ряд одинаковых значков прототипа, под которым не будет содержания.

    На странице-каталоге шаблона над каждой подписью стоит иконка. Если наш
    слайд не раскладывается по этим ячейкам, иконки остаются висеть рядком без
    текста — читается как забытая заготовка. Такой ряд не переносим; одиночный
    значок это правило не трогает, он часть оформления страницы.
    """
    groups = {}
    for shape in source.shapes:
        if shape.has_text_frame and shape.text.strip():
            continue
        if branding and shape_key(shape, width, height) in branding:
            continue
        w, h = shape.width / width, shape.height / height
        if not (0.0002 < w * h < 0.02):
            continue
        if not 0.15 < shape.top / height < 0.85:
            continue
        groups.setdefault((round(w, 2), round(h, 2)), []).append(shape.shape_id)
    orphans = set()
    for members in groups.values():
        if len(members) >= 3:
            orphans.update(members)
    return orphans


def _stage_clutter(source, width, height, branding):
    """Мелкие образцы страницы, которые не переносятся на слайд своей композиции."""
    found = set()
    for shape in source.shapes:
        if shape.is_placeholder or (shape.has_text_frame and shape.text.strip()):
            continue
        box = {
            "x": shape.left / width,
            "y": shape.top / height,
            "w": shape.width / width,
            "h": shape.height / height,
        }
        if stage_clutter(box, branding or ()):
            found.add(shape.shape_id)
    return frozenset(found)


def _resize(element, box, width, height):
    """Переставить скопированную фигуру по исправленной рамке (доли слайда)."""
    transform = element.find(f"{qn('p:spPr')}/{qn('a:xfrm')}")
    if transform is None:
        transform = element.find(f".//{qn('a:xfrm')}")
    if transform is None:
        return
    offset = transform.find(qn("a:off"))
    extent = transform.find(qn("a:ext"))
    if offset is None or extent is None:
        return
    offset.set("x", str(int(box["x"] * width)))
    offset.set("y", str(int(box["y"] * height)))
    extent.set("cx", str(max(1, int(box["w"] * width))))
    extent.set("cy", str(max(1, int(box["h"] * height))))


def _clone_slide(
    deck, source, branding=None, content=(), orphans=frozenset(), fixes=None
):
    target = deck.slides.add_slide(source.slide_layout)
    for shape in list(target.shapes):
        shape._element.getparent().remove(shape._element)
    mapping = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        mapping[rel.rId] = target.part.relate_to(
            rel.target_ref if rel.is_external else rel.target_part,
            rel.reltype,
            is_external=rel.is_external,
        )
    if source._element.cSld.bg is not None:
        target._element.cSld.insert(0, copy.deepcopy(source._element.cSld.bg))
    for shape in source.shapes:
        # Keep artwork and footer/page-number placeholders; remove sample content.
        # The audit reads the same predicate to know what is on the finished slide.
        if not preserved_shape(shape, deck.slide_width, deck.slide_height, branding):
            continue
        # Повторяющийся брендинг остаётся всегда: это оформление страницы.
        repeated = branding and shape_key(
            shape, deck.slide_width, deck.slide_height
        ) in branding
        if not repeated and (
            _buried(shape, content) or shape.shape_id in orphans
        ):
            continue
        element = copy.deepcopy(shape._element)
        # Значок образца внутри группы (кружок + иконка) снимается отдельно:
        # подложка остаётся, пример уходит.
        for picture in element.xpath(".//p:pic"):
            identifier = picture.xpath("./p:nvPicPr/p:cNvPr/@id")
            if identifier and int(identifier[0]) in orphans:
                picture.getparent().remove(picture)
        is_footer = is_footer_placeholder(shape)
        # Groups may contain sample text; do not copy their text into the result.
        for text in element.xpath(".//a:t"):
            if not is_footer:
                text.text = ""
        for node in element.iter():
            for key, value in list(node.attrib.items()):
                if (
                    key.startswith(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                    )
                    and value in mapping
                ):
                    node.set(key, mapping[value])
        fix = (fixes or {}).get(str(shape.shape_id))
        if fix:
            # Аудит уже посчитал честную рамку: применяем её к копии, чтобы
            # картинка перестала быть растянутой и в файле, и в отчёте.
            _resize(element, fix["box"], deck.slide_width, deck.slide_height)
        target.shapes._spTree.insert_element_before(element, "p:extLst")
    return target


def _mix(color, other, ratio):
    """Смесь двух цветов: ratio — доля первого."""
    a = [int(color[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(other[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x * ratio + y * (1 - ratio)):02X}" for x, y in zip(a, b))


def _translucent(shape, color, alpha=0.85):
    """Полупрозрачная заливка: фон шаблона виден, но текст поверх читается."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
    if srgb is not None:
        value = OxmlElement("a:alpha")
        value.set("val", str(int(alpha * 100000)))
        srgb.append(value)
    shape.line.fill.background()
    shape.shadow.inherit = False


def add_scrim(slide, elements, color):
    """Подложка под текстовыми блоками на слайде с фотографией во весь экран.

    Прототип-обложка иногда оказывается лучшим по остальным признакам. Вместо
    того чтобы класть текст прямо на фотографию, кладём под него плашку цвета
    фона шаблона: так слайд остаётся в стиле и остаётся читаемым.
    """
    if not elements:
        return
    # Подложка закрывает всё содержимое слайда: график и таблица на фотографии
    # так же нечитаемы, как текст, а подписи осей рисуются тем же цветом.
    def occupied(element):
        x, y, w, h = element["box"]
        if element["kind"] == "text":
            # Рамка почти всегда выше содержимого: берём фактическую высоту,
            # иначе подложка закрыла бы пол-слайда.
            h = min(h, estimated_text_height(element))
        return x, y, w, h

    boxes = [occupied(e) for e in elements]
    left = min(b[0] for b in boxes)
    top = min(b[1] for b in boxes)
    right = max(b[0] + b[2] for b in boxes)
    bottom = max(b[1] + b[3] for b in boxes)
    pad = 10
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Pt(left - pad),
        Pt(top - pad),
        Pt(right - left + pad * 2),
        Pt(bottom - top + pad * 2),
    )
    _translucent(shape, color)


def diagram_style(font, accent, background, text_color, palette, coverage=None):
    """Кисти для библиотеки диаграмм: заливка, подпись внутри фигуры и под ней."""
    on_accent = best_text_color(accent, palette)

    def paint(shape, role):
        shape.line.fill.background()
        if role == "hole":
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(background)
        elif role == "outline":
            shape.fill.background()
            shape.line.fill.solid()
            shape.line.fill.fore_color.rgb = RGBColor.from_string(accent)
            shape.line.width = Pt(1.25)
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(accent)

    def label(shape, text, size, outline=False):
        """Подпись внутри фигуры схемы.

        Кегль приходит снаружи и один на всю схему: раньше он считался для
        каждой фигуры отдельно, и в одном ряду шевронов «Аудит» оказывался
        крупнее соседей, а на узкой колонке подписи падали до 6 pt.
        """
        text = printable(text, font, coverage)
        frame = shape.text_frame
        frame.word_wrap = True
        # Внутренние поля по умолчанию съедают у фигуры четверть дюйма: для
        # подписи внутри шеврона это половина доступной ширины.
        frame.margin_left = frame.margin_right = Pt(1)
        frame.margin_top = frame.margin_bottom = Pt(1)
        shape.text = text
        colour = text_color if outline else on_accent
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            _font(paragraph.font, font, size, colour)
            for run in paragraph.runs:
                _font(run.font, font, size, colour)

    def caption(shapes, box, text, size=MIN_LABEL_SIZE):
        left, top, width, height = box
        frame = shapes.add_textbox(left, top, width, max(height, Pt(12)))
        frame.text_frame.word_wrap = True
        # Поля текстовой рамки по умолчанию — по 7.2 pt с каждой стороны: на
        # подписи под значком это пятая часть ширины, и слово всё-таки рвалось.
        frame.text_frame.margin_left = frame.text_frame.margin_right = Pt(1)
        frame.text_frame.margin_top = frame.text_frame.margin_bottom = Pt(1)
        frame.text_frame.text = text
        for paragraph in frame.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            _font(paragraph.font, font, size, text_color)
            for run in paragraph.runs:
                _font(run.font, font, size, text_color)

    return {"paint": paint, "label": label, "caption": caption}


# Длинная единица измерения у каждого столбца превращает диаграмму в текст:
# такую единицу несёт только заголовок оси.
MAX_INLINE_UNIT = 10


def value_axis_title(visual):
    """Что написано у оси значений: единица и, при одном ряде, его название.

    Легенда при одном ряде только повторяет сама себя, поэтому её выключаем, —
    но тогда зритель не знает, что именно измерено. Название ряда переезжает
    сюда, к единице измерения: одна короткая строка вместо лишней легенды.
    """
    unit = (visual.get("unit") or "").strip()
    series = visual.get("series") or []
    name = (series[0].get("name") or "").strip() if len(series) == 1 else ""
    if name and unit:
        return f"{name}, {unit}"
    return name or unit


def _horizontal_title(axis):
    """Развернуть заголовок оси в строку.

    По умолчанию PowerPoint кладёт заголовок оси значений на бок. В узкой рамке
    повёрнутое слово переносится по слогам («мин/ут»), поэтому держим его
    горизонтальным: короткая единица читается и сбоку от оси.
    """
    body = axis.axis_title.text_frame._txBody.find(qn("a:bodyPr"))
    if body is None:
        return
    body.set("rot", "0")
    body.set("vert", "horz")


def _number_format(values, unit):
    """Формат подписи значения: число и, если она короткая, единица.

    Дробные значения показываем с одним знаком, целые — без хвоста: «4», а не
    «4,0». Длинная единица остаётся только в заголовке оси, иначе подписи
    закрывают саму диаграмму.
    """
    fractional = any(
        isinstance(v, float) and abs(v - round(v)) > 1e-9 for v in values
    )
    base = "0.#" if fractional else "0"
    unit = unit.strip().replace('"', "")
    if unit and len(unit) <= MAX_INLINE_UNIT:
        return f'{base}" {unit}"'
    return base


def add_value_labels(chart, kind, font, color, unit):
    """Подписи значений на диаграмме — требование Приложения 1.

    Без них столбец «4» рядом со «186» — полоска в пиксель у самой оси: значение
    приходится угадывать по сетке. Подпись ставится снаружи столбца (над точкой
    у линии), чтобы не тонуть в заливке.
    """
    values = [v for series in chart.series for v in series.values if v is not None]
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = _number_format(values, unit)
    labels.number_format_is_linked = False
    labels.position = (
        XL_LABEL_POSITION.OUTSIDE_END if kind == "bar" else XL_LABEL_POSITION.ABOVE
    )
    labels.show_value = True
    _font(labels.font, font, 11, color)
    # Подпись значения PowerPoint переносит по ширине столбца: «186» и «минут»
    # расходятся на две строки и читаются как два числа. Запрещаем перенос —
    # подпись остаётся одной строкой над столбцом.
    body = labels._element.find(f"{qn('c:txPr')}/{qn('a:bodyPr')}")
    if body is not None:
        body.set("wrap", "none")


def stat_colour(accent, background, text_color):
    """Цвет крупной цифры: акцент, если он читается на фоне; иначе цвет текста.

    Для крупного кегля WCAG допускает контраст 3:1 — поэтому фирменный синий на
    белом годится для цифры, хотя для мелкого текста его было бы мало.
    """
    return accent if contrast_ratio(accent, background) >= 3.0 else text_color


def draw_stat(slide, element, font, accent, background, text_color, coverage=None):
    """Фактоид: акцентная черта, крупное число и подпись под ним."""
    x, y, w, h = element["box"]
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Pt(x),
        Pt(y + element["value_size"] * 0.18),
        Pt(STAT_BAR),
        Pt(max(8.0, h - element["value_size"] * 0.18)),
    )
    bar.name = element["id"] + ":bar"
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(accent)
    bar.line.fill.background()
    bar.shadow.inherit = False
    shape = slide.shapes.add_textbox(Pt(x + STAT_INSET), Pt(y), Pt(w - STAT_INSET), Pt(h))
    shape.name = element["id"]
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(6)
    frame.margin_top = frame.margin_bottom = Pt(0)
    first = frame.paragraphs[0]
    first.text = printable(element["value"], font, coverage)
    first.space_after = Pt(2)
    first.line_spacing = 1.0
    colour = stat_colour(accent, background, text_color)
    _font(first.font, font, element["value_size"], colour, True)
    for run in first.runs:
        _font(run.font, font, element["value_size"], colour, True)
    for line in element["text"].splitlines():
        p = frame.add_paragraph()
        p.text = printable(line, font, coverage)
        p.space_after = Pt(3)
        _font(p.font, font, element["font_size"], text_color)
        for run in p.runs:
            _font(run.font, font, element["font_size"], text_color)


def export_pptx(template_path: Path, template, deck_data, output: Path, asset_root=None):
    """asset_root — корень хранилища, где лежат картинки шаблонов и пакетов."""
    deck = Presentation(template_path)
    originals = list(deck.slides)
    original_ids = list(deck.slides._sldIdLst)
    palette = template["tokens"]["colors"]
    # Покрытие гарнитур из разбора шаблона: по нему видно, какой символ
    # шаблонный шрифт не нарисует.
    coverage = template["tokens"].get("font_coverage") or {}
    # Пустой набор — это результат разбора («ничего не повторяется»), а не его
    # отсутствие: отличаем по наличию ключа, иначе включится старое правило.
    branding = (
        {tuple(box) for box in template["branding"]}
        if "branding" in template
        else None
    )
    for slide_data in deck_data["slides"]:
        content_boxes = [
            tuple(Pt(v) for v in element["box"])
            for element in slide_data["elements"]
        ]
        source = originals[slide_data["pattern_index"]]
        orphans = (
            frozenset()
            if any(e.get("from_template") for e in slide_data["elements"])
            else _orphan_icon_row(source, deck.slide_width, deck.slide_height, branding)
        )
        if slide_data.get("drop_shapes"):
            orphans = orphans | frozenset(slide_data["drop_shapes"])
        if slide_data.get("clear_stage"):
            orphans = orphans | _stage_clutter(source, deck.slide_width, deck.slide_height, branding)
        slide = _clone_slide(
            deck,
            source,
            branding,
            content_boxes,
            orphans,
            slide_data.get("shape_fixes"),
        )
        background = slide_data.get("background") or background_color(
            slide, template["tokens"]["theme"]
        )
        # Layout already resolved the readable color; keep both layers in agreement.
        text_color = next(
            (
                e["color"]
                for e in slide_data["elements"]
                if e.get("color") and e["color"] != "auto"
            ),
            best_text_color(background, palette),
        )
        if slide_data.get("needs_scrim") or (
            slide_data.get("image_cover", 0) >= 0.5 and not slide_data.get('native_cover')
        ):
            add_scrim(slide, slide_data["elements"], background)
        roles = color_roles(
            next((e["accent"] for e in slide_data["elements"] if e.get("accent")), palette[0] if palette else "000000"),
            background,
            text_color,
            template["tokens"].get("accents") or [],
        )
        for element in slide_data["elements"]:
            x, y, w, h = [Pt(v) for v in element["box"]]
            kind = element["kind"]
            font, accent = element["font"], element["accent"]
            if kind == "text":
                shape = slide.shapes.add_textbox(x, y, w, h)
                shape.name = element["id"]
                frame = shape.text_frame
                frame.word_wrap = True
                frame.margin_left = frame.margin_right = Pt(6)
                frame.margin_top = frame.margin_bottom = Pt(4)
                align = {
                    "center": PP_ALIGN.CENTER,
                    "right": PP_ALIGN.RIGHT,
                    "justify": PP_ALIGN.JUSTIFY,
                }.get(element.get("align"), PP_ALIGN.LEFT)
                # Начертание и выравнивание — из шаблона: так новый слайд читается
                # как страница той же презентации, а не как вставка.
                bold = element.get("bold", element["role"] == "title")
                # В карточке шаблона первая строка — подзаголовок, остальное текст:
                # так карточка читается как в оригинале, а не как абзац.
                card = element.get("from_template") and element["role"] == "body"
                lines = element["text"].splitlines()
                for j, line in enumerate(lines):
                    p = frame.paragraphs[0] if j == 0 else frame.add_paragraph()
                    p.text = line
                    p.space_after = Pt(3)
                    p.alignment = align
                    line_bold = bold or (card and j == 0 and len(lines) > 1)
                    size = element["font_size"]
                    if card and j > 0:
                        size = max(10.0, size * 0.88)
                    _font(p.font, font, size, text_color, line_bold)
                    for run in p.runs:
                        _font(run.font, font, size, text_color, line_bold)
            elif kind in ("bar", "line"):
                visual = element["data"]

                def safe(value, font=font):
                    """Подпись внутри объекта — тот же текст на слайде."""
                    return printable(str(value), font, coverage)

                data = CategoryChartData()
                data.categories = [safe(c) for c in visual["categories"]]
                for series in visual["series"]:
                    data.add_series(safe(series["name"]), series["values"])
                chart = slide.shapes.add_chart(
                    XL_CHART_TYPE.COLUMN_CLUSTERED
                    if kind == "bar"
                    else XL_CHART_TYPE.LINE,
                    x,
                    y,
                    w,
                    h,
                    data,
                ).chart
                chart.has_title = False
                # Легенда нужна, когда рядов несколько. Один ряд она только
                # повторяет, забирая место у самой диаграммы: его название
                # уходит в заголовок оси значений рядом с единицей.
                chart.has_legend = len(visual["series"]) > 1
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                _font(chart.font, font, 12, text_color)
                # Подпись оси значений — единица измерения из плана. Подписи
                # категорий уже стоят под столбцами: слово «Категория» под ними
                # ничего не добавляет и выдаёт шаблон офисной диаграммы.
                for axis, label in (
                    (chart.value_axis, safe(value_axis_title(visual))),
                    (chart.category_axis, ""),
                ):
                    _font(axis.tick_labels.font, font, 11, text_color)
                    # An axis title is only truthful when the outline supplied a unit;
                    # a missing one is reported by the audit instead of invented here.
                    axis.has_title = bool(label)
                    if not label:
                        continue
                    axis.axis_title.text_frame.text = label
                    for p in axis.axis_title.text_frame.paragraphs:
                        _font(p.font, font, 11, text_color)
                    _horizontal_title(axis)
                # Подписи значений стоят всегда: при разбросе 186 и 4 короткий
                # столбец сливается с осью, и прочитать его иначе нельзя.
                add_value_labels(chart, kind, font, text_color, visual.get("unit", ""))
                accents = template["tokens"].get("accents") or []
                # Сетка офисного серого спорит с палитрой шаблона: делаем её
                # тише текста, а столбцы — шире промежутков.
                grid = chart.value_axis.major_gridlines.format.line
                grid.color.rgb = RGBColor.from_string(_mix(text_color, background, 0.18))
                grid.width = Pt(0.75)
                if kind == "bar":
                    chart.plots[0].gap_width = 60
                for i, series in enumerate(chart.series):
                    color = (
                        accents[i]
                        if i < len(accents)
                        else template["tokens"]["theme"].get(f"accent{i + 1}", accent)
                    )
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = RGBColor.from_string(color)
                    series.format.line.color.rgb = RGBColor.from_string(color)
            elif kind == "table":
                visual = element["data"]
                rows = [visual["columns"], *visual["rows"]]
                table = slide.shapes.add_table(
                    len(rows), len(rows[0]), x, y, w, h
                ).table
                for r, values in enumerate(rows):
                    for c, value in enumerate(values):
                        cell = table.cell(r, c)
                        cell.text = printable(str(value), font, coverage)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(
                            accent if r == 0 else background
                        )
                        color = (
                            best_text_color(accent, palette) if r == 0 else text_color
                        )
                        for p in cell.text_frame.paragraphs:
                            _font(p.font, font, 13, color, r == 0)
            elif kind == "image":
                icon = element.get("role") == "icon"
                draw_asset(
                    slide.shapes,
                    (int(x), int(y), int(w), int(h)),
                    element["asset"],
                    roles,
                    asset_root,
                    element.get("fit", "contain"),
                    # Значок — акцентом, но только если он читается на фоне.
                    tint=(element.get("tint") or stat_colour(accent, background, text_color))
                    if icon
                    else None,
                )
            elif kind == "stat":
                draw_stat(slide, element, font, accent, background, text_color, coverage)
            elif kind in DIAGRAMS:
                style = diagram_style(
                    font, accent, background, text_color, palette, coverage
                )
                icons = element.get("icons")
                if icons:

                    def glyph(shapes, index, box, icons=icons):
                        if index >= len(icons) or not icons[index]:
                            return False
                        return (
                            draw_asset(
                                shapes, box, icons[index], roles, asset_root,
                                tint=stat_colour(accent, background, text_color),
                            )
                            is not None
                        )

                    style["glyph"] = glyph
                DIAGRAMS[kind](
                    slide.shapes,
                    (int(x), int(y), int(w), int(h)),
                    element["data"],
                    style,
                )
        slide.notes_slide.notes_text_frame.text = slide_data["content"]["notes"]
    for identifier in original_ids:
        deck.part.drop_rel(identifier.rId)
        deck.slides._sldIdLst.remove(identifier)
    deck.save(output)
    prune_package(output)


def _rel_targets(folder, data):
    """Куда ведут ссылки одной .rels: пути внутри пакета, без внешних."""
    targets = {}
    for rid, value in re.findall(rb'Id="([^"]+)"[^>]*Target="([^"]+)"', data):
        target = value.decode("utf-8")
        if target.startswith(("http:", "https:", "mailto:")):
            continue
        resolved = os.path.normpath(os.path.join(folder, target)).replace("\\", "/")
        targets[rid.decode("utf-8")] = resolved.lstrip("/")
    return targets


def prune_package(path: Path):
    """Убрать из готового файла то, чем колода не пользуется.

    Клонирование тянет за собой весь шаблон: сорок макетов и иллюстрации всех
    его страниц. Колода из десяти слайдов весила шестнадцать мегабайт, из
    которых пятнадцать — чужие картинки. Выбрасываем макеты, на которые не
    ссылается ни один слайд, и медиа, на которые после этого не осталось ссылок;
    записи в [Content_Types].xml убираются вместе с частями, иначе PowerPoint
    считает файл повреждённым.
    """
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        payload = {name: archive.read(name) for name in names}
    keep_layouts = set()
    for name, data in payload.items():
        if name.startswith("ppt/slides/_rels/"):
            folder = name.rsplit("/_rels/", 1)[0]
            keep_layouts.update(
                target
                for target in _rel_targets(folder, data).values()
                if target.startswith("ppt/slideLayouts/")
            )
    drop = {
        name
        for name in payload
        if name.startswith("ppt/slideLayouts/slideLayout") and name not in keep_layouts
    }
    for name in list(payload):
        # Мастер перечисляет свои макеты дважды: в XML и в .rels. Правим оба.
        if not re.fullmatch(r"ppt/slideMasters/_rels/slideMaster\d+\.xml\.rels", name):
            continue
        folder = name.rsplit("/_rels/", 1)[0]
        targets = _rel_targets(folder, payload[name])
        gone = {rid for rid, target in targets.items() if target in drop}
        if not gone:
            continue
        data = payload[name].decode("utf-8")
        for rid in gone:
            data = re.sub(rf'<Relationship Id="{rid}"[^>]*/>', "", data)
        payload[name] = data.encode("utf-8")
        master = name.replace("/_rels/", "/").removesuffix(".rels")
        xml = payload[master].decode("utf-8")
        for rid in gone:
            xml = re.sub(rf'<p:sldLayoutId[^>]*r:id="{rid}"[^>]*/>', "", xml)
        payload[master] = xml.encode("utf-8")
    for name in drop:
        payload.pop(name, None)
        payload.pop(name.replace("ppt/slideLayouts/", "ppt/slideLayouts/_rels/") + ".rels", None)
    used = set()
    for name, data in payload.items():
        if name.endswith(".rels"):
            used.update(_rel_targets(name.rsplit("/_rels/", 1)[0], data).values())
    for name in [n for n in payload if n.startswith("ppt/media/") and n not in used]:
        payload.pop(name)
    types = "[Content_Types].xml"
    xml = payload[types].decode("utf-8")
    for name in [n for n in names if n not in payload]:
        xml = re.sub(rf'<Override PartName="/{re.escape(name)}"[^>]*/>', "", xml)
    payload[types] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            if name in payload:
                archive.writestr(name, payload[name])
    return len(names) - len(payload)


def _chart_labels(chart, slide_index):
    """Что у диаграммы подписано: легенда, оси, значения.

    Приложение 1 требует подписей осей, единиц и легенды. Проверять это по
    модели вёрстки нечестно: решение принимает экспорт. Смотрим готовый файл.
    """
    try:
        series = len(chart.plots[0].series)
        labels = bool(chart.plots[0].data_labels.show_value)
    except (IndexError, ValueError, AttributeError):
        series, labels = 0, False
    axes = []
    for name in ("value_axis", "category_axis"):
        try:
            axes.append(bool(getattr(chart, name).has_title))
        except (ValueError, AttributeError):
            axes.append(False)
    return {
        "slide": slide_index,
        "series": series,
        "legend": bool(chart.has_legend),
        "axis_title": any(axes),
        "value_labels": labels,
    }


def branding_drift(template_path: Path, deck_data, output: Path, tolerance=0.004):
    """Элементы шаблона, которые уехали с места при клонировании.

    Логотип, колонтитул и декор копируются из страницы-прототипа как есть,
    поэтому сдвинуться они могут только если кто-то тронет экспорт. Приложение 1
    требует это проверять, и проверка обходится дешевле, чем разбор такого бага
    на сцене: сравниваем координаты фигур по их идентификаторам.
    """
    source = Presentation(template_path)
    result = Presentation(output)
    width = float(source.slide_width or 1)
    height = float(source.slide_height or 1)
    moved = []
    originals = list(source.slides)
    for index, slide in enumerate(result.slides):
        pattern = deck_data["slides"][index].get("pattern_index")
        if pattern is None or pattern >= len(originals):
            continue
        # Сопоставляем по паре «номер и имя»: новые блоки получают собственные
        # номера, которые могут совпасть с номерами фигур шаблона.
        before = {
            (shape.shape_id, shape.name): (shape.left, shape.top, shape.width, shape.height)
            for shape in originals[pattern].shapes
        }
        for shape in slide.shapes:
            expected = before.get((shape.shape_id, shape.name))
            if not expected:
                continue
            actual = (shape.left, shape.top, shape.width, shape.height)
            shift = max(
                abs(actual[0] - expected[0]) / width,
                abs(actual[1] - expected[1]) / height,
                abs(actual[2] - expected[2]) / width,
                abs(actual[3] - expected[3]) / height,
            )
            if shift > tolerance:
                moved.append(
                    {
                        "slide": index,
                        "name": shape.name,
                        "shift": round(shift, 4),
                    }
                )
    return moved


def verify_pptx(path: Path, expected_slides: int):
    """Confirm the exported deck opens and carries editable objects.

    The brief rejects slides exported as one flat image, so the service proves the
    opposite on every export instead of asserting it in documentation.
    """
    report = {
        "opens": False,
        "slides": 0,
        "native_objects": 0,
        "raster_slides": [],
        "charts": [],
        "moved_branding": [],
    }
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("Exported PPTX archive is corrupt")
    deck = Presentation(path)
    report["opens"] = True
    report["slides"] = len(deck.slides)
    for index, slide in enumerate(deck.slides):
        editable = 0
        pictures = 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                editable += 1
            elif shape.has_table or shape.has_chart:
                editable += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                editable += 1
            if shape.has_chart:
                report["charts"].append(_chart_labels(shape.chart, index))
        report["native_objects"] += editable
        if not editable and pictures:
            report["raster_slides"].append(index)
    if report["slides"] != expected_slides:
        raise RuntimeError(
            f"Exported {report['slides']} slides, expected {expected_slides}"
        )
    if report["raster_slides"]:
        raise RuntimeError(
            "Slides exported without editable objects: "
            + ", ".join(str(i) for i in report["raster_slides"])
        )
    return report


def convert_pdf(pptx_path: Path, destination: Path):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not executable and mac.exists():
        executable = str(mac)
    if not executable:
        raise RuntimeError(
            "LibreOffice is required for PDF and faithful HTML export; use the backend Docker image"
        )
    with tempfile.TemporaryDirectory(prefix="designer-office-") as temp:
        temp = Path(temp)
        fontconfig = temp / "fonts.conf"
        font_dirs = [
            Path(__file__).parent / "assets/fonts",
            Path("/System/Library/Fonts"),
            Path("/Library/Fonts"),
            Path.home() / "Library/Fonts",
        ]
        fontconfig.write_text(
            '<?xml version="1.0"?><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
            + "".join(
                "<dir>" + html.escape(str(p.resolve())) + "</dir>"
                for p in font_dirs
                if p.exists()
            )
            + "<cachedir>"
            + str(temp / "font-cache")
            + "</cachedir></fontconfig>"
        )
        result = subprocess.run(
            [
                executable,
                "-env:UserInstallation=" + (temp / "profile").as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp),
                str(pptx_path.resolve()),
            ],
            capture_output=True,
            timeout=90,
            check=False,
            env={**os.environ, "FONTCONFIG_FILE": str(fontconfig)},
        )
        converted = temp / (pptx_path.stem + ".pdf")
        if result.returncode or not converted.exists():
            raise RuntimeError("LibreOffice failed to render presentation")
        shutil.copyfile(converted, destination)


def render_slides(pdf_path: Path, width=1280):
    """Слайды как PNG — то, что реально увидит зритель, а не наше представление."""
    import pymupdf

    pages = []
    with pymupdf.open(pdf_path) as document:
        for page in document:
            zoom = width / page.rect.width if page.rect.width else 1
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
            pages.append(pixmap.tobytes("png"))
    return pages


def template_thumbnail(pptx_path: Path, destination: Path, page=0, width=640):
    """Страница шаблона картинкой — чтобы в интерфейсе шаблоны были различимы.

    Рендерит LibreOffice, как и всё остальное: миниатюра совпадает с тем, что
    пользователь увидит в PowerPoint, а не с нашей догадкой о шаблоне.
    """
    with tempfile.TemporaryDirectory(prefix="designer-thumb-") as temp:
        pdf = Path(temp) / "template.pdf"
        convert_pdf(pptx_path, pdf)
        pages = render_slides(pdf, width)
    if not pages:
        raise RuntimeError("Template has no pages to render")
    destination.write_bytes(pages[min(page, len(pages) - 1)])
    return destination


def sample_backgrounds(pptx_path: Path, content_region=(0.05, 0.23, 0.95, 0.87)):
    """Измерить фон каждого слайда шаблона по его реальному рендеру.

    XML врёт: фон может прийти от мастера, от полноэкранной фигуры на макете или
    от картинки. Единственный надёжный источник — то, что видно на странице.
    Возвращает по слайду среднюю светлоту рабочей области и разброс: тёмный фон
    требует светлого текста, пёстрый — подложки под текстом.
    """
    import pymupdf

    with tempfile.TemporaryDirectory(prefix="designer-bg-") as temp:
        pdf = Path(temp) / "template.pdf"
        convert_pdf(pptx_path, pdf)
        stats = []
        with pymupdf.open(pdf) as document:
            for page in document:
                pixmap = page.get_pixmap(dpi=36)
                left = int(pixmap.width * content_region[0])
                right = int(pixmap.width * content_region[2])
                top = int(pixmap.height * content_region[1])
                bottom = int(pixmap.height * content_region[3])
                total = 0.0
                squares = 0.0
                count = 0
                for y in range(top, bottom, 2):
                    for x in range(left, right, 2):
                        r, g, b = pixmap.pixel(x, y)[:3]
                        luma = (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255
                        total += luma
                        squares += luma * luma
                        count += 1
                if not count:
                    stats.append({"luma": 1.0, "spread": 0.0})
                    continue
                mean = total / count
                variance = max(0.0, squares / count - mean * mean)
                stats.append(
                    {"luma": round(mean, 4), "spread": round(variance**0.5, 4)}
                )
    return stats


HTML_EXPORT_MARKER = '<!-- designer-html:live-text-v1 -->'


def export_html(pdf_path: Path, destination: Path, title: str, language: str = "ru"):
    import pymupdf
    from lxml import etree

    with pymupdf.open(pdf_path) as pdf:
        pages = []
        for page in pdf:
            # Keep selectable/searchable SVG text. Serialize numeric character
            # references as UTF-8 too, so Russian phrases can be found in the file.
            svg = page.get_svg_image(text_as_path=False)
            root = etree.fromstring(svg.encode('utf-8'), etree.XMLParser(resolve_entities=False, no_network=True))
            svg = etree.tostring(root, encoding='unicode')
            pages.append('<section class="slide">' + svg + "</section>")
    destination.write_text(
        '<!doctype html>' + HTML_EXPORT_MARKER + '<html lang="' + html.escape(language or 'ru', quote=True)
        + '"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>'
        + html.escape(title)
        + "</title><style>body{margin:0;background:#222}.slide{margin:24px auto;width:min(96vw,1280px);background:white}.slide svg{width:100%;height:auto;display:block}@media print{body{background:white}.slide{margin:0;width:100%;break-after:page}}</style>"
        + "".join(pages)
        + "</html>",
        encoding="utf-8",
    )
