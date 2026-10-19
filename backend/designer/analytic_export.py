"""Native editable charts and KPI objects for the extended data contract."""
import re
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement

MONTHS = ('янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек')

def display_category(value):
    match = re.fullmatch(r'(\d{4})-(0[1-9]|1[0-2])(?:-01)?', str(value))
    return f'{MONTHS[int(match[2])-1]} {match[1][2:]}' if match else str(value)

def display_number(value):
    text = str(value)
    if not re.fullmatch(r'[-+]?\d+(?:[.,]\d+)?', text) or re.match(r'[-+]?0\d', text):
        return text
    integer, dot, fraction = text.replace(',', '.').partition('.')
    sign = '-' if integer.startswith('-') else '+' if integer.startswith('+') else ''
    return sign + f'{int(integer.lstrip("+-")):,}'.replace(',', '\u00a0') + (','+fraction if dot else '')

def number_format(visual):
    values = [abs(v) for s in visual['series'] for v in s['values'] if v]
    if values and min(values) < .01:
        return '0.00E+00'
    maximum = max(values, default=0)
    return '#,##0.0,,," млрд"' if maximum >= 1e9 else '#,##0.0,," млн"' if maximum >= 1e6 else '#,##0.0," тыс."' if maximum >= 1e4 else '#,##0.##'


def add_chart(slide, visual, box, font, palette, text_color):
    x, y, w, h = box
    distinct = []
    def distance(a,b):return sum((int(a[i:i+2],16)-int(b[i:i+2],16))**2 for i in (0,2,4))**.5
    for color in palette:
        if all(distance(color,other)>75 for other in distinct):distinct.append(color)
    if len(distinct)<2 and distinct:
        first=distinct[0]
        distinct.append(''.join(f'{round(int(first[i:i+2],16)*.4+int(text_color[i:i+2],16)*.6):02X}' for i in (0,2,4)))
    needed=len(visual.get('categories',[])) if visual['kind']=='pie' else len(visual.get('series',[]))
    if distinct and len(distinct)<needed:
        first=distinct[0]
        for target,weight in [('FFFFFF',.35),('000000',.25),('FFFFFF',.6),('000000',.7)]:
            color=''.join(f'{round(int(first[i:i+2],16)*(1-weight)+int(target[i:i+2],16)*weight):02X}' for i in (0,2,4))
            if all(distance(color,other)>40 for other in distinct):distinct.append(color)
    palette=distinct or palette
    kind = visual["kind"]
    if kind == "kpi":
        shape = slide.shapes.add_textbox(x, y, w, h)
        frame = shape.text_frame
        frame.word_wrap = True
        for i, series in enumerate(visual["series"]):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            value = series["values"][0]
            rendered = str(int(value)) if value is not None and float(value).is_integer() else str(value)
            if value is not None and abs(value)>=.01 and not float(value).is_integer():
                rendered=f'{value:.2f}'.rstrip('0').rstrip('.')
            if value is not None and 'e' not in rendered.lower():
                rendered = display_number(rendered)
            unit = {'RUB':'₽','rub':'₽'}.get(visual.get('unit'),visual.get('unit',''))
            p.text = rendered + ' ' + unit if value is not None else 'Нет данных'
            p.font.name, p.font.size, p.font.bold = font, Pt(40), True
            p.font.color.rgb = RGBColor.from_string(text_color)
            p = frame.add_paragraph()
            p.text = series["name"]
            p.font.name, p.font.size = font, Pt(18)
            p.font.color.rgb = RGBColor.from_string(text_color)
        return
    if kind == "scatter":
        data = XyChartData()
        for series in visual["series"]:
            item = data.add_series(series["name"])
            for a, b in zip(visual["x_values"], series["values"]):
                if b is not None:
                    item.add_data_point(a, b)
    else:
        data = CategoryChartData()
        data.categories = [display_category(c) for c in visual["categories"]]
        for series in visual["series"]:
            data.add_series(series["name"], series["values"])
    types = dict(bar=XL_CHART_TYPE.COLUMN_CLUSTERED, grouped_bar=XL_CHART_TYPE.COLUMN_CLUSTERED,
                 stacked_bar=XL_CHART_TYPE.COLUMN_STACKED, histogram=XL_CHART_TYPE.COLUMN_CLUSTERED,
                 line=XL_CHART_TYPE.LINE, pie=XL_CHART_TYPE.PIE, scatter=XL_CHART_TYPE.XY_SCATTER)
    horizontal = kind in {"bar", "grouped_bar", "stacked_bar"} and max((len(c) for c in visual["categories"]), default=0) > 22
    if horizontal:
        types[kind] = XL_CHART_TYPE.BAR_STACKED if kind == "stacked_bar" else XL_CHART_TYPE.BAR_CLUSTERED
    chart = slide.shapes.add_chart(types[kind], x, y, w, h, data).chart
    chart.has_title = False
    chart.font.name, chart.font.size = font, Pt(12)
    chart.font.color.rgb = RGBColor.from_string(text_color)
    chart.has_legend = kind == "pie" or len(visual["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name, chart.legend.font.size = font, Pt(12)
        chart.legend.font.color.rgb = RGBColor.from_string(text_color)
    plot = chart.plots[0]
    if kind != "scatter":
        plot.has_data_labels = True
        plot.data_labels.show_value = kind not in {'pie', 'line', 'stacked_bar'} and len(visual['categories']) <= 8 and len(visual['series']) == 1
        plot.data_labels.show_percentage = kind == "pie"
        plot.data_labels.font.name, plot.data_labels.font.size = font, Pt(12)
        plot.data_labels.font.color.rgb = RGBColor.from_string(text_color)
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind == 'pie' else XL_LABEL_POSITION.CENTER if kind == 'stacked_bar' else XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.number_format = "0.0%" if kind == "pie" else number_format(visual)
    if kind == "stacked_bar":
        plot.overlap = 100
    if kind == "histogram":
        plot.gap_width = 0
    elif kind in {'bar', 'grouped_bar', 'stacked_bar'}:
        plot.gap_width = 65
    if kind != "pie":
        for axis, label in ((chart.value_axis, visual.get("y_label") or visual.get("unit", "")), (chart.category_axis, visual.get("x_label", ""))):
            axis.tick_labels.font.name, axis.tick_labels.font.size = font, Pt(12)
            axis.tick_labels.font.color.rgb = RGBColor.from_string(text_color)
            axis.has_title = bool(label)
            if label:
                axis.axis_title.text_frame.text = label
                for p in axis.axis_title.text_frame.paragraphs:
                    p.font.name, p.font.size = font, Pt(11)
                    p.font.color.rgb = RGBColor.from_string(text_color)
        if horizontal:
            chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.number_format = number_format(visual)
        grid = chart.value_axis.major_gridlines.format.line
        grid.color.rgb = RGBColor.from_string('DCE3EB' if text_color != 'FFFFFF' else '526071')
        grid.width = Pt(.5)
        for axis in (chart.category_axis, chart.value_axis):
            axis.format.line.fill.background()
    # LibreOffice imports absent invertIfNegative as absolute-valued bars.
    # Explicit false preserves the sign in PDF as well as the editable workbook.
    for i, series in enumerate(chart.series):
        if kind in {'bar', 'grouped_bar', 'stacked_bar', 'histogram'}:
            series.invert_if_negative = False
        color = palette[i % len(palette)] if palette else "0077FF"
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        series.format.line.color.rgb = RGBColor.from_string(color)
        if horizontal and plot.data_labels.show_value:
            # Outside negative endpoints share the category-label area in Office
            # and LibreOffice. Keep their values inside the bar, with contrast.
            # Tiny bars cannot hold text: the numeric axis still carries scale.
            scale = max((abs(v) for s in visual['series'] for v in s['values'] if v is not None), default=0)
            for j, value in enumerate(visual['series'][i]['values']):
                if value is None:continue
                label = series.points[j].data_label
                label.position = XL_LABEL_POSITION.INSIDE_END if value < 0 else XL_LABEL_POSITION.OUTSIDE_END
                label.font.name, label.font.size = font, Pt(12)
                r, g, b = (int(color[k:k+2], 16) for k in (0, 2, 4))
                label.font.color.rgb = RGBColor.from_string(('FFFFFF' if .2126*r+.7152*g+.0722*b < 155 else '111827') if value < 0 else text_color)
                # LibreOffice does not inherit these settings once any point has
                # an override: write format and visibility for every point.
                node = label._get_or_add_dLbl()
                fmt = OxmlElement('c:numFmt');fmt.set('formatCode', number_format(visual));fmt.set('sourceLinked','0');node.insert(1,fmt)
                show = OxmlElement('c:showVal');show.set('val','1');node.append(show)
                if value < 0 and scale and abs(value) < .15*scale:
                    hidden = OxmlElement('c:delete');hidden.set('val','1');node.insert(1,hidden)
        if kind == 'line':
            series.format.line.width = Pt(2.5)
            series.marker.style = XL_MARKER_STYLE.CIRCLE if len(visual['categories']) <= 12 or any(v is None for v in visual['series'][i]['values']) else XL_MARKER_STYLE.NONE
            series.marker.size = 5
        if kind == "scatter":
            series.format.line.fill.background()
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 5
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.marker.format.line.fill.background()
        if kind == "pie":
            for j, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(palette[j % len(palette)] if palette else "0077FF")
    # Explicit gaps: no interpolation and no synthetic zero at missing points.
    node = OxmlElement("c:dispBlanksAs")
    node.set("val", "gap")
    chart._chartSpace.chart.append(node)
    if kind == "scatter":
        from pptx.oxml.ns import qn
        chart._chartSpace.find('.//' + qn('c:scatterStyle')).set('val', 'marker')


def add_source(slide, visual, box, font, color):
    if not visual.get("source"):
        return
    x, y, w, h = box
    frame = slide.shapes.add_textbox(x, y + h - Pt(24), w, Pt(24)).text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = "Источник: " + re.sub(r'\s*/\s*(?:1|Sheet1)$', '', visual["source"])
    p.font.name, p.font.size = font, Pt(12)
    p.font.color.rgb = RGBColor.from_string(color)


def numeric_cell(value):
    return bool(re.fullmatch(r"[-+]?\d[\d\s,.]*(?:\s*[%₽])?", str(value)))
