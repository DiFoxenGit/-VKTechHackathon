"""PPTX design extraction; never executes embedded content or external links."""

import csv
import io
import json
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pypdf import PdfReader

from .fonts import choose_font, embedded_coverage, font_coverage, resolve_font

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

# Region a designer fills with content, as a fraction of the slide. Shapes that
# carry sample content here are replaced during export; everything else is branding.
CONTENT_REGION = (0.05, 0.23, 0.95, 0.87)

# Версия разбора. Меняется, когда правила извлечения меняют результат: шаблоны,
# разобранные старой версией, переразбираются при старте, иначе экспорт и аудит
# работали бы по устаревшему «паспорту» файла.
PARSER_VERSION = 19


def is_footer_placeholder(shape):
    return shape.is_placeholder and any(
        name in str(shape.placeholder_format.type)
        for name in ("FOOTER", "DATE", "SLIDE_NUMBER")
    )


def shape_key(shape, width, height):
    """Положение фигуры с точностью, достаточной, чтобы узнать её на другом слайде."""
    return (
        round(shape.left / width, 3),
        round(shape.top / height, 3),
        round(shape.width / width, 3),
        round(shape.height / height, 3),
    )


def branding_boxes(slides, width, height, share=0.4):
    """Фигуры, повторяющиеся на многих слайдах: логотип, плашка, колонтитул.

    Всё остальное на конкретном слайде — образец содержания: серые кружки под
    фото, рамки под иконки, демонстрационные иллюстрации. Единственный честный
    признак брендинга — повторяемость, а не координаты: в разных шаблонах
    логотип живёт где угодно.
    """
    counts = Counter()
    for slide in slides:
        seen = set()
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                continue
            if shape.is_placeholder:
                continue
            seen.add(shape_key(shape, width, height))
        counts.update(seen)
    threshold = max(2, int(len(slides) * share))
    return {box for box, count in counts.items() if count >= threshold}


def master_text_defaults(master, theme, font_scheme=None):
    """Типографика уровня мастера: p:txStyles — последний источник правды.

    Стиль заголовка и текста в шаблоне часто не написан ни на слайде, ни на
    макете: он объявлен один раз в мастере. Без него мы бы подставляли свой
    кегль и цвет, то есть ломали дизайн-систему шаблона.
    """
    defaults = {}
    styles = master._element.find(qn("p:txStyles"))
    if styles is None:
        return defaults
    for role, tag in (("title", "p:titleStyle"), ("body", "p:bodyStyle")):
        node = styles.find(qn(tag))
        if node is None:
            continue
        level = node.find(qn("a:lvl1pPr"))
        if level is None:
            continue
        style = {"font": None, "size": None, "bold": None, "color": None, "align": None}
        if level.get("algn"):
            style["align"] = level.get("algn")
        run = level.find(qn("a:defRPr"))
        if run is not None:
            if run.get("sz"):
                style["size"] = round(int(run.get("sz")) / 100, 2)
            if run.get("b") is not None:
                style["bold"] = run.get("b") in ("1", "true")
            latin = run.find(qn("a:latin"))
            if latin is not None and latin.get("typeface"):
                face = latin.get("typeface")
                style["font"] = resolve_font(face, font_scheme)
            srgb = run.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
            scheme = run.find(f"{qn('a:solidFill')}/{qn('a:schemeClr')}")
            if srgb is not None:
                style["color"] = srgb.get("val")
            elif scheme is not None:
                style["color"] = theme.get(scheme.get("val"))
        defaults[role] = style
    return defaults


def frame_defaults(shape, theme, font_scheme=None):
    """Стиль, объявленный в самой рамке: <a:lstStyle> первого уровня.

    В шаблонах формат заголовка обычно живёт здесь — на плейсхолдере макета, а не
    в runs слайда, где текста может не быть вовсе.
    """
    style = {"font": None, "size": None, "bold": None, "color": None, "align": None}
    if not shape.has_text_frame:
        return style
    body = shape.text_frame._txBody
    level = body.find(f"{qn('a:lstStyle')}/{qn('a:lvl1pPr')}")
    if level is None:
        return style
    if level.get("algn"):
        style["align"] = level.get("algn")
    run = level.find(qn("a:defRPr"))
    if run is None:
        return style
    if run.get("sz"):
        style["size"] = round(int(run.get("sz")) / 100, 2)
    if run.get("b") is not None:
        style["bold"] = run.get("b") in ("1", "true")
    latin = run.find(qn("a:latin"))
    if latin is not None and latin.get("typeface"):
        face = latin.get("typeface")
        style["font"] = resolve_font(face, font_scheme)
    srgb = run.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
    scheme = run.find(f"{qn('a:solidFill')}/{qn('a:schemeClr')}")
    if srgb is not None:
        style["color"] = srgb.get("val")
    elif scheme is not None:
        style["color"] = theme.get(scheme.get("val"))
    return style


def text_style(shape, theme, font_scheme=None):
    """Как шаблон оформляет текст в этой рамке.

    Мы не изобретаем типографику: заголовок нового слайда получает кегль, цвет,
    начертание и выравнивание того блока, который занимал его место в шаблоне.
    Берём первый непустой абзац — он задаёт тон всей рамке.
    """
    style = {
        "font": None,
        "size": None,
        "bold": None,
        "color": None,
        "align": None,
        "caps": False,
    }
    for paragraph in shape.text_frame.paragraphs:
        if not paragraph.text.strip() and not paragraph.runs:
            continue
        if paragraph.alignment is not None:
            style["align"] = str(paragraph.alignment).split(" ")[0].lower()
        for font in [*(run.font for run in paragraph.runs), paragraph.font]:
            if style["font"] is None and font.name:
                style["font"] = resolve_font(font.name, font_scheme)
            if style["size"] is None and font.size:
                style["size"] = round(font.size.pt, 2)
            if style["bold"] is None and font.bold is not None:
                style["bold"] = bool(font.bold)
            if style["color"] is None:
                value = color_value(font.color, theme)
                if value:
                    style["color"] = value
        break
    for key, value in frame_defaults(shape, theme, font_scheme).items():
        if style.get(key) is None:
            style[key] = value
    text = shape.text.strip()
    style["caps"] = bool(text) and text == text.upper() and any(c.isalpha() for c in text)
    return style


# Слова, по которым узнаётся назначение страницы шаблона. Роль решает, где
# страницу уместно использовать: обложку — первой, «спасибо» — последней.
ROLE_WORDS = {
    "agenda": ("содержание", "оглавление", "повестка", "agenda", "contents"),
    "closing": ("спасибо", "благодарю", "контакты", "вопросы", "thank", "contacts"),
    "team": ("команда", "о нас", "участники", "спикер", "team", "about us"),
    "section": ("раздел", "часть", "глава", "section", "chapter"),
}


def classify_pattern(pattern, index, total):
    """Назначение страницы шаблона: обложка, содержание, раздел, команда, контент.

    Шаблоны содержат образцы: титульный слайд с именем автора, страницу
    содержания, разделители. Ставить контент на разделитель — то же, что печатать
    текст на обложке книги, поэтому роль важна при подборе.
    """
    title_box = pattern.get("title_box")
    text = " ".join(shape["text"] for shape in pattern["shapes"]
                    if shape.get("box") == title_box).lower()
    for role, words in ROLE_WORDS.items():
        if any(word in text for word in words):
            return role
    slots = pattern.get("text_slots", 0)
    if slots <= 1:
        return "section"
    return "content"


def classify_patterns(patterns):
    """Use the whole template to distinguish early covers from repeat dividers.

    Filenames and layout names are deliberately irrelevant. Empty placeholders
    carry the same typographic evidence as filled ones.
    """
    title_sizes = [slot['style']['size'] for pattern in patterns for slot in pattern['slots']
                   if slot['role'] == 'title' and slot['style'].get('size')]
    large = percentile(title_sizes, 2 / 3) if title_sizes else 28
    eligible = []
    for index, pattern in enumerate(patterns):
        pattern['role'] = classify_pattern(pattern, index, len(patterns))
        title = next((s for s in pattern['slots'] if s['role'] == 'title'), {})
        box = title.get('box') or {}
        size = title.get('style', {}).get('size')
        sparse = pattern['text_slots'] <= 4 and sum(
            s['length'] for s in pattern['slots'] if s['role'] == 'body'
        ) <= 220
        large_title = size >= large if size else box.get('h', 0) >= 0.15
        eligible.append(bool(title and sparse and large_title))
        pattern['role_evidence'] = {'large_title': bool(large_title), 'sparse': sparse,
                                    'opening': index < 3}
    # At most two opening alternatives; later matching pages are dividers when
    # the same composition recurs throughout the file.
    seeds = [i for i, p in enumerate(patterns[:3]) if eligible[i]
             and p['role'] not in ('agenda', 'closing', 'team', 'section')]
    # A single empty title naturally starts as 'section' before the collection
    # supplies context (e.g. a cover placeholder inherited only from its layout).
    seeds += [i for i, p in enumerate(patterns[:3]) if eligible[i]
              and p['role'] == 'section'
              and not any(word in ' '.join(s['text'] for s in p['shapes']).lower()
                          for word in ROLE_WORDS['section'])]
    seeds = sorted(set(seeds))[:2]

    def similar(first, second):
        a, b = first.get('title_box') or {}, second.get('title_box') or {}
        return all(abs(a.get(key, 0) - b.get(key, 0)) <= tolerance
                   for key, tolerance in (('x', .12), ('y', .08), ('w', .2), ('h', .08)))

    for index, pattern in enumerate(patterns):
        matches = [i for i, other in enumerate(patterns) if eligible[i] and similar(pattern, other)]
        opening_match = any(similar(pattern, patterns[i])
                            and pattern['background'] == patterns[i]['background']
                            and pattern['background_kind'] == patterns[i]['background_kind'] for i in seeds)
        if index in seeds:
            pattern['role'] = 'cover'
        elif eligible[index] and pattern['role'] in ('content', 'section') and opening_match:
            if len(matches) >= 3:
                pattern['role'] = 'section'
            elif pattern['role'] == 'content':
                pattern['role'] = 'cover'
        pattern['role_evidence']['similar_pages'] = len(matches)


def branding_score(slide, width, height, background):
    """Reward visible identity outside the content band, not arbitrary artwork."""
    marks = set()
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        for shape in owner.shapes:
            x, y, w, h = shape_key(shape, width, height)
            area = max(0, min(1, x + w) - max(0, x)) * max(0, min(1, y + h) - max(0, y))
            if area <= 0.0001:
                continue
            if owner is slide and is_footer_placeholder(shape) and shape.text.strip():
                marks.add((x, y, w, h))
            elif not shape.is_placeholder and 0.0001 < area < 0.12:
                edge = y + h <= 0.2 or y >= 0.85 or x + w <= 0.05 or x >= 0.95
                visible = shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP)
                visible = visible or bool(fill_colour(shape))
                if edge and visible and (not shape.has_text_frame or not shape.text.strip()):
                    marks.add((x, y, w, h))
    colored = relative_luminance(background['color']) < 0.85
    return round(min(0.6, len(marks) * 0.2) + (0.3 if colored else 0.0), 3)


def fill_colour(shape):
    """Цвет сплошной заливки фигуры, если он есть."""
    try:
        if shape.fill.type == 1 and shape.fill.fore_color.type == 1:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, ValueError, TypeError):
        return None
    return None


def looks_like_placeholder_art(shape, width, height):
    """Серая пустая фигура заметного размера — это рамка под фото или иконку.

    Такие заглушки переезжают на готовый слайд пустыми пятнами. Фирменный декор
    отличается цветом: он взят из палитры шаблона, а не из серой шкалы.
    """
    colour = fill_colour(shape)
    if colour is None:
        return False
    try:
        r, g, b = (int(colour[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    grey = max(r, g, b) - min(r, g, b) <= 24
    area = (shape.width / width) * (shape.height / height)
    # От половины процента слайда: мельче — это точки и линии оформления.
    return grey and 0.005 < area < 0.6


def preserved_shape(shape, width, height, branding=None):
    """True when export keeps this template shape: artwork, footers, page numbers.

    Export and audit must agree on what survives cloning, otherwise the audit
    would check collisions against shapes that are not on the finished slide.
    """
    if shape.has_table or shape.has_chart:
        return False
    # Плейсхолдер — это место под контент, а не украшение. Пустая рамка под фото
    # или иконку, перенесённая в результат, читается как мусор на слайде.
    # Исключение — колонтитул, дата и номер: они часть фирменного оформления.
    if shape.is_placeholder:
        return is_footer_placeholder(shape)
    x, y, w, h = (
        shape.left / width,
        shape.top / height,
        shape.width / width,
        shape.height / height,
    )
    left, top, right, bottom = CONTENT_REGION
    inside = min(y + h, bottom) > max(y, top) and min(x + w, right) > max(x, left)
    # Заполненный текст — всегда образец: его заменяет новый контент.
    if shape.has_text_frame and shape.text.strip():
        return False
    # Дальше — фигуры без текста. Это и есть фирменное оформление страницы:
    # плашки, рамки карточек, линии, иллюстрации. Выбрасывать их целиком значит
    # оставить от шаблона только шрифт и палитру, поэтому убираем лишь заглушки.
    if looks_like_placeholder_art(shape, width, height):
        return False
    return True


def stage_clutter(box, branding=()):
    """Мелкий образец на сцене страницы: стрелка, соединитель, значок, точка.

    Когда слайд собран своей композицией (фактоиды, иконки, картинка), такие
    фигуры образца ни к чему не относятся и читаются как мусор. Брендинг,
    который повторяется по шаблону, и рисунок макета сюда не входят.
    """
    if box.get("layout"):
        return False
    key = tuple(round(box[k], 3) for k in ("x", "y", "w", "h"))
    if key in {tuple(b) for b in branding or ()}:
        return False
    left, top, right, bottom = CONTENT_REGION
    cx, cy = box["x"] + box["w"] / 2, box["y"] + box["h"] / 2
    return box["w"] * box["h"] < 0.05 and left < cx < right and top < cy < bottom


def relative_luminance(value):
    rgb = [int(value[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    rgb = [c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4 for c in rgb]
    return sum(c * k for c, k in zip(rgb, (0.2126, 0.7152, 0.0722)))


def contrast_ratio(first, second):
    """WCAG contrast between two RRGGBB colors; 1.0 when identical, 21.0 black on white."""
    try:
        a, b = relative_luminance(first), relative_luminance(second)
    except (ValueError, IndexError, TypeError):
        return 21.0
    return (max(a, b) + 0.05) / (min(a, b) + 0.05)


def best_text_color(background, palette, minimum=3.0):
    """Читаемый цвет текста из палитры шаблона.

    Одного WCAG мало: на фирменном синем чёрный формально контрастнее белого,
    но так не верстает никто. Поэтому среди цветов, которые проходят порог,
    выбирается тот, что уводит светлоту в сторону от фона: на тёмном — самый
    светлый, на светлом — самый тёмный. Если порог не проходит никто, берём
    максимальный контраст, какой есть.
    """
    options = palette or ["000000", "FFFFFF"]
    fallback = max(options, key=lambda c: contrast_ratio(background, c))
    try:
        luma = relative_luminance(background)
    except (ValueError, IndexError, TypeError):
        return fallback
    # Сторона, в которую уводим светлоту: на тёмном фоне — светлее, на светлом —
    # темнее. Внутри стороны берём самый контрастный вариант палитры.
    lighter = luma < 0.4
    same_side = [
        c
        for c in options
        if (relative_luminance(c) > luma) == lighter
    ]
    if same_side:
        choice = max(same_side, key=lambda c: contrast_ratio(background, c))
        if contrast_ratio(background, choice) >= minimum:
            return choice
    return fallback


def background_info(slide, theme):
    """Чем залит слайд: сплошной цвет, градиент или картинка.

    Для сплошного цвета можно честно посчитать контраст текста. Для градиента и
    фотографии цвет в конкретной точке неизвестен, поэтому вёрстка обязана
    положить под текст подложку, а не надеяться на удачу.
    """
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = owner._element.cSld.bg
        if bg is None:
            continue
        if bg.xpath(".//a:blipFill"):
            return {"color": theme.get("lt1", "FFFFFF"), "kind": "image"}
        rgb = bg.xpath(".//a:solidFill/a:srgbClr/@val")
        if rgb:
            return {"color": rgb[0], "kind": "solid"}
        scheme = bg.xpath(".//a:solidFill/a:schemeClr/@val")
        if scheme and scheme[0] in theme:
            return {"color": theme[scheme[0]], "kind": "solid"}
        stops = bg.xpath(".//a:gsLst/a:gs//a:srgbClr/@val")
        if stops:
            # У градиента берём самый светлый край: подложка всё равно обязательна.
            return {
                "color": max(stops, key=relative_luminance),
                "kind": "gradient",
            }
    return {"color": theme.get("lt1", "FFFFFF"), "kind": "solid"}


def background_color(slide, theme):
    """Цвет фона слайда; для картинки и градиента — приближение."""
    return background_info(slide, theme)["color"]


def title_background(slide, box, width, height, base_color):
    """Sample a raster background at the designer's title slot, without Office.

    A photographic cover's entire canvas can be bright while its text area is
    dark. Sampling that area avoids painting a white panel over the brand.
    """
    if not box:
        return None
    from PIL import Image, ImageStat

    def sample(blob, region_box):
        try:
            with Image.open(io.BytesIO(blob)) as image:
                # Transparent PNG artwork is composited over the slide fill;
                # discarding alpha would measure bright hidden RGB pixels.
                canvas = Image.new('RGBA', image.size, '#' + base_color)
                image = Image.alpha_composite(canvas, image.convert('RGBA')).convert('RGB')
                image.thumbnail((240, 240))
                left, top = max(0, region_box['x']), max(0, region_box['y'])
                right = min(1, region_box['x'] + region_box['w'])
                bottom = min(1, region_box['y'] + region_box['h'])
                if right <= left or bottom <= top:
                    return None
                region = image.crop((int(left * image.width), int(top * image.height),
                                     max(int(left * image.width) + 1, int(right * image.width)),
                                     max(int(top * image.height) + 1, int(bottom * image.height))))
                stats = ImageStat.Stat(region)
                return {'color': ''.join(f'{round(c):02X}' for c in stats.mean),
                        'spread': round(ImageStat.Stat(region.convert('L')).stddev[0] / 255, 4)}
        except (OSError, ValueError, KeyError):
            return None
    owners = (slide, slide.slide_layout, slide.slide_layout.slide_master)
    # A full-slide picture can supply the background while p:bg stays solid.
    for owner in owners:
        for shape in reversed(list(owner.shapes)):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            x, y, w, h = shape_key(shape, width, height)
            if w * h < .85 or not (x <= box['x'] and y <= box['y']
                                   and x + w >= box['x'] + box['w']
                                   and y + h >= box['y'] + box['h']):
                continue
            cw, ch = 1 - shape.crop_left - shape.crop_right, 1 - shape.crop_top - shape.crop_bottom
            return sample(shape.image.blob, {
                'x': shape.crop_left + (box['x'] - x) / w * cw,
                'y': shape.crop_top + (box['y'] - y) / h * ch,
                'w': box['w'] / w * cw, 'h': box['h'] / h * ch,
            })
    for owner in owners:
        bg = owner._element.cSld.bg
        if bg is None:
            continue
        blips = bg.xpath('.//a:blip')
        reference = blips[0].get(qn('r:embed')) if blips else None
        if not reference:
            return None
        try:
            return sample(owner.part.related_part(reference).blob, box)
        except KeyError:
            return None
    return None


# Сколько архив может занять в памяти после распаковки. Python не распакует
# файл больше заявленного в заголовке размера, поэтому сумма заголовков —
# честная верхняя граница. Коэффициент сжатия отсекает архивы, которые
# собраны, чтобы раздуться: у обычных PPTX и картинок он не выше нескольких десятков.
ZIP_MAX_TOTAL = 300 * 1024 * 1024
ZIP_MAX_ENTRY = 100 * 1024 * 1024
ZIP_MAX_ENTRIES = 10000
ZIP_MAX_RATIO = 200
ZIP_RATIO_FLOOR = 1024 * 1024


def check_zip(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > ZIP_MAX_ENTRIES
            or sum(x.file_size for x in entries) > ZIP_MAX_TOTAL
        ):
            raise ValueError("Archive expands beyond 300 MB or 10000 entries")
        if any(x.flag_bits & 1 for x in entries):
            raise ValueError("Encrypted archives are not supported")
        for entry in entries:
            if entry.file_size > ZIP_MAX_ENTRY:
                raise ValueError(f"Archive entry {entry.filename} expands beyond 100 MB")
            if (
                entry.file_size > ZIP_RATIO_FLOOR
                and entry.file_size > ZIP_MAX_RATIO * max(entry.compress_size, 1)
            ):
                raise ValueError(f"Archive entry {entry.filename} is compressed suspiciously well")


def color_value(color, theme):
    try:
        if color.type is not None:
            if color.type == 1:
                return str(color.rgb)
            key = str(color.theme_color).split(" (")[0].lower()
            key = {
                "dark_1": "dk1",
                "dark_2": "dk2",
                "light_1": "lt1",
                "light_2": "lt2",
                "text_1": "dk1",
                "text_2": "dk2",
                "background_1": "lt1",
                "background_2": "lt2",
            }.get(key, key)
            return theme.get(key)
    except (AttributeError, ValueError):
        pass
    return None


def percentile(values, fraction):
    ordered = sorted(values)
    return ordered[min(len(ordered) - 1, int(fraction * (len(ordered) - 1) + 0.5))]


def derive_geometry(patterns):
    """Safe area and vertical guides inferred from where the template puts text.

    The service must fit an unseen template, so margins come from the file rather
    than from constants. Values are fractions of the slide; defaults apply when a
    template carries too little text to measure.
    """
    boxes = [
        shape["box"]
        for pattern in patterns
        for shape in pattern["shapes"]
        if shape["text"].strip() and 0.004 < shape["box"]["w"] * shape["box"]["h"] < 0.6
    ]
    safe_area = {"x": 0.055, "y": 0.055, "w": 0.89, "h": 0.825}
    # Outer limit: no template element sits closer to an edge than this. The safe
    # area is the design rhythm used for placement; margins are the hard boundary
    # the audit enforces, so a title that legitimately rides high is not flagged.
    margins = {"x": 0.04, "y": 0.04, "w": 0.92, "h": 0.9}
    guides = []
    if len(boxes) >= 3:
        outer_left = min(0.1, max(0.02, min(b["x"] for b in boxes)))
        outer_top = min(0.1, max(0.02, min(b["y"] for b in boxes)))
        outer_right = min(0.98, max(0.9, max(b["x"] + b["w"] for b in boxes)))
        outer_bottom = min(0.98, max(0.9, max(b["y"] + b["h"] for b in boxes)))
        margins = {
            "x": round(outer_left, 5),
            "y": round(outer_top, 5),
            "w": round(outer_right - outer_left, 5),
            "h": round(outer_bottom - outer_top, 5),
        }
        left = min(0.14, max(0.03, percentile([b["x"] for b in boxes], 0.1)))
        right = min(0.97, max(0.86, percentile([b["x"] + b["w"] for b in boxes], 0.9)))
        top = min(0.14, max(0.03, percentile([b["y"] for b in boxes], 0.1)))
        bottom = min(0.97, max(0.86, percentile([b["y"] + b["h"] for b in boxes], 0.9)))
        safe_area = {
            "x": round(left, 5),
            "y": round(top, 5),
            "w": round(right - left, 5),
            "h": round(bottom - top, 5),
        }
        # An edge repeated across slides is a guide the designer actually used.
        edges = Counter()
        for pattern in patterns:
            seen = {
                round(value, 3)
                for shape in pattern["shapes"]
                if shape["text"].strip()
                for value in (shape["box"]["x"], shape["box"]["x"] + shape["box"]["w"])
            }
            edges.update(seen)
        # Cluster near-identical edges so one guide is one design decision.
        guides = []
        for value in sorted(v for v in edges if 0 <= v <= 1):
            if guides and value - guides[-1][0] <= 0.008:
                guides[-1][1] += edges[value]
            else:
                guides.append([value, edges[value]])
        guides = [round(value, 4) for value, count in guides if count >= 3]
    return {"safe_area": safe_area, "margins": margins, "guides": guides}


# Ниже этого кегля текст на слайде не читается: это сноски, подписи внутри
# иконок и служебные строки образцов. В основную шкалу они не идут.
MIN_SCALE_SIZE = 9.0
# Столько ступеней держит любая осмысленная типографическая шкала. Всё, что
# длиннее, — это список случайно встреченных значений, а не решение дизайнера.
MAX_SCALE_STEPS = 9
# Кегль выше этой доли высоты слайда — декоративная цифра или буква во весь
# экран (166 pt у VK Tech, 288 pt у Education), а не ступень для текста.
MAX_SCALE_HEIGHT_SHARE = 0.11
# Два кегля, отличающиеся меньше чем на эту долю, — одно и то же решение,
# разъехавшееся при правках (17.5 и 18 pt).
SCALE_TOLERANCE = 0.05


def caption_sizes(counted):
    """Мелкие кегли шаблона: сноски, подписи в иконках, служебные строки.

    Они существуют в файле, поэтому выбрасывать их насовсем нечестно, но и
    верстать ими нельзя: `font_size_off_scale` по такой шкале не отсекает
    ничего.
    """
    return sorted(size for size in counted if size < MIN_SCALE_SIZE)


def type_scale(counted, height_pt=540.0):
    """Типографическая шкала шаблона: кластеры кеглей по объёму текста.

    В `tokens.font_sizes` попадал каждый когда-либо встреченный кегль — у VK Tech
    их набиралось 37 штук от 4.14 до 166 pt. По такой «шкале» проверка
    `font_size_off_scale` не отсекала ничего, а вёрстка выбирала из мусора.

    Близкие значения сливаются в одну ступень: представителем становится тот
    кегль, которым набрано больше текста, — он реально существует в шаблоне.
    Остаются не больше девяти ступеней, но крайние сохраняются всегда: заголовок
    набран немногими знаками и по весу проиграл бы основному тексту, а без
    верхней ступени вёрстка мельчит заголовки и слайд читается как полупустой.
    """
    ceiling = height_pt * MAX_SCALE_HEIGHT_SHARE
    sizes = sorted(s for s in counted if MIN_SCALE_SIZE <= s <= ceiling)
    if not sizes:
        return sorted(counted)
    clusters = [[sizes[0]]]
    for size in sizes[1:]:
        if size - clusters[-1][-1] <= clusters[-1][-1] * SCALE_TOLERANCE:
            clusters[-1].append(size)
        else:
            clusters.append([size])
    steps = [
        (max(cluster, key=lambda s: (counted[s], s)), sum(counted[s] for s in cluster))
        for cluster in clusters
    ]
    while len(steps) > MAX_SCALE_STEPS:
        # Лишняя ступень — та, что ближе всего к соседям: между 17 и 18 pt
        # выбора нет, а между 24 и 48 он есть. Крайние не трогаем: без верхней
        # вёрстка мельчит заголовки, без нижней — теряет подписи. При равном
        # зазоре уходит та, которой набрано меньше текста.
        def redundancy(position):
            size, weight = steps[position]
            gap = min(size / steps[position - 1][0], steps[position + 1][0] / size)
            return (round(gap, 4), weight)

        steps.pop(min(range(1, len(steps) - 1), key=redundancy))
    return [size for size, _ in steps]


def accent_colors(counted, theme, limit=6):
    """Фирменные акценты шаблона — по тому, чем он реально покрашен.

    Тема нередко остаётся офисной по умолчанию (синий 4472C4), а бренд живёт в
    заливках фигур. Поэтому акценты берутся из фактических цветов шаблона по
    частоте: насыщенные, не серые, не почти чёрные и не почти белые. Тема —
    запасной вариант, если красить нечем.
    """
    found = []
    # Цвет, встреченный однажды, обычно пришёл из темы и на слайдах не
    # использован: такой оставляем напоследок.
    ordered = [v for v, n in counted if n > 1] + [v for v, n in counted if n <= 1]
    for value in ordered:
        try:
            r, g, b = (int(value[i : i + 2], 16) for i in (0, 2, 4))
        except (ValueError, IndexError):
            continue
        if max(r, g, b) - min(r, g, b) < 40:
            continue
        luma = (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255
        if not 0.08 < luma < 0.93:
            continue
        if value not in found:
            found.append(value)
    for key in ("accent1", "accent2", "accent3", "accent4"):
        value = theme.get(key)
        if value and value not in found:
            found.append(value)
    return found[:limit]


def _group_transform(group, outer):
    """Как координаты детей группы переводятся в координаты слайда."""
    xfrm = group._element.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    if xfrm is None:
        return outer
    values = {}
    for tag in ("off", "ext", "chOff", "chExt"):
        node = xfrm.find(qn("a:" + tag))
        if node is None:
            return outer
        values[tag] = (
            int(node.get("x", node.get("cx", 0))),
            int(node.get("y", node.get("cy", 0))),
        )
    sx = values["ext"][0] / values["chExt"][0] if values["chExt"][0] else 1.0
    sy = values["ext"][1] / values["chExt"][1] if values["chExt"][1] else 1.0
    ox = values["off"][0] - values["chOff"][0] * sx
    oy = values["off"][1] - values["chOff"][1] * sy
    # Внешнее преобразование применяется поверх внутреннего.
    return (outer[0] + ox * outer[2], outer[1] + oy * outer[3], sx * outer[2], sy * outer[3])


def icon_slots(slide, width, height, branding=(), limit=12):
    """Значки образца страницы: место, где дизайнер поставил тематическую иконку.

    На карточных страницах над каждым тезисом стоит пиктограмма. Сама она —
    пример («лампочка»), а место и цвет — решение дизайнера. Вёрстка ставит на
    это место иконку по смыслу тезиса, а образец снимает. Берутся только
    одноцветные картинки: их можно честно перекрасить в исходный цвет.
    """
    from .assets import inspect_image

    found = []

    def walk(shapes, transform):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, _group_transform(shape, transform))
                continue
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            ox, oy, sx, sy = transform
            box = {
                "x": round((ox + shape.left * sx) / width, 5),
                "y": round((oy + shape.top * sy) / height, 5),
                "w": round(shape.width * sx / width, 5),
                "h": round(shape.height * sy / height, 5),
            }
            area = box["w"] * box["h"]
            aspect = (box["w"] * width) / max(1, box["h"] * height)
            left, top, right, bottom = CONTENT_REGION
            cx, cy = box["x"] + box["w"] / 2, box["y"] + box["h"] / 2
            key = tuple(round(box[k], 3) for k in ("x", "y", "w", "h"))
            if not (0.0003 < area < 0.02 and 0.6 < aspect < 1.67):
                continue
            if not (left < cx < right and top < cy < bottom) or key in branding:
                continue
            try:
                info = inspect_image(shape.image.blob)
            except Exception:  # noqa: BLE001 - нечитаемая картинка — не значок
                continue
            # Глиф, а не плашка: плашка — заливка почти на всю рамку, она
            # остаётся подложкой, а меняется значок поверх неё.
            if info.get("coverage", 1) > 0.6 or info.get("dominance", 0) < 0.7:
                continue
            if not info.get("color"):
                continue
            found.append({"id": shape.shape_id, "box": box, "color": info["color"]})

    walk(slide.shapes, (0.0, 0.0, 1.0, 1.0))
    # Страница-каталог с сотней значков — не карточки, а витрина набора.
    return found if len(found) <= limit else []


def picture_ratio(shape):
    """Пропорции картинки: своя и та, в которую её поставил шаблон.

    Приложение 1 требует ловить растянутую картинку. Сравнивать можно только с
    исходным изображением, причём с учётом обрезки: обрезанная наполовину
    фотография в квадратной рамке не растянута, а кадрирована.
    """
    try:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            return None
        width, height = shape.image.size
    except (AttributeError, ValueError, KeyError, OSError):
        return None
    if not width or not height or not shape.width or not shape.height:
        return None
    visible_w = max(0.01, 1 - (shape.crop_left or 0) - (shape.crop_right or 0))
    visible_h = max(0.01, 1 - (shape.crop_top or 0) - (shape.crop_bottom or 0))
    return {
        "native": round((width * visible_w) / (height * visible_h), 5),
        "frame": round(shape.width / shape.height, 5),
    }


def read_theme(root):
    colors, fonts = {}, {}
    for entry in root.findall(".//a:clrScheme/*", NS):
        if len(entry):
            value = entry[0].get("val") if entry[0].tag.endswith("srgbClr") else entry[0].get("lastClr")
            if value:
                key = etree.QName(entry).localname
                colors[key] = value.upper()
                colors[key.replace("accent", "accent_")] = value.upper()
    for role in ("major", "minor"):
        node = root.find(f".//a:fontScheme/a:{role}Font", NS)
        if node is not None:
            fonts[role] = {etree.QName(child).localname: child.get("typeface") for child in node if etree.QName(child).localname in ("latin", "ea", "cs")}
    return colors, fonts


def parse_template(data: bytes, name: str):
    check_zip(data)
    deck = Presentation(io.BytesIO(data))
    if not deck.slides or len(deck.slides) > 200:
        raise ValueError("Template must contain 1 to 200 slides")
    theme, font_scheme = {}, {}
    role_fonts = {"title": Counter(), "body": Counter()}
    fonts, sizes, colors = Counter(), Counter(), Counter()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        embedded = embedded_coverage(archive)
        for filename in archive.namelist():
            if filename.startswith("ppt/theme/theme") and filename.endswith(".xml"):
                root = etree.fromstring(
                    archive.read(filename),
                    etree.XMLParser(resolve_entities=False, no_network=True),
                )
                if not theme:
                    theme, font_scheme = read_theme(root)
    width, height = int(deck.slide_width), int(deck.slide_height)

    def scan(shapes, collect=True):
        results = []
        for shape in shapes:
            box = {
                "x": round(shape.left / width, 5),
                "y": round(shape.top / height, 5),
                "w": round(shape.width / width, 5),
                "h": round(shape.height / height, 5),
            }
            item = {
                "id": shape.shape_id,
                "name": shape.name,
                "type": str(shape.shape_type),
                "box": box,
                "text": "",
            }
            if shape.has_text_frame:
                item["text"] = shape.text[:2000]
                item["style"] = text_style(shape, theme, font_scheme)
                for p in shape.text_frame.paragraphs:
                    if not collect:
                        break
                    for font in [p.font, *(r.font for r in p.runs)]:
                        value = color_value(font.color, theme)
                        if value:
                            colors[value] += 1
                    # Кегль взвешивается объёмом текста, а не числом run'ов:
                    # иначе сноска в три знака весит столько же, сколько
                    # заголовок страницы, и шкала собирается из случайностей.
                    for run in p.runs:
                        size = run.font.size or p.font.size
                        if size:
                            sizes[round(size.pt, 2)] += max(1, len(run.text))
                    if not p.runs and p.font.size and p.text.strip():
                        sizes[round(p.font.size.pt, 2)] += len(p.text)
            if collect and hasattr(shape, "fill"):
                try:
                    if shape.fill.type == 1:
                        value = color_value(shape.fill.fore_color, theme)
                        if value:
                            colors[value] += 1
                except (AttributeError, ValueError):
                    pass
            native = picture_ratio(shape)
            if native:
                item["picture"] = native
            if shape.shape_type == 6:
                item["children"] = scan(shape.shapes, collect=collect)
            if shape.is_placeholder:
                item["placeholder"] = str(shape.placeholder_format.type)
                item["placeholder_idx"] = shape.placeholder_format.idx
                item["footer"] = is_footer_placeholder(shape)
            results.append(item)
        return results

    patterns = []
    layouts = list(deck.slide_layouts)
    branding = branding_boxes(list(deck.slides), width, height)
    master_defaults = {}
    def inherit_styles(slide, shapes):
        """Стиль плейсхолдера живёт в макете: на слайде он часто пустой.

        Без этого заголовок нового слайда получил бы кегль по умолчанию вместо
        фирменного, хотя шаблон его задаёт — просто уровнем выше.
        """
        layout_styles = {}
        for shape in slide.slide_layout.placeholders:
            if shape.has_text_frame:
                layout_styles[shape.placeholder_format.idx] = text_style(shape, theme, font_scheme)
        for shape in slide.slide_layout.slide_master.placeholders:
            if shape.has_text_frame:
                master_style = text_style(shape, theme, font_scheme)
                role = "title" if "TITLE" in str(shape.placeholder_format.type) else "body"
                master_defaults.setdefault(role, {})
                for key, value in master_style.items():
                    if master_defaults[role].get(key) is None:
                        master_defaults[role][key] = value
        for item, shape in zip(shapes, slide.shapes):
            if not shape.is_placeholder or "style" not in item:
                continue
            role = (
                "title"
                if "TITLE" in str(shape.placeholder_format.type)
                else "body"
            )
            for parent in (
                layout_styles.get(shape.placeholder_format.idx),
                master_defaults.get(role),
            ):
                if not parent:
                    continue
                for key, value in parent.items():
                    if item["style"].get(key) is None:
                        item["style"][key] = value
        return shapes

    for index, slide in enumerate(deck.slides):
        for relationship in slide.slide_layout.slide_master.part.rels.values():
            if relationship.reltype.endswith("/theme"):
                theme, font_scheme = read_theme(etree.fromstring(relationship.target_part.blob, etree.XMLParser(resolve_entities=False, no_network=True)))
                break
        master_defaults = master_text_defaults(slide.slide_layout.slide_master, theme, font_scheme)
        shapes = inherit_styles(slide, scan(slide.shapes))
        # Some templates deliberately omit title/body instances on a slide and
        # leave their empty placeholders only in the layout. They still define
        # a valid content slot; footers and slide numbers never do.
        present = {s.get('placeholder_idx') for s in shapes if 'placeholder' in s}
        inherited = []
        empty_prototype = not any(s['text'].strip() and not s.get('footer') for s in shapes)
        for item in scan(slide.slide_layout.shapes, collect=False):
            if not empty_prototype:
                break
            kind = item.get('placeholder', '')
            if item.get('placeholder_idx') in present or item.get('footer'):
                continue
            if any(word in kind for word in ('TITLE', 'BODY', 'SUBTITLE', 'OBJECT')):
                item['id'] = f"layout:{item['id']}"
                item['text'] = ''  # Layout instructions are not user content.
                item['inherited'] = True
                role = 'title' if 'TITLE' in kind and 'SUBTITLE' not in kind else 'body'
                for key, value in (master_defaults.get(role) or {}).items():
                    if item.get('style', {}).get(key) is None:
                        item.setdefault('style', {})[key] = value
                inherited.append(item)
        shapes.extend(inherited)
        text_shapes = [s for s in shapes if not s.get('footer') and (
            s['text'].strip() or any(word in s.get('placeholder', '')
                                    for word in ('TITLE', 'BODY', 'SUBTITLE', 'OBJECT'))
        )]
        title = next(
            (s for s in text_shapes if "TITLE" in s.get("placeholder", "")
             and 'SUBTITLE' not in s.get('placeholder', '')), None
        )
        if title is None and text_shapes:
            title = min(text_shapes, key=lambda s: s["box"]["y"])
        def count_text(items, source_shapes):
            for item, shape in zip(items, source_shapes):
                if shape.has_text_frame:
                    role = "title" if title and item["id"] == title["id"] else "body"
                    inherited = dict(master_defaults.get(role) or {})
                    inherited.update({key: value for key, value in item.get("style", {}).items() if value is not None})
                    default_face = font_scheme.get("major" if role == "title" else "minor", {}).get("latin")
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            weight = len(run.text.strip())
                            if not weight:
                                continue
                            face = resolve_font(run.font.name or paragraph.font.name, font_scheme) or inherited.get("font") or default_face
                            if face:
                                fonts[face] += weight
                                role_fonts[role][face] += weight
                    item.setdefault("style", {}).setdefault("font", default_face)
                if item.get("children"):
                    count_text(item["children"], shape.shapes)

        count_text(shapes, slide.shapes)
        # Score layout decorations inside the content area, excluding full-page backgrounds.
        decoration_area = 0.0
        decoration_count = 0
        for sh in slide.slide_layout.shapes:
            if sh.is_placeholder:
                continue
            x, y, w, h = (
                sh.left / width,
                sh.top / height,
                sh.width / width,
                sh.height / height,
            )
            if w * h < 0.85:
                covered = max(0, min(x + w, 0.95) - max(x, 0.05)) * max(
                    0, min(y + h, 0.87) - max(y, 0.23)
                )
                decoration_area += covered
                if covered > 0:
                    decoration_count += 1
        # Фотофон: на таком слайде обычный тёмный текст просто не читается.
        background = background_info(slide, theme)
        image_cover = 1.0 if background["kind"] != "solid" else 0.0
        for sh in list(slide.shapes) + list(slide.slide_layout.shapes):
            if sh.shape_type == 13:  # PICTURE
                image_cover = max(
                    image_cover, (sh.width / width) * (sh.height / height)
                )
        # Shapes that survive cloning: artwork, footers, page numbers. Content must
        # not collide with them, so the audit needs their boxes.
        reserved = []
        kept = [
            sh for sh in slide.shapes if preserved_shape(sh, width, height, branding)
        ]
        # Layout artwork is inherited by the cloned slide even though it is not copied.
        inherited = [sh for sh in slide.slide_layout.shapes if not sh.is_placeholder]
        kept += inherited
        inherited_ids = {id(sh) for sh in inherited}
        for sh in kept:
            box = {
                "x": round(sh.left / width, 5),
                "y": round(sh.top / height, 5),
                "w": round(sh.width / width, 5),
                "h": round(sh.height / height, 5),
            }
            left, top, right, bottom = CONTENT_REGION
            covered = max(
                0.0, min(box["x"] + box["w"], right) - max(box["x"], left)
            ) * max(0.0, min(box["y"] + box["h"], bottom) - max(box["y"], top))
            area = box["w"] * box["h"]
            # Рамка, в которую помещается текст, — это контейнер шаблона: карточка
            # или плашка, на ней контенту и место. Всё остальное, что переживает
            # клонирование, — декор: линии, стрелки, иконки, логотипы. Его текст
            # обходит, а аудит ловит наезды.
            # Картинка и группа — это иллюстрация, а не контейнер: текст поверх
            # неё не кладут, даже если рамка большая.
            artwork = sh.shape_type in (6, 13, 14)
            container = (
                not artwork and box["w"] >= 0.15 and box["h"] >= 0.1 and area < 0.6
            )
            if container:
                continue
            if artwork or area < 0.6 or covered < 0.5 * area:
                # Рисунок макета переезжает на слайд вместе с макетом, убрать его
                # нельзя; фигуру самой страницы — можно (см. stage_clutter).
                if id(sh) in inherited_ids:
                    box["layout"] = True
                if box not in reserved:
                    reserved.append(box)
        # Ряд одинаковых значков — заготовка страницы-каталога. Экспорт не
        # переносит его, когда раскладывать по нему нечего, поэтому вёрстка и
        # аудит тоже не должны считать его препятствием.
        icon_sizes = Counter(
            (round(b["w"], 2), round(b["h"], 2))
            for b in reserved
            if 0.0002 < b["w"] * b["h"] < 0.02
        )
        for box in reserved:
            key = (round(box["w"], 2), round(box["h"], 2))
            if icon_sizes.get(key, 0) >= 3 and 0.0002 < box["w"] * box["h"] < 0.02:
                box["icons"] = True
        slots = []
        for item in text_shapes:
            role = "title" if title and item["id"] == title["id"] else "body"
            slots.append(
                {
                    "role": role,
                    "box": item["box"],
                    "style": item.get("style") or {},
                    "length": len(item["text"]),
                }
            )
        # Порядок чтения: сверху вниз, слева направо — как человек смотрит слайд.
        slots.sort(key=lambda s: (s["role"] != "title", s["box"]["y"], s["box"]["x"]))
        branding_keys = {tuple(b) for b in branding}
        patterns.append(
            {
                "slots": slots,
                "background": background["color"],
                "background_kind": background["kind"],
                "reserved": reserved,
                "decoration_area": round(decoration_area, 5),
                "image_cover": round(min(1.0, image_cover), 4),
                "decoration_count": decoration_count,
                "index": index,
                "layout_index": layouts.index(slide.slide_layout)
                if slide.slide_layout in layouts
                else 0,
                "layout_name": slide.slide_layout.name,
                "title_box": title["box"] if title else None,
                "icon_slots": icon_slots(slide, width, height, branding_keys),
                "shapes": shapes,
                "text_slots": len(text_shapes),
                "branding_score": branding_score(slide, width, height, background),
                "title_background": title_background(slide, title['box'] if title else None,
                                                     width, height, background['color']),
            }
        )
    classify_patterns(patterns)
    geometry = derive_geometry(patterns)
    for layout in layouts:
        scan(layout.shapes)
    for master in deck.slide_masters:
        scan(master.shapes)
    if not colors:
        # Шаблон, где ничего не покрашено явно: весь цвет наследуется от темы.
        # Тогда clrScheme и есть его палитра — другого источника просто нет.
        colors.update(dict.fromkeys(theme.values(), 1))
    fallback = [font for font, _ in fonts.most_common()]
    fallback += [value.get("latin") for value in font_scheme.values() if value.get("latin") not in fallback]
    coverage = {name: font_coverage(name, embedded) for name in fallback if name}
    font_choice = {"language": "ru"}
    for role in ("title", "body"):
        font_choice[role] = choose_font(role_fonts[role], fallback, coverage)
        selected = font_choice[role]["selected"]
        coverage.setdefault(selected, font_coverage(selected, embedded))
    for pattern in patterns:
        for slot in pattern["slots"]:
            slot["style"]["font"] = font_choice[slot["role"]]["selected"]
    ordered_fonts = list(dict.fromkeys([font_choice["body"]["selected"], font_choice["title"]["selected"], *fallback]))
    return {
        "name": Path(name).stem,
        "width": width,
        "height": height,
        "tokens": {
            "fonts": ordered_fonts,
            "font_choice": font_choice,
            "font_coverage": coverage,
            "font_sizes": type_scale(sizes, height / 12700),
            "caption_sizes": caption_sizes(sizes),
            "colors": [c for c, _ in colors.most_common()],
            "accents": accent_colors(colors.most_common(), theme),
            "theme": theme,
        },
        "parser": PARSER_VERSION,
        "geometry": geometry,
        # Экспорт должен сохранять ровно то же, что учёл аудит.
        "branding": sorted(branding),
        "layouts": [
            {"index": i, "name": layout.name} for i, layout in enumerate(layouts)
        ],
        "patterns": patterns,
    }


def parse_content(data: bytes, name: str):
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        text = "\n".join(
            p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages
        )
    elif suffix == ".pptx":
        check_zip(data)
        text = "\n".join(
            s.text
            for slide in Presentation(io.BytesIO(data)).slides
            for s in slide.shapes
            if s.has_text_frame
        )
    elif suffix == ".docx":
        check_zip(data)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = etree.fromstring(
                z.read("word/document.xml"),
                etree.XMLParser(resolve_entities=False, no_network=True),
            )
            text = "\n".join(
                root.xpath("//*[local-name()='p']//*[local-name()='t']/text()")
            )
    elif suffix in (".txt", ".md", ".csv", ".json"):
        text = data.decode("utf-8-sig")
        if suffix == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        if suffix == ".csv":
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Supported content: txt, md, csv, json, pdf, docx, pptx")
    if not text.strip():
        raise ValueError("No extractable text; scanned PDFs require OCR before upload")
    if len(text) > 100000:
        raise ValueError("Content exceeds 100000 characters; split the document")
    return text
