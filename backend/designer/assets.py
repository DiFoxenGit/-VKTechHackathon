"""Визуальные ассеты: иконки, иллюстрации и фото для слайдов.

Три источника, в порядке доверия:

- **контент-пакет** пользователя — zip с SVG, PNG и JPEG (и тексты рядом);
- **шаблон** — картинки со страниц-образцов: каталог иконок, 3D-иллюстрации,
  фотографии. Их стиль — стиль бренда, поэтому они лучше любых наших;
- **встроенный набор** — пиктограммы и иллюстрации в SVG, нарисованные ролями
  цвета (акцент, фон, текст). Он есть всегда, поэтому слайд не остаётся без
  картинки на незнакомом шаблоне.

Картинку ставим только туда, где она относится к тексту: подбор идёт по корням
слов подписи. Иконки из шаблона без подписи подбирать не по чему — их подписывает
мультимодальная модель (`tag_assets`), а до тех пор они не участвуют.

SVG не растрируется: пути становятся нативными фигурами PowerPoint, которые
можно перекрасить и отредактировать. Растровые иконки перекрашиваются в акцент.
"""

import functools
import hashlib
import io
import json
import logging
import math
import re
import zipfile
from pathlib import Path, PurePosixPath

from lxml import etree

LOGGER = logging.getLogger("designer.assets")

PACK_DIR = Path(__file__).parent / "assets" / "pack"
KINDS = ("icon", "illustration", "photo")
RASTER = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg"}
# Имя файла ассета в хранилище: хеш содержимого и тип. Другие имена не
# принимаются — ссылка из колоды не может стать путём по диску.
STORED_NAME = re.compile(r"[a-f0-9]{24}\.(png|jpeg|svg)")
# Сколько картинок шаблона имеет смысл держать: каталоги иконок бывают на
# сотни штук, а фотографии весят мегабайты.
HARVEST_LIMITS = {"icon": 320, "illustration": 60, "photo": 40}
MAX_PACK_FILES = 400
MAX_SVG_BYTES = 400_000

# Самое слабое совпадение, при котором картинка считается относящейся к тексту.
MIN_MATCH = 0.55
# Источник, которому верим больше при равном совпадении: пакет пользователя
# собран под эту презентацию, шаблон — под бренд, встроенный набор — запасной.
PRIORITY = {"pack": 3, "template": 2, "builtin": 1}


# ---------------------------------------------------------------- слова и теги


def normalize(text):
    return text.lower().replace("ё", "е")


WORD = re.compile(r"[a-zа-я0-9₽%]+(?:-[a-zа-я0-9]+)?")


def words(text):
    return WORD.findall(normalize(text))


def stem(word):
    """Грубый корень русского слова: окончание отрезается, основа остаётся.

    Нужен, чтобы подпись модели «календарь» совпадала с «календаря» на
    слайде. Короткие слова не режутся: у них окончание — половина слова.
    """
    word = normalize(word).strip()
    if len(word) <= 4:
        return word
    if len(word) == 5:
        return word[:4]
    return word[: max(5, len(word) - 2)]


def tag_roots(tags):
    """Теги ассета как корни: многословный тег даёт корень каждого слова."""
    roots = []
    for tag in tags or []:
        for word in words(str(tag)):
            root = stem(word) if len(word) > 3 else word
            if root and root not in roots:
                roots.append(root)
    return roots


def root_matches(root, word):
    if not word.startswith(root):
        return False
    # Короткий корень («цел», «год») совпадает только с короткими формами, иначе
    # «цел» находит «целиком», а «мин» — «минимум».
    return len(root) > 3 or len(word) <= len(root) + 2


def match_score(tags, text, weight=1.0):
    """Насколько ассет относится к тексту: чем раньше слово, тем больше вес.

    Смысл тезиса обычно держит первое существительное: в «Время сборки
    колоды сократилось» картинка про время, а не про колоду.
    """
    roots = [normalize(t) for t in tags or []]
    if not roots:
        return 0.0
    score = 0.0
    for position, word in enumerate(words(text)):
        matched = [root for root in roots if root_matches(root, word)]
        if matched:
            # Длинный корень точнее короткого: «дизайнер» — про людей, а не
            # про «дизайн» вообще.
            specific = 1 + 0.04 * min(10, max(len(root) for root in matched))
            score += weight * specific / (1 + 0.15 * position)
    return score


# ---------------------------------------------------------------- SVG


def _matrix_multiply(m, n):
    a, b, c, d, e, f = m
    a2, b2, c2, d2, e2, f2 = n
    return (
        a * a2 + c * b2,
        b * a2 + d * b2,
        a * c2 + c * d2,
        b * c2 + d * d2,
        a * e2 + c * f2 + e,
        b * e2 + d * f2 + f,
    )


IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)
NUMBER = re.compile(r"[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?")


def parse_transform(value):
    matrix = IDENTITY
    for name, args in re.findall(r"(\w+)\s*\(([^)]*)\)", value or ""):
        nums = [float(x) for x in NUMBER.findall(args)]
        step = IDENTITY
        if name == "matrix" and len(nums) == 6:
            step = tuple(nums)
        elif name == "translate" and nums:
            step = (1, 0, 0, 1, nums[0], nums[1] if len(nums) > 1 else 0)
        elif name == "scale" and nums:
            step = (nums[0], 0, 0, nums[1] if len(nums) > 1 else nums[0], 0, 0)
        elif name == "rotate" and nums:
            angle = math.radians(nums[0])
            cos, sin = math.cos(angle), math.sin(angle)
            step = (cos, sin, -sin, cos, 0, 0)
            if len(nums) == 3:
                cx, cy = nums[1], nums[2]
                step = _matrix_multiply(
                    _matrix_multiply((1, 0, 0, 1, cx, cy), step), (1, 0, 0, 1, -cx, -cy)
                )
        matrix = _matrix_multiply(matrix, step)
    return matrix


def _apply(m, x, y):
    return (m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5])


NAMED = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "blue": "0000FF", "gray": "808080", "grey": "808080"}


def parse_color(value):
    """Цвет SVG как роль или RRGGBB: none, current, accent… или hex."""
    value = (value or "").strip()
    if not value or value == "inherit":
        return None
    if value == "none" or value == "transparent":
        return "none"
    if value == "currentColor":
        return "current"
    role = re.match(r"var\(--([\w-]+)", value)
    if role:
        return role.group(1)
    if value.startswith("#"):
        hexa = value[1:]
        if len(hexa) == 3:
            hexa = "".join(c * 2 for c in hexa)
        if re.fullmatch(r"[0-9a-fA-F]{6}", hexa[:6]):
            return hexa[:6].upper()
    rgb = re.match(r"rgb\(([^)]*)\)", value)
    if rgb:
        parts = [p.strip() for p in rgb.group(1).split(",")]
        try:
            vals = [
                round(float(p[:-1]) * 2.55) if p.endswith("%") else int(float(p))
                for p in parts[:3]
            ]
            return "".join(f"{max(0, min(255, c)):02X}" for c in vals)
        except ValueError:
            return "current"
    if value.lower() in NAMED:
        return NAMED[value.lower()]
    return "current"


class _PathReader:
    """Разбор атрибута d: команды, числа и слитые флаги дуг («a1 1 0 01 5 5»)."""

    def __init__(self, text):
        self.text = text
        self.pos = 0

    def skip(self):
        while self.pos < len(self.text) and self.text[self.pos] in " \t\r\n,":
            self.pos += 1

    def command(self):
        self.skip()
        if self.pos < len(self.text) and self.text[self.pos].isalpha():
            self.pos += 1
            return self.text[self.pos - 1]
        return None

    def has_number(self):
        self.skip()
        return self.pos < len(self.text) and (
            self.text[self.pos].isdigit() or self.text[self.pos] in "+-."
        )

    def number(self):
        self.skip()
        match = NUMBER.match(self.text, self.pos)
        if not match:
            raise ValueError("bad path number")
        self.pos = match.end()
        return float(match.group())

    def flag(self):
        self.skip()
        char = self.text[self.pos : self.pos + 1]
        if char not in ("0", "1"):
            raise ValueError("bad arc flag")
        self.pos += 1
        return char == "1"


def _cubic(p0, p1, p2, p3, steps=10):
    return [
        (
            (1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * p1[0] + 3 * (1 - t) * t * t * p2[0] + t**3 * p3[0],
            (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * p1[1] + 3 * (1 - t) * t * t * p2[1] + t**3 * p3[1],
        )
        for t in (i / steps for i in range(1, steps + 1))
    ]


def _quad(p0, p1, p2, steps=8):
    return [
        (
            (1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
            (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1],
        )
        for t in (i / steps for i in range(1, steps + 1))
    ]


def _arc(p0, rx, ry, rotation, large, sweep, p1):
    """Дуга SVG (конечные точки) → точки ломаной; формулы из приложения F.6 SVG."""
    if p0 == p1:
        return []
    if not rx or not ry:
        return [p1]
    rx, ry = abs(rx), abs(ry)
    phi = math.radians(rotation)
    cos, sin = math.cos(phi), math.sin(phi)
    dx, dy = (p0[0] - p1[0]) / 2, (p0[1] - p1[1]) / 2
    x1 = cos * dx + sin * dy
    y1 = -sin * dx + cos * dy
    scale = (x1 * x1) / (rx * rx) + (y1 * y1) / (ry * ry)
    if scale > 1:
        rx, ry = rx * math.sqrt(scale), ry * math.sqrt(scale)
    num = rx * rx * ry * ry - rx * rx * y1 * y1 - ry * ry * x1 * x1
    den = rx * rx * y1 * y1 + ry * ry * x1 * x1
    factor = math.sqrt(max(0.0, num / den)) if den else 0.0
    if large == sweep:
        factor = -factor
    cx1, cy1 = factor * rx * y1 / ry, -factor * ry * x1 / rx
    cx = cos * cx1 - sin * cy1 + (p0[0] + p1[0]) / 2
    cy = sin * cx1 + cos * cy1 + (p0[1] + p1[1]) / 2

    def angle(ux, uy, vx, vy):
        dot = ux * vx + uy * vy
        length = math.hypot(ux, uy) * math.hypot(vx, vy)
        value = math.acos(max(-1.0, min(1.0, dot / length))) if length else 0.0
        return -value if ux * vy - uy * vx < 0 else value

    start = angle(1, 0, (x1 - cx1) / rx, (y1 - cy1) / ry)
    delta = angle((x1 - cx1) / rx, (y1 - cy1) / ry, (-x1 - cx1) / rx, (-y1 - cy1) / ry)
    if not sweep and delta > 0:
        delta -= 2 * math.pi
    elif sweep and delta < 0:
        delta += 2 * math.pi
    steps = max(4, math.ceil(abs(delta) / (math.pi / 16)))
    points = []
    for i in range(1, steps + 1):
        t = start + delta * i / steps
        x, y = rx * math.cos(t), ry * math.sin(t)
        points.append((cos * x - sin * y + cx, sin * x + cos * y + cy))
    points[-1] = p1
    return points


def path_contours(d):
    """Атрибут d → список контуров (точки, замкнут ли)."""
    reader = _PathReader(d or "")
    contours, points = [], []
    current = start = (0.0, 0.0)
    last_control = None
    command = None

    def flush(closed=False):
        nonlocal points
        if len(points) > 1:
            contours.append({"points": points, "closed": closed})
        points = []

    while True:
        found = reader.command()
        if found:
            command = found
        elif not reader.has_number() or command is None:
            break
        relative = command.islower()
        op = command.upper()
        base = current if relative else (0.0, 0.0)

        def point():
            x, y = reader.number(), reader.number()
            return (base[0] + x, base[1] + y) if relative else (x, y)

        if op == "Z":
            if points:
                flush(closed=True)
            current = start
            last_control = None
            if not found:
                break
            continue
        if op == "M":
            flush()
            current = start = point()
            points = [current]
            # Следующие пары после M — это L.
            command = "l" if relative else "L"
            last_control = None
            continue
        if not points:
            points = [current]
        if op == "L":
            current = point()
            points.append(current)
            last_control = None
        elif op == "H":
            x = reader.number()
            current = (current[0] + x if relative else x, current[1])
            points.append(current)
            last_control = None
        elif op == "V":
            y = reader.number()
            current = (current[0], current[1] + y if relative else y)
            points.append(current)
            last_control = None
        elif op == "C":
            c1, c2, end = point(), point(), point()
            points.extend(_cubic(current, c1, c2, end))
            last_control, current = c2, end
        elif op == "S":
            c1 = (
                (2 * current[0] - last_control[0], 2 * current[1] - last_control[1])
                if last_control
                else current
            )
            c2, end = point(), point()
            points.extend(_cubic(current, c1, c2, end))
            last_control, current = c2, end
        elif op == "Q":
            c1, end = point(), point()
            points.extend(_quad(current, c1, end))
            last_control, current = c1, end
        elif op == "T":
            c1 = (
                (2 * current[0] - last_control[0], 2 * current[1] - last_control[1])
                if last_control
                else current
            )
            end = point()
            points.extend(_quad(current, c1, end))
            last_control, current = c1, end
        elif op == "A":
            rx, ry, rotation = reader.number(), reader.number(), reader.number()
            large, sweep = reader.flag(), reader.flag()
            end = point()
            points.extend(_arc(current, rx, ry, rotation, large, sweep, end))
            current = end
            last_control = None
        else:
            break
    flush()
    return contours


def _ellipse_points(cx, cy, rx, ry, steps=48):
    return [
        (cx + rx * math.cos(2 * math.pi * i / steps), cy + ry * math.sin(2 * math.pi * i / steps))
        for i in range(steps)
    ]


def _rect_contour(x, y, w, h, rx):
    if not rx:
        return [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]
    rx = min(rx, w / 2, h / 2)
    points = []
    for cx, cy, a0 in (
        (x + w - rx, y + rx, -90),
        (x + w - rx, y + h - rx, 0),
        (x + rx, y + h - rx, 90),
        (x + rx, y + rx, 180),
    ):
        for i in range(7):
            t = math.radians(a0 + 15 * i)
            points.append((cx + rx * math.cos(t), cy + rx * math.sin(t)))
    return points


def _style(node, inherited):
    style = dict(inherited)
    for key in ("fill", "stroke", "stroke-width", "display", "visibility"):
        if node.get(key) is not None:
            style[key] = node.get(key)
    for part in (node.get("style") or "").split(";"):
        if ":" in part:
            key, value = part.split(":", 1)
            key = key.strip()
            if key in ("fill", "stroke", "stroke-width", "display", "visibility"):
                style[key] = value.strip()
    return style


def parse_svg(data):
    """SVG → геометрия в единицах viewBox: фигуры и контуры с ролями цвета.

    Поддерживается то, из чего состоят иконки и плоские иллюстрации: path,
    circle, ellipse, rect, line, polyline, polygon, группы и transform.
    Скрипты, ссылки и внешние сущности не исполняются и не разрешаются.
    """
    if isinstance(data, str):
        data = data.encode("utf-8")
    if len(data) > MAX_SVG_BYTES:
        raise ValueError("SVG is too large")
    root = etree.fromstring(
        data, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    )
    if etree.QName(root).localname != "svg":
        raise ValueError("Not an SVG document")
    box = [float(v) for v in NUMBER.findall(root.get("viewBox") or "")]
    if len(box) != 4 or box[2] <= 0 or box[3] <= 0:
        width = float((NUMBER.findall(root.get("width") or "") or [0])[0])
        height = float((NUMBER.findall(root.get("height") or "") or [0])[0])
        if width <= 0 or height <= 0:
            raise ValueError("SVG has no viewBox or size")
        box = [0.0, 0.0, width, height]
    items = []

    def visit(node, matrix, style):
        if not isinstance(node.tag, str):
            return
        name = etree.QName(node).localname
        if name in ("defs", "clipPath", "mask", "symbol", "style", "script", "title", "desc", "metadata"):
            return
        style = _style(node, style)
        if style.get("display") == "none" or style.get("visibility") == "hidden":
            return
        matrix = _matrix_multiply(matrix, parse_transform(node.get("transform")))
        if name in ("svg", "g", "a"):
            for child in node:
                visit(child, matrix, style)
            return
        fill = parse_color(style.get("fill"))
        stroke = parse_color(style.get("stroke"))
        fill = "000000" if fill is None else fill
        stroke = "none" if stroke is None else stroke
        try:
            width = float((NUMBER.findall(str(style.get("stroke-width", "1"))) or [1])[0])
        except ValueError:
            width = 1.0
        # Масштаб трансформации для толщины линии и радиусов.
        scale = math.sqrt(abs(matrix[0] * matrix[3] - matrix[1] * matrix[2])) or 1.0
        paint = {"fill": fill, "stroke": stroke, "width": width * scale}
        axis_aligned = abs(matrix[1]) < 1e-9 and abs(matrix[2]) < 1e-9

        def num(key, default=0.0):
            found = NUMBER.findall(node.get(key) or "")
            return float(found[0]) if found else default

        contours = []
        if name in ("circle", "ellipse"):
            cx, cy = num("cx"), num("cy")
            rx = num("r") if name == "circle" else num("rx")
            ry = num("r") if name == "circle" else num("ry")
            if rx <= 0 or ry <= 0:
                return
            if axis_aligned:
                x0, y0 = _apply(matrix, cx - rx, cy - ry)
                x1, y1 = _apply(matrix, cx + rx, cy + ry)
                items.append(
                    {"shape": "ellipse", "x": min(x0, x1), "y": min(y0, y1),
                     "w": abs(x1 - x0), "h": abs(y1 - y0), **paint}
                )
                return
            contours = [{"points": _ellipse_points(cx, cy, rx, ry), "closed": True}]
        elif name == "rect":
            x, y, w, h = num("x"), num("y"), num("width"), num("height")
            if w <= 0 or h <= 0:
                return
            rx = num("rx", num("ry"))
            if axis_aligned:
                x0, y0 = _apply(matrix, x, y)
                x1, y1 = _apply(matrix, x + w, y + h)
                items.append(
                    {"shape": "rect", "x": min(x0, x1), "y": min(y0, y1), "w": abs(x1 - x0),
                     "h": abs(y1 - y0), "rx": rx * abs(matrix[0]), **paint}
                )
                return
            contours = [{"points": _rect_contour(x, y, w, h, rx), "closed": True}]
        elif name == "line":
            contours = [{"points": [(num("x1"), num("y1")), (num("x2"), num("y2"))], "closed": False}]
        elif name in ("polyline", "polygon"):
            values = [float(v) for v in NUMBER.findall(node.get("points") or "")]
            points = list(zip(values[0::2], values[1::2]))
            if len(points) < 2:
                return
            contours = [{"points": points, "closed": name == "polygon"}]
        elif name == "path":
            try:
                contours = path_contours(node.get("d"))
            except (ValueError, IndexError):
                return
        else:
            return
        if not contours:
            return
        items.append(
            {
                "shape": "path",
                "contours": [
                    {
                        "points": [
                            [round(v, 3) for v in _apply(matrix, x, y)]
                            for x, y in contour["points"]
                        ],
                        "closed": contour["closed"],
                    }
                    for contour in contours
                ],
                **paint,
            }
        )

    visit(root, IDENTITY, {})
    if not items:
        raise ValueError("SVG has no drawable shapes")
    colors = {
        value
        for item in items
        for value in (item["fill"], item["stroke"])
        if value not in ("none",)
    }
    return {"viewbox": box, "items": items, "colors": sorted(colors)}


@functools.lru_cache(maxsize=512)
def _svg_cached(path: str, mtime: float):
    return parse_svg(Path(path).read_bytes())


def svg_geometry(path: Path):
    return _svg_cached(str(path), path.stat().st_mtime)


# ---------------------------------------------------------------- растр


def inspect_image(blob):
    """Что за картинка: размер, прозрачность, одноцветность, пестрота.

    По этим признакам шаблонная картинка раскладывается на иконку, иллюстрацию,
    фотографию или декор, который брать не нужно (градиентные полосы, фон).
    """
    from .render import image_from_bytes

    image = image_from_bytes(blob)
    width, height = image.size
    side = 40
    scale = side / max(width, height)
    small = image.resize(
        (max(1, round(width * scale)), max(1, round(height * scale))), resample=3
    )
    alpha = small.mode == "RGBA"
    opaque, solid = [], []
    total = 0
    for pixel_value in small.getdata():
        total += 1
        if alpha and pixel_value[3] < 48:
            continue
        pixel = tuple(pixel_value[:3])
        opaque.append(pixel)
        # Полупрозрачные края сглаживания темнее самого цвета: для оценки
        # одноцветности берём только плотные пиксели.
        if not alpha or pixel_value[3] >= 200:
            solid.append(pixel)
    coverage = len(opaque) / total if total else 0.0
    info = {"w": width, "h": height, "alpha": alpha, "coverage": round(coverage, 3)}
    if not opaque:
        return {**info, "mono": False, "dominance": 0.0, "color": None, "spread": 0.0, "colors": 0}
    core = solid if len(solid) >= 8 else opaque
    ordered = sorted(core, key=lambda c: sum(c))
    median = ordered[len(ordered) // 2]
    near = sum(
        1 for c in core if max(abs(c[0] - median[0]), abs(c[1] - median[1]), abs(c[2] - median[2])) < 56
    )
    lumas = [(0.2126 * r + 0.7152 * g + 0.0722 * b) / 255 for r, g, b in opaque]
    mean = sum(lumas) / len(lumas)
    spread = math.sqrt(max(0.0, sum((v - mean) ** 2 for v in lumas) / len(lumas)))
    buckets = {(r >> 5, g >> 5, b >> 5) for r, g, b in opaque}
    return {
        **info,
        "mono": near / len(core) >= 0.88,
        # Доля пикселей основного цвета: у сглаженного белого глифа края
        # серые, и строгая одноцветность его не узнаёт.
        "dominance": round(near / len(core), 3),
        "color": "".join(f"{c:02X}" for c in median),
        "spread": round(spread, 3),
        "colors": len(buckets),
    }


def classify_image(info, frame_area=None):
    """Иконка, иллюстрация, фото — или None, если картинка не годится в слайд."""
    width, height = info["w"], info["h"]
    if min(width, height) < 16:
        return None
    ratio = width / height
    coverage = info["coverage"]
    if frame_area is not None and frame_area > 0.6:
        # Картинка во весь слайд — фон страницы, а не иллюстрация к тезису.
        return None
    if (
        info["alpha"]
        and info["mono"]
        and 0.03 < coverage < 0.85
        and 0.6 <= ratio <= 1.67
        and (frame_area is None or frame_area < 0.03)
    ):
        return "icon"
    if (
        info["alpha"]
        and not info["mono"]
        and 0.12 < coverage < 0.95
        and info["colors"] >= 6
        and min(width, height) >= 96
        and 0.4 <= ratio <= 2.5
    ):
        # Квадрат с закруглёнными углами и почти без прозрачности — значок
        # приложения или логотип: как иллюстрация к тезису он не читается.
        if 0.85 <= ratio <= 1.18 and coverage > 0.8:
            return None
        return "illustration"
    if (
        (not info["alpha"] or coverage > 0.97)
        and min(width, height) >= 200
        and info["spread"] > 0.12
        and info["colors"] >= 24
        and 0.4 <= ratio <= 2.6
    ):
        return "photo"
    return None


def tint_png(blob, color):
    """Одноцветная иконка в заданный цвет: прозрачность сохраняется."""
    from PIL import Image

    from .render import image_from_bytes, png

    image = image_from_bytes(blob)
    if image.mode != "RGBA":
        return blob
    r, g, b = (int(color[i : i + 2], 16) for i in (0, 2, 4))
    tinted = Image.new("RGBA", image.size, (r, g, b, 255))
    tinted.putalpha(image.getchannel("A"))
    return png(tinted)


def image_format(blob):
    if blob[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if blob[:3] == b"\xff\xd8\xff":
        return "jpeg"
    return None


def asset_id(blob):
    return hashlib.sha256(blob).hexdigest()[:24]


# ---------------------------------------------------------------- сбор из PPTX


def _pictures(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from _pictures(shape.shapes)
            continue
        try:
            image = shape.image
        except (AttributeError, ValueError, KeyError):
            continue
        yield shape, image


def harvest_pptx(data, branding=None):
    """Картинки со страниц PPTX, разложенные по видам. Возвращает [(meta, blob)].

    Повторяющийся брендинг (логотип на каждой странице) не берётся: это не
    иллюстрация к тезису. Одинаковые картинки собираются один раз.
    """
    from pptx import Presentation

    from .parsing import shape_key

    deck = Presentation(io.BytesIO(data))
    width, height = deck.slide_width, deck.slide_height
    found, seen = [], set()
    counts = dict.fromkeys(KINDS, 0)
    for index, slide in enumerate(deck.slides):
        for shape, image in _pictures(slide.shapes):
            if branding and shape_key(shape, width, height) in branding:
                continue
            blob = image.blob
            fmt = image_format(blob)
            if not fmt:
                continue
            identifier = asset_id(blob)
            if identifier in seen:
                continue
            seen.add(identifier)
            try:
                info = inspect_image(blob)
            except Exception:  # noqa: BLE001 - битая картинка шаблона не должна ронять импорт
                continue
            area = (shape.width or 0) * (shape.height or 0) / (width * height or 1)
            kind = classify_image(info, area)
            if not kind or counts[kind] >= HARVEST_LIMITS[kind]:
                continue
            counts[kind] += 1
            found.append(
                (
                    {
                        "id": identifier,
                        "kind": kind,
                        "format": fmt,
                        "file": f"{identifier}.{fmt}",
                        "ratio": round(info["w"] / info["h"], 4),
                        "mono": bool(info["mono"]) and kind == "icon",
                        "color": info["color"],
                        "slide": index,
                        "tags": [],
                        "label": [],
                    },
                    blob,
                )
            )
    return found


def save_assets(folder: Path, harvested):
    """Записать картинки на диск; вернуть метаданные для записи в хранилище."""
    folder.mkdir(parents=True, exist_ok=True)
    metas = []
    for meta, blob in harvested:
        target = folder / meta["file"]
        if not target.exists():
            target.write_bytes(blob)
        metas.append(meta)
    return metas


# ---------------------------------------------------------------- контент-пакет zip


TEXT_SUFFIXES = (".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".pptx")
IMAGE_SUFFIXES = (".svg", ".png", ".jpg", ".jpeg")
FOLDER_KINDS = (
    ("icon", "icon"),
    ("иконк", "icon"),
    ("пиктограм", "icon"),
    ("illustr", "illustration"),
    ("иллюстр", "illustration"),
    ("photo", "photo"),
    ("фото", "photo"),
    ("screen", "photo"),
    ("скрин", "photo"),
    ("image", "photo"),
)


def _kind_from_path(path):
    lowered = normalize(path)
    for hint, kind in FOLDER_KINDS:
        if hint in lowered:
            return kind
    return None


def _svg_kind(geometry):
    """Иконка — одноцветный рисунок; иллюстрация — несколько цветов."""
    colors = [c for c in geometry["colors"] if c != "none"]
    return "icon" if len(set(colors)) <= 1 else "illustration"


def parse_pack_zip(data, parse_text):
    """Zip контент-пакета: тексты для плана и картинки для вёрстки.

    manifest.json (необязателен) задаёт вид и теги: {"assets": [{"file":
    "icons/rocket.svg", "kind": "icon", "tags": ["запуск", "ракета"]}]}. Без
    него вид берётся из папки (icons/, illustrations/, photos/) или по самой
    картинке, а теги — из имени файла: «запуск-ракета.svg».
    Возвращает (текст, [(meta, blob)]).
    """
    from .parsing import check_zip

    check_zip(data)
    texts, assets = [], []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = [
            n
            for n in archive.namelist()
            if not n.endswith("/")
            and not n.startswith("__MACOSX/")
            and not PurePosixPath(n).name.startswith(".")
        ]
        if len(names) > MAX_PACK_FILES:
            raise ValueError(f"Content pack holds more than {MAX_PACK_FILES} files")
        manifest = {}
        if "manifest.json" in names:
            try:
                raw = json.loads(archive.read("manifest.json").decode("utf-8-sig"))
                manifest = {
                    str(item.get("file")): item
                    for item in raw.get("assets", [])
                    if isinstance(item, dict) and item.get("file")
                }
            except (ValueError, AttributeError) as exc:
                raise ValueError("manifest.json is not valid JSON") from exc
        seen = set()
        for name in sorted(names):
            if name == "manifest.json":
                continue
            suffix = PurePosixPath(name).suffix.lower()
            payload = archive.read(name)
            if suffix in TEXT_SUFFIXES:
                try:
                    texts.append(parse_text(payload, name))
                except ValueError:
                    LOGGER.info("Skip unreadable text %s in content pack", name)
                if suffix == ".pptx":
                    assets.extend(harvest_pptx(payload))
                continue
            if suffix not in IMAGE_SUFFIXES:
                continue
            entry = manifest.get(name, {})
            stem_words = re.split(r"[-_\s.]+", PurePosixPath(name).stem)
            label = [str(t) for t in entry.get("tags") or stem_words if str(t).strip()]
            identifier = asset_id(payload)
            if identifier in seen:
                continue
            seen.add(identifier)
            kind = entry.get("kind") if entry.get("kind") in KINDS else _kind_from_path(name)
            if suffix == ".svg":
                try:
                    geometry = parse_svg(payload)
                except (ValueError, etree.XMLSyntaxError):
                    LOGGER.info("Skip unreadable SVG %s in content pack", name)
                    continue
                box = geometry["viewbox"]
                kind = kind or _svg_kind(geometry)
                meta = {
                    "format": "svg",
                    "ratio": round(box[2] / box[3], 4),
                    "mono": _svg_kind(geometry) == "icon",
                    "color": None,
                }
            else:
                fmt = image_format(payload)
                if not fmt:
                    continue
                try:
                    info = inspect_image(payload)
                except Exception:  # noqa: BLE001 - испорченная картинка пропускается
                    continue
                kind = kind or classify_image(info) or "photo"
                meta = {
                    "format": fmt,
                    "ratio": round(info["w"] / info["h"], 4),
                    "mono": bool(info["mono"]) and kind == "icon",
                    "color": info["color"],
                }
            meta.update(
                id=identifier,
                kind=kind,
                file=f"{identifier}.{meta['format']}",
                name=name,
                label=label,
                tags=tag_roots(label),
            )
            assets.append((meta, payload))
    return "\n\n".join(t for t in texts if t.strip()), assets


def list_roots(tags):
    """Теги из манифеста уже корни («команд»): берём как есть, в нижнем регистре."""
    return [normalize(str(t)).strip() for t in tags if str(t).strip()]


# ---------------------------------------------------------------- библиотека


@functools.lru_cache(maxsize=1)
def builtin_manifest():
    path = PACK_DIR / "manifest.json"
    if not path.exists():
        return {"assets": []}
    return json.loads(path.read_text(encoding="utf-8"))


def builtin_assets():
    manifest = builtin_manifest()
    neutral = manifest.get("neutral_icon")
    result = []
    for item in manifest.get("assets", []):
        geometry = svg_geometry(PACK_DIR / item["file"])
        box = geometry["viewbox"]
        result.append(
            {
                "id": "builtin-" + PurePosixPath(item["file"]).stem,
                "kind": item["kind"],
                "format": "svg",
                "source": "builtin",
                "family": "builtin",
                "ref": "builtin:" + item["file"],
                "ratio": round(box[2] / box[3], 4),
                "mono": item["kind"] == "icon",
                "tags": list_roots(item.get("tags")),
                "neutral": item["file"] == neutral,
            }
        )
    return result


def library(template=None, packs=()):
    """Всё, что можно поставить на слайд этой колоды: пакеты, шаблон, встроенное."""
    items = []
    for pack in packs:
        for meta in pack.get("assets") or []:
            items.append(
                {
                    **meta,
                    "source": "pack",
                    "family": "pack:" + pack["id"],
                    "ref": f"pack:{pack['id']}:{meta['file']}",
                }
            )
    if template:
        for meta in template.get("assets") or []:
            items.append(
                {
                    **meta,
                    "source": "template",
                    "family": "template:" + template.get("id", ""),
                    "ref": f"template:{template.get('id', '')}:{meta['file']}",
                }
            )
    items.extend(builtin_assets())
    return items


def _best(candidates, text, used, weight=1.0):
    scored = []
    for asset in candidates:
        score = match_score(asset.get("tags"), text, weight)
        if score < MIN_MATCH:
            continue
        # Повтор той же картинки на соседнем слайде — не ошибка, но хуже.
        penalty = 0.5 if asset["id"] in used else 0.0
        scored.append((score - penalty, PRIORITY.get(asset["source"], 0), asset["id"], asset))
    if not scored:
        return None
    scored.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return scored[0][3]


def pick_icons(labels, assets, used=()):
    """Иконка на каждую подпись — из одного набора, чтобы ряд был в одном стиле.

    Возвращает (иконки, сколько подобрано по смыслу). Набор выбирается тот, что
    покрывает больше подписей; неподобранные получают нейтральный значок того
    же набора. Набор без нейтрального значка берётся, только если в нём нашлось
    всё. Одна иконка на слайде не повторяется.
    """
    families = {}
    for asset in assets:
        if asset["kind"] == "icon":
            families.setdefault(asset["family"], []).append(asset)
    best = None
    for family, members in families.items():
        neutral = next((a for a in members if a.get("neutral")), None)
        chosen, matched, taken = [], 0, set()
        for label in labels:
            pool = [a for a in members if a["id"] not in taken and not a.get("neutral")]
            icon = _best(pool, label, used)
            if icon:
                matched += 1
                taken.add(icon["id"])
            chosen.append(icon)
        if any(icon is None for icon in chosen):
            if neutral is None:
                continue
            chosen = [icon or neutral for icon in chosen]
        rank = (matched, PRIORITY.get(members[0]["source"], 0))
        if best is None or rank > best[0]:
            best = (rank, chosen, matched)
    if best is None:
        return [None] * len(labels), 0
    return best[1], best[2]


def pick_picture(text, assets, kinds=("illustration", "photo"), used=(), title=""):
    """Иллюстрация или фото к слайду: заголовок весит больше тезисов."""
    candidates = [a for a in assets if a["kind"] in kinds]
    scored = []
    for asset in candidates:
        score = match_score(asset.get("tags"), title, 1.5) + match_score(
            asset.get("tags"), text, 1.0
        )
        if score < MIN_MATCH:
            continue
        penalty = 0.8 if asset["id"] in used else 0.0
        scored.append((score - penalty, PRIORITY.get(asset["source"], 0), asset["id"], asset))
    if not scored:
        return None
    scored.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return scored[0][3]


def asset_ref(asset):
    """Что кладётся в модель слайда: ссылка и всё, что нужно экспорту."""
    return {
        key: asset.get(key)
        for key in ("id", "ref", "kind", "format", "ratio", "mono", "source", "color")
    }


def resolve(ref, root=None):
    """Ссылка ассета → путь к файлу. Чужой путь из ссылки не собирается."""
    source, _, rest = (ref or "").partition(":")
    if source == "builtin":
        path = (PACK_DIR / rest).resolve()
        if PACK_DIR.resolve() not in path.parents:
            raise ValueError("Asset outside of the built-in pack")
        return path
    owner, _, name = rest.partition(":")
    if root is None or not re.fullmatch(r"[a-f0-9]{32}", owner) or not STORED_NAME.fullmatch(name):
        raise ValueError("Unknown asset reference")
    kind = {
        "template": "templates",
        "pack": "content_packs",
        "generated": "illustrations",
    }.get(source)
    if not kind:
        raise ValueError("Unknown asset source")
    return Path(root) / kind / owner / "assets" / name


# ---------------------------------------------------------------- отрисовка


def _mix(color, other, ratio):
    a = [int(color[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(other[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x * ratio + y * (1 - ratio)):02X}" for x, y in zip(a, b))


def color_roles(accent, background, text, accents=()):
    """Роли цвета иллюстраций → цвета палитры этого слайда."""
    from .parsing import contrast_ratio

    second = next(
        (c for c in accents if c.upper() != accent.upper() and contrast_ratio(c, background) >= 1.6),
        None,
    ) or _mix(accent, text, 0.55)
    return {
        "accent": accent,
        "accent2": second,
        "soft": _mix(accent, background, 0.16),
        "muted": _mix(text, background, 0.3),
        "paper": background,
        "ink": text,
        "current": accent,
    }


def _paint(shape, fill, stroke, width_emu, colors):
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu

    def value(role):
        if role in (None, "none"):
            return None
        if re.fullmatch(r"[0-9A-F]{6}", role):
            return colors.get("mono_override") or role
        return colors.get(role) or colors["current"]

    fill_color, stroke_color = value(fill), value(stroke)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill_color)
    else:
        shape.fill.background()
    if stroke_color:
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = RGBColor.from_string(stroke_color)
        shape.line.width = Emu(max(3175, int(width_emu)))
        line = shape.line._get_or_add_ln()
        line.set("cap", "rnd")
        if line.find(qn("a:round")) is None:
            line.append(OxmlElement("a:round"))
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def draw_vector(shapes, box, geometry, colors, name="illustration"):
    """Нарисовать SVG нативными фигурами в рамке (EMU), без искажения пропорций."""
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE

    x, y, w, h = box
    vx, vy, vw, vh = geometry["viewbox"]
    scale = min(w / vw, h / vh)
    ox = x + (w - vw * scale) / 2 - vx * scale
    oy = y + (h - vh * scale) / 2 - vy * scale
    group = shapes.add_group_shape()
    group.name = name
    members = group.shapes
    # Иконка в одном цвете: чужие цвета SVG приводятся к акценту слайда.
    palette = {**colors}
    if colors.get("tint_all"):
        palette["mono_override"] = colors["current"]
    for item in geometry["items"]:
        width = item["width"] * scale
        if item["shape"] in ("ellipse", "rect"):
            left, top = ox + item["x"] * scale, oy + item["y"] * scale
            cx, cy = item["w"] * scale, item["h"] * scale
            if cx < 1 or cy < 1:
                continue
            preset = SHAPE.OVAL
            if item["shape"] == "rect":
                preset = SHAPE.ROUNDED_RECTANGLE if item.get("rx") else SHAPE.RECTANGLE
            shape = members.add_shape(preset, int(left), int(top), int(cx), int(cy))
            if item["shape"] == "rect" and item.get("rx"):
                shape.adjustments[0] = min(0.5, item["rx"] * scale / max(1, min(cx, cy)))
            _paint(shape, item["fill"], item["stroke"], width, palette)
            continue
        contours = [
            [(int(ox + px * scale), int(oy + py * scale)) for px, py in contour["points"]]
            for contour in item["contours"]
        ]
        contours = [c for c in zip(item["contours"], contours) if len(c[1]) > 1]
        if not contours:
            continue
        first = contours[0][1]
        builder = members.build_freeform(first[0][0], first[0][1], scale=1.0)
        for index, (source, points) in enumerate(contours):
            if index:
                builder.move_to(points[0][0], points[0][1])
            builder.add_line_segments(points[1:], close=source["closed"])
        shape = builder.convert_to_shape()
        _paint(shape, item["fill"], item["stroke"], width, palette)
    group._element.recalculate_extents()
    return group


def draw_raster(shapes, box, blob, ratio, fit="contain", name="image"):
    """Картинка в рамке: contain — вписать целиком, cover — заполнить обрезкой.

    Растягивать нельзя (Приложение 1 ТЗ), поэтому рамка либо подгоняется под
    пропорции картинки, либо лишнее обрезается.
    """
    x, y, w, h = box
    frame = w / h if h else 1.0
    if fit == "cover":
        picture = shapes.add_picture(io.BytesIO(blob), int(x), int(y), int(w), int(h))
        if ratio > frame:
            cut = (1 - frame / ratio) / 2
            picture.crop_left = picture.crop_right = cut
        elif ratio < frame:
            cut = (1 - ratio / frame) / 2
            picture.crop_top = picture.crop_bottom = cut
    else:
        if ratio > frame:
            cw, ch = w, w / ratio
        else:
            cw, ch = h * ratio, h
        picture = shapes.add_picture(
            io.BytesIO(blob), int(x + (w - cw) / 2), int(y + (h - ch) / 2), int(cw), int(ch)
        )
    picture.name = name
    return picture


def draw_asset(shapes, box, asset, colors, root=None, fit="contain", tint=None):
    """Поставить ассет на слайд. Возвращает фигуру или None, если файла нет."""
    try:
        path = resolve(asset["ref"], root)
        if not path.exists():
            return None
    except ValueError:
        LOGGER.warning("Asset %s cannot be resolved", asset.get("ref"))
        return None
    name = f"{asset['kind']}:{asset['id']}"
    if asset.get("format") == "svg":
        geometry = svg_geometry(path)
        palette = dict(colors)
        if tint:
            palette["current"] = tint
        if asset["kind"] == "icon":
            palette["tint_all"] = True
        return draw_vector(shapes, box, geometry, palette, name)
    blob = path.read_bytes()
    if asset.get("mono") and asset["kind"] == "icon":
        try:
            blob = tint_png(blob, tint or colors["current"])
        except Exception:  # noqa: BLE001 - без перекраски иконка всё равно уместна
            LOGGER.warning("Cannot tint %s", asset.get("ref"))
    return draw_raster(shapes, box, blob, asset.get("ratio") or 1.0, fit, name)


# ---------------------------------------------------------------- подписи моделью


TAG_BATCH = 16


def svg_raster(blob, size):
    """SVG картинкой для контакт-листа: наш же разбор геометрии, залитый серым.

    Модели нужно узнать, что изображено, а не оценить цвета, поэтому хватает
    силуэта. Сторонний растеризатор SVG для этого не нужен.
    """
    from PIL import Image, ImageDraw

    geometry = parse_svg(blob)
    vx, vy, vw, vh = geometry["viewbox"]
    scale = size / max(vw or 1, vh or 1)
    image = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    ink = (40, 40, 40, 255)
    ox = (size - vw * scale) / 2 - vx * scale
    oy = (size - vh * scale) / 2 - vy * scale
    for item in geometry["items"]:
        if item["shape"] in ("ellipse", "rect"):
            box = (
                ox + item["x"] * scale,
                oy + item["y"] * scale,
                ox + (item["x"] + item["w"]) * scale,
                oy + (item["y"] + item["h"]) * scale,
            )
            if box[2] - box[0] < 1 or box[3] - box[1] < 1:
                continue
            (draw.ellipse if item["shape"] == "ellipse" else draw.rectangle)(box, fill=ink)
            continue
        for contour in item["contours"]:
            points = [(ox + px * scale, oy + py * scale) for px, py in contour["points"]]
            if len(points) < 2:
                continue
            if contour["closed"] and len(points) > 2:
                draw.polygon(points, fill=ink)
            else:
                draw.line(points, fill=ink, width=max(1, round(item["width"] * scale)))
    return image


def contact_sheet(blobs, cell=150, columns=4):
    """Пронумерованный лист картинок для мультимодальной модели."""
    from PIL import Image, ImageDraw

    from .render import image_from_bytes, png

    rows = math.ceil(len(blobs) / columns)
    sheet = Image.new("RGB", (columns * cell, max(1, rows) * cell), (255, 255, 255))
    draw = ImageDraw.Draw(sheet)
    inner = cell - 34
    for index, (blob, mono) in enumerate(blobs):
        col, row = index % columns, index // columns
        x0, y0 = col * cell, row * cell
        draw.rectangle(
            (x0 + 2, y0 + 2, x0 + cell - 3, y0 + cell - 3),
            fill=(245, 245, 245),
            outline=(204, 204, 204),
        )
        try:
            if blob[:5] == b"<?xml" or b"<svg" in blob[:512]:
                picture = svg_raster(blob, inner)
            else:
                if mono:
                    try:
                        blob = tint_png(blob, "111111")
                    except Exception:  # noqa: BLE001
                        pass
                picture = image_from_bytes(blob)
                picture.thumbnail((inner, inner))
        except Exception:  # noqa: BLE001 - картинку, которую не открыть, модель не увидит
            continue
        left = x0 + 22 + (inner - picture.width) // 2
        top = y0 + 22 + (inner - picture.height) // 2
        sheet.paste(picture, (left, top), picture if picture.mode == "RGBA" else None)
        draw.text((x0 + 6, y0 + 4), str(index + 1), fill=(204, 26, 26))
    return png(sheet)


# ---------------------------------------------------------------- генерация картинок


# Что просим у генератора всегда: картинка идёт на слайд рядом с текстом, поэтому
# букв на ней быть не должно, а композиция — спокойной и в цветах шаблона.
IMAGE_STYLE = (
    "минималистичная векторная иллюстрация, плоские заливки, без текста, "
    "без букв, без цифр, без логотипов, без рамок, спокойная композиция, "
    "много свободного места, фон {background}, основные цвета {accent} и {secondary}"
)


def illustration_prompt(title, bullets, palette):
    """Описание картинки для слайда: тема из заголовка, цвета — из шаблона.

    Промпт собирается детерминированно, а не моделью: лишний вызов на каждый
    слайд стоит денег и времени, а тема слайда уже сформулирована в заголовке.
    """
    idea = " ".join([title or ""] + [b for b in (bullets or [])][:2]).strip()
    idea = re.sub(r"\s+", " ", idea)[:220]
    style = IMAGE_STYLE.format(
        background=palette.get("background", "#FFFFFF"),
        accent=palette.get("accent", "#1478F4"),
        secondary=palette.get("secondary", palette.get("accent", "#1478F4")),
    )
    return f"{idea}. {style}"


def illustration_targets(slides, limit):
    """Слайды, которым картинка нужнее всего.

    Берём те, где нет своей визуализации и мало текста: именно туда вёрстка
    ставит иллюстрацию. Обложку и финал пропускаем — там композицию задаёт
    страница шаблона.
    """
    chosen = []
    total = len(slides)
    for index, content in enumerate(slides):
        if index == 0 or (total > 2 and index == total - 1):
            continue
        if (content.get("visual") or {}).get("kind", "none") != "none":
            continue
        bullets = content.get("bullets") or []
        if not bullets or len(bullets) > 4 or sum(len(b) for b in bullets) > 320:
            continue
        chosen.append((index, content))
        if len(chosen) >= limit:
            break
    return chosen


async def generate_illustrations(slides, palette, folder: Path, owner, limit=3, size=None):
    """Нарисовать иллюстрации к слайдам и сохранить их как ассеты колоды.

    Возвращает метаданные в том же виде, в каком их отдаёт библиотека: вёрстка
    подбирает картинку по тегам, а теги здесь — слова самого слайда, поэтому
    сгенерированная картинка встаёт именно на тот слайд, для которого нарисована.
    Ошибка генератора не роняет колоду: слайд просто останется без картинки.
    """
    import asyncio

    from .generation import draw_image, image_available

    if not image_available():
        return []
    targets = illustration_targets(slides, limit)
    if not targets:
        return []

    async def one(index, content):
        prompt = illustration_prompt(content.get("title"), content.get("bullets"), palette)
        try:
            blob = await draw_image(prompt, size)
            info = inspect_image(blob)
        except Exception as exc:  # noqa: BLE001 - без картинки слайд собирается как раньше
            LOGGER.warning("Illustration for slide %s failed: %s", index, exc)
            return None
        identifier = hashlib.sha256(blob).hexdigest()[:24]
        label = [w for w in words(content.get("title") or "") if len(w) > 2][:5]
        return (
            {
                "id": identifier,
                "kind": "illustration",
                "format": "png",
                "file": f"{identifier}.png",
                "ratio": round((info["w"] or 1) / (info["h"] or 1), 4),
                "mono": False,
                "color": info.get("color"),
                "slide": index,
                "label": label,
                "tags": tag_roots(label),
                "prompt": prompt,
            },
            blob,
        )

    drawn = await asyncio.gather(*(one(index, content) for index, content in targets))
    harvested = [item for item in drawn if item]
    if not harvested:
        return []
    metas = save_assets(folder, harvested)
    for meta in metas:
        meta["source"] = "generated"
        meta["family"] = "generated:" + owner
        meta["ref"] = f"generated:{owner}:{meta['file']}"
    return metas


async def drop_lettered(metas, folder: Path):
    """Убрать картинки, на которых генератор всё-таки написал текст.

    Буквы на иллюстрации — это чужой язык, опечатки и бессмысленные слова рядом
    с выверенным текстом слайда. Проверяет мультимодальная модель, та же, что
    смотрит слайды в аудите; без неё картинки остаются как есть.
    """
    import asyncio

    from .generation import PROMPTS, vision_available, vision_completion, workflow

    if not metas or not vision_available():
        return metas
    agent = workflow()["agents"].get("image_check")
    if not agent:
        return metas
    prompt = (PROMPTS / agent).read_text(encoding="utf-8")

    async def check(meta):
        try:
            answer = await vision_completion(
                prompt, {"id": meta["id"]}, (folder / meta["file"]).read_bytes()
            )
        except Exception as exc:  # noqa: BLE001 - проверка не должна ронять колоду
            LOGGER.warning("Cannot check illustration %s: %s", meta["id"], exc)
            return True
        return not bool(answer.get("has_text"))

    keep = await asyncio.gather(*(check(meta) for meta in metas))
    clean = []
    for meta, ok in zip(metas, keep):
        if ok:
            clean.append(meta)
            continue
        LOGGER.info("Illustration %s dropped: generator wrote text on it", meta["id"])
        (folder / meta["file"]).unlink(missing_ok=True)
    return clean


async def tag_assets(assets, load, budget=None):
    """Подписать картинки шаблона: что на них изображено, по-русски.

    Без подписи иконку не подобрать к тезису. Модель видит лист из 16
    пронумерованных картинок и возвращает по два-четыре слова на каждую.
    Картинки, которые не удалось подписать, помечаются и больше не отправляются.
    `load(meta)` возвращает байты картинки.
    """
    import asyncio

    from .generation import PROMPTS, vision_available, vision_completion, workflow

    if not vision_available():
        return 0
    pending = [a for a in assets if not a.get("tags") and not a.get("tag_tried")]
    if budget is not None:
        pending = pending[:budget]
    if not pending:
        return 0
    prompt = (PROMPTS / workflow()["agents"]["assets"]).read_text()
    limit = asyncio.Semaphore(4)
    tagged = 0

    async def batch(items):
        nonlocal tagged
        blobs = []
        for meta in items:
            try:
                blobs.append((load(meta), meta.get("mono")))
            except OSError:
                blobs.append((b"", False))
        sheet = contact_sheet(blobs)
        async with limit:
            try:
                answer = await vision_completion(prompt, {"count": len(items)}, sheet)
            except Exception as exc:  # noqa: BLE001 - подписи не критичны для колоды
                LOGGER.warning("Asset tagging failed: %s", exc)
                answer = {}
        entries = answer.get("items") if isinstance(answer, dict) else None
        by_number = {}
        for entry in entries or []:
            if isinstance(entry, dict) and isinstance(entry.get("n"), int):
                by_number[entry["n"]] = [str(t)[:40] for t in entry.get("tags") or []][:5]
        for number, meta in enumerate(items, start=1):
            meta["tag_tried"] = True
            label = by_number.get(number) or []
            if label:
                meta["label"] = label
                meta["tags"] = tag_roots(label)
                tagged += 1

    await asyncio.gather(
        *(batch(pending[i : i + TAG_BATCH]) for i in range(0, len(pending), TAG_BATCH))
    )
    return tagged
