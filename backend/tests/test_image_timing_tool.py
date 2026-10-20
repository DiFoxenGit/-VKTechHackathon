import importlib.util
import json
import io
import hashlib
from pathlib import Path
import sys

import httpx
import pytest
from PIL import Image
from pptx import Presentation

script = Path(__file__).resolve().parents[2] / 'tools/check_image_timing.py'
spec = importlib.util.spec_from_file_location('check_image_timing', script)
tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tool)


def test_default_is_dry_run_without_network(monkeypatch, capsys):
    monkeypatch.setattr(sys, 'argv', [str(script)])
    monkeypatch.setattr(httpx, 'get', lambda *a, **kw: pytest.fail('Unexpected paid/network call'))
    tool.main()
    output = json.loads(capsys.readouterr().out)
    assert output['dry_run'] is True and output['requested_new_images'] == 2


def test_existing_marker_prevents_resubmission(monkeypatch, tmp_path):
    path = tmp_path / 'report.json'
    path.write_text('{"status":"accepted"}')
    monkeypatch.setattr(sys, 'argv', [str(script), '--live', '--output', str(path)])
    monkeypatch.setattr(httpx, 'get', lambda *a, **kw: pytest.fail('Unexpected network call'))
    with pytest.raises(SystemExit, match='inspect accepted jobs'):
        tool.main()


def test_concurrent_run_cannot_overwrite_reservation(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, 'argv', [str(script), '--live', '--output', str(tmp_path / 'report.json')])
    monkeypatch.setattr(httpx, 'get', lambda *a, **kw: pytest.fail('Unexpected network call'))
    def reserved(*args, **kwargs):
        raise FileExistsError()
    monkeypatch.setattr(tool.os, 'open', reserved)
    with pytest.raises(SystemExit, match='inspect accepted jobs'):
        tool.main()


def test_failed_job_does_not_repeat_post():
    requests = []
    def handler(request):
        requests.append(request)
        return httpx.Response(202, json={'id': 'job'}) if request.method == 'POST' else httpx.Response(200, json={'status': 'failed', 'stage': 'failed'})
    with httpx.Client(base_url='http://test/api/v1', transport=httpx.MockTransport(handler)) as client:
        entries = []
        saved = []
        with pytest.raises(RuntimeError, match='no automatic resubmission'):
            tool.run_case(client, {'generate_images': True, 'outline': tool.OUTLINE}, entries,
                          'with_image', lambda: saved.append(json.dumps(entries)))
    assert sum(r.method == 'POST' for r in requests) == 1
    assert entries[0]['job_id'] == 'job'
    assert 'submitting' in saved[0] and 'accepted' in saved[1]


@pytest.mark.parametrize('expected_slide', [1, 2])
def test_native_picture_verification_detects_wrong_slide(expected_slide):
    png = io.BytesIO()
    Image.new('RGB', (16, 16), 'blue').save(png, format='PNG')
    digest = hashlib.sha256(png.getvalue()).hexdigest()[:24]
    pptx = Presentation()
    pptx.slides.add_slide(pptx.slide_layouts[6]).shapes.add_picture(io.BytesIO(png.getvalue()), 0, 0)
    pptx.slides.add_slide(pptx.slide_layouts[6])
    blob = io.BytesIO(); pptx.save(blob)
    slides = [{'elements': []}, {'elements': []}]
    slides[expected_slide-1]['elements'] = [{'asset': {'ref': f'generated:owner:{digest}.png'}}]
    if expected_slide == 1:
        assert tool.verify_generated_picture(blob.getvalue(), {'deck': {'slides': slides}}) == [(1, digest)]
    else:
        with pytest.raises(AssertionError):
            tool.verify_generated_picture(blob.getvalue(), {'deck': {'slides': slides}})
