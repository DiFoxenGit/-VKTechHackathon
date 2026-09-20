import io
import os
import zipfile
from collections import Counter
from pathlib import Path

import pytest
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from designer.fonts import embedded_coverage, missing_glyphs, resolve_font
from designer.parsing import NS, parse_template


def pptx_bytes(deck):
    stream = io.BytesIO()
    deck.save(stream)
    # Set both theme roles explicitly so this fixture never depends on which
    # Office version supplied python-pptx's default template.
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("ppt/theme/theme") and item.filename.endswith(".xml"):
                root = etree.fromstring(data)
                for latin in root.findall(".//a:fontScheme//a:latin", NS):
                    latin.set("typeface", "Calibri")
                data = etree.tostring(root)
            target.writestr(item, data)
    return output.getvalue()


def add_text(slide, text, face, size, y=2):
    shape = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(8), Inches(0.5))
    shape.text = text
    shape.text_frame.paragraphs[0].runs[0].font.name = face
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return shape


def test_theme_font_resolution_and_character_weighted_russian_choice(monkeypatch):
    monkeypatch.setattr("designer.fonts._system_fonts", lambda: {})
    deck = Presentation()
    for index in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Заголовок по теме " * 6
        slide.shapes.title.text_frame.paragraphs[0].font.name = "+mj-lt"
        slide.placeholders[1].text = "Основной текст по теме " * 40
        slide.placeholders[1].text_frame.paragraphs[0].font.name = "+mn-lt"
    prototype = deck.slides.add_slide(deck.slide_layouts[6])
    add_text(prototype, "Prototype", "Poppins Light", 32)
    template = parse_template(pptx_bytes(deck), "synthetic.pptx")
    assert template["tokens"]["fonts"][0] == "Calibri"
    for role in ("title", "body"):
        choice = template["tokens"]["font_choice"][role]
        assert choice["selected"] == "Calibri"
        assert choice["candidates"][0]["characters"] > 100
    assert all(slot["style"]["font"] == "Calibri" for pattern in template["patterns"] for slot in pattern["slots"])
    assert resolve_font("+mn-cs", {"minor": {"cs": "", "latin": "Calibri"}}) == "Calibri"
    assert resolve_font("+mj-lt", {"major": {"latin": "Arial"}}) == "Arial"


def test_latin_only_dominant_font_falls_back_with_recorded_reason(monkeypatch):
    monkeypatch.setattr("designer.fonts._system_fonts", lambda: {})
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_text(slide, "Prototype " * 100, "Poppins Light", 36, 0.4)
    add_text(slide, "Body specimen " * 100, "Poppins Light", 20)
    add_text(slide, "Русский образец", "+mn-lt", 20, 4)
    template = parse_template(pptx_bytes(deck), "fallback.pptx")
    for role in ("title", "body"):
        choice = template["tokens"]["font_choice"][role]
        assert choice["selected"] == "Calibri"
        assert choice["reason"] == "cyrillic_fallback"
        assert "Poppins Light" in choice["rejected"]
    assert missing_glyphs("Poppins Light", "ABC Яя", template["tokens"]["font_coverage"]) == ["Я", "я"]


def test_embedded_font_character_map_overrides_fallback_name():
    builder = FontBuilder(1000, isTTF=True)
    glyphs = [".notdef", "A", "uni0410"]
    builder.setupGlyphOrder(glyphs)
    builder.setupCharacterMap({65: "A", 1040: "uni0410"})
    builder.setupGlyf({name: TTGlyphPen(None).glyph() for name in glyphs})
    builder.setupHorizontalMetrics({name: (500, 0) for name in glyphs})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({"familyName": "Fixture Font", "styleName": "Regular", "fullName": "Fixture Font", "psName": "FixtureFont"})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
    builder.setupPost()
    buffer = io.BytesIO()
    builder.save(buffer)
    archive_buffer = io.BytesIO()
    with zipfile.ZipFile(archive_buffer, "w") as archive:
        archive.writestr("ppt/fonts/font1.fntdata", buffer.getvalue())
    with zipfile.ZipFile(io.BytesIO(archive_buffer.getvalue())) as archive:
        coverage = embedded_coverage(archive)
    assert coverage["fixture font"]["source"] == "embedded"
    assert missing_glyphs("Fixture Font", "AА Б", coverage) == ["Б"]


def test_russian_export_uses_theme_font_and_audit_detects_missing_glyphs(tmp_path, monkeypatch):
    from designer.audit import audit
    from designer.exporting import export_pptx
    from designer.layout import compose
    from designer.models import Outline
    monkeypatch.setattr("designer.fonts._system_fonts", lambda: {})
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = 'Русский заголовок'
    slide.placeholders[1].text = 'Основной текст по теме ' * 20
    slide.placeholders[1].text_frame.paragraphs[0].font.name = '+mn-lt'
    add_text(slide, 'Specimen', 'Poppins Light', 18, 4)
    data = pptx_bytes(source)
    template = parse_template(data, 'unknown.pptx')
    plan = Outline.model_validate({'title': 'Колода', 'slides': [
        {'title': 'Русский заголовок', 'bullets': ['Русский текст слайда'], 'source_refs': ['brief']}
    ]}).model_dump()
    deck = compose(plan, template, 'classic')
    original = tmp_path / 'template.pptx'
    original.write_bytes(data)
    result = tmp_path / 'result.pptx'
    export_pptx(original, template, deck, result)
    with zipfile.ZipFile(result) as archive:
        part = next(name for name in archive.namelist() if name.startswith('ppt/slides/slide') and name.endswith('.xml'))
        xml = etree.fromstring(archive.read(part))
        assert 'Poppins Light' not in xml.xpath('//a:r/a:rPr/a:latin/@typeface', namespaces=NS)
        assert 'Calibri' in xml.xpath('//a:r/a:rPr/a:latin/@typeface', namespaces=NS)
    deck['slides'][0]['elements'][0]['font'] = 'Poppins Light'
    findings = audit(deck, template, [{'id': 'brief', 'text': 'Русский текст'}])['issues']
    assert any(i['code'] == 'font_missing_glyphs' and i['severity'] == 'error' for i in findings)


