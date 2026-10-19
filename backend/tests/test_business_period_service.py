"""Explicit inputs must survive HTTP generation, replay and revisions."""
import copy
import io
import json
import os
from pathlib import Path
import time

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from designer.app import create_app
from designer.analytics_bridge import extract, build_outline, verify_outline
from designer.business_bridge import prepare_business

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'samples/business-acceptance/data-v1'
EXT=ROOT/'samples/business-acceptance/period-context-v1'


def read(path):return json.loads(path.read_text(encoding='utf8'))


def inputs(version='baseline'):
    folder=EXT/version
    sources=[{'id':'brief','text':'Прибыль и маржа, план-факт'}]
    for name in ('sales','plan','products','regions'):
        path=(folder if name in {'sales','plan'} else DATA)/(name+'.csv')
        sources.append({'id':name,'text':'','datasets':extract(path.read_bytes(),path.name)})
    sources.append({'id':'context','text':(folder/'context.txt').read_text(encoding='utf8')})
    sources[0]['business_request']={'settings':read(EXT/'settings.json'),'conditions':read(folder/'conditions.json'),'context_source_id':'context'}
    return sources


def upload(client, version='baseline'):
    folder=EXT/version;ids=[]
    for name in ('sales.csv','plan.csv','products.csv','regions.csv','context.txt'):
        path=(DATA if name in {'products.csv','regions.csv'} else folder)/name
        response=client.post('/api/v1/content-packs',files={'file':(name,path.read_bytes())})
        assert response.status_code==201,response.text
        ids.append(response.json()['id'])
    return {'brief':'Прибыль и маржа, план-факт по выбранным периодам','slide_count':12,'template_id':'analysis',
        'content_pack_ids':ids,'business_settings':read(EXT/'settings.json'),
        'business_conditions':read(folder/'conditions.json'),'business_context_id':ids[-1]}


def test_bridge_periods_exclusions_conditions_and_source_changes():
    totals=[];tokens=[]
    for version,capacity in [('baseline',2),('version-b',3)]:
        sources=inputs(version);results=prepare_business(sources,sources[0]['text']);payload=sources[0]['_business']
        assert payload['analysis']['facts']['q3_net']['period']==['2026-04','2026-06']
        assert len(sources[1]['datasets'][0]['contract']['rows'])==1512
        assert len(payload['analysis']['source_scope']['excluded_sales_rows'])==216
        assert payload['story']['capacity']['engineers']==capacity
        for result in results:
            assert '2026-04' in result['business_explanation']['period']
            assert '216' in result['business_explanation']['exclusions']
        outline=build_outline(results,sources,12)
        assert verify_outline(outline,results,sources)
        trend=next(s for s in outline.slides if s.analysis_id.startswith('business:trend:'))
        assert trend.visual.categories[-1]=='2026-06-01'
        assert not any('III квартал' in s.title or 'II кварталу' in s.title or '90 дней' in s.title for s in outline.slides)
        totals.append(int(payload['analysis']['facts']['q3_net']['exact']));tokens.append(outline.slides[0].analysis_id)
    assert totals==[93563698,94823698] and tokens[0]!=tokens[1]


@pytest.mark.parametrize('target',['settings','conditions','missing_request','old_version'])
def test_trusted_request_cannot_be_replaced_by_result_metadata(target):
    sources=inputs();results=prepare_business(sources,sources[0]['text']);outline=build_outline(results,sources,12)
    if target=='settings':sources[0]['business_request']['settings']['current']=['2026-07','2026-09']
    if target=='conditions':sources[0]['business_request']['conditions']['capacity']['engineers']['value']=3
    if target=='missing_request':del sources[0]['business_request']
    if target=='old_version':sources[0]['_business']['version']='business-service-1'
    with pytest.raises(ValueError):verify_outline(outline,results,sources)


def test_http_analysis_outline_and_explained_refusals(tmp_path,monkeypatch):
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    with TestClient(create_app(tmp_path)) as client:
        request=upload(client)
        analysis=client.post('/api/v1/analysis',json=request)
        assert analysis.status_code==200,analysis.text
        assert len(analysis.json()['items'])==12
        outline_request={k:v for k,v in request.items() if k!='template_id'}
        assert client.post('/api/v1/outlines',json=outline_request).status_code==200
        bad=copy.deepcopy(request);bad['business_conditions']['capacity']['engineers']['value']=3
        response=client.post('/api/v1/analysis',json=bad)
        assert response.status_code==422 and 'не подтверждено' in response.text
        bad=copy.deepcopy(request);bad['business_settings']['previous']=['2026-01','2026-02']
        response=client.post('/api/v1/analysis',json=bad)
        assert response.status_code==422 and 'одинаковой длины' in response.text
        bad=copy.deepcopy(request);bad['business_context_id']='missing'
        assert client.post('/api/v1/analysis',json=bad).status_code==422
        bad=copy.deepcopy(request);bad['business_conditions']=None;bad['business_context_id']=None
        assert client.post('/api/v1/analysis',json=bad).status_code==422
        # A fresh valid request succeeds after each refusal, without stale server state.
        assert client.post('/api/v1/analysis',json=request).status_code==200


def test_generation_save_reopen_native_exports_and_preview(tmp_path,monkeypatch):
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    monkeypatch.delenv('DESIGNER_IMAGE_BASE_URL',raising=False)
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text='Тема';slide.placeholders[1].text='Образец'
    stream=io.BytesIO();prs.save(stream)
    with TestClient(create_app(tmp_path)) as client:
        template=client.post('/api/v1/templates',files={'file':('template.pptx',stream.getvalue())}).json()
        request=upload(client);request['template_id']=template['id'];request['generate_images']=False
        response=client.post('/api/v1/generations',json=request)
        assert response.status_code==202,response.text
        job=response.json();deadline=time.monotonic()+90
        while job['status'] not in {'completed','failed'} and time.monotonic()<deadline:
            time.sleep(.05);job=client.get('/api/v1/jobs/'+job['id']).json()
        assert job['status']=='completed',job
        identifier=job['presentation_ids'][0]
        record=client.get('/api/v1/presentations/'+identifier).json()
        assert record['business_request']['settings']==request['business_settings']
        assert record['business_request']['conditions']==request['business_conditions']
        slides=[s['content'] for s in record['deck']['slides']]
        changed=copy.deepcopy(slides);changed[0]['title']='Подменённый вывод'
        edit={'revision':1,'title':'Сохранённый отчёт Q2','variant':'focus','slides':changed}
        assert client.patch('/api/v1/presentations/'+identifier,json=edit).status_code==422
        edit['slides']=slides[1:]+slides[:1]
        response=client.patch('/api/v1/presentations/'+identifier,json=edit)
        assert response.status_code==200,response.text
        assert response.json()['revision']==2
        assert client.patch('/api/v1/presentations/'+identifier,json=edit).status_code==409
        pptx=client.get(f'/api/v1/presentations/{identifier}/export/pptx?revision=2')
        native=Presentation(io.BytesIO(pptx.content))
        assert len(native.slides)==12 and sum(s.has_chart for p in native.slides for s in p.shapes)==5
        assert sum(s.has_table for p in native.slides for s in p.shapes)==2
        if os.getenv('DESIGNER_SOFFICE'):
            from pypdf import PdfReader
            pdf=client.get(f'/api/v1/presentations/{identifier}/export/pdf?revision=2')
            assert pdf.status_code==200,pdf.text[:200] if pdf.status_code!=200 else ''
            assert len(PdfReader(io.BytesIO(pdf.content)).pages)==12
            preview=client.get(f'/api/v1/presentations/{identifier}/slides/0/thumbnail?revision=2&width=640')
            assert preview.status_code==200 and preview.content.startswith(b'\x89PNG')
            assert client.get(f'/api/v1/presentations/{identifier}/slides/0/thumbnail?revision=1').status_code==409
    with TestClient(create_app(tmp_path)) as client:
        reopened=client.get('/api/v1/presentations/'+identifier).json()
        assert reopened['revision']==2 and reopened['title']=='Сохранённый отчёт Q2'
        assert reopened['business_request']==record['business_request']
        assert reopened['deck']['slides'][0]['content']==edit['slides'][0]
