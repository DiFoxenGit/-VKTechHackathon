"""Vendor-independent regressions from held-out template acceptance."""
import io
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from designer.exporting import _adapt_inherited_layout, _set_notes


def test_reflow_does_not_change_other_slides_shared_layout():
    p=Presentation()
    first=p.slides.add_slide(p.slide_layouts[6])
    second=p.slides.add_slide(p.slide_layouts[6])
    layout=first.slide_layout
    native=first.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(250),Pt(600),Pt(1))
    native.name='Native divider'
    native._element.xpath('./p:nvSpPr/p:cNvPr')[0].set('id','1001')
    layout.shapes._spTree.insert_element_before(native._element,'p:extLst')
    logo=first.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(600),Pt(20),Pt(50),Pt(20))
    logo.name='Native logo'
    logo._element.xpath('./p:nvSpPr/p:cNvPr')[0].set('id','1002')
    layout.shapes._spTree.insert_element_before(logo._element,'p:extLst')
    elements=[{'kind':'text','text':'The new paragraph crosses the old divider',
               'box':[40,230,600,100],'font_size':20}]
    _adapt_inherited_layout(first,elements,p.slide_width,p.slide_height)
    out=io.BytesIO();p.save(out);out.seek(0);reopened=Presentation(out)
    names=lambda slide:{s.name for s in slide.slide_layout.shapes}
    assert 'Native divider' not in names(reopened.slides[0])
    assert 'Native divider' in names(reopened.slides[1])
    assert all('Native logo' in names(s) for s in reopened.slides)
    assert reopened.slides[0].slide_layout.part != reopened.slides[1].slide_layout.part


def test_missing_notes_body_does_not_break_editable_notes():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);notes=s.notes_slide
    body=notes.notes_placeholder
    body._element.getparent().remove(body._element)
    assert notes.notes_text_frame is None
    _set_notes(s,'Evidence and caveats remain editable')
    out=io.BytesIO();p.save(out);out.seek(0);saved=Presentation(out)
    assert saved.slides[0].notes_slide.notes_text_frame.text=='Evidence and caveats remain editable'


def test_theme_background_reference_honours_master_color_map():
    from designer.parsing import background_info
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);master=s.slide_layout.slide_master
    master._element.find(qn('p:clrMap')).set('bg1','dk1')
    bg=master._element.cSld.get_or_add_bg()
    for child in list(bg):bg.remove(child)
    ref=OxmlElement('p:bgRef');ref.set('idx','1001')
    color=OxmlElement('a:schemeClr');color.set('val','bg1');ref.append(color);bg.append(ref)
    assert background_info(s,{'lt1':'FFFFFF','dk1':'0A1828'})=={'color':'0A1828','kind':'solid'}


def test_divider_follows_title_without_striking_new_picture():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);layout=s.slide_layout
    rule=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(40),Pt(300),Pt(600),Pt(1))
    rule.name='Divider';layout.shapes._spTree.insert_element_before(rule._element,'p:extLst')
    elements=[{'role':'title','kind':'text','text':'A short title','box':[40,40,600,50],'font_size':24},
              {'role':'picture','kind':'image','box':[400,140,200,200]}]
    _adapt_inherited_layout(s,elements,p.slide_width,p.slide_height)
    out=io.BytesIO();p.save(out);out.seek(0);saved=Presentation(out)
    moved=next(x for x in saved.slides[0].slide_layout.shapes if x.name=='Divider')
    assert 70 < moved.top/12700 < 140
    original=next(x for x in layout.shapes if x.name=='Divider')
    assert original.top==Pt(300)


def test_picture_stage_keeps_usable_height_around_specimen_rule():
    from designer.parsing import parse_template
    from designer.models import Outline
    from designer.layout import compose
    from designer.illustration_plan import target_key
    from tests.test_pipeline import template_bytes
    t=parse_template(template_bytes(),'unseen.pptx')
    for pattern in t['patterns']:
        pattern['role']='content'
        pattern['reserved'].append({'x':.04,'y':.65,'w':.9,'h':.003,'layout':True})
    content=Outline.model_validate({'title':'Test','slides':[{'role':'content','title':'A decision supported by evidence',
        'bullets':['A verified result from the selected period.','A limitation of the available data.','An action to verify the assumption.']}]}).model_dump()
    asset={'source':'generated','kind':'illustration','id':'picture','ref':'generated:'+('a'*32)+':picture.png',
           'format':'png','ratio':1,'target_key':target_key(content['slides'][0]),'placement':'right','tags':[],'label':[]}
    for variant in ['classic','split','focus']:
        deck=compose(content,t,variant,assets=[asset])
        picture=next(e for e in deck['slides'][0]['elements'] if e.get('role')=='picture')
        assert picture['box'][3]>=deck['height']*.35
        assert picture['box'][1]<deck['height']*.5
