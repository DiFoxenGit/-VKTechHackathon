"""Rules run without model transport; pipeline calls each dependency explicitly."""
import asyncio
import copy
import io
import subprocess
import sys

import pytest

from designer.algorithms.context import chunk_text, select_context
from designer.algorithms.balance import balance_plan, apply_condensation, finish_balance
from designer.models import GenerateRequest, Outline
from designer.pipeline import run_generation
from pptx import Presentation
from tests.test_pipeline import client, template_bytes, wait_job


def plan():
    return Outline.model_validate({"title": "Пилот", "slides": [
        {"title": "Пилот для 20 команд", "bullets": ["Подробности внедрения. " * 40], "source_refs": ["brief"]}
    ]}).model_dump()


def test_deterministic_modules_do_not_load_ai_transport():
    code = """
import sys
from designer.algorithms import context, content, balance, validation
from designer import audit, layout, assets
assert 'designer.ai' not in sys.modules
assert 'designer.pipeline' not in sys.modules
assert 'httpx' not in sys.modules
"""
    subprocess.run([sys.executable, "-c", code], check=True)


def test_chunking_bounds_order_and_progress_without_spaces():
    text = 'Начало\n\n' + 'x' * 5000 + '\n\nКонец'
    chunks = chunk_text(text, 80)
    assert all(0 < len(chunk) <= 80 for chunk in chunks)
    assert ''.join(chunks).replace('\n', '') == text.replace('\n', '')
    with pytest.raises(ValueError):
        chunk_text(text, 0)


@pytest.mark.parametrize('budget', [0, 3, 50, 1000, 6000])
def test_context_never_exceeds_budget_across_many_sources(budget):
    sources = [{"id": str(index), "text": 'Контекст источника. ' * 300} for index in range(20)]
    original = copy.deepcopy(sources)
    selected = select_context(sources, 'источника', budget)
    assert sum(len(source['text']) for source in selected) <= budget
    assert [source['id'] for source in selected] == [source['id'] for source in sources]
    if budget >= len(sources):
        assert all(source['text'] for source in selected)
    assert sources == original


@pytest.mark.parametrize('answer', [
    None, [], {}, {'slides': [None]},
    {'slides': [{'index': True, 'bullets': ['Тезис']}]},
    {'slides': [{'index': 10, 'bullets': ['Тезис']}]},
    {'slides': [{'index': 0, 'bullets': [None]}]},
    {'slides': [{'index': 0, 'bullets': ['Доход 99999 рублей']}]},
    {'slides': []},
    {'slides': [{'index': 0, 'bullets': ['Первый']}, {'index': 0, 'bullets': ['Второй']}]},
])
def test_invalid_condensation_is_rejected_without_mutation(answer):
    original = plan()
    snapshot = copy.deepcopy(original)
    with pytest.raises(ValueError):
        apply_condensation(original, answer, balance_plan(original, 200))
    assert original == snapshot


def test_condensation_preserves_sources_visuals_and_original_text_in_notes():
    original = plan()
    payload = balance_plan(original, 200)
    result = apply_condensation(original, {'slides': [{'index': 0, 'bullets': ['Пилот для 20 команд']}]}, payload)
    warnings = []
    balanced = finish_balance(original, result, payload, warnings)
    assert balanced['slides'][0]['source_refs'] == ['brief']
    assert balanced['slides'][0]['visual'] == original['slides'][0]['visual']
    assert original['slides'][0]['bullets'][0] in balanced['slides'][0]['notes']
    assert balanced['slides'][0]['bullets'] == ['Пилот для 20 команд']
    assert warnings


def test_notes_overflow_keeps_source_and_warns_instead_of_losing_text():
    original = plan()
    original['slides'][0]['bullets'] = ['Подробности. ' * 1000]
    warnings = []
    balanced = finish_balance(original, original, balance_plan(original, 200), warnings)
    assert balanced == original
    assert len(warnings) == 1 and 'не сокращён' in warnings[0]


def test_condensation_reapplies_typography_rules_after_the_model():
    original = plan()
    result = apply_condensation(original, {'slides': [
        {'index': 0, 'title': 'пилот для 20 команд.', 'bullets': ['дизайн‑система для 20\u202fкоманд.']}
    ]}, balance_plan(original, 200))
    assert result['slides'][0]['title'] == 'Пилот для 20 команд'
    assert result['slides'][0]['bullets'] == ['Дизайн-система для 20 команд']


def test_invalid_optional_model_profile_keeps_deterministic_fallback(client, monkeypatch):
    from designer.pipeline import balance_outline
    monkeypatch.setenv('DESIGNER_LLM_PROVIDER', 'openai')
    monkeypatch.setenv('DESIGNER_LLM_TIMEOUT', 'invalid')
    response = client.get('/api/v1/templates')
    assert response.status_code == 200
    assert response.json()['llm_configured'] is False
    warnings = []
    original = plan()
    balanced = asyncio.run(balance_outline(original, {}, capacity=200, warnings=warnings))
    assert original['slides'][0]['bullets'][0] in balanced['slides'][0]['notes']
    assert any('ИИ не сократил' in warning for warning in warnings)


def test_pipeline_combines_ai_and_rules_for_each_renderer():
    events, rendered, published = [], [], []
    request = GenerateRequest(brief='Пилот для 20 команд', template_id='template', slide_count=1, contextual_audit=True)
    async def generate(request, sources, warnings):
        events.append('ai:outline')
        return Outline.model_validate(plan())
    async def balance(outline, template, warnings):
        events.append('rules:balance')
        outline['slides'][0]['bullets'] = ['20 команд участвуют']
        return outline
    async def label(*args): events.append('ai:assets')
    async def draw(deck, template):
        events.append('ai:images')
        assert deck['slides'][0]['content']['bullets'] == ['20 команд участвуют']
        return 'owner', [{'id': 'image'}]
    async def check(deck, sources):
        events.append('ai:audit')
        return [{'code': 'review'}]
    async def render(request, outline, template, sources, variant, findings, images):
        rendered.append(copy.deepcopy(outline))
        assert findings == [{'code': 'review'}]
        assert images[0] == 'owner'
        outline['slides'].clear()  # A renderer cannot corrupt subsequent variants.
        return {'id': variant}
    result = asyncio.run(run_generation(
        request, {}, [], outline_factory=generate, balance=balance, render_variant=render,
        variants=[{'id': name} for name in ['classic', 'split', 'focus']],
        on_stage=lambda *args: None, on_outline=lambda outline: None,
        on_variant=published.append, warnings=[], label_assets=label,
        draw_illustrations=draw, contextual_check=check,
    ))
    assert events == ['ai:outline', 'rules:balance', 'ai:assets', 'ai:images', 'ai:audit']
    assert len(rendered) == 3 and all(outline == rendered[0] for outline in rendered)
    assert result == published


def test_provided_outline_runs_without_ai_and_cancellation_stops_rendering():
    request = GenerateRequest(brief='Исходное содержание', template_id='template', slide_count=1, outline=plan())
    async def never(*args, **kwargs):
        raise AssertionError('Unexpected AI or rendering call')
    async def cancel(*args, **kwargs):
        raise asyncio.CancelledError()
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(run_generation(
            request, {}, [], outline_factory=never, balance=cancel, render_variant=never,
            variants=[{'id': 'classic'}], on_stage=lambda *args: None,
            on_outline=never, on_variant=never, warnings=[],
        ))
    assert request.outline.model_dump() == plan()


def test_api_offline_pipeline_keeps_original_in_exported_notes(client, monkeypatch):
    monkeypatch.delenv('DESIGNER_LLM_MODEL', raising=False)
    monkeypatch.delenv('DESIGNER_LLM_BASE_URL', raising=False)
    monkeypatch.setenv('DESIGNER_LLM_PROVIDER', 'openai')
    assert client.get('/api/v1/templates').json()['llm_configured'] is False
    template = client.post('/api/v1/templates', files={'file': ('sample.pptx', template_bytes())}).json()
    original = plan()
    response = client.post('/api/v1/generations', json={
        'brief': 'Пилот для 20 команд', 'template_id': template['id'],
        'slide_count': 1, 'outline': original,
    })
    assert response.status_code == 202, response.text
    job = wait_job(client, response.json()['id'])
    assert any('заметках' in warning for warning in job['warnings'])
    assert len(job['presentation_ids']) == 3
    for identifier in job['presentation_ids']:
        record = client.get(f'/api/v1/presentations/{identifier}').json()
        blob = client.get(record['exports']['pptx']).content
        pptx = Presentation(io.BytesIO(blob))
        assert original['slides'][0]['bullets'][0] in pptx.slides[0].notes_slide.notes_text_frame.text
    monkeypatch.setenv('DESIGNER_LLM_BASE_URL', 'https://inference.example/v1')
    monkeypatch.setenv('DESIGNER_LLM_MODEL', 'test-model')
    assert client.get('/api/v1/templates').json()['llm_configured'] is True
