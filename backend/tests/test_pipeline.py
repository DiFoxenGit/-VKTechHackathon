import json
import io
import os
import time
import zipfile
from pathlib import Path
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
from designer.app import create_app
from designer.audit import audit
from designer.layout import compose
from designer.models import Visual
from designer.parsing import parse_template

# Corporate templates are not committed: they belong to the case organisers.
# Point DESIGNER_TEMPLATE_DIR at a folder with .pptx files, or drop them into
# <repo>/templates. Tests that need real templates skip when the folder is empty.
TEMPLATES = Path(
    os.getenv("DESIGNER_TEMPLATE_DIR")
    or Path(__file__).resolve().parents[2] / "templates"
)


def sample_templates():
    return sorted(TEMPLATES.glob("*.pptx")) if TEMPLATES.is_dir() else []


def require_template(name_hint=""):
    files = sample_templates()
    if not files:
        pytest.skip(f"No .pptx templates in {TEMPLATES}")
    for path in files:
        if name_hint and name_hint.lower() in path.stem.lower():
            return path
    return files[0]


def template_bytes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Sample title"
    slide.placeholders[1].text = "Sample content"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def outline(count=3):
    return {
        "title": "Проверка пайплайна",
        "slides": [
            {
                "title": f"Вывод {i + 1}",
                "bullets": ["Первый тезис", "Второй тезис"],
                "source_refs": ["brief"],
                "visual": visual,
            }
            for i, visual in enumerate(
                [
                    {
                        "kind": "bar",
                        "categories": ["А", "Б"],
                        "series": [{"name": "Продажи", "values": [10, 20]}],
                        "unit": "млн руб.",
                    },
                    {
                        "kind": "table",
                        "columns": ["Этап", "Срок"],
                        "rows": [["А", "10"], ["Б", "20"]],
                    },
                    {"kind": "process", "steps": ["План", "Работа", "Итог"]},
                ][:count]
            )
        ],
    }


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("DESIGNER_API_KEY", raising=False)
    with TestClient(create_app(tmp_path)) as c:
        yield c


def wait_job(client, identifier):
    for _ in range(200):
        job = client.get("/api/v1/jobs/" + identifier).json()
        if job["status"] in ("completed", "failed"):
            assert job["status"] == "completed", job
            return job
        time.sleep(0.05)
    pytest.fail("generation timed out")


def generate(client, data=None):
    response = client.post(
        "/api/v1/templates", files={"file": ("template.pptx", data or template_bytes())}
    )
    assert response.status_code == 201, response.text
    template = response.json()
    response = client.post(
        "/api/v1/generations",
        json={
            "template_id": template["id"],
            "brief": "Первый тезис. Второй тезис. Продажи А 10 Б 20 млн руб. План Работа Итог. Выводы 1 2 3.",
            "slide_count": 3,
            "outline": outline(),
        },
    )
    assert response.status_code == 202, response.text
    return template, wait_job(client, response.json()["id"])


def test_end_to_end_native_variants_and_revision(client):
    template, job = generate(client)
    assert len(job["presentation_ids"]) == 3
    positions = []
    for identifier in job["presentation_ids"]:
        record = client.get("/api/v1/presentations/" + identifier).json()
        positions.append([e["box"] for e in record["deck"]["slides"][0]["elements"]])
        response = client.get(record["exports"]["pptx"])
        assert response.status_code == 200
        deck = Presentation(io.BytesIO(response.content))
        assert len(deck.slides) == 3
        assert any(s.has_chart for s in deck.slides[0].shapes)
        assert any(s.has_table for s in deck.slides[1].shapes)
        assert any(
            s.has_text_frame and "Вывод" in s.text for s in deck.slides[0].shapes
        )
        assert "Sample content" not in "\n".join(
            s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame
        )
    assert len({str(p) for p in positions}) == 3
    identifier = job["presentation_ids"][0]
    response = client.patch(
        f"/api/v1/presentations/{identifier}/slides/0",
        json={
            "revision": 1,
            "content": {
                "title": "Новый вывод",
                "bullets": ["Подтвержденный тезис"],
                "source_refs": ["brief"],
            },
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["revision"] == 2
    assert (
        client.patch(
            f"/api/v1/presentations/{identifier}/slides/0",
            json={"revision": 1, "content": {"title": "Конфликт"}},
        ).status_code
        == 409
    )
    old = client.get(f"/api/v1/presentations/{identifier}/export/pptx?revision=1")
    assert old.status_code == 200
    assert (
        client.get(
            f"/api/v1/presentations/{identifier}/export/pptx?revision=999"
        ).status_code
        == 404
    )


@pytest.mark.parametrize(
    "path", sample_templates(), ids=lambda p: p.stem
)
def test_supplied_templates(client, path):
    template, job = generate(client, path.read_bytes())
    assert template["tokens"]["colors"]
    assert template["patterns"]
    for identifier in job["presentation_ids"]:
        response = client.get(f"/api/v1/presentations/{identifier}/export/pptx")
        deck = Presentation(io.BytesIO(response.content))
        assert len(deck.slides) == 3
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert archive.testzip() is None


def test_content_upload_invalid_inputs_and_dedup(client, monkeypatch):
    data = template_bytes()
    a = client.post("/api/v1/templates", files={"file": ("a.pptx", data)}).json()
    b = client.post("/api/v1/templates", files={"file": ("b.pptx", data)}).json()
    assert a["id"] == b["id"]
    assert (
        client.post(
            "/api/v1/templates", files={"file": ("x.pptx", b"broken")}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/v1/content-packs", files={"file": ("source.csv", "name,value\nA,10")}
        ).status_code
        == 201
    )
    assert (
        client.post(
            "/api/v1/content-packs", files={"file": ("source.txt", b"")}
        ).status_code
        == 422
    )
    assert client.get("/api/v1/templates/unknown").status_code == 404
    monkeypatch.delenv("DESIGNER_LLM_BASE_URL", raising=False)
    assert (
        client.post(
            "/api/v1/generations", json={"template_id": a["id"], "brief": "test brief"}
        ).status_code
        == 503
    )
    assert (
        client.post(
            "/api/v1/generations",
            json={
                "template_id": a["id"],
                "brief": "test brief",
                "slide_count": 2,
                "outline": outline(),
            },
        ).status_code
        == 422
    )


def test_selective_fixes(client):
    _, job = generate(client)
    identifier = job["presentation_ids"][0]
    store = client.app.state.store
    record = store.get("presentations", identifier)
    record["deck"]["slides"][0]["elements"][0]["box"][0] = -100
    record["deck"]["slides"][1]["elements"][0]["box"][0] = -100
    record["audit"] = audit(
        record["deck"], store.get("templates", record["template_id"]), record["sources"]
    )
    store.put("presentations", record)
    findings = [i for i in record["audit"]["issues"] if i["code"] == "out_of_bounds"]
    response = client.post(
        f"/api/v1/presentations/{identifier}/fixes",
        json={"revision": 1, "issue_ids": [findings[0]["id"]]},
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["deck"]["slides"][0]["elements"][0]["box"][0] == 0
    assert data["deck"]["slides"][1]["elements"][0]["box"][0] == -100
    assert (
        client.post(
            f"/api/v1/presentations/{identifier}/fixes",
            json={"revision": 2, "issue_ids": ["unknown"]},
        ).status_code
        == 422
    )


def test_auth_cors_and_recovery(tmp_path, monkeypatch):
    monkeypatch.setenv("DESIGNER_API_KEY", "test-token")
    app = create_app(tmp_path)
    app.state.store.put("jobs", {"id": "a" * 32, "status": "running"})
    with TestClient(app) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/api/v1/templates").status_code == 401
        assert (
            client.get(
                "/api/v1/templates", headers={"Authorization": "Bearer test-token"}
            ).status_code
            == 200
        )
        response = client.options(
            "/api/v1/generations",
            headers={
                "Origin": "http://localhost:3000",
                "Access-Control-Request-Method": "POST",
                "Access-Control-Request-Headers": "Authorization,Content-Type",
            },
        )
        assert response.status_code == 200
        assert (
            response.headers["access-control-allow-origin"] == "http://localhost:3000"
        )
        assert app.state.store.get("jobs", "a" * 32)["status"] == "failed"


def test_visual_validation():
    with pytest.raises(ValueError):
        Visual(
            kind="bar",
            categories=["A"],
            series=[{"name": "x", "values": [1, 2]}],
            unit="units",
        )
    with pytest.raises(ValueError):
        Visual(kind="table", columns=["A"], rows=[["1", "2"]])


def test_openapi(client):
    spec = client.get("/openapi.json").json()
    assert "/api/v1/generations" in spec["paths"]
    assert "/api/v1/presentations/{presentation_id}/fixes" in spec["paths"]


def test_provider_contract_and_failures(client, monkeypatch):
    import httpx
    from designer import generation

    monkeypatch.setenv("DESIGNER_LLM_BASE_URL", "https://inference.example/v1")
    monkeypatch.setenv("DESIGNER_LLM_MODEL", "test-open-model")
    captured = []
    real_client = httpx.AsyncClient

    def handler(request):
        import json

        payload = json.loads(request.content)
        captured.append(payload)
        return httpx.Response(
            200, json={"choices": [{"message": {"content": json.dumps(outline())}}]}
        )

    monkeypatch.setattr(
        generation.httpx,
        "AsyncClient",
        lambda **kwargs: real_client(transport=httpx.MockTransport(handler), **kwargs),
    )
    response = client.post(
        "/api/v1/outlines", json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 3}
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["slides"]) == 3
    assert captured[0]["model"] == "test-open-model"
    assert captured[0]["response_format"] == {"type": "json_object"}
    # Модель отдаёт три слайда вместо четырёх: после трёх попыток берём что есть,
    # но структура остаётся валидной.
    response = client.post(
        "/api/v1/outlines", json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 4}
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["slides"]) == 3

    def bad_handler(request):
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "not json"}}]}
        )

    monkeypatch.setattr(
        generation.httpx,
        "AsyncClient",
        lambda **kwargs: real_client(
            transport=httpx.MockTransport(bad_handler), **kwargs
        ),
    )
    assert (
        client.post(
            "/api/v1/outlines", json={"brief": "Данные для проверки: 1, 2, 3"}
        ).status_code
        == 502
    )


def test_failed_generation_visible(client, monkeypatch):
    from designer import app as module

    monkeypatch.setenv("DESIGNER_LLM_BASE_URL", "https://inference.example/v1")
    monkeypatch.setenv("DESIGNER_LLM_MODEL", "test-open-model")

    async def fail(*args):
        raise RuntimeError("provider disconnected")

    monkeypatch.setattr(module, "generate_outline", fail)
    template = client.post(
        "/api/v1/templates", files={"file": ("template.pptx", template_bytes())}
    ).json()
    job = client.post(
        "/api/v1/generations",
        json={"template_id": template["id"], "brief": "Проверка ошибки"},
    ).json()
    for _ in range(50):
        result = client.get("/api/v1/jobs/" + job["id"]).json()
        if result["status"] == "failed":
            break
        time.sleep(0.02)
    assert result["status"] == "failed"
    assert result["error"]


def test_pdf_html_export(client):
    if not client.get("/health").json()["pdf_available"]:
        pytest.skip("LibreOffice not installed")
    _, job = generate(client)
    identifier = job["presentation_ids"][0]
    response = client.get(f"/api/v1/presentations/{identifier}/export/pdf")
    assert response.status_code == 200, response.text
    from pypdf import PdfReader

    assert len(PdfReader(io.BytesIO(response.content)).pages) == 3
    response = client.get(f"/api/v1/presentations/{identifier}/export/html")
    assert response.status_code == 200, response.text
    assert response.text.count('<section class="slide">') == 3
    assert "<svg" in response.text
    response = client.get(
        f"/api/v1/presentations/{identifier}/slides/0/preview?highlight=true"
    )
    assert response.status_code == 200, response.text
    assert "image/svg+xml" in response.headers["content-type"]
    assert response.headers["x-presentation-revision"] == "1"


def test_twelve_slide_target(client):
    from copy import deepcopy

    template = client.post(
        "/api/v1/templates", files={"file": ("unknown.pptx", template_bytes())}
    ).json()
    contents = outline()
    contents["slides"] = [deepcopy(contents["slides"][i % 3]) for i in range(12)]
    for i, slide in enumerate(contents["slides"]):
        slide["title"] = f"Уникальный вывод {i + 1}"
    response = client.post(
        "/api/v1/generations",
        json={
            "template_id": template["id"],
            "brief": "Источник: 10, 20. Тестовые данные.",
            "slide_count": 12,
            "outline": contents,
        },
    )
    job = wait_job(client, response.json()["id"])
    assert len(job["presentation_ids"]) == 3
    for identifier in job["presentation_ids"]:
        response = client.get(f"/api/v1/presentations/{identifier}/export/pptx")
        assert len(Presentation(io.BytesIO(response.content)).slides) == 12


def test_dark_template_color_and_title_geometry():
    from designer.parsing import best_text_color

    assert best_text_color("000000", ["FFFFFF", "000000", "0077FF"]) == "FFFFFF"
    template = parse_template(template_bytes(), "unknown.pptx")
    content = outline(1)
    content["slides"][0]["title"] = (
        "Длинный заголовок должен помещаться в рамку и не накладываться на основной текст"
    )
    from designer.models import Outline

    deck = compose(Outline.model_validate(content).model_dump(), template, "classic")
    title = deck["slides"][0]["elements"][0]
    body = deck["slides"][0]["elements"][1]
    assert title["box"][1] + title["box"][3] < body["box"][1]


def test_contextual_audit_inside_generation(client, monkeypatch):
    from designer import app as module

    monkeypatch.setenv("DESIGNER_LLM_BASE_URL", "https://inference.example/v1")
    monkeypatch.setenv("DESIGNER_LLM_MODEL", "test-open-model")
    called = []

    async def fake_audit(deck, sources):
        from designer.audit import issue

        called.append(deck)
        return [
            issue(
                0,
                "context_0",
                "title_mismatch",
                "Проверить заголовок",
                category="contextual",
            )
        ]

    monkeypatch.setattr(module, "contextual_audit", fake_audit)
    template = client.post(
        "/api/v1/templates", files={"file": ("template.pptx", template_bytes())}
    ).json()
    response = client.post(
        "/api/v1/generations",
        json={
            "template_id": template["id"],
            "brief": "Данные 10 и 20",
            "slide_count": 3,
            "outline": outline(),
            "contextual_audit": True,
        },
    )
    assert response.status_code == 202
    job = wait_job(client, response.json()["id"])
    assert len(called) == 1
    for identifier in job["presentation_ids"]:
        report = client.get(f"/api/v1/presentations/{identifier}/audit").json()
        assert report["contextual"]["status"] == "completed"
        assert any(i["category"] == "contextual" for i in report["issues"])


def audited(template, slide_mutator=None, content_mutator=None, count=1):
    """Compose a deck from an unknown template, break it on purpose, then audit."""
    from designer.models import Outline

    content = outline(count)
    if content_mutator:
        content_mutator(content)
    deck = compose(Outline.model_validate(content).model_dump(), template, "classic")
    if slide_mutator:
        slide_mutator(deck["slides"][0])
    report = audit(deck, template, [{"id": "brief", "text": "Первый тезис 10 20"}])
    return deck, report, {i["code"] for i in report["issues"]}


def test_clean_deck_on_supplied_templates_has_no_layout_findings():
    """The checks must stay quiet on decks the service itself composed."""
    noisy = {
        "out_of_bounds",
        "margin_encroachment",
        "misaligned",
        "branding_overlap",
        "low_contrast",
        "font_not_in_template",
        "font_size_off_scale",
        "color_not_in_palette",
        "font_variety",
        "layout_not_from_template",
        "text_overflow",
    }
    for path in sample_templates() or pytest.skip(f"No .pptx templates in {TEMPLATES}"):
        template = parse_template(path.read_bytes(), path.name)
        for variant in ("classic", "split", "focus"):
            from designer.models import Outline

            deck = compose(
                Outline.model_validate(outline(3)).model_dump(), template, variant
            )
            report = audit(deck, template, [{"id": "brief", "text": "Первый тезис"}])
            found = {i["code"] for i in report["issues"]} & noisy
            assert not found, (path.name, variant, found)


def test_layout_checks_fire_on_broken_geometry():
    template = parse_template(template_bytes(), "unknown.pptx")

    def off_slide(slide):
        slide["elements"][0]["box"][1] = -80

    assert "out_of_bounds" in audited(template, off_slide)[2]

    def into_margin(slide):
        slide["elements"][0]["box"][0] = 1.0

    codes = audited(template, into_margin)[2]
    assert "margin_encroachment" in codes and "out_of_bounds" not in codes

    def nudged(slide):
        slide["elements"][1]["box"][0] += 17

    assert "misaligned" in audited(template, nudged)[2]


def test_template_token_checks_fire():
    # A real template carries a type scale; the synthetic one has none, and the
    # check correctly stays silent when the template defines no sizes.
    branded = parse_template(
        require_template("VK Tech").read_bytes(), "vk.pptx"
    )
    assert branded["tokens"]["font_sizes"]

    def off_scale(slide):
        slide["elements"][0]["font_size"] = 37.5

    assert "font_size_off_scale" in audited(branded, off_scale)[2]
    assert "font_size_off_scale" not in audited(branded)[2]

    template = parse_template(template_bytes(), "unknown.pptx")

    def off_palette(slide):
        slide["elements"][0]["color"] = "123456"

    assert "color_not_in_palette" in audited(template, off_palette)[2]

    def unreadable(slide):
        slide["background"], slide["elements"][0]["color"] = "FFFFFF", "FEFEFE"

    assert "low_contrast" in audited(template, unreadable)[2]

    def third_font(slide):
        slide["elements"][0]["font"] = "Comic Sans MS"
        slide["elements"][1]["font"] = "Courier New"

    # Two replaced fonts plus the base font on the other slides make three families.
    assert "font_variety" in audited(template, third_font, count=3)[2]

    def foreign_layout(slide):
        slide["layout_index"] = 99

    assert "layout_not_from_template" in audited(template, foreign_layout)[2]


def test_density_and_chart_checks_fire():
    template = parse_template(template_bytes(), "unknown.pptx")

    def wide_table(content):
        content["slides"][0]["visual"] = {
            "kind": "table",
            "columns": [f"К{i}" for i in range(6)],
            "rows": [[f"з{i}" for i in range(6)] for _ in range(8)],
        }

    assert "table_size" in audited(template, content_mutator=wide_table)[2]

    def many_series(content):
        content["slides"][0]["visual"] = {
            "kind": "bar",
            "categories": ["А", "Б"],
            "series": [{"name": f"С{i}", "values": [1, 2]} for i in range(6)],
            "unit": "шт.",
        }

    assert "chart_series" in audited(template, content_mutator=many_series)[2]

    def unlabelled(content):
        content["slides"][0]["visual"] = {
            "kind": "bar",
            "categories": ["А", "Б"],
            "series": [{"name": "С", "values": [1, 2]}],
        }

    assert "chart_labels" in audited(template, content_mutator=unlabelled)[2]

    def sparse(content):
        content["slides"][1]["bullets"] = ["Мало"]
        content["slides"][1]["visual"] = {"kind": "none"}

    # Первый слайд — обложка в рамках шаблона, её заливку задаёт дизайнер;
    # разреженность ловится на обычном слайде.
    _, report, _ = audited(template, content_mutator=sparse, count=2)
    assert any(
        i["code"] == "fill_ratio" and i["slide_index"] == 1 for i in report["issues"]
    )


def test_new_fixes_repair_the_deck():
    from designer.audit import apply_fixes

    template = parse_template(template_bytes(), "unknown.pptx")

    def broken(slide):
        slide["elements"][0]["box"][0] = 1.0
        slide["elements"][0]["font_size"] = 37.5
        slide["elements"][0]["color"] = "123456"
        slide["elements"][1]["box"][0] += 17

    deck, report, codes = audited(template, broken)
    repairable = ("margin_encroachment", "font_size_off_scale", "color_not_in_palette")
    selected = [i["id"] for i in report["issues"] if i["code"] in repairable]
    assert {"margin_encroachment", "color_not_in_palette", "misaligned"} <= codes
    fixed = apply_fixes(deck, report, selected, template)
    element = fixed["slides"][0]["elements"][0]
    assert element["box"][0] > 1.0
    assert element["color"] in template["tokens"]["colors"]
    after = audit(fixed, template, [{"id": "brief", "text": "x"}])
    assert not [
        i
        for i in after["issues"]
        if i["code"] == "margin_encroachment" and i["element_id"] == "title"
    ]
    # An unselected finding must survive: the user chooses what gets changed.
    assert [i for i in after["issues"] if i["code"] == "misaligned"]


def test_export_verification_proves_native_objects(client):
    _, job = generate(client)
    record = client.app.state.store.get("presentations", job["presentation_ids"][0])
    check = record["export_check"]
    assert check["opens"] and check["slides"] == 3
    assert check["native_objects"] >= 3 and check["raster_slides"] == []


def test_base_path_makes_urls_proxy_safe(tmp_path, monkeypatch):
    """Behind a reverse proxy sub-path, returned URLs must carry the prefix."""
    monkeypatch.setenv("DESIGNER_BASE_PATH", "/presentations")
    monkeypatch.delenv("DESIGNER_API_KEY", raising=False)
    with TestClient(create_app(tmp_path)) as client:
        template = client.post(
            "/api/v1/templates", files={"file": ("t.pptx", template_bytes())}
        ).json()
        accepted = client.post(
            "/api/v1/generations",
            json={
                "template_id": template["id"],
                "brief": "Первый тезис. Второй тезис. Продажи А 10 Б 20 млн руб.",
                "slide_count": 3,
                "outline": outline(),
            },
        )
        assert accepted.status_code == 202, accepted.text
        # status_url is only returned by the 202 response, not by the job resource.
        assert accepted.json()["status_url"].startswith("/presentations/api/v1/jobs/")
        job = wait_job(client, accepted.json()["id"])
        record = client.get(
            "/api/v1/presentations/" + job["presentation_ids"][0]
        ).json()
        for url in record["exports"].values():
            assert url.startswith("/presentations/api/v1/presentations/")
        assert client.get("/openapi.json").json()["servers"] == [
            {"url": "/presentations"}
        ]


def mock_provider(monkeypatch, responses):
    """Serve scripted model answers; returns the list of captured requests."""
    import httpx

    from designer import generation

    monkeypatch.setenv("DESIGNER_LLM_BASE_URL", "https://inference.example/v1")
    monkeypatch.setenv("DESIGNER_LLM_MODEL", "test-open-model")
    monkeypatch.setenv("DESIGNER_LLM_RETRY_DELAY", "0")
    monkeypatch.setattr(generation, "RETRY_DELAY", 0.0)
    captured = []
    real_client = httpx.AsyncClient

    def handler(request):
        payload = json.loads(request.content)
        captured.append(payload)
        answer = responses[min(len(captured) - 1, len(responses) - 1)]
        return httpx.Response(200, json={"choices": [{"message": {"content": answer}}]})

    monkeypatch.setattr(
        generation.httpx,
        "AsyncClient",
        lambda **kwargs: real_client(transport=httpx.MockTransport(handler), **kwargs),
    )
    return captured


def test_retry_recovers_from_unusable_model_answers(client, monkeypatch):
    """A weaker model gets its answer back with the reason and a second chance."""
    captured = mock_provider(
        monkeypatch,
        [
            "сейчас соберу структуру",  # no JSON at all
            json.dumps(outline(2)),  # valid JSON, wrong slide count
            "```json\n" + json.dumps(outline(3)) + "\n```",  # fenced, correct
        ],
    )
    response = client.post(
        "/api/v1/outlines", json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 3}
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["slides"]) == 3
    assert len(captured) == 3
    # The repair turn carries the failed answer and the reason, not just the prompt.
    roles = [m["role"] for m in captured[-1]["messages"]]
    assert roles == ["system", "user", "assistant", "user"]
    assert "не подошёл" in captured[-1]["messages"][-1]["content"]


def test_retry_gives_up_with_a_readable_reason(client, monkeypatch):
    captured = mock_provider(monkeypatch, ["not json"])
    response = client.post("/api/v1/outlines", json={"brief": "Данные для проверки: 1, 2, 3"})
    assert response.status_code == 502
    assert "Модель не вернула корректный ответ" in response.json()["detail"]
    assert len(captured) == 3


def test_unknown_source_refs_are_retried(client, monkeypatch):
    bad = outline(1)
    bad["slides"][0]["source_refs"] = ["не-существует"]
    captured = mock_provider(monkeypatch, [json.dumps(bad), json.dumps(outline(1))])
    response = client.post(
        "/api/v1/outlines", json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 1}
    )
    assert response.status_code == 200, response.text
    assert len(captured) == 2
    assert "source_refs" in captured[-1]["messages"][-1]["content"]


def synthetic_template(patterns):
    """Minimal template dict with the fields the layout scorer reads."""
    return {
        "name": "synthetic",
        "width": 9144000,
        "height": 5143500,
        "tokens": {
            "fonts": ["Arial"],
            "font_sizes": [18, 24, 32],
            "colors": ["1C1D22", "FFFFFF"],
            "theme": {"lt1": "FFFFFF", "accent1": "FF0053"},
        },
        "geometry": {
            "safe_area": {"x": 0.06, "y": 0.06, "w": 0.88, "h": 0.8},
            "margins": {"x": 0.04, "y": 0.04, "w": 0.92, "h": 0.9},
            "guides": [],
        },
        "layouts": [{"index": 0, "name": "Blank"}],
        "patterns": patterns,
    }


def pattern(index, decoration=0.0, reserved=(), text_slots=3):
    return {
        "index": index,
        "layout_index": 0,
        "layout_name": "Blank",
        "background": "FFFFFF",
        "decoration_area": decoration,
        "reserved": [dict(zip(("x", "y", "w", "h"), box)) for box in reserved],
        "title_box": {"x": 0.06, "y": 0.08, "w": 0.8, "h": 0.14},
        "shapes": [],
        "text_slots": text_slots,
    }


def test_layout_picks_the_roomiest_pattern_for_a_chart():
    from designer.models import Outline

    busy = pattern(0, decoration=0.22, reserved=[(0.5, 0.2, 0.45, 0.6)])
    roomy = pattern(1, decoration=0.01)
    content = outline(1)  # first slide carries a bar chart
    deck = compose(
        Outline.model_validate(content).model_dump(),
        synthetic_template([busy, roomy]),
        "classic",
    )
    assert deck["slides"][0]["pattern_index"] == 1
    assert deck["slides"][0]["pattern_choice"]["free_area"] > 0.9


def test_layout_does_not_repeat_one_pattern_across_the_deck():
    from designer.models import Outline

    template = synthetic_template([pattern(0), pattern(1), pattern(2)])
    deck = compose(
        Outline.model_validate(outline(3)).model_dump(), template, "classic"
    )
    used = [s["pattern_index"] for s in deck["slides"]]
    assert len(set(used)) == 3, used
    # Neighbouring slides never share a pattern.
    assert all(a != b for a, b in zip(used, used[1:]))


def test_layout_choice_is_deterministic():
    from designer.models import Outline

    template = synthetic_template([pattern(0), pattern(1), pattern(2)])
    plan = Outline.model_validate(outline(3)).model_dump()
    first = [s["pattern_index"] for s in compose(plan, template, "classic")["slides"]]
    second = [s["pattern_index"] for s in compose(plan, template, "classic")["slides"]]
    assert first == second


def test_layout_avoids_patterns_whose_branding_sits_in_the_content_band():
    """A corner logo is fine; a panel across the content area is not."""
    from designer.layout import branding_in_band
    from designer.models import Outline

    corner_logo = pattern(0, reserved=[(0.86, 0.02, 0.12, 0.08)])
    mid_panel = pattern(1, reserved=[(0.1, 0.3, 0.8, 0.4)])
    assert branding_in_band(corner_logo) == 0
    assert branding_in_band(mid_panel) > 0.3
    deck = compose(
        Outline.model_validate(outline(1)).model_dump(),
        synthetic_template([mid_panel, corner_logo]),
        "classic",
    )
    assert deck["slides"][0]["pattern_index"] == 0


def test_long_content_pack_is_cut_to_the_model_window():
    """A 32k-token model cannot swallow a whole content pack; keep what matters."""
    from designer.generation import chunk_text, select_context

    noise = ("Посторонний текст про парковку и столовую. " * 20 + "\n\n") * 30
    fact = "Конверсия выросла на 18 процентов после запуска пилота.\n\n"
    text = noise + fact + noise
    assert len(chunk_text(text)) > 10
    selected = select_context(
        [{"id": "pack", "text": text}], "Как выросла конверсия после пилота", budget=6000
    )
    kept = selected[0]["text"]
    assert len(kept) <= 6000
    assert "Конверсия выросла на 18" in kept
    assert selected[0]["id"] == "pack"


def test_short_sources_are_passed_through_untouched():
    from designer.generation import select_context

    sources = [{"id": "brief", "text": "Короткий бриф"}, {"id": "p1", "text": "Материал"}]
    assert select_context(sources, "бриф", budget=6000) == sources


def test_every_source_survives_selection_so_refs_stay_valid():
    """Dropping a document entirely would make its source_refs unresolvable."""
    from designer.generation import select_context

    sources = [
        {"id": "brief", "text": "Конверсия и пилот"},
        {"id": "pack", "text": "Совершенно другая тема. " * 2000},
    ]
    selected = select_context(sources, "конверсия пилот", budget=3000)
    assert [s["id"] for s in selected] == ["brief", "pack"]
    assert all(s["text"] for s in selected)


def test_foreign_chart_labels_are_reported():
    """ТЗ требует одну языковую колоду; подписи осей — место, где модель срывается."""
    from designer.audit import audit, foreign_labels
    from designer.models import Outline

    assert foreign_labels(["Before launch"], True) == ["Before launch"]
    assert foreign_labels(["Пилот", "Q1", "KPI", "%"], True) == []
    assert foreign_labels(["Before launch"], False) == []

    plan = outline(1)
    plan["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Before launch", "After launch"],
        "series": [{"name": "Teams", "values": [10.0, 20.0]}],
        "unit": "команд",
    }
    template = synthetic_template([pattern(0), pattern(1)])
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    report = audit(deck, template, [{"id": "brief", "text": "Пилот 10 20 команд"}])
    found = [i for i in report["issues"] if i["code"] == "foreign_language"]
    assert found, [i["code"] for i in report["issues"]]
    assert "Before launch" in found[0]["message"]


def test_english_deck_keeps_english_labels():
    from designer.audit import audit
    from designer.models import Outline

    plan = outline(1)
    plan["title"] = "Pilot results"
    plan["slides"][0]["title"] = "Teams doubled after launch"
    plan["slides"][0]["bullets"] = ["Pilot covered 10 teams", "Launch reached 20 teams"]
    plan["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Before", "After"],
        "series": [{"name": "Teams", "values": [10.0, 20.0]}],
        "unit": "teams",
    }
    template = synthetic_template([pattern(0), pattern(1)])
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    report = audit(deck, template, [{"id": "brief", "text": "Pilot 10 20 teams"}])
    assert not [i for i in report["issues"] if i["code"] == "foreign_language"]


def test_split_variant_keeps_short_slides_in_one_column():
    """Две колонки по одной строке читаются как пустой слайд."""
    from designer.audit import audit
    from designer.models import Outline

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = ["Единственный тезис о результате пилота"]
    template = synthetic_template([pattern(0), pattern(1)])
    deck = compose(Outline.model_validate(plan).model_dump(), template, "split")
    bodies = [e for e in deck["slides"][0]["elements"] if e.get("role") == "body"]
    assert len(bodies) == 1, [e["id"] for e in bodies]
    # Текст занимает всю ширину рабочей области, а не половину.
    title = next(e for e in deck["slides"][0]["elements"] if e["role"] == "title")
    assert bodies[0]["box"][2] > title["box"][2] * 0.9
    audit(deck, template, [{"id": "brief", "text": "Пилот"}])


def test_split_variant_still_uses_two_columns_for_dense_slides():
    from designer.models import Outline

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = [f"Тезис номер {i} о результате" for i in range(5)]
    deck = compose(
        Outline.model_validate(plan).model_dump(),
        synthetic_template([pattern(0), pattern(1)]),
        "split",
    )
    ids = {e["id"] for e in deck["slides"][0]["elements"]}
    assert {"body_left", "body_right"} <= ids, ids


def test_layout_shrinks_text_before_the_audit_sees_it():
    """Вёрстка сама подбирает кегль из шкалы, а не оставляет это пользователю."""
    from designer.audit import audit
    from designer.models import Outline

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = [
        "Очень длинный тезис о результатах пилота, который занимает много места",
        "Второй столь же длинный тезис о работе команд и планах развития сервиса",
        "Третий длинный тезис про экономию времени дизайнера на каждой колоде",
        "Четвёртый длинный тезис про масштабирование сервиса на всю компанию",
    ]
    template = synthetic_template([pattern(0), pattern(1)])
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    body = next(e for e in deck["slides"][0]["elements"] if e.get("role") == "body")
    assert body["font_size"] in template["tokens"]["font_sizes"]
    report = audit(deck, template, [{"id": "brief", "text": "Пилот"}])
    assert not [i for i in report["issues"] if i["code"] == "text_overflow"]


def test_invented_numbers_send_the_model_back_for_another_try(client, monkeypatch):
    """ТЗ: все цифры со слайда есть в материалах. Выдуманное число — повод переспросить."""
    derived = outline(1)
    derived["slides"][0]["title"] = "Участие выросло на 100%"
    good = outline(1)
    good["slides"][0]["title"] = "Команд стало 20 вместо 10"
    captured = mock_provider(monkeypatch, [json.dumps(derived), json.dumps(good)])
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Пилот: 10 команд, после запуска 20 команд", "slide_count": 1},
    )
    assert response.status_code == 200, response.text
    assert response.json()["slides"][0]["title"] == "Команд стало 20 вместо 10"
    assert len(captured) == 2
    assert "100" in captured[-1]["messages"][-1]["content"]


def test_focus_variant_fills_the_slide_with_larger_type():
    """Один тезис на слайд — это крупный кегль, а не четверть пустой страницы."""
    from designer.layout import ink_area
    from designer.models import Outline

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = ["Пилот охватил десять команд"]
    template = synthetic_template([pattern(0), pattern(1)])
    data = Outline.model_validate(plan).model_dump()
    classic = compose(data, template, "classic")
    focus = compose(data, template, "focus")
    body_size = lambda deck: next(
        e["font_size"] for e in deck["slides"][0]["elements"] if e.get("role") == "body"
    )
    # «Фокус» отвечает на пустой слайд крупным кеглем, а не наполнителем.
    assert body_size(focus) >= body_size(classic)
    area = focus["width"] * focus["height"]
    assert sum(ink_area(e) for e in focus["slides"][0]["elements"]) / area > 0.1
    for element in focus["slides"][0]["elements"]:
        if element["kind"] == "text":
            assert element["font_size"] in template["tokens"]["font_sizes"]


def test_english_labels_send_the_model_back(client, monkeypatch):
    """Русская колода с английскими подписями осей не доходит до вёрстки."""
    english = outline(1)
    english["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Before launch", "After launch"],
        "series": [{"name": "Teams", "values": [10.0, 20.0]}],
        "unit": "teams",
    }
    russian = outline(1)
    russian["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Пилот", "Запуск"],
        "series": [{"name": "Команды", "values": [10.0, 20.0]}],
        "unit": "команд",
    }
    captured = mock_provider(monkeypatch, [json.dumps(english), json.dumps(russian)])
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Пилот: 10 команд, запуск 20 команд. Разделы 1, 2, 3.", "slide_count": 1},
    )
    assert response.status_code == 200, response.text
    assert response.json()["slides"][0]["visual"]["categories"] == ["Пилот", "Запуск"]
    assert len(captured) == 2
    assert "Before launch" in captured[-1]["messages"][-1]["content"]


def test_every_diagram_kind_exports_as_native_shapes(client, tmp_path):
    """Схемы и пиктограммы должны быть фигурами PowerPoint, а не картинкой."""
    from pptx import Presentation as Deck

    from designer.exporting import export_pptx, verify_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    kinds = {
        "icon": {"kind": "icon", "steps": ["Рост выручки", "Команда", "Сроки"]},
        "process": {"kind": "process", "steps": ["Анализ", "Пилот", "Запуск"]},
        "cycle": {"kind": "cycle", "steps": ["План", "Работа", "Оценка"]},
        "pyramid": {"kind": "pyramid", "steps": ["Цель", "Метрики", "Данные"]},
        "timeline": {"kind": "timeline", "steps": ["Q1", "Q2", "Q3"]},
        "comparison": {
            "kind": "comparison",
            "columns": ["Было", "Стало"],
            "rows": [["Ручная вёрстка", "Автоматическая"], ["4 часа", "20 секунд"]],
        },
    }
    plan = {
        "title": "Все виды схем",
        "slides": [
            {
                "title": f"Схема {name}",
                "bullets": ["Проверка нативности"],
                "notes": "",
                "source_refs": ["brief"],
                "visual": visual,
            }
            for name, visual in kinds.items()
        ],
    }
    data = Outline.model_validate(plan).model_dump()
    source = template_bytes()
    template = parse_template(source, "unknown.pptx")
    deck = compose(data, template, "classic")
    source_path = tmp_path / "template.pptx"
    source_path.write_bytes(source)
    output = tmp_path / "deck.pptx"
    export_pptx(source_path, template, deck, output)

    report = verify_pptx(output, len(kinds))
    assert report["raster_slides"] == []
    assert report["native_objects"] >= len(kinds) * 3
    # Подписи шагов доезжают до файла как редактируемый текст.
    text = " ".join(
        shape.text
        for slide in Deck(output).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    for word in ("Анализ", "План", "Метрики", "Было", "Стало", "Q1"):
        assert word in text, word


def test_visual_audit_reads_the_slide_image(monkeypatch):
    """Контекстные проверки Приложения 1 идут по картинке, а не по тексту."""
    import asyncio

    from designer import audit as audit_module

    seen = {}

    async def fake_vision(prompt, payload, image):
        seen["prompt"] = prompt
        seen["payload"] = payload
        seen["image"] = image
        return {
            "issues": [
                {"code": "title_conclusion", "message": "Заголовок называет тему"},
                {"code": "readable", "message": "Текст наезжает на логотип"},
                "мусор, который модель прислала не по схеме",
            ]
        }

    monkeypatch.setattr(audit_module, "vision_completion", fake_vision)
    deck = {
        "slides": [
            {
                "content": {
                    "title": "Конверсия",
                    "bullets": ["Тезис"],
                    "visual": {"kind": "none"},
                }
            }
        ]
    }
    findings = asyncio.run(
        audit_module.visual_audit(deck, [{"id": "brief", "text": "Источник"}], [b"PNG"])
    )
    assert [f["code"] for f in findings] == ["title_conclusion", "readable"]
    assert all(f["category"] == "contextual" for f in findings)
    assert all(f["element_id"] == "slide_image" for f in findings)
    assert seen["image"] == b"PNG"
    assert seen["payload"]["title"] == "Конверсия"
    assert "картинк" in seen["prompt"].lower() or "изображени" in seen["prompt"].lower()


def test_visual_audit_survives_a_broken_model_answer(monkeypatch):
    """Сбой проверки по картинке не должен ронять весь аудит."""
    import asyncio

    from designer import audit as audit_module

    async def broken(prompt, payload, image):
        raise audit_module.InvalidCompletion("модель вернула мусор")

    monkeypatch.setattr(audit_module, "vision_completion", broken)
    deck = {"slides": [{"content": {"title": "Т", "bullets": [], "visual": {"kind": "none"}}}]}
    findings = asyncio.run(audit_module.visual_audit(deck, [{"id": "brief", "text": "и"}], [b"PNG"]))
    assert findings == []


def test_reasoning_models_answer_is_recovered():
    """Модель в режиме рассуждений кладёт ответ мимо content — читаем всё равно."""
    from designer.generation import InvalidCompletion, message_text

    assert message_text({"content": " {\"a\":1} "}) == '{"a":1}'
    assert message_text({"content": None, "reasoning_content": "{}"}) == "{}"
    assert message_text({"content": [{"type": "text", "text": "{}"}]}) == "{}"
    try:
        message_text({"content": None})
    except InvalidCompletion:
        pass
    else:
        raise AssertionError("пустой ответ должен считаться непригодным")


def test_layout_prefers_slides_without_template_clutter():
    """Пустые заглушки макета переезжают на готовый слайд, поэтому их избегаем."""
    from designer.layout import candidate_patterns

    clutter = [dict(pattern(i), decoration_count=6, decoration_area=0.2) for i in range(3)]
    clean = [dict(pattern(i + 3), decoration_count=0) for i in range(3)]
    chosen = candidate_patterns(clutter + clean)
    assert {p["index"] for p in chosen} == {3, 4, 5}
    # Если чистых страниц мало, берём что есть — пустая колода хуже украшений.
    only_two_clean = clutter + clean[:2]
    assert len(candidate_patterns(only_two_clean)) == 5


def test_templates_are_reparsed_after_a_parser_upgrade(tmp_path, monkeypatch):
    """Шаблон в хранилище не должен остаться с устаревшим разбором."""
    from designer.app import create_app, refresh_templates
    from designer.parsing import PARSER_VERSION

    with TestClient(create_app(data_dir=tmp_path)) as client:
        created = client.post(
            "/api/v1/templates",
            files={"file": ("t.pptx", template_bytes())},
        )
        assert created.status_code == 201, created.text
        template_id = created.json()["id"]
        store = client.app.state.store
        stale = store.get("templates", template_id)
        stale["parser"] = 0
        stale["geometry"] = {}
        store.put("templates", stale)

        assert refresh_templates(store) == 1
        fresh = store.get("templates", template_id)
        assert fresh["parser"] == PARSER_VERSION
        assert fresh["geometry"]["safe_area"]
        assert fresh["id"] == template_id
        # Повторный вызов ничего не делает: версия уже актуальна.
        assert refresh_templates(store) == 0


def test_sample_artwork_is_not_copied_into_the_result(tmp_path):
    """Фигуры-образцы с прототипа не должны переезжать на готовый слайд."""
    import io

    from pptx import Presentation as Deck
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Emu

    from designer.exporting import export_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    source = Deck()
    for index in range(3):
        slide = source.slides.add_slide(source.slide_layouts[1])
        slide.shapes.title.text = "Образец заголовка"
        slide.placeholders[1].text = "Образец текста"
        # Серый кружок-заглушка под фото: на каждом слайде в своём месте.
        stub = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Emu(900000 + index * 400000),
            Emu(3000000),
            Emu(700000),
            Emu(700000),
        )
        stub.fill.solid()
        stub.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    # Цветная плашка — фирменный декор, он должен остаться на результате.
    for slide in source.slides:
        plate = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(100000), Emu(100000), Emu(400000), Emu(200000)
        )
        plate.fill.solid()
        plate.fill.fore_color.rgb = RGBColor(0x00, 0x77, 0xFF)
    buffer = io.BytesIO()
    source.save(buffer)
    data = buffer.getvalue()

    template = parse_template(data, "sample.pptx")

    plan = {
        "title": "Проверка",
        "slides": [
            {
                "title": "Вывод",
                "bullets": ["Тезис"],
                "notes": "",
                "source_refs": ["brief"],
                "visual": {"kind": "none"},
            }
        ],
    }
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    source_path = tmp_path / "template.pptx"
    source_path.write_bytes(data)
    output = tmp_path / "out.pptx"
    export_pptx(source_path, template, deck, output)

    shapes = list(Deck(output).slides[0].shapes)
    ovals = [s for s in shapes if s.shape_type == 1 and s.width == Emu(700000)]
    plates = [s for s in shapes if s.shape_type == 1 and s.width == Emu(400000)]
    assert not ovals, "серая заглушка под фото попала в результат"
    assert plates, "фирменная плашка шаблона потерялась"


def test_photo_backgrounds_are_avoided_and_flagged():
    """Текст на фотографии не читается: такой прототип избегаем, а если взяли — помечаем."""
    from designer.audit import audit
    from designer.layout import compose
    from designer.models import Outline

    photo = dict(pattern(0), image_cover=0.95)
    plain = dict(pattern(1), image_cover=0.0)
    plan = Outline.model_validate(outline(1)).model_dump()

    deck = compose(plan, synthetic_template([photo, plain]), "classic")
    assert deck["slides"][0]["pattern_index"] == 1

    only_photo = synthetic_template([photo])
    forced = compose(plan, only_photo, "classic")
    assert forced["slides"][0]["image_cover"] == 0.95
    report = audit(forced, only_photo, [{"id": "brief", "text": "10 20 А Б Продажи"}])
    assert [i for i in report["issues"] if i["code"] == "text_over_image"]


def test_measured_background_decides_text_colour():
    """Тёмный фон меняет цвет текста, пёстрый — включает подложку."""
    from designer.layout import compose
    from designer.models import Outline

    plan = Outline.model_validate(outline(1)).model_dump()

    dark = synthetic_template([dict(pattern(0), bg_luma=0.05, bg_spread=0.01)])
    dark["tokens"]["colors"] = ["1C1D22", "FFFFFF"]
    light_text = compose(plan, dark, "classic")["slides"][0]["elements"][0]["color"]
    assert light_text == "FFFFFF"

    light = synthetic_template([dict(pattern(1), bg_luma=0.95, bg_spread=0.01)])
    light["tokens"]["colors"] = ["1C1D22", "FFFFFF"]
    dark_text = compose(plan, light, "classic")["slides"][0]["elements"][0]["color"]
    assert dark_text == "1C1D22"

    busy = synthetic_template([dict(pattern(2), bg_luma=0.4, bg_spread=0.35)])
    slide = compose(plan, busy, "classic")["slides"][0]
    assert slide["needs_scrim"] is True


def test_busy_backgrounds_lose_to_calm_ones():
    from designer.layout import compose
    from designer.models import Outline

    busy = dict(pattern(0), bg_luma=0.4, bg_spread=0.4)
    calm = dict(pattern(1), bg_luma=0.95, bg_spread=0.02)
    deck = compose(
        Outline.model_validate(outline(1)).model_dump(),
        synthetic_template([busy, calm]),
        "classic",
    )
    assert deck["slides"][0]["pattern_index"] == 1


def test_layout_follows_the_template_typography():
    """Кегль, шрифт, начертание и выравнивание берутся у слайда-прототипа."""
    from designer.layout import compose
    from designer.models import Outline

    styled = dict(
        pattern(0),
        slots=[
            {
                "role": "title",
                "box": {"x": 0.08, "y": 0.08, "w": 0.8, "h": 0.16},
                "style": {
                    "font": "Play",
                    "size": 40.0,
                    "bold": True,
                    "color": "FF0053",
                    "align": "ctr",
                },
                "length": 20,
            },
            {
                "role": "body",
                "box": {"x": 0.08, "y": 0.34, "w": 0.8, "h": 0.5},
                "style": {
                    "font": "Play",
                    "size": 20.0,
                    "bold": False,
                    "color": "1C1D22",
                    "align": "l",
                },
                "length": 80,
            },
        ],
    )
    template = synthetic_template([styled, pattern(1)])
    template["tokens"]["font_sizes"] = [14, 20, 28, 40]
    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    title = deck["slides"][0]["elements"][0]
    body = next(e for e in deck["slides"][0]["elements"] if e.get("role") == "body")
    assert title["font"] == "Play"
    assert title["font_size"] == 40.0
    assert title["bold"] is True
    assert title["align"] == "center"
    # Кегль тела остаётся в пределах шаблонного: расти он может, но немного и
    # только по шкале самого шаблона.
    assert 20.0 <= body["font_size"] <= 44.0
    assert body["font_size"] in template["tokens"]["font_sizes"]
    assert body["align"] == "left"
    # Контент начинается там, где его держит прототип.
    assert body["box"][1] >= deck["height"] * 0.3


def test_first_slide_uses_a_cover_page_of_the_template():
    """У шаблона для обложки свои страницы: крупный заголовок, мало рамок."""
    from designer.layout import compose
    from designer.models import Outline

    cover = dict(pattern(0, text_slots=2))
    cover["title_box"] = {"x": 0.1, "y": 0.35, "w": 0.8, "h": 0.22}
    content = dict(pattern(1, text_slots=5))
    plan = outline(3)
    # Титульный слайд — название и одна строка под ним, без списка тезисов.
    plan["slides"][0]["bullets"] = ["Отдел дизайна, сентябрь 2026"]
    plan["slides"][0]["visual"] = {"kind": "none"}
    deck = compose(
        Outline.model_validate(plan).model_dump(),
        synthetic_template([cover, content]),
        "classic",
    )
    assert deck["slides"][0]["pattern_index"] == 0
    # Дальше идут обычные контентные страницы.
    assert deck["slides"][1]["pattern_index"] == 1


def test_sparse_content_is_pulled_to_the_optical_centre():
    from designer.layout import compose
    from designer.models import Outline

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = ["Один короткий тезис"]
    template = synthetic_template([pattern(0), pattern(1)])
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    title, body = deck["slides"][0]["elements"][0], deck["slides"][0]["elements"][1]
    gap = body["box"][1] - (title["box"][1] + title["box"][3])
    assert gap > 0
    # Блок опущен ниже верхней кромки области, но не улетел за её пределы.
    assert body["box"][1] + body["box"][3] <= deck["height"]


def test_template_roles_drive_slide_placement():
    """Обложка шаблона достаётся первому слайду, «спасибо» — последнему."""
    from designer.layout import compose
    from designer.models import Outline

    cover = dict(pattern(0, text_slots=2), role="cover")
    content = dict(pattern(1, text_slots=4), role="content")
    closing = dict(pattern(2, text_slots=2), role="closing")
    plan = outline(3)
    plan["slides"][0]["bullets"] = ["Отдел дизайна, сентябрь 2026"]
    plan["slides"][0]["visual"] = {"kind": "none"}
    deck = compose(
        Outline.model_validate(plan).model_dump(),
        synthetic_template([cover, content, closing]),
        "classic",
    )
    used = [s["pattern_index"] for s in deck["slides"]]
    assert used[0] == 0, used
    assert used[-1] == 2, used


def test_poor_template_still_produces_a_readable_deck(tmp_path):
    """Шаблон может быть плохим: без стилей, тем и образцов. Колода всё равно нужна."""
    import io

    from pptx import Presentation as Deck

    from designer.audit import audit
    from designer.exporting import export_pptx, verify_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    poor = Deck()
    slide = poor.slides.add_slide(poor.slide_layouts[6])  # пустой макет
    slide.shapes.add_textbox(0, 0, 100, 50).text_frame.text = "Текст"
    buffer = io.BytesIO()
    poor.save(buffer)
    data = buffer.getvalue()

    template = parse_template(data, "poor.pptx")
    plan = outline(3)
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")

    width, height = deck["width"], deck["height"]
    for slide_data in deck["slides"]:
        assert slide_data["elements"], "слайд без содержимого"
        for element in slide_data["elements"]:
            x, y, w, h = element["box"]
            assert x >= 0 and y >= 0
            assert x + w <= width + 1 and y + h <= height + 1
            if element["kind"] == "text":
                assert element["font_size"] >= 12

    report = audit(deck, template, [{"id": "brief", "text": "10 20 А Б Продажи Этап Срок"}])
    assert report["counts"]["errors"] == 0, [
        i["code"] for i in report["issues"] if i["severity"] == "error"
    ]

    source = tmp_path / "poor.pptx"
    source.write_bytes(data)
    output = tmp_path / "out.pptx"
    export_pptx(source, template, deck, output)
    assert verify_pptx(output, len(deck["slides"]))["raster_slides"] == []


def test_bullets_fill_the_card_grid_of_the_template():
    """Если у прототипа ряд одинаковых блоков — тезисы ложатся в них."""
    from designer.layout import compose
    from designer.models import Outline

    cards = [
        {
            "role": "body",
            "box": {"x": 0.06 + i * 0.31, "y": 0.4, "w": 0.26, "h": 0.3},
            "style": {"font": "Play", "size": 16.0, "align": "l"},
            "length": 40,
        }
        for i in range(3)
    ]
    grid = dict(
        pattern(0, text_slots=4),
        role="content",
        slots=[
            {
                "role": "title",
                "box": {"x": 0.06, "y": 0.08, "w": 0.8, "h": 0.14},
                "style": {"size": 32.0, "align": "l"},
                "length": 20,
            },
            *cards,
        ],
    )
    plan = outline(2)
    # Первый слайд — обложка, карточки проверяем на втором, контентном.
    for slide in plan["slides"]:
        slide["visual"] = {"kind": "none"}
        slide["bullets"] = ["Первый", "Второй", "Третий"]
    deck = compose(
        Outline.model_validate(plan).model_dump(),
        synthetic_template([grid, dict(pattern(1, text_slots=2), role="cover")]),
        "classic",
    )
    content_slide = deck["slides"][1]
    ids = [e["id"] for e in content_slide["elements"]]
    assert ids == ["title", "card_0", "card_1", "card_2"], ids
    # Карточки стоят там же, где они в шаблоне.
    xs = [e["box"][0] for e in content_slide["elements"] if e["id"].startswith("card")]
    assert xs == sorted(xs) and len(set(xs)) == 3


def test_overloaded_slides_are_condensed_to_the_template_capacity(monkeypatch):
    """Текст подгоняется под место, которое даёт шаблон, до вёрстки."""
    import asyncio

    from designer.generation import balance_outline, outline_overflow
    from designer.layout import slide_capacity
    from designer.parsing import parse_template

    template = parse_template(template_bytes(), "unknown.pptx")
    capacity = slide_capacity(template)
    assert capacity > 200

    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = ["Очень длинный тезис про пилот сервиса. " * 30]
    assert outline_overflow(plan, capacity)

    async def fake_condense(agent, payload, validate=None):
        assert agent == "condense"
        assert payload["slides"][0]["limit_chars"] > 0
        return {"slides": [{"index": 0, "bullets": ["Пилот сервиса завершён"]}]}

    monkeypatch.setattr("designer.generation.completion", fake_condense)
    balanced = asyncio.run(balance_outline(plan, template))
    assert balanced["slides"][0]["bullets"] == ["Пилот сервиса завершён"]
    assert not outline_overflow(balanced, capacity)


def test_without_a_model_overflow_is_trimmed_deterministically(monkeypatch):
    import asyncio

    from designer.generation import balance_outline, outline_overflow
    from designer.layout import slide_capacity
    from designer.parsing import parse_template

    template = parse_template(template_bytes(), "unknown.pptx")
    plan = outline(1)
    plan["slides"][0]["visual"] = {"kind": "none"}
    plan["slides"][0]["bullets"] = ["Длинный тезис про внедрение сервиса. " * 20]

    async def unavailable(agent, payload, validate=None):
        raise HTTPException(503, "модель не настроена")

    monkeypatch.setattr("designer.generation.completion", unavailable)
    balanced = asyncio.run(balance_outline(plan, template))
    assert balanced["slides"][0]["bullets"], "текст не должен исчезать целиком"
    assert not outline_overflow(balanced, slide_capacity(template) * 1.2)


def test_stubborn_model_still_gets_a_deck(client, monkeypatch):
    """Модель настаивает на своём числе слайдов — лучше колода, чем ошибка."""
    short = outline(1)
    captured = mock_provider(monkeypatch, [json.dumps(short)])
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Пилот: 10 и 20 команд. Разделы 1, 2, 3.", "slide_count": 5},
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["slides"]) == 1
    # Сначала переспросили столько раз, сколько положено.
    assert len(captured) == 3


def test_narrative_frame_depends_on_the_kind_of_deck(monkeypatch):
    """У продукта и проекта разный состав слайдов; команда — не по умолчанию."""
    from designer.generation import narrative

    product = narrative("product")
    project = narrative("project")
    feature = narrative("feature")
    assert product and project and feature
    assert product["beats"] != project["beats"]
    assert feature["team_slide"] == "no"
    assert all("команд" not in beat.lower() for beat in product["beats"])


def test_planner_receives_the_narrative(client, monkeypatch):
    captured = mock_provider(monkeypatch, [json.dumps(outline(1))])
    response = client.post(
        "/api/v1/outlines",
        json={
            "brief": "Данные: 10 и 20. Разделы 1, 2, 3.",
            "slide_count": 1,
            "purpose": "feature",
        },
    )
    assert response.status_code == 200, response.text
    payload = json.loads(captured[0]["messages"][1]["content"])
    assert payload["narrative"]["label"] == "Фича"
    assert payload["narrative"]["beats"][0].startswith("Боль")


def test_repeated_titles_are_sent_back_to_the_model(client, monkeypatch):
    """Два слайда с одним заголовком — это один слайд, разрезанный пополам."""
    doubled = outline(2)
    doubled["slides"][1]["title"] = doubled["slides"][0]["title"]
    good = outline(2)
    captured = mock_provider(monkeypatch, [json.dumps(doubled), json.dumps(good)])
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 2},
    )
    assert response.status_code == 200, response.text
    titles = [s["title"] for s in response.json()["slides"]]
    assert len(set(titles)) == 2
    assert "повторяются" in captured[-1]["messages"][-1]["content"].lower()


def test_thousands_separator_counts_as_the_same_number():
    """«1 400 рублей» в материалах и «1400» на слайде — одно и то же число."""
    from designer.generation import known_numbers, unsupported_numbers

    known = known_numbers([{"text": "Стоимость пилота — 1 400 рублей за 68 колод."}])
    assert unsupported_numbers("1400 рублей, 68 колод", known) == []
    assert unsupported_numbers("2100 рублей", known) == ["2100"]


def test_single_digits_do_not_fail_generation():
    """«три команды» и «3 команды» — не повод заваливать всю колоду."""
    from designer.generation import known_numbers, unsupported_numbers

    known = known_numbers([{"text": "Пилот шёл на двух командах, собрано 68 колод."}])
    assert unsupported_numbers("3 команды", known) == []
    assert unsupported_numbers("186 минут", known) == ["186"]


def test_foreign_labels_drop_the_visual_on_the_last_attempt(client, monkeypatch):
    """Колода без одной диаграммы лучше, чем 502 вместо колоды."""
    stubborn = outline(2)
    stubborn["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Auto", "Manual"],
        "series": [{"name": "Speed", "values": [10, 20]}],
        "unit": "minutes",
    }
    answers = [json.dumps(stubborn)] * 3
    mock_provider(monkeypatch, answers)
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 2},
    )
    assert response.status_code == 200, response.text
    assert response.json()["slides"][0]["visual"]["kind"] == "none"


def test_bullets_are_tidied_before_layout():
    """Строчная буква и точка в конце тезиса — разнобой, видимый на слайде."""
    from designer.generation import tidy_line

    assert tidy_line(" принимает шаблон pptx.") == "Принимает шаблон pptx"
    assert tidy_line("Сроки  и   ресурсы") == "Сроки и ресурсы"
    assert tidy_line("Что дальше...") == "Что дальше..."


def test_export_drops_media_nobody_references(tmp_path):
    """Колода не должна таскать иллюстрации всех страниц шаблона."""
    import zipfile

    from designer.exporting import export_pptx, verify_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    for path in sample_templates() or pytest.skip("нет шаблонов"):
        template = parse_template(path.read_bytes(), path.name)
        deck = compose(Outline.model_validate(outline(3)).model_dump(), template, "classic")
        result = tmp_path / "deck.pptx"
        export_pptx(path, template, deck, result)
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(result) as made:
            before = len([n for n in source.namelist() if n.startswith("ppt/media/")])
            after = len([n for n in made.namelist() if n.startswith("ppt/media/")])
        assert after <= before
        assert verify_pptx(result, len(deck["slides"]))["opens"]


def test_dense_first_slide_still_uses_cover_without_losing_content():
    """P0-2: первый слайд всегда использует обложку, содержание сохраняется."""
    from designer.layout import compose
    from designer.models import Outline

    cover = dict(pattern(0, text_slots=2), role="cover")
    content = dict(pattern(1, text_slots=5), role="content")
    plan = outline(3)
    plan["slides"][0]["bullets"] = ["Раз", "Два", "Три", "Четыре"]
    deck = compose(
        Outline.model_validate(plan).model_dump(),
        synthetic_template([cover, content]),
        "classic",
    )
    assert deck["slides"][0]["pattern_index"] == 0
    assert deck["slides"][0]["content"]['bullets'] == ["Раз", "Два", "Три", "Четыре"]
    assert any(e['kind'] == 'bar' for e in deck['slides'][0]['elements'])


def test_parked_shapes_never_become_a_card_grid():
    """Блоки за краем страницы — заготовки дизайнера, а не карточки."""
    from designer.layout import card_slots

    parked = [
        {"role": "body", "box": {"x": 0.68, "y": 0.33, "w": 0.42, "h": 0.12}},
        {"role": "body", "box": {"x": 0.68, "y": 0.48, "w": 0.42, "h": 0.12}},
        {"role": "body", "box": {"x": 0.68, "y": 0.63, "w": 0.42, "h": 0.12}},
    ]
    assert card_slots(parked, 960, 540) == []


def test_invented_number_is_shipped_flagged_on_the_last_attempt(client, monkeypatch):
    """Упрямая модель не должна оставлять пользователя без колоды."""
    stubborn = outline(2)
    stubborn["slides"][0]["bullets"] = ["Рост составил 95 процентов"]
    mock_provider(monkeypatch, [json.dumps(stubborn)] * 3)
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 2},
    )
    assert response.status_code == 200, response.text
    assert "95" in response.json()["slides"][0]["bullets"][0]


def test_chart_carries_value_labels_and_a_readable_axis_title(tmp_path):
    """Приложение 1: у диаграммы есть подписи значений и единица по оси.

    Без подписей значений столбец «4» рядом со «186» — полоска у оси: прочитать
    его можно только по сетке. Заголовок оси PowerPoint по умолчанию кладёт на
    бок, и в узкой рамке слово рвётся по слогам, поэтому он должен быть
    горизонтальным.
    """
    from lxml import etree

    from designer.exporting import export_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    plan = {
        "title": "Диаграмма",
        "slides": [
            {
                "title": "Колода за 4 минуты вместо 186",
                "bullets": ["Замер на пилоте"],
                "notes": "",
                "source_refs": ["brief"],
                "visual": {
                    "kind": "bar",
                    "unit": "минут",
                    "categories": ["Вручную", "Сервис"],
                    "series": [{"name": "Сборка колоды", "values": [186, 4]}],
                },
            }
        ],
    }
    source = template_bytes()
    template = parse_template(source, "unknown.pptx")
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    source_path = tmp_path / "template.pptx"
    source_path.write_bytes(source)
    output = tmp_path / "deck.pptx"
    export_pptx(source_path, template, deck, output)

    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    }
    with zipfile.ZipFile(output) as archive:
        name = next(n for n in archive.namelist() if n.startswith("ppt/charts/chart"))
        chart = etree.fromstring(archive.read(name))
    labels = chart.find(".//c:plotArea//c:dLbls", namespaces)
    assert labels is not None, "подписей значений нет вовсе"
    assert labels.find("c:showVal", namespaces).get("val") in ("1", "true")
    # Единица стоит в самой подписи, а перенос запрещён: «186 минут» одной строкой.
    assert "минут" in labels.find("c:numFmt", namespaces).get("formatCode")
    assert labels.find("c:txPr/a:bodyPr", namespaces).get("wrap") == "none"
    # Ось значений подписана единицей и названием ряда — легенда при одном ряде
    # только повторяла бы его.
    title = chart.find(".//c:valAx/c:title", namespaces)
    assert title is not None
    text = "".join(title.itertext())
    assert "минут" in text and "Сборка колоды" in text
    assert title.find(".//a:bodyPr", namespaces).get("rot") == "0"
    assert chart.find("c:chart/c:legend", namespaces) is None


def test_a_squeezed_chart_is_reported_and_can_be_grown():
    """chart_too_small ловит диаграмму, сжатую до полоски, и умеет её починить."""
    template = parse_template(template_bytes(), "unknown.pptx")

    def squeeze(slide):
        visual = next(e for e in slide["elements"] if e["id"] == "visual")
        visual["box"][3] = 40.0

    def chart(content):
        content["slides"][0]["visual"] = {
            "kind": "bar",
            "unit": "минут",
            "categories": ["Вручную", "Сервис"],
            "series": [{"name": "Сборка", "values": [186, 4]}],
        }

    deck, report, codes = audited(template, squeeze, chart)
    assert "chart_too_small" in codes
    from designer.audit import apply_fixes

    finding = next(i for i in report["issues"] if i["code"] == "chart_too_small")
    assert finding["fixable"]
    apply_fixes(deck, report, [finding["id"]], template)
    visual = next(e for e in deck["slides"][0]["elements"] if e["id"] == "visual")
    assert visual["box"][3] >= deck["height"] * 0.25
    assert "chart_too_small" not in {
        i["code"]
        for i in audit(deck, template, [{"id": "brief", "text": "Первый тезис 10 20"}])[
            "issues"
        ]
    }


def test_bullets_move_beside_a_chart_instead_of_squeezing_it():
    """Диаграмма не ужимается под текст: тезисы уходят в колонку слева."""
    from designer.layout import MIN_VISUAL_SHARE, compose
    from designer.models import Outline

    plan = {
        "title": "Диаграмма и текст",
        "slides": [
            {
                "title": "Колода за 4 минуты вместо 186",
                "bullets": ["Замер на пилоте, колода из 12 слайдов " * 6] * 4,
                "notes": "",
                "source_refs": ["brief"],
                "visual": {
                    "kind": "bar",
                    "unit": "минут",
                    "categories": ["Вручную", "Сервис"],
                    "series": [{"name": "Сборка", "values": [186, 4]}],
                },
            }
        ],
    }
    template = parse_template(template_bytes(), "unknown.pptx")
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    slide = deck["slides"][0]
    visual = next(e for e in slide["elements"] if e["id"] == "visual")
    body = next(e for e in slide["elements"] if e["id"] == "body")
    assert visual["box"][3] >= deck["height"] * MIN_VISUAL_SHARE * 0.9
    # Блоки стоят рядом, а не друг под другом.
    assert body["box"][0] + body["box"][2] <= visual["box"][0] + 1


def test_type_scale_is_clustered_and_bounded():
    """Шкала кеглей — ступени дизайнера, а не список всех встреченных значений."""
    from collections import Counter

    from designer.parsing import MAX_SCALE_STEPS, MIN_SCALE_SIZE, type_scale

    # Так выглядит реальный шаблон: основной текст весит много, декоративная
    # цифра во весь экран — три знака, сноски набраны 4–8 pt.
    counted = Counter(
        {
            4.14: 40, 6.0: 30, 8.0: 20,
            9.0: 2650, 9.92: 308, 10.13: 95, 10.5: 214, 11.0: 45,
            12.0: 1148, 12.25: 263, 13.22: 84, 14.0: 299, 14.06: 263,
            16.0: 225, 16.88: 98, 18.0: 96, 18.03: 24,
            24.0: 22, 24.38: 48, 27.05: 18, 32.0: 11, 48.0: 3,
            96.0: 3, 166.0: 3,
        }
    )
    scale = type_scale(counted, height_pt=540.0)
    assert 5 <= len(scale) <= MAX_SCALE_STEPS
    assert min(scale) >= MIN_SCALE_SIZE
    # Декоративные 96 и 166 pt — не ступень для текста.
    assert max(scale) <= 540.0 * 0.11
    # Близкие значения слились: 14.0 и 14.06 не могут быть двумя ступенями.
    assert not any(
        b - a <= a * 0.05 for a, b in zip(scale, scale[1:])
    ), scale
    # Мелкие кегли не пропали, но живут отдельно.
    from designer.parsing import caption_sizes

    assert caption_sizes(counted) == [4.14, 6.0, 8.0]


def test_palette_drops_office_theme_colours_nobody_uses():
    """В палитру идут цвета, которыми в шаблоне что-то покрашено."""
    import io

    from pptx import Presentation as Deck
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Pt

    from designer.parsing import parse_template

    deck = Deck()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Заголовок"
    slide.placeholders[1].text = "Текст"
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(10), Pt(10), Pt(80), Pt(40)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("0077FF")
    out = io.BytesIO()
    deck.save(out)

    palette = parse_template(out.getvalue(), "brand.pptx")["tokens"]["colors"]
    assert "0077FF" in palette
    # Офисная тема по умолчанию: эти цвета объявлены в clrScheme, но ими в
    # шаблоне ничего не покрашено.
    assert not {"ED7D31", "A5A5A5", "FFC000", "70AD47"} & set(palette)


def test_a_template_without_explicit_colour_keeps_its_theme_palette():
    """Если явных цветов нет вовсе, палитрой становится тема: другого нет."""
    from designer.parsing import parse_template

    tokens = parse_template(template_bytes(), "unknown.pptx")["tokens"]
    assert tokens["colors"], "палитра не может быть пустой"


def test_layout_font_sizes_stay_on_the_template_scale():
    """Вёрстка не встаёт между ступенями, иначе её же аудит это находит."""
    from designer.models import Outline

    for name in ("VK Tech", "Education", "WorkSpace"):
        path = require_template(name)
        template = parse_template(path.read_bytes(), path.name)
        scale = template["tokens"]["font_sizes"]
        for variant in ("classic", "split", "focus"):
            deck = compose(
                Outline.model_validate(outline(4)).model_dump(), template, variant
            )
            report = audit(deck, template, [{"id": "brief", "text": "Первый тезис"}])
            assert not [
                i for i in report["issues"] if i["code"] == "font_size_off_scale"
            ], (path.name, variant, scale)


def test_icons_are_drawn_not_left_as_empty_circles(tmp_path):
    """У каждого значка внутри круга есть рисунок, а не пустая заливка."""
    from pptx import Presentation as Deck
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from designer.exporting import export_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template
    from designer.visuals import ICONS, pick_icon

    # Подпись, которая не попадает ни в одно ключевое слово: раньше такой значок
    # рисовался одним закрашенным кругом.
    assert pick_icon("Квазар") == "пункт"
    assert len(ICONS["пункт"]) >= 3

    plan = {
        "title": "Схема",
        "slides": [
            {
                "title": "Шаблон разбирается в дизайн-систему",
                "bullets": ["Имён макетов в коде нет"],
                "notes": "",
                "source_refs": ["brief"],
                "visual": {"kind": "icon", "steps": ["Квазар", "Пульсар", "Гарнитуры"]},
            }
        ],
    }
    source = template_bytes()
    template = parse_template(source, "unknown.pptx")
    deck = compose(Outline.model_validate(plan).model_dump(), template, "classic")
    source_path = tmp_path / "template.pptx"
    source_path.write_bytes(source)
    output = tmp_path / "deck.pptx"
    export_pptx(source_path, template, deck, output)
    # Значок — группа нативных фигур из библиотеки иконок; у неподобранной
    # подписи — нейтральный знак, но тоже рисунок, а не одна заливка.
    icons = [
        s
        for s in Deck(output).slides[0].shapes
        if s.shape_type == MSO_SHAPE_TYPE.GROUP and s.name.startswith("icon:")
    ]
    assert len(icons) == 3
    assert all(len(icon.shapes) >= 2 for icon in icons)
    assert any("palette" in icon.name for icon in icons)


def test_diagram_labels_share_one_size_and_never_break_a_word():
    """Единый кегль на схему, не мельче 10 pt, и слово целиком."""
    from designer.visuals import (
        EMU_PER_PT,
        MIN_LABEL_SIZE,
        icon_grid,
        label_plan,
        label_size,
        word_fits,
    )

    steps = ["Гарнитуры", "Шкала кеглей", "Палитра", "Поля", "Страницы"]
    # Узкая колонка варианта split: пять значков в ряд не помещаются целыми
    # словами, поэтому рядов становится два.
    rows, per_row = icon_grid(steps, 318.0)
    assert (rows, per_row) == (2, 3)
    plan = label_plan("icon", [0, 0, 318.0 * EMU_PER_PT, 132.0 * EMU_PER_PT], {"steps": steps})
    assert plan["size"] >= MIN_LABEL_SIZE
    assert not plan["word_break"]

    # Кегль не опускается ниже ориентира даже на совсем узкой схеме.
    assert label_size(["Непереносимоедлинноеслово"], 20.0) == MIN_LABEL_SIZE
    assert not word_fits(["Непереносимоедлинноеслово"], 20.0)


def test_process_stacks_vertically_when_labels_do_not_fit_in_a_row():
    """Схема получает другую раскладку, а не кегль 6 pt."""
    from designer.visuals import EMU_PER_PT, MIN_LABEL_SIZE, label_plan

    steps = ["Шаблон", "Материалы", "План колоды", "Вёрстка", "Аудит"]
    narrow = [0, 0, 300.0 * EMU_PER_PT, 250.0 * EMU_PER_PT]
    plan = label_plan("process", narrow, {"steps": steps})
    assert plan["size"] >= MIN_LABEL_SIZE
    assert not plan["word_break"]


def test_small_text_and_word_breaks_are_reported_and_fixable():
    """Новые коды ловят мелкий кегль и перенос внутри слова."""
    from designer.audit import apply_fixes
    from designer.visuals import MIN_LABEL_SIZE

    template = parse_template(template_bytes(), "unknown.pptx")

    def shrink(slide):
        slide["elements"][1]["font_size"] = 7.5

    deck, report, codes = audited(template, shrink)
    assert "text_too_small" in codes
    finding = next(i for i in report["issues"] if i["code"] == "text_too_small")
    apply_fixes(deck, report, [finding["id"]], template)
    assert deck["slides"][0]["elements"][1]["font_size"] >= MIN_LABEL_SIZE

    def squeeze(slide):
        slide["elements"][1]["text"] = "Непереносимоедлинноесловонастраницу"
        slide["elements"][1]["box"][2] = 60.0

    assert "word_break" in audited(template, squeeze)[2]


def test_focus_variant_keeps_one_accent_and_equal_bullets():
    """В focus акцент крупнее, остальные тезисы одинаковы и не вылезают."""
    from designer.layout import compose, estimated_text_height
    from designer.models import Outline

    plan = {
        "title": "Фокус",
        "slides": [
            {
                "title": "Три риска и как мы их закрываем",
                "bullets": [
                    "Выдуманные цифры: аудит сверяет каждое число с материалом",
                    "Слабый шаблон: вёрстка опирается на рабочую область файла",
                    "Долгий рендер: колода собирается за четыре минуты",
                ],
                "notes": "",
                "source_refs": ["brief"],
                "visual": {"kind": "none"},
            }
        ],
    }
    template = parse_template(template_bytes(), "unknown.pptx")
    deck = compose(Outline.model_validate(plan).model_dump(), template, "focus")
    elements = deck["slides"][0]["elements"]
    lead = next(e for e in elements if e["id"] == "lead")
    bullets = [
        e
        for e in elements
        if e["kind"] == "text" and e["role"] == "body" and e["id"] != "lead"
    ]
    assert bullets, "тезисы под акцентом должны остаться"
    # Акцент — самый крупный блок содержания, а не самый мелкий.
    assert lead["font_size"] >= max(e["font_size"] for e in bullets)
    # Все обычные тезисы набраны одинаково: и кеглем, и маркером.
    assert len({e["font_size"] for e in bullets}) == 1
    assert len({e["text"].lstrip().startswith("•") for e in bullets}) == 1
    # Ничего не выходит за свою рамку.
    for element in [lead, *bullets]:
        assert estimated_text_height(element) <= element["box"][3] + 0.01, element["id"]


def test_focus_bullets_do_not_leave_their_box_on_supplied_templates():
    """На шаблонах кейса вариант focus не выпускает текст за рамку."""
    from designer.layout import compose, estimated_text_height
    from designer.models import Outline

    for name in ("VK Tech", "Education", "WorkSpace"):
        path = require_template(name)
        template = parse_template(path.read_bytes(), path.name)
        deck = compose(
            Outline.model_validate(outline(6)).model_dump(), template, "focus"
        )
        for slide in deck["slides"]:
            for element in slide["elements"]:
                if element["kind"] != "text" or element.get("from_template"):
                    continue
                assert estimated_text_height(element) <= element["box"][3] + 0.01, (
                    path.name,
                    slide["index"],
                    element["id"],
                )


def test_english_slide_in_a_russian_deck_is_reported():
    """deck_language ловит слайд, целиком написанный на другом языке."""
    from designer.audit import deck_language_strays

    deck = {
        "slides": [
            {
                "index": 0,
                "content": {
                    "title": "Сервис автоматической вёрстки",
                    "bullets": ["Колода собирается за четыре минуты"],
                },
            },
            {
                "index": 1,
                "content": {
                    "title": "Revenue grew 500 percent",
                    "bullets": ["The pilot team shipped the first release"],
                },
            },
        ]
    }
    assert deck_language_strays(deck) == ("ru", [2])
    # Запрошенный язык важнее большинства.
    assert deck_language_strays(deck, "en") == ("en", [1])


def test_product_names_do_not_make_a_slide_foreign():
    """vCPU, gpt-oss-20b и OpenAI живут в русском слайде на своих местах."""
    from designer.audit import deck_language_strays
    from designer.language import script_of

    assert script_of("Инференс на 8 vCPU, модель gpt-oss-20b от OpenAI") == "ru"
    assert script_of("gpt-oss-20b, vCPU, OpenAI, VK Tech") is None
    deck = {
        "slides": [
            {
                "index": 0,
                "content": {
                    "title": "Модели и железо",
                    "bullets": ["gpt-oss-20b на 8 vCPU"],
                },
            },
            {
                "index": 1,
                "content": {"title": "gpt-oss-20b, vCPU, OpenAI", "bullets": []},
            },
        ]
    }
    assert deck_language_strays(deck) is None


def test_stretched_picture_finding_and_fix():
    """Растянутая картинка страницы шаблона находится и чинится."""
    from designer.audit import restore_picture, stretched_pictures

    pattern = {
        "index": 0,
        "shapes": [
            {
                "id": 7,
                "box": {"x": 0.1, "y": 0.1, "w": 0.4, "h": 0.4},
                # Рамка 1:1 при собственных пропорциях 16:9 — картинка сплющена.
                "picture": {"native": 1.778, "frame": 1.0},
            },
            {
                # Значок в пару пикселей: искажение там не читается.
                "id": 8,
                "box": {"x": 0.9, "y": 0.9, "w": 0.02, "h": 0.02},
                "picture": {"native": 1.0, "frame": 2.0},
            },
        ],
    }
    found = stretched_pictures(pattern)
    assert [s["id"] for s in found] == [7]

    slide = {"pattern_index": 0, "elements": []}
    restore_picture(slide, pattern, {"element_id": "picture_7"})
    fix = slide["shape_fixes"]["7"]
    assert fix["ratio"] == 1.778
    # Короткая сторона на месте, длинная подрезана.
    assert fix["box"]["w"] == pytest.approx(0.4)
    assert fix["box"]["h"] < 0.4
    # После правки находки больше нет.
    assert stretched_pictures(pattern, slide["shape_fixes"]) == []


def test_text_overflow_fix_either_repairs_or_explains_itself():
    """POST /fixes больше не плодит ревизии с той же находкой."""
    from designer.audit import apply_fixes

    template = parse_template(template_bytes(), "unknown.pptx")
    # Шаблон без объявленной шкалы: ступени мельче текущей просто нет.
    template["tokens"]["font_sizes"] = [18.0]

    def squeeze(slide):
        element = slide["elements"][1]
        element["font_size"] = 18.0
        element["text"] = "Очень длинный тезис, который заведомо не помещается " * 6
        element["box"][3] = 30.0

    deck, report, codes = audited(template, squeeze)
    assert "text_overflow" in codes
    finding = next(i for i in report["issues"] if i["code"] == "text_overflow")
    assert finding["fixable"]
    apply_fixes(deck, report, [finding["id"]], template)
    again = audit(deck, template, [{"id": "brief", "text": "Первый тезис 10 20"}])
    repeat = [i for i in again["issues"] if i["code"] == "text_overflow"]
    # Либо находки нет, либо она честно помечена неисправимой с причиной.
    for item in repeat:
        assert not item["fixable"], item["message"]
        assert "автоисправление невозможно" in item["message"]


def test_model_typography_is_normalised_before_layout(client, monkeypatch):
    """Неразрывный дефис и узкий пробел модели — пустые прямоугольники на слайде."""
    plan = outline(2)
    plan["title"] = "Дизайн‑система"
    plan["slides"][0]["title"] = "Пилот стоил 1 400 рублей"
    plan["slides"][0]["bullets"] = ["дизайн‑система за 1 400 рублей."]
    plan["slides"][0]["visual"] = {
        "kind": "bar",
        "categories": ["Было‑раньше", "Стало"],
        "series": [{"name": "Сборка колоды", "values": [10, 20]}],
        "unit": "мин на колоду",
    }
    mock_provider(monkeypatch, [json.dumps(plan)])
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 2},
    )
    assert response.status_code == 200, response.text
    body = json.dumps(response.json(), ensure_ascii=False)
    for char in ("‑", " ", " "):
        assert char not in body, f"в ответе остался {char!r}"
    slide = response.json()["slides"][0]
    assert slide["bullets"][0] == "Дизайн-система за 1 400 рублей"
    assert slide["visual"]["categories"][0] == "Было-раньше"
    assert slide["visual"]["series"][0]["name"] == "Сборка колоды"


def test_short_deck_is_sent_back_with_a_way_to_grow(client, monkeypatch):
    """Недобор слайдов — повод попросить модель разбить слайды, а не молча принять."""
    # Хелпер отдаёт не больше трёх слайдов, поэтому запрашиваем три.
    captured = mock_provider(
        monkeypatch, [json.dumps(outline(2)), json.dumps(outline(3))]
    )
    response = client.post(
        "/api/v1/outlines",
        json={"brief": "Данные: 10 и 20. Разделы 1, 2, 3.", "slide_count": 3},
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["slides"]) == 3
    nudge = captured[-1]["messages"][-1]["content"].lower()
    assert "разбей" in nudge and ("таблиц" in nudge or "схем" in nudge)


def test_stubborn_short_deck_is_accepted_with_a_warning(monkeypatch):
    """Если материала правда мало, колода выходит, но пользователь об этом знает."""
    import asyncio

    from designer.generation import generate_outline
    from designer.models import Brief

    mock_provider(monkeypatch, [json.dumps(outline(2))] * 3)
    warnings: list[str] = []
    request = Brief(brief="Данные: 10 и 20. Разделы 1, 2, 3.", slide_count=5)
    result = asyncio.run(
        generate_outline(request, [{"id": "brief", "text": request.brief}], warnings)
    )
    assert len(result.slides) == 2
    assert warnings and "2 слайдов вместо 5" in warnings[0]


def test_deck_never_exceeds_the_requested_size(monkeypatch):
    """Верхняя граница — это граница: лишние слайды не уезжают пользователю."""
    import asyncio

    from designer.generation import generate_outline
    from designer.models import Brief

    mock_provider(monkeypatch, [json.dumps(outline(5))] * 3)
    request = Brief(brief="Данные: 10 и 20. Разделы 1, 2, 3.", slide_count=3)
    result = asyncio.run(
        generate_outline(request, [{"id": "brief", "text": request.brief}], [])
    )
    assert len(result.slides) == 3


def test_file_level_findings_reach_the_report():
    """Приложение 1: «файл не открывается» и «слайд оказался картинкой»."""
    from designer.audit import file_issues, merge_file_issues

    broken = file_issues({"opens": False})
    assert [i["code"] for i in broken] == ["file_broken"]
    assert broken[0]["severity"] == "error"

    found = file_issues(
        {
            "opens": True,
            "raster_slides": [2],
            "moved_branding": [{"slide": 1, "name": "Logo", "shift": 0.05}],
            "charts": [{"slide": 3, "series": 2, "legend": False, "axis_title": True, "value_labels": True}],
        }
    )
    codes = {i["code"]: i for i in found}
    assert codes["raster_slide"]["slide_index"] == 2
    assert codes["branding_moved"]["severity"] == "error"
    assert "Logo" in codes["branding_moved"]["message"]
    assert "легенды" in codes["chart_labels"]["message"]

    report = {"issues": [], "counts": {"errors": 0, "warnings": 0}}
    merge_file_issues(report, {"opens": True, "raster_slides": [0]})
    assert report["counts"]["errors"] == 1 and len(report["issues"]) == 1


def test_rerun_audit_keeps_file_level_findings(client):
    """Повторный аудит собирает отчёт заново — проверки файла должны в нём остаться.

    Кнопки «Проверить содержание» и «Проверить по изображению» зовут этот же
    эндпоинт, и без слияния находки, видные только в pptx, молча исчезали.
    """
    _, job = generate(client)
    identifier = job["presentation_ids"][0]
    store = client.app.state.store
    record = store.get("presentations", identifier)
    record["export_check"] = {"opens": True, "raster_slides": [0], "charts": []}
    store.put("presentations", record)

    report = client.post(f"/api/v1/presentations/{identifier}/audit").json()
    assert "raster_slide" in {i["code"] for i in report["issues"]}
    assert report["counts"]["errors"] >= 1
    # Отчёт сохраняется, значит и в карточке презентации находка на месте.
    saved = client.get(f"/api/v1/presentations/{identifier}/audit").json()
    assert "raster_slide" in {i["code"] for i in saved["issues"]}


def test_clean_chart_leaves_no_file_findings():
    """У подписанной диаграммы претензий быть не должно."""
    from designer.audit import file_issues

    assert file_issues(
        {
            "opens": True,
            "charts": [{"slide": 1, "series": 1, "legend": False, "axis_title": True, "value_labels": True}],
        }
    ) == []


def test_export_keeps_template_branding_in_place(tmp_path):
    """«Логотип или колонтитул сдвинуты» — проверяем по записанному файлу."""
    from designer.exporting import branding_drift, export_pptx, verify_pptx
    from designer.layout import compose
    from designer.models import Outline
    from designer.parsing import parse_template

    for path in sample_templates() or pytest.skip("нет шаблонов"):
        template = parse_template(path.read_bytes(), path.name)
        deck = compose(Outline.model_validate(outline(3)).model_dump(), template, "classic")
        result = tmp_path / f"{path.stem[:8]}.pptx"
        export_pptx(path, template, deck, result)
        assert branding_drift(path, deck, result) == []
        check = verify_pptx(result, len(deck["slides"]))
        # Диаграммы разбираются: у каждой видно легенду, подписи осей и значений.
        assert all({"legend", "axis_title", "value_labels"} <= set(c) for c in check["charts"])
        assert check["charts"], "в плане есть диаграмма, она должна попасть в отчёт"


# ---------------------------------------------------------------- визуальные ассеты


def png_glyph(color=(0, 119, 255), size=64, fill=0.45):
    """Одноцветный значок на прозрачном фоне — как иконка из шаблона."""
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, size, size), True)
    pix.clear_with()
    margin = int(size * (1 - fill) / 2)
    pix.set_rect(pymupdf.IRect(margin, margin, size - margin, size - margin), (*color, 255))
    return pix.tobytes("png")


def png_photo(width=400, height=300):
    """Непрозрачная пёстрая картинка — как фотография."""
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    stripe = width // 40
    for i in range(40):
        colour = ((i * 37) % 256, (i * 91) % 256, (i * 53) % 256)
        pix.set_rect(pymupdf.IRect(i * stripe, 0, (i + 1) * stripe, height), colour)
    return pix.tobytes("png")


def test_svg_paths_become_contours_with_colour_roles():
    from designer.assets import parse_svg

    geometry = parse_svg(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" stroke="currentColor" fill="none">'
        '<g transform="translate(2 1)"><path d="M2 2h10v10z m2 2a3 3 0 01 6 0"/></g>'
        '<circle cx="12" cy="12" r="4" fill="var(--accent)" stroke="none"/>'
        '<rect x="1" y="1" width="4" height="4" rx="1" fill="#abc"/>'
        '<script>alert(1)</script></svg>'
    )
    assert geometry["viewbox"] == [0.0, 0.0, 24.0, 24.0]
    path, circle, rect = geometry["items"]
    assert path["shape"] == "path" and path["stroke"] == "current"
    square, arc = path["contours"]
    assert square["closed"] and square["points"][0] == [4.0, 3.0]
    # Дуга «a3 3 0 01 6 0» со слитыми флагами — полуокружность из многих точек.
    assert not arc["closed"] and len(arc["points"]) > 8
    assert arc["points"][-1] == [12.0, 5.0]
    assert circle["shape"] == "ellipse" and circle["fill"] == "accent"
    assert rect["shape"] == "rect" and rect["fill"] == "AABBCC" and rect["rx"] == 1.0


def test_asset_references_cannot_leave_storage(tmp_path):
    from designer.assets import resolve

    for ref in (
        "builtin:../../app.py",
        "template:../../etc:passwd",
        f"template:{'a' * 32}:../source.pptx",
        "elsewhere:x:y",
    ):
        with pytest.raises(ValueError):
            resolve(ref, tmp_path)
    good = resolve(f"pack:{'b' * 32}:{'c' * 24}.png", tmp_path)
    assert good == tmp_path / "content_packs" / ("b" * 32) / "assets" / ("c" * 24 + ".png")


def test_stat_values_are_taken_from_the_text_verbatim():
    from designer.layout import stat_plan, stat_value

    nb = " "
    assert stat_value("Время сборки сократилось с 186 до 4 минут") == f"4{nb}мин"
    assert stat_value("Колода за 4 минуты вместо 186") == f"4{nb}мин"
    assert stat_value("Нарушений стиля 2 из 68") == f"2{nb}из{nb}68"
    assert stat_value("Инференс стоил 1 400 рублей") == f"1{nb}400{nb}₽"
    assert stat_value("Доля выросла до 23%") == "23%"
    assert stat_value("Запросы закрыты 38 вместо 11") == "38"
    assert stat_value("Без чисел вовсе") is None
    # Две цифры и фраза без чисел — фактоиды; одна цифра — нет.
    assert stat_plan(["С 186 до 4 минут", "2 из 68 колод", "Рутина ушла"]) is not None
    assert stat_plan(["С 186 до 4 минут", "Рутина ушла", "Дизайн свободен"]) is None


def asset_outline():
    return {
        "title": "Сервис вёрстки",
        "slides": [
            {"title": "Сервис вёрстки", "bullets": [], "source_refs": ["brief"]},
            {
                "title": "Пилот сэкономил время",
                "bullets": [
                    "Время сборки сократилось с 186 до 4 минут",
                    "Нарушений стиля 2 из 68 колод",
                    "Дизайнеры закрыли 38 запросов вместо 11",
                ],
                "source_refs": ["brief"],
            },
            {
                "title": "Что нужно для запуска",
                "bullets": [
                    "Команда из четырёх дизайнеров",
                    "Сервер с двумя ядрами",
                    "Бюджет на инференс",
                ],
                "source_refs": ["brief"],
            },
            {
                "title": "Три риска и как мы их закрываем",
                "bullets": ["Риск выдуманных цифр закрывает аудит"],
                "source_refs": ["brief"],
            },
            {"title": "Спасибо", "bullets": ["Вопросы"], "source_refs": ["brief"]},
        ],
    }


ASSET_SOURCES = [
    {
        "id": "brief",
        "text": "С 186 до 4 минут, 2 из 68 колод, 38 запросов вместо 11.",
    }
]


@pytest.mark.parametrize("variant", ["classic", "split", "focus"])
def test_slides_get_stats_icons_and_pictures(tmp_path, variant):
    from pptx import Presentation as Deck
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from designer.exporting import export_pptx, verify_pptx
    from designer.models import Outline

    source = template_bytes()
    template = parse_template(source, "plain.pptx")
    plan = Outline.model_validate(asset_outline()).model_dump()
    deck = compose(plan, template, variant)
    kinds = [{e["kind"] for e in s["elements"]} for s in deck["slides"]]
    # Цифры в тезисах стали фактоидами во всех вариантах.
    stats = [e for e in deck["slides"][1]["elements"] if e["kind"] == "stat"]
    assert stats and stats[0]["value"] == "4 мин"
    if variant == "classic":
        icons = [e for e in deck["slides"][2]["elements"] if e.get("role") == "icon"]
        assert [i["asset"]["id"] for i in icons] == ["builtin-team", "builtin-server", "builtin-money"]
    if variant == "split":
        assert "image" in kinds[3]
        picture = next(e for e in deck["slides"][3]["elements"] if e["id"] == "picture")
        # Иллюстрация вписана без искажения: рамка в пропорциях картинки.
        assert abs(picture["box"][2] / picture["box"][3] - picture["asset"]["ratio"]) < 0.01
    report = audit(deck, template, ASSET_SOURCES, "ru")
    assert not [i for i in report["issues"] if i["severity"] == "error"], report["issues"]
    path = tmp_path / "template.pptx"
    path.write_bytes(source)
    output = tmp_path / "deck.pptx"
    export_pptx(path, template, deck, output)
    assert verify_pptx(output, len(deck["slides"]))["raster_slides"] == []
    groups = [
        s.name
        for slide in Deck(output).slides
        for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.GROUP
    ]
    # SVG встроенного набора лёг нативными фигурами, а не картинкой.
    assert any(name.startswith(("icon:", "illustration:")) for name in groups) or variant == "focus"


def test_visual_content_pack_zip_is_imported_and_used(client):
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr(
            "manifest.json",
            json.dumps({"assets": [{"file": "icons/rocket.svg", "kind": "icon", "tags": ["запуск", "старт"]}]}),
        )
        z.writestr(
            "icons/rocket.svg",
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24"><path d="M4 20L12 3L20 20z" fill="#123456"/></svg>',
        )
        z.writestr("icons/команда-люди.png", png_glyph())
        z.writestr("photos/офис.png", png_photo())
        z.writestr("notes.md", "Запуск пилота: команда из четырёх человек.")
        z.writestr("icons/broken.svg", "<svg><path d='M0 0'")
    response = client.post(
        "/api/v1/content-packs", files={"file": ("pack.zip", archive.getvalue(), "application/zip")}
    )
    assert response.status_code == 201, response.text
    pack = response.json()
    assert pack["asset_counts"]["icon"] == 2 and pack["asset_counts"]["photo"] == 1
    assert "Запуск пилота" in pack["text"]
    listing = client.get(f"/api/v1/content-packs/{pack['id']}/assets").json()
    rocket = next(a for a in listing["items"] if a["name"] == "icons/rocket.svg")
    assert rocket["tags"] == ["запус", "стар"]
    served = client.get(rocket["url"].removeprefix(""))
    assert served.status_code == 200 and served.headers["content-type"].startswith("image/svg")
    assert "default-src 'none'" in served.headers["content-security-policy"]

    template = client.post(
        "/api/v1/templates", files={"file": ("template.pptx", template_bytes())}
    ).json()
    plan = {
        "title": "Пилот",
        "slides": [
            {"title": "Пилот", "bullets": [], "source_refs": ["brief"]},
            {
                "title": "Как мы запускаем",
                "bullets": ["Запуск через месяц", "Команда из четырёх человек"],
                "source_refs": ["brief"],
            },
            {"title": "Итог", "bullets": ["Готовы"], "source_refs": ["brief"]},
        ],
    }
    job = client.post(
        "/api/v1/generations",
        json={
            "template_id": template["id"],
            "brief": "Запуск пилота",
            "slide_count": 3,
            "outline": plan,
            "content_pack_ids": [pack["id"]],
        },
    ).json()
    done = wait_job(client, job["id"])
    classic = next(
        client.get(f"/api/v1/presentations/{i}").json()
        for i in done["presentation_ids"]
        if client.get(f"/api/v1/presentations/{i}").json()["variant"] == "classic"
    )
    icons = [e for e in classic["deck"]["slides"][1]["elements"] if e.get("role") == "icon"]
    sources = [i["asset"]["source"] for i in icons]
    # Пакет пользователя важнее встроенного набора: он покрыл все тезисы, и
    # ряд собран из его значков, а не вперемешку со встроенными.
    assert sources == ["pack", "pack"], icons
    pptx = client.get(f"/api/v1/presentations/{classic['id']}/export/pptx")
    deck = Presentation(io.BytesIO(pptx.content))
    names = [s.name for s in deck.slides[1].shapes]
    assert any(name == f"icon:{rocket['id']}" for name in names), names


def test_single_picture_and_empty_upload(client):
    ok = client.post("/api/v1/content-packs", files={"file": ("рост-выручки.svg", b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="10" height="10" fill="#f00"/></svg>')})
    assert ok.status_code == 201, ok.text
    assert ok.json()["asset_counts"]["icon"] == 1
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w") as z:
        z.writestr("readme.bin", b"\x00\x01")
    bad = client.post("/api/v1/content-packs", files={"file": ("empty.zip", empty.getvalue())})
    assert bad.status_code == 422


def pictured_template():
    """Шаблон с карточками, значком над каждой и фотографией на отдельной странице."""
    from pptx.util import Emu

    deck = Presentation()
    width, height = deck.slide_width, deck.slide_height
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Заголовок страницы с карточками"
    colours = [(0, 119, 255), (255, 57, 133), (255, 255, 255)]
    for i in range(3):
        left = int(width * (0.06 + i * 0.31))
        slide.shapes.add_picture(
            io.BytesIO(png_glyph(colours[i])), left, int(height * 0.3), Emu(int(width * 0.06)), Emu(int(width * 0.06))
        )
        box = slide.shapes.add_textbox(left, int(height * 0.42), int(width * 0.26), int(height * 0.3))
        box.text_frame.text = "Тезис под тематической иконкой"
    extra = deck.slides.add_slide(deck.slide_layouts[5])
    extra.shapes.title.text = "Фото"
    extra.shapes.add_picture(io.BytesIO(png_photo()), int(width * 0.5), int(height * 0.3), int(width * 0.4), int(height * 0.4))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_template_pictures_are_harvested_and_labelled(client, monkeypatch):
    import designer.generation as generation

    response = client.post("/api/v1/templates", files={"file": ("pictured.pptx", pictured_template())})
    assert response.status_code == 201, response.text
    template = response.json()
    listing = client.get(f"/api/v1/templates/{template['id']}/assets").json()
    assert listing["counts"]["icon"] == 3 and listing["counts"]["photo"] == 1
    assert client.get(listing["items"][0]["url"]).status_code == 200
    # Без мультимодальной модели подписывать нечем — честный 503.
    monkeypatch.delenv("DESIGNER_VLM_MODEL", raising=False)
    assert client.post(f"/api/v1/templates/{template['id']}/assets/tags").status_code == 503

    monkeypatch.setenv("DESIGNER_VLM_MODEL", "vlm")
    monkeypatch.setenv("DESIGNER_VLM_BASE_URL", "http://vlm.invalid/v1")
    calls = []

    async def fake_vision(prompt, payload, image, attempts=2):
        calls.append(payload["count"])
        assert image.startswith(b"\x89PNG")
        return {"items": [{"n": n, "tags": ["команда", "люди"]} for n in range(1, payload["count"] + 1)]}

    monkeypatch.setattr(generation, "vision_completion", fake_vision)
    tagged = client.post(f"/api/v1/templates/{template['id']}/assets/tags").json()
    assert tagged["tagged"] == 4 and calls == [4]
    # Повторно не спрашиваем: всё подписано.
    assert client.post(f"/api/v1/templates/{template['id']}/assets/tags").json()["tagged"] == 0


def test_template_sample_icons_are_replaced_in_place(tmp_path):
    from pptx import Presentation as Deck

    from designer.exporting import export_pptx
    from designer.models import Outline

    source = pictured_template()
    template = parse_template(source, "pictured.pptx")
    card_page = template["patterns"][0]
    assert len(card_page["icon_slots"]) == 3
    assert {s["color"] for s in card_page["icon_slots"]} == {"0077FF", "FF3985", "FFFFFF"}
    plan = Outline.model_validate(
        {
            "title": "Карточки",
            "slides": [
                {"title": "Обложка", "bullets": [], "source_refs": ["brief"]},
                {
                    "title": "Что даёт сервис",
                    "bullets": ["Команда свободна от рутины", "Бюджет меньше", "Сервер один"],
                    "source_refs": ["brief"],
                },
                {"title": "Итог", "bullets": ["Готовы"], "source_refs": ["brief"]},
            ],
        }
    ).model_dump()
    deck = compose(plan, template, "classic")
    slide = deck["slides"][1]
    assert slide["pattern_index"] == 0
    swaps = [e for e in slide["elements"] if e["id"].startswith("swap_")]
    assert [s["asset"]["id"] for s in swaps] == ["builtin-team", "builtin-money", "builtin-server"]
    # Цвет значка — цвет образца, место — его рамка.
    assert [s["tint"] for s in swaps] == ["0077FF", "FF3985", "FFFFFF"]
    path = tmp_path / "t.pptx"
    path.write_bytes(source)
    output = tmp_path / "d.pptx"
    export_pptx(path, template, deck, output)
    shapes = Deck(output).slides[1].shapes
    assert not [s for s in shapes if s.shape_type == 13], "образцы значков остались на слайде"
    assert len([s for s in shapes if s.name.startswith("icon:builtin-")]) == 3


def _deck_on_background(**background):
    """Колода на странице шаблона с заданным фоном."""
    from designer.layout import compose
    from designer.models import Outline

    page = dict(pattern(0, text_slots=4), background_kind="image", image_cover=1.0, **background)
    template = synthetic_template([page, dict(page, index=1)])
    deck = compose(Outline.model_validate(outline(2)).model_dump(), template, "classic")
    return deck, template


def test_calm_measured_background_is_not_a_finding():
    """Фоновая подложка шаблона — не повод для находки, если фон спокойный."""
    deck, template = _deck_on_background(bg_luma=0.92, bg_spread=0.04)
    codes = [i["code"] for i in audit(deck, template, [{"id": "brief", "text": "тест"}])["issues"]]
    assert "text_over_image" not in codes


def test_dark_busy_background_is_reported_with_measurement():
    """Тёмный пёстрый фон: вёрстка кладёт подложку, аудит просит проверить."""
    deck, template = _deck_on_background(bg_luma=0.18, bg_spread=0.4)
    found = [
        i for i in audit(deck, template, [{"id": "brief", "text": "тест"}])["issues"]
        if i["code"] == "text_over_image"
    ]
    assert found, "на тёмном пёстром фоне находка обязана быть"
    assert "40%" in found[0]["message"], found[0]["message"]


def test_unmeasured_background_still_warns():
    """Без измерения фон неизвестен — честнее предупредить."""
    deck, template = _deck_on_background()
    found = [
        i for i in audit(deck, template, [{"id": "brief", "text": "тест"}])["issues"]
        if i["code"] == "text_over_image"
    ]
    assert found and "не измерен" in found[0]["message"], found


def test_cover_fill_is_left_to_the_template():
    """Обложка в рамках шаблона не судится по заливке, обычный слайд — судится."""
    template = parse_template(template_bytes(), "unknown.pptx")

    def sparse(content):
        for slide in content["slides"]:
            slide["bullets"] = ["Мало"]
            slide["visual"] = {"kind": "none"}

    deck, report, _ = audited(template, content_mutator=sparse, count=2)
    assert deck["slides"][0]["native_cover"]
    flagged = {i["slide_index"] for i in report["issues"] if i["code"] == "fill_ratio"}
    assert 0 not in flagged and 1 in flagged

    # Мера та же: сними с обложки признак — и находка вернётся.
    deck["slides"][0]["native_cover"] = False
    report = audit(deck, template, [{"id": "brief", "text": "Мало"}])
    assert any(
        i["code"] == "fill_ratio" and i["slide_index"] == 0 for i in report["issues"]
    )


def test_card_text_frame_grows_down_to_the_next_obstacle():
    from designer.layout import card_depth

    top_row = [10.0, 100.0, 200.0, 40.0]
    below = [10.0, 300.0, 200.0, 40.0]
    # Следующая карточка в колонке — предел.
    assert card_depth(top_row, [top_row, below], [], 500.0) == 300.0 - 8.0 - 100.0
    # Фигура-подложка карточки — тоже.
    assert card_depth(top_row, [top_row], [(0.0, 90.0, 220.0, 150.0)], 500.0) == 140.0
    # Значок ниже рамки — тоже.
    assert card_depth(top_row, [top_row], [(50.0, 200.0, 30.0, 30.0)], 500.0) == 92.0
    # Свободно до низа рабочей области; но рамка никогда не уменьшается.
    assert card_depth(top_row, [top_row], [], 500.0) == 400.0
    assert card_depth(top_row, [top_row], [], 120.0) == 40.0


def test_title_backdrop_is_the_plate_under_the_template_title():
    from designer.layout import title_backdrop

    title_box = {"x": 0.1, "y": 0.066, "w": 0.6, "h": 0.066}
    # Плашка, в которой начинается рамка заголовка образца.
    assert title_backdrop((75.0, 32.0, 580.0, 49.0), title_box, 960, 540)
    # Крупная иллюстрация рядом с заголовком — не плашка.
    assert not title_backdrop((474.0, 54.0, 486.0, 486.0), title_box, 960, 540)
    # Ряд декора ниже заголовка — тоже.
    assert not title_backdrop((120.0, 107.0, 118.0, 110.0), title_box, 960, 540)


def test_title_on_its_plate_is_not_a_branding_overlap():
    """Заголовок, поставленный на плашку образца, аудит не считает наездом."""
    template = parse_template(template_bytes(), "unknown.pptx")
    content = outline(2)
    from designer.models import Outline

    deck = compose(Outline.model_validate(content).model_dump(), template, "classic")
    slide = deck["slides"][1]
    title = next(e for e in slide["elements"] if e["id"] == "title")
    pattern = next(p for p in template["patterns"] if p["index"] == slide["pattern_index"])
    width, height = deck["width"], deck["height"]
    x, y, w, h = title["box"]
    pattern["title_box"] = {"x": x / width, "y": y / height, "w": w / width, "h": h / height}
    plate = {"x": (x - 4) / width, "y": (y - 4) / height, "w": (w + 8) / width, "h": (h + 8) / height}
    pattern["reserved"] = [plate]

    def overlaps(report):
        return any(
            i["code"] == "branding_overlap" and i["slide_index"] == 1 for i in report["issues"]
        )

    assert not overlaps(audit(deck, template, [{"id": "brief", "text": "10 20"}]))
    # Та же фигура, но не под рамкой заголовка образца, — наезд.
    pattern["title_box"] = {"x": 0.9, "y": 0.9, "w": 0.05, "h": 0.05}
    assert overlaps(audit(deck, template, [{"id": "brief", "text": "10 20"}]))
