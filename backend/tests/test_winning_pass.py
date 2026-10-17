"""Regression checks for the final story, faithful edits and readable evidence."""
import copy
import io
import time
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from designer.app import create_app
from designer.analytics_bridge import extract, analyse_sources, build_outline, verify_outline
from designer.brandbook import apply_style


def test_story_keeps_context_and_verified_data_and_checks_budget():
    table=extract('Регион,Продажи\nСевер,10\nЮг,20'.encode(),'regions.csv')[0]
    sources=[{'id':'brief','text':'Сравнить продажи','_story':{'goal':'Сравнить продажи'}},
        {'id':'data','datasets':[table]}, {'id':'context','text':'Запуск пилота ограничен двумя регионами.\n\nСледующий шаг: проверить повторные покупки.'}]
    results=analyse_sources(sources,'Сравнить')
    outline=build_outline(results,sources,10)
    assert [s.title for s in outline.slides][0]=='Задача анализа'
    assert outline.slides[-1].title=='Следующие действия'
    assert any('Запуск пилота' in ' '.join(s.bullets) for s in outline.slides)
    assert verify_outline(outline,results,sources)
    with pytest.raises(ValueError,match='истории'):
        build_outline(results,sources,1)
    reordered=outline.model_dump();reordered['slides'].reverse()
    assert verify_outline(reordered,results,sources)
    analytical=next(s for s in reordered['slides'] if s['analysis_id'])
    reordered['slides'].append(copy.deepcopy(analytical))
    with pytest.raises(ValueError,match='продублирован'):
        verify_outline(reordered,results,sources)


def test_partial_brandbook_inherits_dark_background():
    template={'tokens':{'colors':['FFFFFF','0077FF'],'theme':{'lt1':'FFFFFF'}},
              'patterns':[{'role':'content','background':'111827','slots':[]}]}
    result=apply_style(template,{'font':'Arial'})
    assert result['brandbook']['background']=='111827'
    assert result['brandbook']['text_color']=='FFFFFF'
    assert template['patterns'][0]['background']=='111827'


def test_long_reference_cells_paginate_without_dropping_content():
    import json
    rows=[[f'R{i}', 'Длинное описание товара ' * 8] for i in range(12)]
    table=extract(json.dumps({'columns':['Код','Описание'],'rows':rows}).encode(),'long.json')[0]
    source={'id':'data','datasets':[table],'analysis_requests':{table['dataset_id']:{'task':'reference'}}}
    result=analyse_sources([source],'Таблица')
    outline=build_outline(result,[source],30)
    assert max(len(s.visual.rows) for s in outline.slides)<=3
    assert [row for s in outline.slides for row in s.visual.rows]==rows


def test_edit_rebuilds_same_template_all_variants_and_rejects_stale_revision(tmp_path,monkeypatch):
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Заголовок';slide.placeholders[1].text='Содержание'
    stream=io.BytesIO();deck.save(stream)
    with TestClient(create_app(tmp_path)) as client:
        template=client.post('/api/v1/templates',files={'file':('original.pptx',stream.getvalue())}).json()
        data=client.post('/api/v1/content-packs',files={'file':('data.csv','Регион,Продажи\nСевер,10\nЮг,20'.encode())}).json()
        job=client.post('/api/v1/generations',json={'brief':'Сравнить продажи','analysis_goal':'Сравнить','narrative':True,'template_id':template['id'],'content_pack_ids':[data['id']],'slide_count':10}).json()
        for _ in range(500):
            job=client.get(f"/api/v1/jobs/{job['id']}").json()
            if job['status'] in {'completed','failed'}:break
            time.sleep(.01)
        assert job['status']=='completed',job
        records=[client.get(f'/api/v1/presentations/{id}').json() for id in job['presentation_ids']]
        frames=[next(e['box'] for s in r['deck']['slides'] if s['content']['analysis_id'] for e in s['elements'] if e['id']=='visual') for r in records]
        assert len({tuple(box) for box in frames})==3
        record=records[0];url=f"/api/v1/presentations/{record['id']}"
        slides=[s['content'] for s in record['deck']['slides']]
        slides[0]['title']='Вопрос для команды'
        body={'revision':1,'title':'Новая версия','variant':'focus','slides':slides}
        response=client.patch(url,json=body)
        assert response.status_code==200,response.text
        saved=response.json();assert saved['revision']==2 and saved['variant']=='focus'
        assert saved['template_id']==template['id']
        exported=Presentation(io.BytesIO(client.get(url+'/export/pptx').content))
        assert exported.slide_width==deck.slide_width and exported.slide_height==deck.slide_height
        assert any('Вопрос для команды' in s.text for s in exported.slides[0].shapes if s.has_text_frame)
        assert client.patch(url,json=body).status_code==409
        body['revision']=2
        next(s for s in slides if s['analysis_id'])['visual']['series'][0]['values'][0]=99999
        assert client.patch(url,json=body).status_code==422
        assert client.get(url).json()['revision']==2


def test_analytic_export_removes_empty_layout_panel_but_keeps_artwork(tmp_path):
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    from designer.parsing import parse_template
    from designer.layout import compose
    from designer.exporting import export_pptx
    deck=Presentation();layout=deck.slide_layouts[1]
    placeholder=layout.placeholders[1]
    placeholder.left=Inches(5);placeholder.top=0;placeholder.width=Inches(5);placeholder.height=deck.slide_height
    placeholder.fill.solid();placeholder.fill.fore_color.rgb=RGBColor(255,255,255)
    slide=deck.slides.add_slide(layout);slide.shapes.title.text='Тема';slide.placeholders[1].text='Образец'
    source=tmp_path/'template.pptx';deck.save(source)
    template=parse_template(source.read_bytes(),'template.pptx')
    table=extract(b'Category,Value\nA,10\nB,20','data.csv')[0]
    sources=[{'id':'data','datasets':[table]}];results=analyse_sources(sources,'comparison')
    outline=build_outline(results,sources,10).model_dump();composed=compose(outline,template,'classic')
    target=tmp_path/'final.pptx';export_pptx(source,template,composed,target)
    exported=Presentation(target)
    assert not any(s.is_placeholder and s.width*s.height>=deck.slide_width*deck.slide_height*.25 for s in exported.slides[0].slide_layout.shapes)
    assert len(Presentation(source).slide_layouts[1].placeholders)==len(layout.placeholders)
