"""Semantic cover budgeting and selective geometry fixes with real exports."""
import copy
import io
import pytest
from pptx import Presentation
from designer.analytics_bridge import build_outline, verify_outline
from designer.business_bridge import prepare_business
from designer.layout import compose, content_demand
from designer.models import Outline
from designer.parsing import parse_template
from designer.audit import issue
from tests.test_business_period_service import inputs
from tests.test_covers import cover_template, save
from tests.test_pipeline import client, generate


@pytest.mark.parametrize('count', [12, 13])
def test_cover_budget_and_verified_story_survive_reorder(count):
    sources = inputs()
    results = prepare_business(sources, sources[0]['text'])
    original = build_outline(results, sources, 30)
    sources[0]['_business_presentation'] = {'slide_count':count,'audience':'Руководитель продукта'}
    outline = build_outline(results, sources, count)
    assert len(outline.slides) == count
    assert outline.slides[0].role == 'cover'
    assert outline.slides[0].bullets == ['II кв. 2026\nРуководитель продукта']
    assert all(s.role == 'content' for s in outline.slides[1:])
    text = '\n'.join([s.title+'\n'+'\n'.join(s.bullets) for s in outline.slides])
    for slide in original.slides:
        for bullet,role in zip(slide.bullets,slide.block_roles):
            if count == 12 and slide.analysis_id.startswith('business:kpi:') and role == 'evidence':
                assert outline.slides[1].visual == slide.visual  # duplicate KPI prose is represented natively
            else:
                assert bullet in text
    actual = outline.model_dump()
    actual['slides'] = actual['slides'][1:] + actual['slides'][:1]
    assert verify_outline(actual,results,sources)
    # Moving the cover must not turn either first content or last cover into a
    # different semantic page. The explicit cover uses its native template slot.
    template = parse_template(save(cover_template()), 'arbitrary.pptx')
    deck = compose(actual, template, 'classic', assets=[])
    assert deck['slides'][-1]['native_cover']
    assert not deck['slides'][0]['native_cover']
    damaged = copy.deepcopy(actual)
    damaged['slides'][0]['role'] = 'cover'
    with pytest.raises(ValueError):verify_outline(damaged,results,sources)


def test_explicit_content_at_zero_is_not_a_cover_and_cover_rejects_evidence():
    content = Outline.model_validate({'title':'Test','slides':[{'role':'content','title':'Проверенный вывод','bullets':['Факт']}]}).model_dump()['slides'][0]
    assert not content_demand(content,0,1)['cover']
    with pytest.raises(ValueError):
        Outline.model_validate({'title':'Test','slides':[{'role':'cover','title':'Тема','bullets':['Один','Два']} ]})


def test_selected_fix_retains_other_issue_exports_and_undo(client):
    _,job=generate(client)
    identifier=job['presentation_ids'][0]
    store=client.app.state.store
    record=store.get('presentations',identifier)
    first=record['deck']['slides'][0]['elements'][0]
    first['box'][0]=-40
    other=record['deck']['slides'][1]['elements'][0]
    other['box'][0]=-35
    semantic=issue(1,'context_0','title_mismatch','Явная тестовая смысловая находка',category='contextual')
    from designer.audit import audit
    record['audit']=audit(record['deck'],store.get('templates',record['template_id']),record['sources'])
    record['audit']['issues'].append(semantic)
    store.put('presentations',record)
    store.put('revisions',{**copy.deepcopy(record),'id':identifier+'_1'})
    finding=next(i for i in record['audit']['issues'] if i['code']=='out_of_bounds' and i['slide_index']==0)
    before=copy.deepcopy(record['deck']['slides'][1])
    path=f'/api/v1/presentations/{identifier}'
    fixed=client.post(path+'/fixes',json={'revision':1,'issue_ids':[finding['id']]})
    assert fixed.status_code==200,fixed.text
    fixed=fixed.json()
    assert fixed['revision']==2
    assert fixed['deck']['slides'][1]==before
    assert finding['id'] not in {i['id'] for i in fixed['audit']['issues']}
    assert semantic in fixed['audit']['issues']
    assert not fixed['audit'].get('model_checks')
    exported=client.get(fixed['exports']['pptx'])
    assert exported.status_code==200
    native=Presentation(io.BytesIO(exported.content))
    assert len(native.slides)==len(fixed['deck']['slides'])
    assert client.post(path+'/fixes',json={'revision':1,'issue_ids':[finding['id']]}).status_code==409
    restored=client.post(path+'/revisions/1/restore',json={'revision':2})
    assert restored.status_code==200,restored.text
    assert restored.json()['revision']==3
    assert restored.json()['deck']==record['deck']
    assert semantic in restored.json()['audit']['issues']
    assert not restored.json()['audit'].get('model_checks')
    assert client.get(path).json()['revision']==3
