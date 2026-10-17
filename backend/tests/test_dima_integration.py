"""Integration contract, independent replay and slide pagination boundaries."""
import copy
import io
import json
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from designer.analytics_bridge import extract, extract_pack, analyse, analyse_sources, build_outline, verify_outline
from designer.app import create_app

SAMPLES = Path(__file__).resolve().parents[2]/"samples"/"analytics"/"contract"


@pytest.mark.parametrize("name",["time-series","categories","shares"])
def test_published_contract_examples_survive_slide_adapter(name):
    table=extract((SAMPLES/f"{name}.csv").read_bytes(),f"{name}.csv")[0]
    request=json.loads((SAMPLES/f"{name}.request.json").read_text(encoding="utf-8"))
    source={"id":"pack", "datasets":[table], "analysis_requests":{table["dataset_id"]:request}}
    results=analyse_sources([source],request["goal"])
    assert results[0]["contract"]==json.loads((SAMPLES/f"{name}.result.json").read_text(encoding="utf-8"))
    slides=build_outline(results,[source],30)
    assert verify_outline(slides,results,[source])
    visual=slides.slides[0].visual
    assert visual.categories==results[0]["contract"]["data"]["categories"]
    assert visual.series[0].values==results[0]["contract"]["data"]["series"][0]["values"]
    altered=copy.deepcopy(results)
    altered[0]["contract"]["facts"][0]["evidence"]["row_ids"]=[999]
    with pytest.raises(ValueError):
        verify_outline(slides,altered,[source])
    # Mutating the stored cells must invalidate both the graph and its conclusion.
    source["datasets"][0]["contract"]["rows"][0][-1]=99999
    with pytest.raises(ValueError):
        verify_outline(slides,results,[source])


def test_category_pages_and_wide_filtered_table_are_complete():
    table=extract(("Категория,Значение\n"+"\n".join(f"К{i},{i}" for i in range(25))).encode(),"categories.csv")[0]
    source={"id":"pack","datasets":[table]}
    results=analyse_sources([source],"Сравнение")
    outline=build_outline(results,[source],30)
    assert len(outline.slides)==2
    assert sum(len(s.visual.categories) for s in outline.slides)==25
    assert verify_outline(outline,results,[source])
    matrix={"columns":[f"К{i}" for i in range(12)],"rows":[[r]*12 for r in range(31)]}
    table=extract(json.dumps(matrix).encode(),"wide.json")[0]
    source={"id":"wide","datasets":[table],"analysis_requests":{table["dataset_id"]:{"task":"reference","filters":{"c1":[0,2,4,6,8,10]}}}}
    results=analyse_sources([source],"Таблица")
    outline=build_outline(results,[source],30)
    assert len(outline.slides)==2
    assert all(len(s.visual.rows)==6 for s in outline.slides)
    assert outline.slides[1].visual.columns==["К0","К7","К8","К9","К10","К11"]
    assert verify_outline(outline,results,[source])


def test_zip_paths_rejections_and_unsafe_asset_paths():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as archive:
        archive.writestr("nested/data.csv","Категория,Число\nА,1\nБ,2")
        archive.writestr("bad.json","{}")
    tables,report=extract_pack(stream.getvalue(),"pack.zip")
    assert tables[0]["contract"]["ref"]["archive_path"]=="nested/data.csv"
    assert next(r for r in report if r["name"]=="bad.json")["status"]=="rejected"
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as archive:
        archive.writestr("../asset.png",b"image")
    with pytest.raises(ValueError):
        extract_pack(stream.getvalue(),"unsafe.zip")


def test_series_pages_keep_every_category_metric_pair():
    columns=["Category"]+[f"Metric {i}" for i in range(12)]
    matrix={"columns":columns,"rows":[[f"C{r}"]+[r*100+i for i in range(12)] for r in range(25)]}
    table=extract(json.dumps(matrix).encode(),"series.json")[0]
    source={"id":"pack","datasets":[table]}
    results=analyse_sources([source],"comparison")
    outline=build_outline(results,[source],30)
    assert len(outline.slides)==6
    assert all(len(s.visual.series)<=5 for s in outline.slides)
    observed={(category,s.name):value for slide in outline.slides for s in slide.visual.series for category,value in zip(slide.visual.categories,s.values)}
    expected={(c,s["name"]):v for s in results[0]["contract"]["data"]["series"] for c,v in zip(results[0]["contract"]["data"]["categories"],s["values"])}
    assert observed==expected
    assert verify_outline(outline,results,[source])


def test_xlsx_empty_sheet_is_reported_without_losing_valid_sheet():
    from openpyxl import Workbook
    workbook=Workbook();workbook.active.title="Data"
    workbook.active.append(["Category","Value"]);workbook.active.append(["North",10])
    workbook.create_sheet("Empty")
    stream=io.BytesIO();workbook.save(stream)
    tables,report=extract_pack(stream.getvalue(),"data.xlsx")
    assert len(tables)==1 and tables[0]["sheet"]=="Data"
    assert any(r["status"]=="rejected" and "Empty" in r["name"] for r in report)
    assert len(extract(stream.getvalue(),"data.xlsx"))==1


def test_explicit_api_request_resolves_ambiguity_without_changing_contract(tmp_path,monkeypatch):
    monkeypatch.delenv("DESIGNER_TEMPLATE_DIR",raising=False)
    with TestClient(create_app(tmp_path)) as client:
        response=client.post("/api/v1/content-packs",files={"file":("data.csv","Категория,Сумма,Срок\nА,10,2\nБ,20,4".encode())})
        assert response.status_code==201,response.text
        pack=response.json();dataset=pack["datasets"][0]["dataset_id"]
        preview=client.get(f"/api/v1/content-packs/{pack['id']}/analysis",params={"goal":"Распределение"}).json()
        assert preview["analysis_results"][0]["status"]=="refused"
        selection={"goal":"Распределение","requests":{dataset:{"task":"distribution","metrics":["c2"]}}}
        preview=client.post(f"/api/v1/content-packs/{pack['id']}/analysis",json=selection)
        assert preview.status_code==200,preview.text
        result=preview.json()["analysis_results"][0]
        assert result["visualization"]["selected"]=="histogram"
        assert next(f for f in result["facts"] if f["evidence"]["operation"]=="median")["value"]==15
        assert client.post(f"/api/v1/content-packs/{pack['id']}/analysis",json={"requests":None}).status_code==422


@pytest.mark.parametrize("field,value", [
    ("bullets", ["Прибыль 999999 рублей"]),
    ("source_refs", ["other-pack"]),
    ("notes", "Подменённые доказательства"),
])
def test_verified_slide_rejects_tampered_narrative_and_provenance(field, value):
    table = extract(b"Category,Value\nA,10\nB,20", "data.csv")[0]
    source = {"id": "pack", "datasets": [table]}
    results = analyse_sources([source], "comparison")
    outline = build_outline(results, [source], 30).model_dump()
    outline["slides"][0][field] = value
    with pytest.raises(ValueError):
        verify_outline(outline, results, [source])


def test_reference_keeps_literal_missing_like_category_names():
    table = extract(json.dumps({"columns": ["Region"], "rows": [["NA"], ["None"], ["-"]]}).encode(), "regions.json")[0]
    source = {"id": "pack", "datasets": [table], "analysis_requests": {table["dataset_id"]: {"task": "reference"}}}
    results = analyse_sources([source], "reference")
    assert build_outline(results, [source], 30).slides[0].visual.rows == [["NA"], ["None"], ["-"]]


def test_summary_labels_each_metric_and_histogram_keeps_exact_intervals():
    table = extract(b"Category,Revenue,Cost\nA,10,2\nB,20,4", "data.csv")[0]
    source = {"id": "pack", "datasets": [table], "analysis_requests": {table["dataset_id"]: {"task": "summary"}}}
    results = analyse_sources([source], "summary")
    outline = build_outline(results, [source], 30)
    assert [(s.visual.series[0].name, s.visual.series[0].values) for s in outline.slides] == [("Revenue", [30]), ("Cost", [6])]
    table = extract(b"Value\n1000001\n1000002\n1000003\n1000004", "values.csv")[0]
    source = {"id": "pack", "datasets": [table], "analysis_requests": {table["dataset_id"]: {"task": "distribution", "bins": 3}}}
    results = analyse_sources([source], "distribution")
    visual = build_outline(results, [source], 30).slides[0].visual
    assert visual.categories == ["[1000001; 1000002)", "[1000002; 1000003)", "[1000003; 1000004]"]
    assert len(set(visual.categories)) == 3


def test_histogram_repeating_decimal_labels_stay_readable():
    table = extract(b"Value\n1\n2\n3\n5", "values.csv")[0]
    source = {"id": "pack", "datasets": [table], "analysis_requests": {table["dataset_id"]: {"task": "distribution", "bins": 3}}}
    result = analyse_sources([source], "distribution")
    labels = build_outline(result, [source], 30).slides[0].visual.categories
    assert all(len(label) < 40 for label in labels)
    assert any("≈" in label for label in labels)
    assert len(set(labels)) == 3
