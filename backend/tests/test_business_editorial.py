"""Controlled transport failures and persisted, grounded model contributions."""
import asyncio
import copy
import io
import json
import time

import httpx
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation

from designer import ai
from designer.analytics_bridge import build_outline, verify_outline
from designer.app import create_app
from designer.business_bridge import prepare_business
from designer.business_editorial import EDITORIAL_VERSION, editor_payload, replay_editorial, validate_editorial
from designer.inference_trace import capture_inference, validated_content
from test_business_period_service import inputs, upload


@pytest.fixture
def model_env(monkeypatch):
    monkeypatch.setenv('DESIGNER_LLM_PROVIDER', 'vk')
    monkeypatch.setenv('DESIGNER_VK_BASE_URL', 'https://model.invalid/v1')
    monkeypatch.setenv('DESIGNER_VK_MODEL', 'qwen-fixture-not-live')
    monkeypatch.setenv('DESIGNER_VK_API_KEY', 'test-secret-never-persist')
    monkeypatch.setenv('DESIGNER_VK_PROJECT_ID', 'fixture-project')
    monkeypatch.setattr(ai, 'RETRY_DELAY', 0)
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR', raising=False)


def response_for(payload):
    pages = []
    for page in payload['pages']:
        key = page['fact_ids'][0]
        name = page['id']
        title = page['approved_titles'][0] if name in {'risks', 'actions'} else f'Приоритет {name}: {{{{fact:{key}}}}}'
        pages.append(dict(id=name, title=title,
            interpretation=f'{{{{interpretation:{name}}}}}',
            interpretation_kind='observation',
            next_step=f'{{{{next_step:{name}}}}}',
            fact_ids=page['fact_ids'], context_refs=page['context_refs'], action_ref=page['action_ref']))
    return dict(input_sha256=payload['input_sha256'], pages=pages)


def draft_for(payload):
    return dict(input_sha256=payload['input_sha256'], pages=[
        {'id': page['id'], 'title': page['title'], 'hypothesis': None}
        for page in response_for(payload)['pages']])


def install_transport(monkeypatch, handler):
    real_client = httpx.AsyncClient
    monkeypatch.setattr(ai.httpx, 'AsyncClient', lambda **kwargs: real_client(
        transport=httpx.MockTransport(handler), **kwargs))


@pytest.mark.parametrize('failure,expected_attempts,code', [
    ('auth', 1, 'authorization'), ('rate', 2, 'rate_limit'),
    ('timeout', 2, 'timeout'), ('empty', 2, 'invalid_response'),
    ('truncated', 2, 'invalid_response'), ('json', 2, 'invalid_response')])
def test_failures_are_bounded_and_never_claim_success(monkeypatch, model_env, failure, expected_attempts, code):
    def handler(request):
        if failure == 'auth': return httpx.Response(401, json={'error': 'denied'})
        if failure == 'rate': return httpx.Response(429, json={'error': 'busy'})
        if failure == 'timeout': raise httpx.ReadTimeout('fixture timeout', request=request)
        return httpx.Response(200, json={'choices': [{'finish_reason': 'length' if failure == 'truncated' else 'stop',
            'message': {'content': '' if failure == 'empty' else '{partial' if failure == 'json' else '{"ok":true}'}}]})
    install_transport(monkeypatch, handler)
    async def run():
        with capture_inference() as trace:
            with pytest.raises(HTTPException):
                await ai.complete_json([], attempts=2)
            return trace
    trace = asyncio.run(run())
    assert trace['attempts'] == expected_attempts == len(trace['calls'])
    assert trace['attempted'] and not trace['completed']
    assert trace['successful_calls'] == 0 and trace['validation'] == 'failed'
    assert trace['error_code'] == code
    assert all(call['validation'] == 'failed' for call in trace['calls'])
    assert 'test-secret' not in json.dumps(trace)


def test_actual_response_metadata_requires_validation(monkeypatch, model_env):
    install_transport(monkeypatch, lambda request: httpx.Response(200, json={
        'id': 'fixture-response', 'usage': {'prompt_tokens': 11, 'completion_tokens': 7},
        'choices': [{'message': {'content': '{"ok":true}'}}]}))
    async def run():
        with capture_inference() as trace:
            await ai.complete_json([], validate=lambda value: value if value['ok'] else None)
            return trace
    trace = asyncio.run(run())
    assert trace['completed'] and trace['validation'] == 'passed' and trace['successful_calls'] == 1
    assert trace['response_id'] == 'fixture-response' and trace['usage']['prompt_tokens'] == 11
    assert len(trace['response_sha256']) == 64
    assert trace['calls'][0]['status_code'] == 200 and trace['calls'][0]['validation'] == 'passed'


def prepared():
    sources = inputs()
    results = prepare_business(sources, sources[0]['text'])
    outline = build_outline(results, sources, 12)
    return sources, results, outline


@pytest.mark.parametrize('tamper', ['number', 'fact', 'period', 'cause', 'binding', 'extra', 'stale'])
def test_editorial_rejects_fabrication_and_stale_bindings(tamper):
    sources, results, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    if tamper == 'number': response['pages'][0]['title'] += ' рост 999%'
    if tamper == 'fact': response['pages'][0]['title'] += ' {{fact:invented}}'
    if tamper == 'period': response['pages'][0]['interpretation'] += ' III квартал 2026'
    if tamper == 'cause': response['pages'][0]['interpretation'] += ' Выручка снизилась из-за партнёров.'
    if tamper == 'binding': response['pages'][0]['context_refs'] = ['foreign#4']
    if tamper == 'extra': response['pages'][0]['visual'] = {'kind': 'none'}
    if tamper == 'stale': response['input_sha256'] = '0' * 64
    with pytest.raises(ValueError): validate_editorial(response, sources, outline)


def test_editorial_replays_and_preserves_numeric_data():
    sources, results, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    edited, saved = validate_editorial(response, sources, outline)
    sources[0]['_business_editorial'] = dict(version=EDITORIAL_VERSION, response=saved)
    assert all(old.visual == new.visual and old.source_refs == new.source_refs and old.analysis_id == new.analysis_id
               for old, new in zip(outline.slides, edited.slides))
    assert all(len(old.bullets) == len(new.bullets) for old, new in zip(outline.slides, edited.slides))
    assert edited.slides[-1].bullets == outline.slides[-1].bullets
    assert edited.slides[-2].bullets[1] == outline.slides[-2].bullets[1]
    assert verify_outline(edited, results, sources)
    changed = edited.model_copy(deep=True); changed.slides[0].title += ' 999'
    with pytest.raises(ValueError): verify_outline(changed, results, sources)
    updated = inputs('version-b'); updated_results = prepare_business(updated, updated[0]['text'])
    updated[0]['_business_editorial'] = copy.deepcopy(sources[0]['_business_editorial'])
    with pytest.raises(ValueError): build_outline(updated_results, updated, 12)


def test_protected_narrative_preserves_current_verified_text_and_card_count():
    sources, _, outline = prepared()
    payload = editor_payload(sources, outline, sources[0]['text'])
    response = response_for(payload)
    for page in payload['pages']:
        original = next(item for item in sources[0]['_business']['story']['slides'] if item['id'] == page['id'])
        from designer.business_semantics import protected_interpretation
        assert page['verified_interpretation'] == protected_interpretation(original)
        assert all(token.startswith('fact:') for token in page['allowed_tokens'])
    edited, _ = validate_editorial(response, sources, outline)
    assert [slide.bullets for slide in edited.slides[1:]] == [slide.bullets for slide in outline.slides[1:]]
    assert edited.slides[0].bullets[1] == 'Причины изменения показателей не установлены.'
    assert edited.slides[0].bullets[2:] == outline.slides[0].bullets[2:]
    assert all('{{' not in slide.title for slide in edited.slides)


@pytest.mark.parametrize('kind', ['observation', 'caveat'])
@pytest.mark.parametrize('text', [
    'Причина роста установлена: клиенты массово перешли от конкурентов к нам.',
    '{{interpretation:decision}} Это гарантирует повышение прибыли компании.',
    '{{interpretation:channels}}',
])
def test_unprotected_observations_and_caveats_cannot_claim_validated_provenance(kind, text):
    sources, _, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    response['pages'][0].update(interpretation_kind=kind, interpretation=text)
    with pytest.raises(ValueError, match='защищённую интерпретацию'):
        validate_editorial(response, sources, outline)


@pytest.mark.parametrize('text', [
    'Уволить команду региональных продаж для гарантированного повышения прибыли.',
    '{{next_step:decision}} Затем закрыть региональные продажи.',
    '{{next_step:channels}}',
    '{{action:partners}}',
])
def test_next_step_cannot_add_actions_or_rebind_another_page(text):
    sources, _, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    response['pages'][0]['next_step'] = text
    with pytest.raises(ValueError, match='защищённый токен'):
        validate_editorial(response, sources, outline)


def test_new_explanation_requires_explicit_hypothesis_and_keeps_verified_next_step():
    sources, _, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    text = 'Гипотеза (не доказана): изменение спроса могло повлиять на динамику продаж.'
    response['pages'][0].update(interpretation_kind='hypothesis', interpretation=text)
    edited, _ = validate_editorial(response, sources, outline)
    assert edited.slides[0].bullets[1] == text
    assert edited.slides[0].bullets[2] == outline.slides[0].bullets[2]
    response['pages'][0]['interpretation'] = text.removeprefix('Гипотеза (не доказана): ')
    with pytest.raises(ValueError, match='Гипотезе нужна явная оговорка'):
        validate_editorial(response, sources, outline)


@pytest.mark.parametrize('title', [
    'Причина роста установлена: {{fact:qoq_net_growth}}',
    'Рост выручки объясняется рекламой: {{fact:qoq_net_growth}}',
    'Выручка — следствие рекламы: {{fact:qoq_net_growth}}',
])
def test_title_cannot_turn_unproven_cause_into_a_verified_fact(title):
    sources, _, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    response['pages'][0]['title'] = title
    with pytest.raises(ValueError, match='Неподтверждённая причина'):
        validate_editorial(response, sources, outline)


@pytest.mark.parametrize('direction', ['снизилась', 'падает', 'упала', 'уменьшилась', 'сократилась', 'ниже'])
def test_positive_growth_cannot_be_described_as_decline(direction):
    sources, _, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    response['pages'][0]['title'] = 'Выручка ' + direction + ': {{fact:qoq_net_growth}}'
    with pytest.raises(ValueError, match='Направление изменения'):
        validate_editorial(response, sources, outline)


@pytest.mark.parametrize('value,direction', [
    ('-2', 'выросла'), ('-2', 'растёт'), ('-2', 'увеличилась'), ('-2', 'повысилась'),
    ('0', 'выросла'), ('0', 'снизилась'), ('2', 'снизилась'),
])
def test_signed_margin_changes_reject_opposite_or_zero_direction(value, direction):
    from designer.business_editorial import _validate_directions
    with pytest.raises(ValueError, match='Направление изменения'):
        _validate_directions('Маржа ' + direction + ': {{fact:qoq_margin_change}}',
                             {'qoq_margin_change': {'exact': value}})


def test_direction_guard_binds_each_metric_and_excludes_absolute_values_and_attainment():
    from designer.business_editorial import _validate_directions
    facts = {'qoq_net_growth': {'exact': '24'}, 'qoq_margin_change': {'exact': '-0.4'},
             'q3_net': {'exact': '93563698'}, 'q3_net_plan_attainment': {'exact': '94'},
             'worst_region_plan_gap': {'exact': '-100'}}
    for text in (
        'Выручка выросла {{fact:qoq_net_growth}}, маржа снизилась {{fact:qoq_margin_change}}',
        'Рост выручки {{fact:qoq_net_growth}} при снижении маржи {{fact:qoq_margin_change}}',
        'Чистая выручка: {{fact:q3_net}}',
        'Выполнение плана: {{fact:q3_net_plan_attainment}}',
        'Отклонение от плана: {{fact:worst_region_plan_gap}}',
    ):
        _validate_directions(text, facts)


@pytest.mark.parametrize('title', [
    'Выручка снизилась до {{fact:q3_net}}',
    'Маржа снизилась: {{fact:qoq_net_growth}}',
    'Прибыль выросла: {{fact:qoq_net_growth}}',
    'План вырос: {{fact:qoq_net_growth}}',
    'Выполнение повысилось: {{fact:qoq_net_growth}}',
    'Выполнение плана повысилось до {{fact:q3_net_plan_attainment}}',
    'Рост выручки; значение {{fact:qoq_net_growth}}',
])
def test_directional_title_requires_matching_signed_change_in_same_clause(title):
    from designer.business_editorial import _validate_directions
    facts = {'qoq_net_growth': {'exact': '24'}, 'q3_net': {'exact': '93563698'},
             'q3_net_plan_attainment': {'exact': '94'}}
    with pytest.raises(ValueError, match='нужен факт изменения той же метрики'):
        _validate_directions(title, facts)


def test_unavailable_change_cannot_support_directional_claim():
    from designer.business_editorial import _validate_directions
    with pytest.raises(ValueError, match='нет рассчитанного значения'):
        _validate_directions('Рост выручки: {{fact:qoq_net_growth}}', {'qoq_net_growth': {'exact': None}})


def test_legacy_saved_editorial_requires_regeneration_instead_of_silent_migration():
    sources, _, outline = prepared()
    sources[0]['_business_editorial'] = dict(version='business-editorial-1', response={})
    with pytest.raises(ValueError, match='создайте презентацию заново из исходников'):
        replay_editorial(sources, outline)


@pytest.mark.parametrize('period', ['current', 'previous', 'year_ago'])
def test_model_period_tokens_cannot_relabel_current_facts(period):
    sources, results, outline = prepared()
    payload = editor_payload(sources, outline, sources[0]['text'])
    assert payload['selected_periods'] == sources[0]['_business']['analysis']['period_settings']
    assert not any(token.startswith('period:') for page in payload['pages'] for token in page['allowed_tokens'])
    response = response_for(payload)
    response['pages'][1]['title'] = 'Выручка за {{period:' + period + '}}: {{fact:q3_net}}'
    with pytest.raises(ValueError): validate_editorial(response, sources, outline)


@pytest.mark.parametrize('literal', [
    'двадцать процентов', 'двадцати', 'двадцатью', 'тридцать', 'сорока',
    'двенадцати', 'пятнадцать', 'четырёхсот', 'семьюдесятью', 'восьмидесяти',
    'девяноста', 'триста', 'одной', 'двумя', 'тремя', 'полтора', 'вдвое',
    'twenty percent', 'процентов', '%', 'рублей', 'май', 'прошлый год', 'II квартал'])
def test_number_words_units_and_period_literals_cannot_add_unsupported_claims(literal):
    sources, results, outline = prepared()
    response = response_for(editor_payload(sources, outline, sources[0]['text']))
    response['pages'][1]['title'] = 'Ориентир ' + literal + ': {{fact:q3_net}}'
    with pytest.raises(ValueError): validate_editorial(response, sources, outline)


def test_outline_success_survives_optional_condense_authorization_failure(monkeypatch, model_env):
    from designer.models import GenerateRequest, Outline
    from designer.pipeline import run_generation, balance_outline
    plan = {'title': 'План', 'slides': [{'title': 'Пилот',
        'bullets': ['Модель сформулировала этот длинный текст. ' * 35], 'source_refs': ['brief']}]}
    calls, snapshots, warnings = [], [], []
    def handler(request):
        calls.append(request)
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(plan)}}]}) if len(calls) == 1 else httpx.Response(401, json={'error': 'denied'})
    install_transport(monkeypatch, handler)
    async def factory(*args): return await ai.complete_json([], validate=Outline.model_validate)
    async def balance(value, template, warnings): return await balance_outline(value, template, capacity=200, warnings=warnings)
    async def render(request, outline, template, sources, *args):
        assert sources[0]['id'] == 'brief'  # Empty sources remain supported.
        snapshots.append(copy.deepcopy(sources[0]['_inference']))
        assert outline['slides'][0]['bullets'][0].startswith('Модель сформулировала')
        return {'id': 'fixture'}
    asyncio.run(run_generation(GenerateRequest(brief='Пилот и исходные материалы', template_id='fixture', slide_count=1),
        {}, [], outline_factory=factory, balance=balance, render_variant=render,
        variants=[{'id': 'classic'}, {'id': 'split'}, {'id': 'focus'}],
        on_stage=lambda *args: None, on_outline=lambda *args: None, on_variant=lambda *args: None, warnings=warnings))
    trace = snapshots[0]
    assert len(calls) == 2 and trace['attempts'] == trace['successful_calls'] == 1
    assert validated_content(trace) and trace['status'] == 'completed'
    assert trace['validation'] == 'passed' and not trace['fallback_used'] and not trace['error_code']
    assert trace['content_reused_for_variants'] == 3 and all(snapshot == trace for snapshot in snapshots)
    assert trace['condense']['attempts'] == 1 and trace['condense']['error_code'] == 'authorization'
    assert trace['condense']['fallback_used'] and any('ИИ не сократил' in warning for warning in warnings)


@pytest.mark.parametrize('primary_model,asset_status', [(True, 401), (False, 200)])
def test_optional_asset_calls_cannot_replace_primary_content_provenance(monkeypatch, model_env, primary_model, asset_status):
    from designer.models import GenerateRequest, Outline
    from designer.pipeline import run_generation
    plan = {'title': 'План', 'slides': [{'title': 'Пилот', 'bullets': ['Проверить ограничения']}]}
    snapshots, warnings, calls = [], [], []
    def handler(request):
        calls.append(request)
        if primary_model and len(calls) == 1:
            return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(plan)}}]})
        if asset_status != 200:
            return httpx.Response(asset_status, json={'error': 'fixture failure'})
        return httpx.Response(200, json={'choices': [{'message': {'content': '{"items":[]}'}}]})
    install_transport(monkeypatch, handler)
    async def factory(*args): return await ai.complete_json([], validate=Outline.model_validate)
    async def balance(value, template, warnings): return value
    async def label(*args): return await ai.complete_json([])
    async def render(request, outline, template, sources, *args):
        snapshots.append(copy.deepcopy(sources[0]['_inference']))
        return {'id': 'fixture'}
    request = GenerateRequest(brief='Пилот и исходные материалы', template_id='fixture', slide_count=1,
                              outline=None if primary_model else Outline.model_validate(plan))
    asyncio.run(run_generation(request, {}, [], outline_factory=factory, balance=balance, render_variant=render,
        variants=[{'id': 'classic'}], on_stage=lambda *args: None, on_outline=lambda *args: None,
        on_variant=lambda *args: None, warnings=warnings, label_assets=label))
    trace = snapshots[0]
    assert validated_content(trace) is primary_model
    assert trace['successful_calls'] == int(primary_model)
    assert trace['asset_tagging']['attempts'] == 1
    if asset_status == 200:
        assert trace['asset_tagging']['successful_calls'] == 1 and not trace['requested']
    else:
        assert trace['asset_tagging']['fallback_used'] and warnings


def test_browser_job_shares_one_actual_validated_answer_and_safe_retry(tmp_path, monkeypatch, model_env):
    calls = []
    failing = False
    def handler(request):
        calls.append(request)
        if failing: return httpx.Response(401, json={'error': 'denied'})
        payload = json.loads(json.loads(request.content)['messages'][1]['content'])
        return httpx.Response(200, json={'id': 'business-fixture', 'usage': {'total_tokens': 22},
            'choices': [{'message': {'content': json.dumps(draft_for(payload), ensure_ascii=False)}}]})
    install_transport(monkeypatch, handler)
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Обложка'; slide.placeholders[1].text = 'Образец'
    stream = io.BytesIO(); prs.save(stream)
    with TestClient(create_app(tmp_path)) as client:
        template = client.post('/api/v1/templates', files={'file': ('template.pptx', stream.getvalue())}).json()
        request = upload(client); request.update(template_id=template['id'], generate_images=False, model_content_generation=True)
        def generate():
            queued = client.post('/api/v1/generations', json=request)
            assert queued.status_code == 202, queued.text
            job = queued.json(); deadline = time.monotonic() + 80
            while job['status'] not in {'completed', 'failed'} and time.monotonic() < deadline:
                time.sleep(.03); job = client.get('/api/v1/jobs/' + job['id']).json()
            assert job['status'] == 'completed', job
            return job
        job = generate()
        assert len(calls) == 1 and job['inference']['successful_calls'] == 1
        assert job['inference']['content_reused_for_variants'] == 3
        records = [client.get('/api/v1/presentations/' + identifier).json() for identifier in job['presentation_ids']]
        original = records[0]
        snapshot = client.get(original['exports']['pptx']).content
        assert all(record['generation']['mode'] == 'llm' and record['generation']['inference'] == job['inference'] for record in records)
        exported = Presentation(io.BytesIO(snapshot))
        assert any('Приоритет decision' in shape.text for slide in exported.slides for shape in slide.shapes if shape.has_text_frame)
        slides = [slide['content'] for slide in original['deck']['slides']]
        edit = dict(revision=1, title=original['title'], variant=original['variant'], slides=slides)
        assert client.patch('/api/v1/presentations/' + original['id'], json=edit).status_code == 200
        assert client.get('/api/v1/presentations/' + original['id'] + '/export/pptx?revision=1').content == snapshot
        failing = True
        fallback = generate()
        assert len(calls) == 2 and fallback['inference']['fallback_used'] and not fallback['inference']['completed']
        assert fallback['inference']['error_code'] == 'authorization'
        failed_record = client.get('/api/v1/presentations/' + fallback['presentation_ids'][0]).json()
        assert failed_record['generation']['mode'] == 'deterministic_analytics'
        assert client.get('/api/v1/presentations/' + original['id'] + '/export/pptx?revision=1').content == snapshot
    with TestClient(create_app(tmp_path)) as client:
        reopened = client.get('/api/v1/presentations/' + original['id']).json()
        assert reopened['revision'] == 2 and reopened['generation']['inference']['response_id'] == 'business-fixture'
        assert reopened['deck']['slides'][0]['content']['title'] == slides[0]['title']
