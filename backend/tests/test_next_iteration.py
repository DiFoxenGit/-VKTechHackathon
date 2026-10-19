"""Mixed decks must keep canonical evidence and use pack/generated assets."""
import copy
import io
import time
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from designer.app import create_app
from designer.analytics_bridge import extract, analyse_sources, build_outline, verify_outline
from designer.assets import library
from designer.layout import compose
from designer.models import Outline
from designer.parsing import parse_template


def test_business_request_cannot_be_silently_reduced_to_gross_sum():
    from designer.analytics_bridge import validate_business_scope
    source={'datasets':[{'columns':['gross_revenue_rub','discount_rub','refund_rub','cogs_rub']}]}
    with pytest.raises(ValueError,match='D2'):
        validate_business_scope([source],'Рассчитай маржу и выполнение план-факта')
    validate_business_scope([source],'Покажи сумму gross_revenue_rub')


def test_dense_line_with_gaps_keeps_observations_visible():
    from designer.analytic_export import add_chart, display_number
    from pptx.util import Inches
    from pptx.enum.chart import XL_MARKER_STYLE
    values=[10 if i%3==0 else None for i in range(21)]
    visual={'kind':'line','categories':[str(i) for i in range(21)],'series':[{'name':'Выручка','values':values}],'unit':'RUB'}
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    add_chart(slide,visual,[Inches(.5),Inches(.5),Inches(8),Inches(4)],'Arial',['0077FF'],'243149')
    chart=next(s.chart for s in slide.shapes if s.has_chart)
    assert chart.series[0].marker.style==XL_MARKER_STYLE.CIRCLE
    assert list(chart.series[0].values)==values
    assert display_number('00123')=='00123'
    assert display_number('12989284.25')=='12\u00a0989\u00a0284,25'
    assert display_number('-0.25')=='-0,25'
    assert display_number('+1000.50')=='+1\u00a0000,50'


def template_bytes():
    deck = Presentation()
    deck.slide_width, deck.slide_height = 12192000, 6858000
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Образец заголовка'
    slide.placeholders[1].text = 'Чужие числа 999999'
    stream = io.BytesIO(); deck.save(stream)
    return stream.getvalue()


def test_mixed_layout_keeps_verified_objects_and_pack_provenance():
    template = parse_template(template_bytes(), 'template.pptx')
    table = extract('Регион,Выручка\nСевер,1200\nЮг,2100'.encode(), 'sales.csv')[0]
    sources = [{'id':'brief','_story':{'goal':'Сравнить выручку'}}, {'id':'data','datasets':[table]},
        {'id':'context','text':'Команда согласует роли.\n\nКоманда проверяет условия.\n\nКоманда учитывает ограничения.'}]
    results = analyse_sources(sources, 'Сравнить')
    outline = build_outline(results, sources, 12).model_dump()
    before = copy.deepcopy(outline)
    assets = library(template)
    # Reuse actual team artwork as a pack entry. No synthetic image substitutes.
    team = next(a for a in assets if a.get('ref','').endswith('illustrations/team.svg'))
    pack = {**team, 'source':'pack', 'family':'pack:test', 'id':'team-from-pack'}
    assets.insert(0, pack)
    for variant in ['classic','split','focus']:
        deck = compose(outline, template, variant, assets)
        assert verify_outline({'slides':[s['content'] for s in deck['slides']]}, results, sources)
        assert any(e.get('asset',{}).get('source')=='pack' for s in deck['slides'] for e in s['elements']), variant
        assert any(s.get('asset_usage') for s in deck['slides'])
        assert next(e['data'] for s in deck['slides'] if s['content'].get('analysis_id') for e in s['elements'] if e['id']=='visual') == next(s['visual'] for s in outline['slides'] if s.get('analysis_id'))
    assert outline == before


def test_generated_asset_stays_on_semantic_page_after_reorder():
    template = parse_template(template_bytes(), 'template.pptx')
    slides = Outline.model_validate({'title':'Колода','slides':[
        {'title':'Обложка','bullets':['Контекст']},
        {'title':'Команда продукта','bullets':['Команда согласует критерии пилота.']},
        {'title':'Ограничения','bullets':['Нельзя выводить эффект инвестиций без затрат.']},
        {'title':'Итог','bullets':['Решение после проверки.']}]}).model_dump()['slides']
    asset = {**next(a for a in library() if a.get('ref','').endswith('illustrations/team.svg')), 'source':'generated', 'id':'generated-team', 'slide':1, 'target_title':'Команда продукта'}
    slides[1],slides[2] = slides[2],slides[1]
    for variant in ['classic','split','focus']:
        deck = compose({'slides':slides},template,variant,[asset])
        page = next(s for s in deck['slides'] if s['content']['title']=='Команда продукта')
        assert any(e.get('asset',{}).get('id')=='generated-team' for e in page['elements'])


def test_image_unavailable_is_visible_retryable_and_persisted(tmp_path,monkeypatch):
    monkeypatch.delenv('DESIGNER_IMAGE_MODEL',raising=False)
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    with TestClient(create_app(tmp_path)) as client:
        template = client.post('/api/v1/templates',files={'file':('test.pptx',template_bytes())}).json()
        data = client.post('/api/v1/content-packs',files={'file':('data.csv',b'Category,Value\nA,10\nB,20')}).json()
        job = client.post('/api/v1/generations',json={'brief':'Сравнить данные','generate_images':True,'template_id':template['id'],'content_pack_ids':[data['id']]}).json()
        for _ in range(300):
            job=client.get('/api/v1/jobs/'+job['id']).json()
            if job['status'] in {'completed','failed'}: break
            time.sleep(.01)
        assert job['status']=='completed',job
        assert any('провайдер' in w for w in job['warnings'])
        for id in job['presentation_ids']:
            record=client.get('/api/v1/presentations/'+id).json()
            assert record['image_generation']['status']=='unavailable'
            retry=client.post(f'/api/v1/presentations/{id}/illustrations/retry',json={'revision':1})
            assert retry.status_code==200 and retry.json()['image_generation']['status']=='unavailable'
            assert client.get(f'/api/v1/presentations/{id}/slides/0/thumbnail?revision=99').status_code==409
            assert client.post(f'/api/v1/presentations/{id}/illustrations/retry',json={'revision':99}).status_code==409
    with TestClient(create_app(tmp_path)) as client:
        assert client.get('/api/v1/presentations/'+job['presentation_ids'][0]).json()['image_generation']['status']=='unavailable'
