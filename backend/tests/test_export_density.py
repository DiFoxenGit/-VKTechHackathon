"""Regressions found by rendering a real eight-slide model-generated deck."""
import shutil

import pytest

from designer.exporting import convert_pdf, export_pptx, value_axis_title
from designer.layout import compose
from designer.models import Outline
from designer.parsing import parse_template
from designer.render import text_runs
from tests.test_pipeline import template_bytes


def test_axis_keeps_unit_instead_of_repeating_a_slide_conclusion():
    visual = {'unit': 'мин', 'series': [{'name': 'Автоматическая вёрстка повышает эффективность и соблюдение стиля, как показал пилот'}]}
    assert value_axis_title(visual) == 'мин'
    assert value_axis_title({'unit': 'мин', 'series': [{'name': 'Сборка'}]}) == 'Сборка, мин'


@pytest.mark.skipif(not (shutil.which('libreoffice') or shutil.which('soffice')), reason='LibreOffice required')
def test_dense_table_stays_inside_its_frame_after_pdf_conversion(tmp_path):
    source = tmp_path / 'template.pptx'
    source.write_bytes(template_bytes())
    template = parse_template(source.read_bytes(), 'sample.pptx')
    outline = Outline.model_validate({'title': 'Dense table', 'slides': [{
        'title': 'Seven rows fit the planned frame', 'source_refs': ['brief'],
        'visual': {'kind': 'table', 'columns': ['Metric', 'Value'],
                   'rows': [[f'Metric {index}', f'Value {index}'] for index in range(6)]},
    }]})
    deck = compose(outline.model_dump(), template, 'classic')
    visual = next(element for element in deck['slides'][0]['elements'] if element['kind'] == 'table')
    visual['box'] = [36, 200, 600, 133]
    pptx, pdf = tmp_path / 'dense.pptx', tmp_path / 'dense.pdf'
    export_pptx(source, template, deck, pptx)
    convert_pdf(pptx, pdf)
    runs = [run for run in text_runs(pdf, 0) if 'Metric 5' in run['text'] or 'Value 5' in run['text']]
    assert runs, 'All table rows must remain visible'
    assert max(run['y'] + run['height'] for run in runs) <= 334


def test_long_category_names_all_survive_real_pdf(tmp_path):
    """Office must render every category, not merely keep it in chart XML."""
    import os
    from pathlib import Path
    from pptx import Presentation
    from pptx.util import Pt
    from pypdf import PdfReader
    from designer.analytic_export import add_chart
    if not (os.getenv('DESIGNER_SOFFICE') or shutil.which('libreoffice') or shutil.which('soffice')
            or Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists()):
        pytest.skip('LibreOffice required')
    labels = ['Москва и Московская область', 'Санкт-Петербург и Ленинградская область',
              'Поволжье', 'Уральский федеральный округ',
              'Сибирь и Дальневосточный федеральный округ',
              'Южный и Северо-Кавказский федеральные округа']
    values = [600000, -2100000, -2100000, -2300000, -2400000, -7200000]
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(720), Pt(405)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_chart(slide, {'kind':'bar', 'categories':labels, 'series':[{'name':'План-факт', 'values':values}],
                     'unit':'Рубли'}, tuple(Pt(n) for n in (40,117,640,154)),
              'Arial', ['0077FF'], '111111', 'FFFFFF')
    source, pdf = tmp_path/'categories.pptx', tmp_path/'categories.pdf'
    deck.save(source)
    reopened = Presentation(source).slides[0].shapes[0].chart
    assert list(reopened.series[0].values) == values
    assert [c.label for c in reopened.plots[0].categories] == labels
    convert_pdf(source, pdf)
    text = ' '.join(PdfReader(pdf).pages[0].extract_text().split())
    assert all(label in text for label in labels), text
