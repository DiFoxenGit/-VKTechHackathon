"""Acceptance evidence must reject configuration-only and fallback false positives."""
import copy
import io
import json
from pathlib import Path
import sys

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
from acceptance_evidence import (response_summary, sanitized_inference, validate_inference,
    verify_business, verify_independent_jobs, verify_office_exports)
from check_inference import probe_profile, validate_probe
from check_parallel_generation import business_inputs


def valid_trace():
    return dict(requested=True, attempted=True, completed=True, provider="vk", model="qwen38-27b-fp8",
        attempts=1, successful_calls=1, duration_seconds=1.2, validation="passed", fallback_used=False,
        status="completed", error_code=None, response_id="response-one", content_reused_for_variants=3,
        calls=[dict(attempt=1, status_code=200, duration_seconds=1.2, validation="passed", response_id="response-one")])


def test_probe_fallback_requires_opt_in_and_exact_configured_model(monkeypatch):
    from designer.inference import Inference
    primary = Inference(profile='vk', url='https://primary.invalid/v1', model='qwen-fixture', key='fixture')
    monkeypatch.setenv('DESIGNER_FALLBACK_PROVIDER', 'genapi')
    monkeypatch.setenv('DESIGNER_FALLBACK_MODEL', 'gemma-4-31b')
    monkeypatch.setenv('DESIGNER_FALLBACK_API_KEY', 'fixture-not-live')
    trace = {'provider': 'genapi', 'model': 'gemma-4-31b'}
    assert probe_profile(primary, trace) is primary
    assert probe_profile(primary, trace, True).profile == 'genapi'
    assert probe_profile(primary, {**trace, 'model': 'unconfigured-model'}, True) is primary
    monkeypatch.delenv('DESIGNER_FALLBACK_PROVIDER')
    assert probe_profile(primary, trace, True) is primary


@pytest.mark.parametrize('arguments', [
    ['--accept-fallback', '--attempts', '2'],
    ['--accept-fallback', '--service', 'http://localhost', '--template-id', 'fixture'],
])
def test_budgeted_fallback_probe_rejects_modes_with_unbounded_extra_calls(monkeypatch, arguments):
    import check_inference
    monkeypatch.setattr(sys, 'argv', ['check_inference.py', *arguments])
    monkeypatch.setattr(check_inference, 'load_env', lambda *args: pytest.fail('must reject before loading credentials'))
    with pytest.raises(SystemExit) as error:
        check_inference.main()
    assert error.value.code == 2


@pytest.mark.parametrize("change", [
    {"attempted": False}, {"completed": False}, {"attempts": 0}, {"successful_calls": 0},
    {"validation": "not_run"}, {"fallback_used": True}, {"status": "fallback"},
    {"error_code": "rate_limit"}, {"calls": []}, {"duration_seconds": float("nan")},
    {"content_reused_for_variants": 0}, {"calls": [{"status_code": 500, "validation": "passed"}]},
    {"calls": [{"status_code": 200, "validation": "not_run"}]},
    {"successful_calls": True}, {"duration_seconds": True},
])
def test_live_false_positives_are_rejected(change):
    trace = valid_trace()
    trace.update(change)
    with pytest.raises(ValueError):
        validate_inference(trace, live=True)


def test_live_requires_real_selected_profile_but_allows_validated_retries():
    trace = valid_trace()
    trace["calls"].insert(0, dict(attempt=1, status_code=429, duration_seconds=.2, validation="failed", error_code="rate_limit"))
    trace["attempts"] = 2
    trace["calls"][1]["attempt"] = 2
    assert validate_inference(trace, live=True, provider="vk", model="qwen38-27b-fp8")["attempts"] == 2
    with pytest.raises(ValueError):
        validate_inference(trace, live=True, provider="openai")


@pytest.mark.parametrize("value", [{"ok": 1}, {"ok": True, "extra": "ignored"}, [True], {"ok": "true"}])
def test_transport_probe_requires_exact_boolean_schema(value):
    with pytest.raises(ValueError):
        validate_probe(value)
    assert validate_probe({"ok": True}) == {"ok": True}


def test_configured_model_name_does_not_prove_call_and_secrets_are_not_exported():
    with pytest.raises(ValueError):
        validate_inference({"model": "qwen38-27b-fp8", "llm_configured": True}, live=True)
    trace = valid_trace()
    trace.update(headers={"Authorization": "SECRET"}, messages=["SECRET"], api_key="SECRET")
    trace["calls"][0]["response_body"] = "SECRET"
    assert "SECRET" not in json.dumps(sanitized_inference(trace))


def test_response_summary_counts_jobs_once_and_keeps_unknown_identity_honest():
    one = valid_trace()
    two = copy.deepcopy(one)
    two["calls"][0]["response_id"] = "response-two"
    summary = response_summary([one, two])
    assert summary["validated_model_calls"] == 2
    assert summary["distinct_provider_response_ids"] == 2
    assert summary["distinct_response_content_hashes"] is None
    two["calls"][0].pop("response_id")
    assert response_summary([one, two])["distinct_provider_response_ids"] is None


def test_explicit_deterministic_matrix_cannot_hide_model_use():
    from designer.inference_trace import new_trace
    assert validate_inference(new_trace(False), live=False)["attempts"] == 0
    with pytest.raises(ValueError):
        validate_inference(valid_trace(), live=False)


def test_independent_jobs_do_not_require_optional_provider_ids_or_unique_text():
    jobs = []
    for number in range(7):
        trace = valid_trace()
        trace["calls"][0].pop("response_id")
        trace["response_id"] = None
        trace["calls"][0]["response_sha256"] = "a" * 64  # Same valid text may be returned independently.
        jobs.append({"job_id": str(number), "inference": trace})
    result = verify_independent_jobs(jobs, expected_count=7)
    assert result["validated_model_calls"] == 7
    assert result["distinct_provider_response_ids"] is None
    assert result["distinct_response_content_hashes"] == 1
    assert result["provider_response_identity"] == "unavailable_or_partial"
    jobs[-1]["job_id"] = "0"
    with pytest.raises(ValueError):
        verify_independent_jobs(jobs, expected_count=7)


def test_duplicate_returned_provider_ids_cannot_pass_as_independent_responses():
    jobs = [{"job_id": str(number), "inference": valid_trace()} for number in range(7)]
    with pytest.raises(ValueError, match="response ID was reused"):
        verify_independent_jobs(jobs, expected_count=7)


def test_all_fifteen_oracle_facts_and_source_evidence_are_checked():
    folder, settings, _, _, expected = business_inputs()
    proof = json.loads((folder / "analysis.json").read_text(encoding="utf8"))
    record = {"business_proof": proof}
    assert verify_business(record, expected, settings)["independent_numeric_controls"] == 15
    proof["facts"]["q3_net"]["exact"] = "94823698"
    with pytest.raises(ValueError, match="Independent fact mismatch"):
        verify_business(record, expected, settings)
    proof["facts"]["q3_net"]["exact"] = "93563698"
    proof["facts"]["q3_net"]["evidence"][0]["row_ids"] = []
    with pytest.raises(ValueError, match="source rows absent"):
        verify_business(record, expected, settings)


@pytest.mark.parametrize("version", ["baseline", "version-b"])
def test_alternate_period_has_self_consistent_context_and_independent_oracle(version):
    from designer.analytics.business_conditions import ContextConditions
    folder, settings, conditions, context, expected = business_inputs(version, "alternate")
    ContextConditions.model_validate(conditions)
    assert settings["current"] == ["2026-07", "2026-09"]
    assert conditions["reporting_period"]["quote"] in context
    assert conditions["reporting_period"]["value"] == settings["current"]
    assert all(answer["period"] == settings["current"] for answer in expected["answers"])
    assert "октябрь — декабрь 2026" in context


def test_alternate_period_prepares_with_production_contract_without_any_model_call():
    from designer.analytics_bridge import extract
    from designer.business_bridge import prepare_business
    from check_parallel_generation import DATA
    folder, settings, conditions, context, expected = business_inputs("baseline", "alternate")
    sources = [{"id": "brief", "text": "Прибыль и маржа, план-факт", "business_request": {
        "settings": settings, "conditions": conditions, "context_source_id": "context"}}]
    for name in ("sales.csv", "plan.csv", "products.csv", "regions.csv"):
        path = (DATA if name in {"products.csv", "regions.csv"} else folder) / name
        sources.append({"id": name, "text": "", "datasets": extract(path.read_bytes(), name)})
    sources.append({"id": "context", "text": context})
    assert len(prepare_business(sources, sources[0]["text"])) == 12
    proof = sources[0]["_business"]["analysis"]
    assert verify_business({"business_proof": proof}, expected, settings)["independent_numeric_controls"] == 15
    assert proof["facts"]["q3_net"]["exact"] == "116241720"


def export_fixture(values):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Revenue", values)
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
    table = slide.shapes.add_table(2, 1, Inches(6), Inches(1), Inches(2), Inches(2)).table
    table.cell(0, 0).text = "Revenue"
    table.cell(1, 0).text = "93\u00a0563\u00a0698,25"
    raw = io.BytesIO()
    presentation.save(raw)
    pdf = io.BytesIO()
    document = PdfWriter()
    page = document.add_blank_page(width=800, height=600)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): font})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 20 Td (Revenue 93563698.25) Tj ET")
    page[NameObject("/Contents")] = content
    document.write(pdf)
    return raw.getvalue(), pdf.getvalue()


def test_native_chart_numbers_are_compared_instead_of_only_object_counts():
    record = {"deck": {"slides": [{"content": {"visual": {
        "kind": "bar", "series": [{"values": [100, -50]}]}}}]}}
    pptx, pdf = export_fixture([100, -50])
    assert verify_office_exports(pptx, pdf, record)["native_charts"] == 1
    record["deck"]["slides"][0]["content"]["visual"]["series"][0]["values"][0] = 101
    with pytest.raises(ValueError, match="chart numbers differ"):
        verify_office_exports(pptx, pdf, record)
    record["deck"]["slides"][0]["content"]["visual"] = {"kind": "table", "rows": [["93563698.25"]]}
    assert verify_office_exports(pptx, pdf, record)["native_tables"] == 1
