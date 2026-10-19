"""Business evidence must survive real service adaptation and persisted replay."""
import copy
from pathlib import Path
import pytest
from designer.analytics_bridge import extract, build_outline, verify_outline
from designer.business_bridge import prepare_business

DATA=Path(__file__).resolve().parents[2]/'samples/business-acceptance/data-v1'


def sources(directory=DATA):
    result=[{'id':'brief','text':'Развитие продуктов: маржа и план-факт'}]
    for name in ('sales','plan','products','regions'):
        path=directory/(name+'.csv')
        if not path.exists():path=DATA/path.name
        result.append({'id':name,'datasets':extract(path.read_bytes(),path.name),'text':''})
    result.append({'id':'context','text':(DATA/'business-context.txt').read_text(encoding='utf8')})
    return result


def test_full_business_sources_yield_twelve_pages_and_replay():
    inputs=sources();results=prepare_business(inputs,inputs[0]['text'])
    outline=build_outline(results,inputs,12)
    assert len(outline.slides)==12
    assert all(s.analysis_id.startswith('business:') for s in outline.slides)
    assert verify_outline(outline,results,inputs)
    assert inputs[0]['_business']['analysis']['facts']['q3_net']['exact']=='116241720'
    trend=next(s for s in outline.slides if s.analysis_id.startswith('business:trend:'))
    assert len(trend.visual.categories)==12 and 'последние 12' in trend.notes
    assert len(inputs[0]['_business']['analysis']['views']['monthly_trend']['dataset']['rows'])==21


@pytest.mark.parametrize('target',['source','fact','proof','story','result','slide'])
def test_tampering_cannot_cross_business_boundary(target):
    inputs=sources();results=prepare_business(inputs,inputs[0]['text']);outline=build_outline(results,inputs,12).model_dump()
    if target=='source':inputs[1]['datasets'][0]['contract']['rows'][0][4]='999999'
    if target=='fact':inputs[0]['_business']['analysis']['facts']['q3_net']['exact']='999999'
    if target=='proof':inputs[0]['_business']['sources']['sales']['source_sha256']='0'*64
    if target=='story':inputs[0]['_business']['story']['slides'][0]['answer']='Подменённый вывод'
    if target=='result':results[0]['contract']['presentation']['title']='Подменённый вывод'
    if target=='slide':outline['slides'][0]['title']='Подменённый вывод'
    with pytest.raises(ValueError):verify_outline(outline,results,inputs)


def test_duplicate_equivalent_inputs_and_missing_sources_are_explained():
    inputs=sources();inputs.append(copy.deepcopy(inputs[1]))
    with pytest.raises(ValueError,match='ровно одна таблица'):prepare_business(inputs,inputs[0]['text'])
    inputs=sources();inputs.pop(2)
    with pytest.raises(ValueError,match='plan'):prepare_business(inputs,inputs[0]['text'])


def test_source_update_requires_fresh_business_plan():
    first=sources();results=prepare_business(first,first[0]['text']);outline=build_outline(results,first,12)
    updated=sources(DATA/'updated');new_results=prepare_business(updated,updated[0]['text'])
    assert updated[0]['_business']['analysis']['facts']['q3_net']['exact']=='117501720'
    with pytest.raises(ValueError):verify_outline(outline,new_results,updated)


def test_http_full_materials_analysis_and_raw_download(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from designer.app import create_app
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    with TestClient(create_app(tmp_path)) as client:
        ids=[]
        # Large JSON exceeds the old text-only parser's limit; complete cells are retained.
        for name in ('sales.json','plan.csv','products.csv','regions.csv','business-context.txt'):
            data=(DATA/name).read_bytes()
            response=client.post('/api/v1/content-packs',files={'file':(name,data)})
            assert response.status_code==201,response.text
            pack=response.json();ids.append(pack['id'])
            assert client.get(f"/api/v1/content-packs/{pack['id']}/source").content==data
            if name=='sales.json':
                assert len(pack['datasets'][0]['rows'])==1512
                assert len(pack['text'])<10000 and 'предпросмотр' in pack['text']
        response=client.post('/api/v1/analysis',json={'brief':'Прибыль и маржа, план-факт','template_id':'analysis','content_pack_ids':ids})
        assert response.status_code==200,response.text
        result=response.json()
        assert result['business'] and len(result['items'])==12
        outline=client.post('/api/v1/outlines',json={'brief':'Прибыль и маржа, план-факт','slide_count':12,'content_pack_ids':ids})
        assert outline.status_code==200,outline.text
        assert len(outline.json()['slides'])==12
        assert all(r['business_explanation']['formula'] and r['packId']==ids[0] for r in result['items'])


def test_negative_bar_sign_survives_pdf_render(tmp_path):
    import shutil
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from designer.analytic_export import add_chart
    from designer.exporting import convert_pdf,render_slides
    if not (shutil.which('soffice') or shutil.which('libreoffice')):pytest.skip('LibreOffice required for sign geometry regression')
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    visual={'kind':'bar','categories':['Long positive category for horizontal','Long negative category for horizontal'],
        'series':[{'name':'Gap','values':[3000000,-4000000]}]}
    add_chart(slide,visual,[Inches(.5),Inches(.5),Inches(9),Inches(5)],'Arial',['0077FF'],'000000')
    chart=next(s.chart for s in slide.shapes if s.has_chart)
    assert chart.series[0].invert_if_negative is False
    from pptx.enum.chart import XL_LABEL_POSITION
    assert chart.series[0].points[1].data_label.position == XL_LABEL_POSITION.OUTSIDE_END
    assert chart.value_axis.minimum_scale < -4000000
    from pptx.oxml.ns import qn
    for point in chart.series[0].points:
        label=point.data_label._dLbl
        assert label.find(qn('c:numFmt')).get('sourceLinked')=='0'
        assert label.find(qn('c:showVal')).get('val')=='1'
    path=tmp_path/'negative.pptx';prs.save(path);convert_pdf(path,path.with_suffix('.pdf'))
    image=Image.open(BytesIO(render_slides(path.with_suffix('.pdf'),1000)[0])).convert('RGB')
    blue=[]
    for y in range(image.height):
        xs=[x for x in range(image.width) if (lambda p:p[2]>180 and p[2]-p[0]>70 and p[1]>80)(image.getpixel((x,y)))]
        if len(xs)>30:blue.append((y,sum(xs)/len(xs)))
    groups=[]
    for y,x in blue:
        if not groups or y>groups[-1][-1][0]+1:groups.append([])
        groups[-1].append((y,x))
    assert len(groups)==2
    # Horizontal charts list the last (negative) category above the positive one.
    assert sum(x for _,x in groups[0])/len(groups[0]) < sum(x for _,x in groups[1])/len(groups[1])


def test_generated_fixture_is_shared_placed_saved_and_reopened(tmp_path,monkeypatch):
    import io,time
    from fastapi.testclient import TestClient
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from designer.app import create_app
    from designer import ai
    monkeypatch.delenv('DESIGNER_TEMPLATE_DIR',raising=False)
    monkeypatch.setenv('DESIGNER_IMAGE_BASE_URL','https://example.invalid/v1')
    monkeypatch.setenv('DESIGNER_IMAGE_MODEL','fixture-only-not-a-live-provider')
    monkeypatch.setenv('DESIGNER_IMAGE_LIMIT','1')
    calls=[]
    async def draw(prompt,size=None):
        calls.append(prompt)
        image=Image.new('RGB',(160,100),'#0077FF');stream=io.BytesIO();image.save(stream,format='PNG');return stream.getvalue()
    monkeypatch.setattr(ai,'draw_image',draw)
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text='Тема';slide.placeholders[1].text='Образец'
    stream=io.BytesIO();prs.save(stream)
    with TestClient(create_app(tmp_path)) as client:
        template=client.post('/api/v1/templates',files={'file':('test.pptx',stream.getvalue())}).json()
        ids=[]
        for name in ('sales.csv','plan.csv','products.csv','regions.csv','business-context.txt'):
            response=client.post('/api/v1/content-packs',files={'file':(name,(DATA/name).read_bytes())})
            assert response.status_code==201,response.text
            ids.append(response.json()['id'])
        job=client.post('/api/v1/generations',json={'brief':'Маржа и прибыль: сгенерируй иллюстрацию команды','slide_count':12,
            'template_id':template['id'],'content_pack_ids':ids,'generate_images':True}).json()
        deadline=time.monotonic()+60
        while job['status'] not in {'completed','failed'} and time.monotonic()<deadline:
            time.sleep(.02);job=client.get('/api/v1/jobs/'+job['id']).json()
        assert job['status']=='completed',job
        assert len(calls)==1
        owners=set()
        for identifier in job['presentation_ids']:
            record=client.get('/api/v1/presentations/'+identifier).json()
            assert record['image_generation']['status']=='ready'
            owners.add(record['illustrations_id'])
            assert any(e.get('asset',{}).get('source')=='generated' for s in record['deck']['slides'] for e in s['elements'])
            response=client.patch('/api/v1/presentations/'+identifier,json={'revision':1,'title':record['title'],'variant':record['variant'],
                'slides':[s['content'] for s in record['deck']['slides']]})
            assert response.status_code==200,response.text
            assert response.json()['revision']==2
        assert len(owners)==1
    with TestClient(create_app(tmp_path)) as client:
        record=client.get('/api/v1/presentations/'+identifier).json()
        assert record['image_generation']['status']=='ready' and record['revision']==2
        assert any(e.get('asset',{}).get('source')=='generated' for s in record['deck']['slides'] for e in s['elements'])
