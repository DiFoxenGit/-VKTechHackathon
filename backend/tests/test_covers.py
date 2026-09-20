"""P0-2: cover recognition, brand-aware selection and native cover export."""

import copy
import io
import os
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from designer.audit import audit
from designer.exporting import export_pptx
from designer.layout import candidate_patterns, choose_pattern, compose
from designer.models import Outline
from designer.parsing import branding_score, classify_patterns, parse_template


def save(deck):
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def cover_template(layout_only=False, image_background=False):
    deck = Presentation()
    layout = deck.slide_layouts[0]
    title, subtitle = layout.placeholders[0], layout.placeholders[1]
    title.left, title.top, title.width, title.height = [Inches(v) for v in (.6, 3.9, 6, 1.4)]
    subtitle.left, subtitle.top, subtitle.width, subtitle.height = [Inches(v) for v in (.6, 5.5, 6, .6)]
    for shape, size in ((title, 44), (subtitle, 20)):
        shape.text_frame.paragraphs[0].font.size = Pt(size)
    slide = deck.slides.add_slide(layout)
    logo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(.4), Inches(1), Inches(.3))
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    logo._element.nvSpPr.cNvPr.set('id', str(max(s.shape_id for s in layout.shapes) + 1))
    layout.shapes._spTree.insert_element_before(logo._element, 'p:extLst')
    if layout_only:
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('41006B')
    if image_background:
        image = io.BytesIO()
        Image.new('RGB', (160, 120), '#41006B').save(image, format='PNG')
        if image_background == 'picture':
            slide.shapes.add_picture(io.BytesIO(image.getvalue()), 0, 0, deck.slide_width, deck.slide_height)
            return deck
        _, rid = slide.part.get_or_add_image_part(io.BytesIO(image.getvalue()))
        bg = slide._element.cSld.bg
        for node in list(bg):
            bg.remove(node)
        props, fill, blip = OxmlElement('p:bgPr'), OxmlElement('a:blipFill'), OxmlElement('a:blip')
        from pptx.oxml.ns import qn
        blip.set(qn('r:embed'), rid)
        fill.append(blip)
        stretch = OxmlElement('a:stretch')
        stretch.append(OxmlElement('a:fillRect'))
        fill.append(stretch)
        props.append(fill)
        bg.append(props)
    return deck


def plan():
    return Outline.model_validate({'title': 'Тест', 'slides': [
        {'title': 'Название инициативы', 'bullets': ['Защита перед руководством'],
         'visual': {'kind': 'none'}, 'source_refs': []}
    ]}).model_dump()


@pytest.mark.parametrize('layout_only', [False, True])
def test_empty_cover_slots_inherit_geometry_and_type(layout_only):
    template = parse_template(save(cover_template(layout_only)), 'unseen-name.pptx')
    cover = template['patterns'][0]
    assert cover['role'] == 'cover'
    assert [s['role'] for s in cover['slots']] == ['title', 'body']
    assert cover['slots'][0]['style']['size'] == 44
    assert cover['title_box']['y'] == pytest.approx(3.9 / 7.5)
    assert all(s['length'] == 0 for s in cover['slots'])


def test_speaker_in_subtitle_does_not_turn_cover_into_team_page():
    deck = cover_template()
    deck.slides[0].shapes.title.text = 'Новая платформа'
    deck.slides[0].placeholders[1].text = 'Имя спикера, команда продукта'
    assert parse_template(save(deck), 'anything.pptx')['patterns'][0]['role'] == 'cover'


def test_repeated_dividers_are_not_reusable_covers():
    cover = parse_template(save(cover_template()), 'any.pptx')['patterns'][0]
    content = copy.deepcopy(cover)
    content['title_box'] = {'x': .06, 'y': .1, 'w': .88, 'h': .1}
    content['slots'][0]['box'] = content['title_box']
    content['slots'][0]['style']['size'] = 24
    patterns = [copy.deepcopy(p) for p in (cover, content, content, cover, content, cover)]
    for i, pattern in enumerate(patterns):
        pattern['index'] = i
    classify_patterns(patterns)
    assert [p['role'] for p in patterns] == ['cover', 'content', 'content', 'section', 'content', 'section']


def test_visible_brand_wins_equal_capacity_and_survives_candidate_filter():
    base = parse_template(save(cover_template()), 'any.pptx')['patterns'][0]
    base.update(role='content', branding_score=0, decoration_count=0, reserved=[])
    base['title_box'] = {'x': .06, 'y': .1, 'w': .88, 'h': .1}
    patterns = [dict(base, index=i) for i in range(4)]
    # A tiny logo must not disappear just because three plain pages are clean.
    patterns[-1].update(branding_score=.2, decoration_count=1)
    pool = candidate_patterns(patterns)
    assert patterns[-1] in pool
    selected, _ = choose_pattern(pool, plan()['slides'][0], [], position=1, total=3)
    assert selected['index'] == 3


def test_two_slide_deck_uses_cover_and_closing():
    template = parse_template(save(cover_template()), 'any.pptx')
    template['patterns'].append(dict(template['patterns'][0], index=1, role='closing'))
    outline = plan()
    outline['slides'].append(dict(outline['slides'][0], title='Просим решение о запуске'))
    assert [s['pattern_index'] for s in compose(outline, template, 'classic')['slides']] == [0, 1]


def test_transparent_art_is_sampled_over_the_slide_fill():
    source = cover_template()
    pixels = io.BytesIO()
    Image.new('RGBA', (160, 120), (255, 255, 255, 0)).save(pixels, format='PNG')
    source.slides[0].shapes.add_picture(io.BytesIO(pixels.getvalue()), 0, 0, source.slide_width, source.slide_height)
    template = parse_template(save(source), 'anything.pptx')
    assert template['patterns'][0]['title_background']['color'] == '41006B'
    assert compose(plan(), template, 'classic')['slides'][0]['elements'][0]['color'] == 'FFFFFF'


def test_branding_ignores_empty_frames_and_off_canvas_art():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.1), Inches(.1), Inches(1), Inches(.3))
    off = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(-1), Inches(1), Inches(.2))
    off.fill.solid()
    off.fill.fore_color.rgb = RGBColor.from_string('FF0000')
    args = (slide, deck.slide_width, deck.slide_height, {'color': 'FFFFFF'})
    assert branding_score(*args) == 0
    off.top = Inches(.1)
    assert branding_score(*args) > 0


@pytest.mark.parametrize('variant', ['classic', 'split', 'focus'])
@pytest.mark.parametrize('image_background', [True, 'picture'])
def test_image_cover_keeps_native_art_and_editable_title(tmp_path, variant, image_background):
    data = save(cover_template(layout_only=True, image_background=image_background))
    template = parse_template(data, 'unknown.pptx')
    deck = compose(plan(), template, variant)
    slide = deck['slides'][0]
    assert slide['native_cover'] and not slide['needs_scrim']
    assert slide['background'] == '41006B'
    assert slide['elements'][0]['box'][1] == pytest.approx(3.9 * 72)
    assert all(e['color'] == 'FFFFFF' for e in slide['elements'])
    assert not [i for i in audit(deck, template, [])['issues'] if i['severity'] == 'error']
    source, output = tmp_path / 'source.pptx', tmp_path / 'result.pptx'
    source.write_bytes(data)
    export_pptx(source, template, deck, output)
    result = Presentation(output).slides[0]
    assert result._element.cSld.bg.xpath('.//a:blip') or image_background == 'picture'
    assert len(result.shapes) == len(slide['elements']) + (image_background == 'picture')  # no scrim
    assert any(s.has_text_frame and s.text == plan()['slides'][0]['title'] for s in result.shapes)
    assert len(result.slide_layout.shapes) == len(cover_template().slide_layouts[0].shapes)


CORPORATE_DIR = Path(os.getenv('DESIGNER_TEMPLATE_DIR') or Path(__file__).resolve().parents[2] / 'templates')
CORPORATE_FILES = sorted(CORPORATE_DIR.glob('*.pptx'))


@pytest.mark.parametrize('path', CORPORATE_FILES or [None], ids=lambda p: p.name if p else 'no-templates')
def test_supplied_template_has_cover_and_uses_it(path):
    if path is None:
        pytest.skip(f'No .pptx templates in {CORPORATE_DIR}')
    # Deliberately hide the real filename from classification.
    template = parse_template(path.read_bytes(), 'unknown.pptx')
    assert any(p['role'] == 'cover' for p in template['patterns']), path.name
    for variant in ('classic', 'split', 'focus'):
        deck = compose(plan(), template, variant)
        first = template['patterns'][deck['slides'][0]['pattern_index']]
        assert first['role'] == 'cover'
