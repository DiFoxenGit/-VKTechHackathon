"""Slot capacity and signed-chart regressions from the final business audit."""
import copy
import io

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from designer.analytic_export import add_chart
from designer.exporting import export_pptx
from designer.layout import card_slots, compose
from designer.models import Outline
from designer.parsing import parse_template


def card_template(counts=(0, 2, 3, 4)):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for count in counts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.6), Inches(.5), Inches(12), Inches(.8))
        title.text = 'Заголовок образца'
        title.text_frame.paragraphs[0].font.size = Pt(30)
        for i in range(count or 1):
            width = 12 / (count or 1)
            x = .6 + i * width
            if count:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2), Inches(width-.15), Inches(3.4))
                card.name = f'card-background-{i}'
                card.fill.solid()
                card.fill.fore_color.rgb = RGBColor.from_string('DDEEFF')
            body = slide.shapes.add_textbox(Inches(x+.08), Inches(2.1), Inches(width-.31), Inches(3))
            body.text = f'Пример блока {i}'
            body.text_frame.paragraphs[0].font.size = Pt(20)
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()


def content(count):
    outline = Outline.model_validate({'title': 'Решения', 'slides': [{
        'title': 'Условия реализации',
        'bullets': [f'Самостоятельная мысль номер {i}' for i in range(count)],
    }]}).model_dump()
    outline['content_pages'] = True
    return outline


@pytest.mark.parametrize('count', [1, 2, 3, 4])
@pytest.mark.parametrize('variant', ['classic', 'split', 'focus'])
def test_block_count_selects_filled_grid_before_variant_columns(count, variant):
    template = parse_template(card_template(), 'sample.pptx')
    for pattern in template['patterns']:
        pattern['role'] = 'content'
    outline = content(count)
    before = copy.deepcopy(outline)
    page = compose(outline, template, variant, assets=[])['slides'][0]
    pattern = template['patterns'][page['pattern_index']]
    cards = [e for e in page['elements'] if e['id'].startswith('card_')]
    if count >= 2:
        assert len(card_slots(pattern['slots'], 1, 1)) == count
        assert len(cards) == count
        assert [e['text'] for e in cards] == outline['slides'][0]['bullets']
    else:
        assert not card_slots(pattern['slots'], 1, 1)
    assert outline == before


@pytest.mark.parametrize('count', [1, 2, 4])
@pytest.mark.parametrize('variant', ['classic', 'split', 'focus'])
def test_unavoidable_wrong_capacity_drops_only_card_containers(tmp_path, count, variant):
    source = tmp_path / 'source.pptx'
    source.write_bytes(card_template((3,)))
    template = parse_template(source.read_bytes(), 'sample.pptx')
    deck = compose(content(count), template, variant, assets=[])
    target = tmp_path / 'result.pptx'
    export_pptx(source, template, deck, target)
    slide = Presentation(target).slides[0]
    assert not any(s.name.startswith('card-background-') for s in slide.shapes)
    text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    for bullet in content(count)['slides'][0]['bullets']:
        assert text.count(bullet) == 1


@pytest.mark.parametrize('variant', ['classic', 'split', 'focus'])
def test_dense_actions_choose_a_readable_page_over_narrow_exact_card_count(variant):
    template = parse_template(card_template((0, 4)), 'sample.pptx')
    for pattern in template['patterns']:
        pattern['role'] = 'content'
    outline = content(4)
    outline['slides'][0]['bullets'] = [
        ('Проверить готовность внедрения аналитики. Ответственный руководитель продукта. '
         'Согласовать очередь внедрений и подтвердить доступную инженерную мощность. ')*3
        for _ in range(4)
    ]
    page = compose(outline, template, variant, assets=[])['slides'][0]
    selected = template['patterns'][page['pattern_index']]
    assert not card_slots(selected['slots'], 1, 1)


def test_short_negative_native_labels_remain_visible_and_signed(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    visual = {'kind': 'bar', 'categories': ['Длинное название положительного региона', 'Длинное название отрицательного региона', 'Небольшое отрицательное отклонение'],
              'series': [{'name': 'Отклонение', 'values': [3_000_000, -4_000_000, -120_000]}], 'unit': 'RUB'}
    add_chart(slide, visual, [Inches(.5), Inches(.5), Inches(9), Inches(5)], 'Arial', ['0077FF'], '111111')
    target = tmp_path / 'native.pptx'
    prs.save(target)
    chart = next(s.chart for s in Presentation(target).slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [3_000_000, -4_000_000, -120_000]
    assert chart.series[0].invert_if_negative is False
    assert chart.value_axis.minimum_scale < -4_000_000
    for point in chart.series[0].points:
        label = point.data_label
        assert label.font.size.pt >= 12
        assert label._dLbl.find(qn('c:delete')) is None
        assert label._dLbl.find(qn('c:showVal')).get('val') == '1'


def test_plain_text_fallback_materializes_original_layout_footer_with_image_relation(tmp_path):
    from PIL import Image
    prs = Presentation(io.BytesIO(card_template((0,))))
    slide = prs.slides[0]
    blob = io.BytesIO()
    Image.new('RGB', (16, 8), '#0077ff').save(blob, format='PNG')
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(blob.getvalue()), Inches(.4), Inches(7), Inches(1), Inches(.3))
    # A layout owns its own relationship IDs. The explicit footer must remap
    # these, rather than referencing a nonexistent image relation on the slide.
    picture = group.shapes[0]
    rid = slide.slide_layout.part.relate_to(slide.part.related_part(picture._element.blipFill.blip.rEmbed),
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
    picture._element.blipFill.blip.rEmbed = rid
    slide.slide_layout.shapes._spTree.insert_element_before(group._element, 'p:extLst')
    source = tmp_path / 'source.pptx'
    prs.save(source)
    template = parse_template(source.read_bytes(), 'sample.pptx')
    deck = compose(content(1), template, 'split', assets=[])
    target = tmp_path / 'fallback.pptx'
    export_pptx(source, template, deck, target)
    exported = Presentation(target).slides[0]
    copied = [s for s in exported.shapes if s.name.startswith('template-footer-')]
    assert len(copied) == 1
    assert copied[0].shapes[0].image.blob == blob.getvalue()
    assert copied[0].left == group.left and copied[0].top == group.top
